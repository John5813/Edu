"""Saytda taqdimotni QO'LDA tahrirlash (bepul): matn, diagramma raqamlari, sahifalar tartibi.

Haqiqiy baza (vaqtinchalik) va haqiqiy brauzer (sahifalar chiziladi, PPTX yig'iladi); AI soxta va
tahrirlashda umuman chaqirilmasligi tekshiriladi.

    python test_sayt_qolda.py
"""
import asyncio, io, os, sys, tempfile, time, types

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import database.database as dbmod
from database.database import Database, init_db
from database import web_store
from services import web_jobs, web_kinds, web_decks  # noqa
from services.premium_presentation import html_render, html_slides, llm_client, manual_edit
import webapp
from aiohttp.test_utils import TestClient, TestServer

TMP = tempfile.mkdtemp()
dbmod.DATABASE_FILE = os.path.join(TMP, "test.db")
web_jobs.RESULTS_DIR = os.path.join(TMP, "results")


class FakeBot:
    async def send_document(self, *a, **kw): return types.SimpleNamespace()
    async def send_message(self, *a, **kw): return types.SimpleNamespace(message_id=1)
webapp.BOT, webapp.BOT_USERNAME = FakeBot(), "Edufayl_bot"


def head(title):
    return f'<section class="slide"><div class="head"><h2 class="title">{title}</h2><div class="rule"></div></div>'

TITLES = ["Raqamli iqtisodiyot", "Taqdimot rejasi", "Kirish va tushunchalar", "Onlayn to'lovlar hajmi",
          "Asosiy ko'rsatkichlar", "Xulosa"]
BODIES = [
    f'<section class="slide dark"><div class="body"><h1 class="title big">{TITLES[0]}</h1><p class="lead">Mavzu haqida</p></div></section>',
    head(TITLES[1]) + '<div class="body"><p class="lead">Reja</p></div></section>',
    head(TITLES[2]) + '<div class="body"><p class="lead">Asosiy fikr</p><div class="list">'
    '<div class="item"><span class="item-dot"></span><div class="item-text"><b>Birinchi.</b> Band matni</div></div>'
    '<div class="item"><span class="item-dot"></span><div class="item-text">Ikkinchi band</div></div></div></div></section>',
    head(TITLES[3]) + '<div class="body"><p class="lead">To\'lovlar o\'smoqda</p>'
    '<div class="chart" data-kind="bar" data-labels="2020,2021,2022" data-series="Trln so\'m: 38,61,94" data-unit="trln so\'m"></div></div></section>',
    head(TITLES[4]) + '<div class="body"><div class="cols cols-2"><div class="kpi"><div class="kpi-value">78%</div>'
    '<div class="kpi-label">Internet</div><div class="kpi-note">Foydalanuvchilar ulushi.</div></div>'
    '<div class="kpi"><div class="kpi-value">400+</div><div class="kpi-label">E-xizmatlar</div>'
    '<div class="kpi-note">Davlat xizmatlari.</div></div></div></div></section>',
    head(TITLES[5]) + '<div class="body"><p class="lead">Yakuniy fikr</p></div></section>',
]

AI = {"calls": 0, "armed": False}
def fake_write(topic, count, theme, language="uz", level=2, preferences="", source_text="", author="",
               progress_cb=None, outline_out=None, plan_cb=None):
    if outline_out is not None:
        outline_out.update(family="umumiy", outline=[{"title": t, "brief": t, "category": "kartalar"} for t in TITLES])
    return html_slides.build_pages(BODIES, theme, language)
html_slides.write_slides = fake_write
def no_ai(*a, **k):
    # Taqdimot yaratilayotganda (tuzatuvchi) javob bermaydi; qo'lda tahrirlashda esa umuman chaqirilmasligi kerak.
    if AI["armed"]:
        AI["calls"] += 1
    raise RuntimeError("AI yo'q")
llm_client._call_openrouter = no_ai
llm_client._call_openrouter_text = no_ai


async def wait_job(client, job_id, limit=240.0):
    end = time.time() + limit
    while time.time() < end:
        data = await (await client.get(f"/api/v1/jobs/{job_id}")).json()
        if data["job"]["status"] in ("done", "failed"):
            return data
        await asyncio.sleep(0.3)
    return data


async def login(client, origin, telegram_id):
    start = await (await client.post("/api/v1/auth/start", headers=origin)).json()
    await web_store.confirm_login(start["token"], telegram_id)
    await client.get("/api/v1/auth/poll", params={"token": start["token"]})


def blocks(page_html):
    """Brauzer ko'radigan tahrirlanadigan bloklar matni (sayt ham aynan shu ro'yxatni tuzadi)."""
    return manual_edit._browser_eval([(page_html, f"(sel) => ({manual_edit._LIST_JS})(sel).map((e) => e.textContent.trim())",
                                       manual_edit.EDITABLE)])[0]


def pptx_text(path):
    from pptx import Presentation
    return " ".join(sh.text_frame.text for sl in Presentation(path).slides for sh in sl.shapes if sh.has_text_frame)


async def main():
    if not html_render.available():
        check("brauzer o'rnatilgan", False, "Chromium kerak"); return
    from webapp.server import create_web_app
    await init_db()
    await Database.create_user(100, "ali", "Ali", "uz")
    await Database.create_user(200, "vali", "Vali", "uz")
    await Database.update_user_balance(100, 20_000)
    server = TestServer(create_web_app()); client = TestClient(server); await client.start_server()
    origin = {"Origin": f"http://{server.host}:{server.port}"}
    await login(client, origin, 100)

    print("1) Taqdimot")
    r = await client.post("/api/v1/jobs", json={"kind": "premium_presentation", "params": {"topic": "Raqamli iqtisodiyot", "slide_count": 4}}, headers=origin)
    job_id = (await r.json())["job"]["id"]
    job = await wait_job(client, job_id)
    check("taqdimot tayyor", job["job"]["status"] == "done", job)
    info = await (await client.get(f"/api/v1/jobs/{job_id}/deck")).json()
    deck = info["deck"]
    check("6 sahifa, reja bor, diagramma 4-sahifada", deck["count"] == 6 and deck["has_plan"]
          and [s["charts"] for s in deck["slides"]] == [0, 0, 0, 1, 0, 0], deck)
    balance0 = (await Database.get_user(100)).balance
    AI["armed"] = True

    print("2) Matnni joyida tahrirlash")
    page = await (await client.get(f"/api/v1/jobs/{job_id}/page/3")).json()
    check("sahifa HTML i va qoidasi beriladi", page["ok"] and "<section" in page["html"] and page["selector"] == manual_edit.EDITABLE
          and not page["locked"], {k: page.get(k) for k in ("ok", "locked", "selector")})
    texts = await asyncio.to_thread(blocks, page["html"])
    check("bloklar: sarlavha, asosiy fikr, bandlar", texts[:4] == ["Kirish va tushunchalar", "Asosiy fikr", "Birinchi. Band matni", "Ikkinchi band"], texts)
    plan_img = await (await client.get(f"/api/v1/jobs/{job_id}/slide/2")).read()
    edits = [{"k": 0, "html": "Kirish: yangi sarlavha"}, {"k": 2, "html": "<b>Birinchi.</b> Tuzatilgan <script>alert(1)</script>band"}]
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "text", "n": 3, "version": deck["version"], "count": len(texts), "edits": edits}, headers=origin)
    data = await r.json()
    check("saqlandi, versiya oshdi", r.status == 200 and data["deck"]["version"] == deck["version"] + 1, data)
    check("sarlavha ro'yxatda yangilandi", data["deck"]["slides"][2]["title"] == "Kirish: yangi sarlavha", data["deck"]["slides"][2])
    stored = web_decks.load(job_id)
    new_page = stored["pages"][2]
    check("sahifada yangi matn, qalin qism saqlandi, skript yo'q", "Tuzatilgan" in new_page and "<b>Birinchi.</b>" in new_page and "alert(1)" not in new_page)
    check("asl slayd ham yangilandi (AI qayta yozsa tuzatish yo'qolmaydi)", "Kirish: yangi sarlavha" in html_slides.source_of(new_page)
          and "Tuzatilgan" in html_slides.source_of(new_page), html_slides.source_of(new_page)[:300])
    check("reja sahifasi yangi sarlavhadan qayta yig'ildi", "Kirish: yangi sarlavha" in stored["pages"][1]
          and plan_img != await (await client.get(f"/api/v1/jobs/{job_id}/slide/2")).read())
    check("bepul: balans o'zgarmadi, AI chaqirilmadi", (await Database.get_user(100)).balance == balance0 and AI["calls"] == 0, AI)
    path = (await web_store.get_job(job_id))["result_path"]
    check("PPTX keyinroq yig'iladi (belgi qo'yildi)", stored.get("pptx_stale") is True)
    r = await client.get(f"/api/v1/jobs/{job_id}/file")
    body = await r.read()
    check("yuklab olishda PPTX yangi matn bilan", r.status == 200 and "Kirish: yangi sarlavha" in pptx_text(path), pptx_text(path)[:200])
    check("belgi olib tashlandi", web_decks.load(job_id).get("pptx_stale") is False)

    v = data["deck"]["version"]
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "text", "n": 3, "version": v - 1, "count": len(texts), "edits": edits}, headers=origin)
    check("eski versiyadan saqlash — 409", r.status == 409, r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "text", "n": 3, "version": v, "count": len(texts) + 5, "edits": edits}, headers=origin)
    check("bloklar soni mos kelmasa — rad etiladi", r.status in (400, 409), r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "text", "n": 2, "version": v, "count": 3, "edits": edits}, headers=origin)
    check("reja sahifasini qo'lda tahrirlab bo'lmaydi", r.status == 400, r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "text", "n": 3, "version": v, "count": len(texts), "edits": [{"k": 1, "html": "a" * 900}]}, headers=origin)
    check("juda uzun matn rad etiladi", r.status == 400, r.status)

    print("3) Diagramma raqamlari")
    page = await (await client.get(f"/api/v1/jobs/{job_id}/page/4")).json()
    charts = page["charts"]
    check("diagramma ma'lumoti o'qildi", len(charts) == 1 and charts[0]["labels"] == ["2020", "2021", "2022"]
          and charts[0]["series"][0]["values"] == [38, 61, 94] and charts[0]["kind"] == "bar", charts)
    img4 = await (await client.get(f"/api/v1/jobs/{job_id}/slide/4")).read()
    spec = {"kind": "line", "labels": ["2020", "2021", "2022", "2023"], "unit": "trln so'm", "source": "Markaziy bank",
            "series": [{"name": "To'lovlar", "values": [40, 65, 99, 142]}]}
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "chart", "n": 4, "k": 0, "version": v, "chart": spec}, headers=origin)
    data = await r.json()
    check("diagramma saqlandi", r.status == 200 and data["deck"]["version"] == v + 1, data)
    v = data["deck"]["version"]
    page = await (await client.get(f"/api/v1/jobs/{job_id}/page/4")).json()
    c = page["charts"][0]
    check("yangi qiymatlar va turi saqlandi", c["kind"] == "line" and c["labels"][-1] == "2023" and c["series"][0]["values"] == [40, 65, 99, 142]
          and c["source"] == "Markaziy bank", c)
    check("sahifadagi diagramma qayta chizildi", "142" in page["html"] and img4 != await (await client.get(f"/api/v1/jobs/{job_id}/slide/4")).read())
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "chart", "n": 4, "k": 0, "version": v,
                          "chart": {"kind": "bar", "labels": ["A"], "series": [{"name": "", "values": [1]}]}}, headers=origin)
    check("bitta qiymatli diagramma rad etiladi", r.status == 400, r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "chart", "n": 4, "k": 0, "version": v,
                          "chart": {"kind": "bar", "labels": ["A", "B"], "series": [{"name": "", "values": [1, "x"]}]}}, headers=origin)
    check("raqam bo'lmagan qiymat rad etiladi", r.status == 400, r.status)

    print("4) Sahifalar tartibi")
    before = [s["title"] for s in (await (await client.get(f"/api/v1/jobs/{job_id}/deck")).json())["deck"]["slides"]]
    img5 = await (await client.get(f"/api/v1/jobs/{job_id}/slide/5")).read()
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "order", "version": v, "order": [5, 3, 4, 4, 6]}, headers=origin)
    data = await r.json()
    titles = [s["title"] for s in data["deck"]["slides"]] if r.status == 200 else []
    check("tartib o'zgardi, nusxa qo'shildi", r.status == 200 and titles == [before[0], before[1], before[4], before[2], before[3], before[3], before[5]], (titles, before))
    v = data["deck"]["version"]
    check("suratlar ham ko'chdi", await (await client.get(f"/api/v1/jobs/{job_id}/slide/3")).read() == img5)
    stored = web_decks.load(job_id)
    check("reja yangi tartibda", stored["pages"][1].find(before[4]) < stored["pages"][1].find(before[2]) and before[4] in stored["pages"][1])
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "order", "version": v, "order": [3, 6]}, headers=origin)
    data = await r.json()
    check("sahifalar o'chirildi", r.status == 200 and data["deck"]["count"] == 4 and data["deck"]["has_plan"], data)
    v = data["deck"]["version"]
    check("ortiqcha suratlar o'chdi", (await client.get(f"/api/v1/jobs/{job_id}/slide/5")).status == 404)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "order", "version": v, "order": [3]}, headers=origin)
    check("reja bilan kamida 2 sahifa qoladi", r.status == 400, r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "order", "version": v, "order": [1, 3]}, headers=origin)
    check("muqovani ko'chirib bo'lmaydi", r.status == 400, r.status)
    check("hammasi bepul, AI chaqirilmadi", (await Database.get_user(100)).balance == balance0 and AI["calls"] == 0, AI)
    r = await client.get(f"/api/v1/jobs/{job_id}/file"); await r.read()
    check("PPTX 4 sahifa", len(__import__("pptx").Presentation(path).slides) == 4)

    print("5) Boshqa odam")
    other = TestClient(TestServer(create_web_app())); await other.start_server()
    o_origin = {"Origin": f"http://{other.server.host}:{other.server.port}"}
    await login(other, o_origin, 200)
    r = await other.get(f"/api/v1/jobs/{job_id}/page/3")
    check("boshqa odam sahifani ololmaydi", r.status == 404, r.status)
    r = await other.post(f"/api/v1/jobs/{job_id}/edit", json={"op": "order", "version": v, "order": [4, 3]}, headers=o_origin)
    check("boshqa odam tahrirlay olmaydi", r.status == 404, r.status)
    await other.close()
    await client.close()

    print()
    print("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}")


asyncio.run(main())
sys.exit(1 if FAILS else 0)
