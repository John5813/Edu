"""Rejadagi varaq raqamlarini va O'zbekiston qoidasini tekshiradi.

Ikki narsa sinaladi:

1. Reja qatorlari nuqtali chiziq bilan tugaydi va oxiridagi varaq
   raqami haqiqatga mos — hujjat PDF'ga aylantirilib, har sarlavha
   qaysi varaqqa tushgani bilan solishtiriladi.

2. Mavzu O'zbekiston iqtisodi yoki siyosatiga tegishli bo'lsa, kirish
   Prezident so'zlaridan boshlanadi, o'sha abzatsga snoska qo'yiladi va
   adabiyotlar ro'yxati talab qilingan tartibda bo'ladi.

Ishga tushirish:

    python test_reja_va_ozbekiston.py
"""

import asyncio
import glob
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, ".")
os.environ.setdefault("BOT_TOKEN", "test")

from services import uzbekistan  # noqa: E402
from services.doc_toc import TocPlan, _locate, _normalize  # noqa: E402

FAILS = []


def check(name, condition, detail=""):
    if condition:
        print(f"  ok   {name}")
    else:
        print(f"  XATO {name} — {detail}")
        FAILS.append(name)


# ────────────────────────────────────────────── 1. mavzuni aniqlash

def check_topic_detection():
    print("\n1) O'zbekiston mavzusini aniqlash")
    cases = [
        ("O'zbekiston Respublikasining milliy boyligi", True),
        ("Oʻzbekiston iqtisodiyotini rivojlantirish", True),
        ("Ozbekiston siyosiy tizimi", True),
        ("Узбекистан и мировая экономика", True),
        ("Uzbekistan's tax policy", True),
        ("Jahon savdosi va O'zbekiston eksporti", True),
        ("Alisher Navoiy ijodi", False),
        ("Marketing strategiyalari", False),
        ("Fotosintez jarayoni", False),
        ("Buxoro shahri tarixi", False),
    ]
    for topic, want in cases:
        got = uzbekistan.is_uzbek_topic(topic)
        check(f"«{topic[:38]}» -> {want}", got == want, f"aniqlangani {got}")


# ──────────────────────────────────────── 2. adabiyotlar tartibi

def check_reference_order():
    print("\n2) Adabiyotlar tartibi")
    refs = [
        "1. Karimov B. Iqtisodiyot asoslari. — Toshkent: Fan, 2019. — 200 b.",
        "2. O'zbekiston Respublikasining Konstitutsiyasi. — Toshkent, 2023.",
        "3. Sattorov D. Moliya // Jurnal. — 2024. — №2.",
        "4. Mirziyoyev Sh.M. Yangi O'zbekiston strategiyasi. — Toshkent, 2022.",
        "5. Yusupov A. Bank tizimi. — Toshkent, 2021.",
    ]
    ordered = uzbekistan.order_references(refs, "uz")
    check("birinchisi Prezident asari", "Mirziyoyev" in ordered[0], ordered[0])
    check("ikkinchisi Konstitutsiya", "Konstitutsiya" in ordered[1], ordered[1])
    years = [uzbekistan._year_of(r) for r in ordered[2:]]
    check("qolganlari yangidan eskiga", years == sorted(years, reverse=True), str(years))
    check("eski raqamlash olib tashlangan",
          not any(re.match(r"^\d+\.", r) for r in ordered), str(ordered[:1]))

    # Ro'yxatda Prezident ham, Konstitutsiya ham bo'lmasa — qo'shiladi.
    bare = uzbekistan.order_references(
        ["Yusupov A. Bank tizimi. — Toshkent, 2021."], "uz",
        "Mirziyoyev Sh.M. Oliy Majlisga Murojaatnoma. — Toshkent, 2025.")
    check("Prezident manbasi qo'shildi", "Mirziyoyev" in bare[0], bare[0])
    check("Konstitutsiya qo'shildi", "Konstitutsiya" in bare[1], bare[1])
    check("uchta yozuv", len(bare) == 3, str(len(bare)))


# ──────────────────────────────── 3. rejadagi varaq raqamlari

def _headings_in_pdf(pdf_path, headings):
    """Har sarlavha PDF'ning nechanchi varag'ida turganini qaytaradi."""
    import pymupdf

    found = {}
    with pymupdf.open(pdf_path) as document:
        pages = [[_normalize(l) for l in page.get_text().splitlines()]
                 for page in document]
    entries = [{"heading": _normalize(h)} for h in headings]
    for heading, number in zip(headings, _locate(entries, pages)):
        found[heading] = number
    return found


async def check_toc_numbers():
    print("\n3) Rejadagi varaq raqamlari")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    titles = [
        "1. Milliy boylik tushunchasi va uning tarkibi",
        "2. Iqtisodiyotni rivojlantirishda milliy boylikning o'rni",
        "3. Milliy boylikni ko'paytirish omillari",
    ]
    headings = ["KIRISH"] + titles + ["XULOSA VA TAKLIFLAR"]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.59)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run("TITUL VARAQ").bold = True
    doc.add_page_break()

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.add_run("REJA").bold = True
    doc.add_paragraph()

    plan = TocPlan(width_in=8.5 - 1.18 - 0.59)
    plan.line(doc, "Kirish", "KIRISH")
    for title in titles:
        plan.line(doc, title, title)
    plan.line(doc, "Xulosa va takliflar", "XULOSA VA TAKLIFLAR")
    doc.add_page_break()

    for heading, pages in zip(headings, (2, 4, 3, 4, 2)):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run(heading).bold = True
        for _ in range(pages * 21):
            doc.add_paragraph("Milliy boylik mamlakat iqtisodiyotining asosidir. " * 2)

    work_dir = tempfile.mkdtemp(prefix="toc_test_")
    path = os.path.join(work_dir, "reja.docx")
    doc.save(path)
    filled = await plan.fill(doc, path)
    check("raqamlar qo'yildi", filled)

    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", work_dir, path], capture_output=True, timeout=300)
    pdf = glob.glob(os.path.join(work_dir, "*.pdf"))
    if not pdf:
        check("PDF yaratildi", False, "soffice javob bermadi")
        return

    import pymupdf
    with pymupdf.open(pdf[0]) as document:
        toc_text = document[1].get_text()

    check("nuqtali chiziq bor", "......" in toc_text, toc_text[:60])

    written = {}
    for line in toc_text.splitlines():
        match = re.match(r"^(.*?)\.{4,}(\d+)$", line.strip())
        if match:
            written[_normalize(match.group(1))] = int(match.group(2))
    check("har bir qatorda raqam bor", len(written) == len(headings),
          f"{len(written)} / {len(headings)}")

    real = _headings_in_pdf(pdf[0], headings)
    for heading in headings:
        key = _normalize(heading)
        check(f"«{heading[:34]}» varag'i to'g'ri",
              written.get(key) == real.get(heading),
              f"rejada {written.get(key)}, aslida {real.get(heading)}")


# ──────────────────────────── 4. kirish Prezident so'zlaridan

async def check_presidential_opening():
    print("\n4) Kirishdagi Prezident abzatsi")
    from services.document_service import DocumentService

    opening = ("O'zbekiston Respublikasi Prezidenti Shavkat Mirziyoyev Oliy "
               "Majlisga Murojaatnomasida milliy boylikdan oqilona foydalanish "
               "masalasiga alohida to'xtaldi. ") * 3
    source = "Mirziyoyev Sh.M. Oliy Majlisga Murojaatnoma. — Toshkent, 2025."
    body = "Milliy boylik iqtisodiyotning moddiy asosini tashkil etadi. " * 20

    content = {
        "introduction": "Mavzuning dolzarbligi ortib bormoqda. " * 14,
        "intro_points": {f"point_{i}": f"{i}-band mazmuni." for i in range(1, 7)},
        "presidential_opening": {"text": opening.strip(), "source": source},
        "chapters": [{
            "title": "Milliy boylik tushunchasi",
            "subsections": [
                {"number": "1.1", "title": "Mohiyati", "content": body},
                {"number": "1.2", "title": "Tarkibi", "content": body},
            ],
        }],
        "conclusion": "Xulosalar chiqarildi. " * 14,
        "references": uzbekistan.order_references([
            "Karimov B. Iqtisodiyot. — Toshkent: Fan, 2019. — 200 b.",
            "Sattorov D. Moliya // Jurnal. — 2024. — №2.",
        ], "uz", source),
    }

    service = DocumentService()
    path = await service.create_course_work(
        "O'zbekiston milliy boyligi", content, "Toshmatov T.", "uz")
    check("hujjat yaratildi", os.path.exists(path), path)

    work_dir = tempfile.mkdtemp(prefix="uz_test_")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", work_dir, path], capture_output=True, timeout=300)
    pdf = glob.glob(os.path.join(work_dir, "*.pdf"))
    if not pdf:
        check("PDF yaratildi", False, "soffice javob bermadi")
        return

    import pymupdf
    with pymupdf.open(pdf[0]) as document:
        intro_page = document[2].get_text()

    lines = [l.strip() for l in intro_page.splitlines() if l.strip()]
    check("kirish sarlavhasi birinchi", lines[0].upper().startswith("KIRISH"), lines[0])
    check("Prezident abzatsi sarlavhadan keyin darhol",
          "Prezidenti Shavkat Mirziyoyev" in " ".join(lines[1:4]), lines[1][:70])
    check("abzatsdan keyin snoska belgisi",
          re.search(r"to['’]xtaldi\.\d", intro_page.replace("\n", " ")) is not None)
    check("snoska varaq pastida", source.split(".")[0] in intro_page.splitlines()[-2],
          intro_page.splitlines()[-2][:70])

    os.remove(path)


async def _audit_toc(path):
    """Rejadagi raqamlarni hujjatning o'zi bilan solishtiradi."""
    work_dir = tempfile.mkdtemp(prefix="audit_")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", work_dir, path], capture_output=True, timeout=300)
    produced = glob.glob(os.path.join(work_dir, "*.pdf"))
    if not produced:
        return None, {}

    import pymupdf
    with pymupdf.open(produced[0]) as document:
        pages = [page.get_text() for page in document]

    written = {}
    for line in pages[1].splitlines():
        match = re.match(r"^(.*?)\.{4,}(\d+)$", line.strip())
        if match:
            written[_normalize(match.group(1))] = int(match.group(2))

    real = {}
    for number, text in enumerate(pages[2:], 3):
        for line in text.splitlines():
            flat = _normalize(line)
            for key in written:
                if key not in real and flat.startswith(key):
                    real[key] = number

    wrong = {k: (v, real.get(k)) for k, v in written.items() if v != real.get(k)}
    return written, wrong


async def check_short_documents():
    """Qisqa ishda ham raqam to'g'ri chiqsinmi.

    Ilgari qisqa referatda bir nechta bo'lim bitta varaqqa sig'ib
    qolardi va o'sha varaq reja deb qabul qilinib, raqamlar bir varaqqa
    surilib ketardi.
    """
    print("\n5) Qisqa hujjatlarda raqamlar")
    from services.document_service import DocumentService

    body = "Milliy boylik iqtisodiyotning moddiy asosini tashkil etadi. " * 20
    service = DocumentService()
    for sections in (2, 4):
        content = {
            "language": "uz",
            "sections": (
                [{"title": "Kirish", "content": "Mavzuning dolzarbligi ortmoqda. " * 12}]
                + [{"title": f"{i}-bo'lim mavzusi", "content": body}
                   for i in range(1, sections + 1)]
                + [{"title": "Xulosa", "content": "Xulosalar chiqarildi. " * 12}]
            ),
            "references": ["Karimov B. Iqtisodiyot. — Toshkent, 2019."],
        }
        path = await service.create_referat("Milliy boylik", content)
        written, wrong = await _audit_toc(path)
        check(f"{sections} bo'limli referat rejasi to'liq",
              written is not None and len(written) == sections + 3,
              f"{len(written or {})} qator")
        check(f"{sections} bo'limli referat raqamlari to'g'ri", not wrong, str(wrong))
        os.remove(path)


async def main():
    check_topic_detection()
    check_reference_order()
    await check_toc_numbers()
    await check_presidential_opening()
    await check_short_documents()

    print()
    if FAILS:
        print(f"❌ {len(FAILS)} ta tekshiruv o'tmadi: {', '.join(FAILS)}")
        return 1
    print("✅ Reja raqamlari ham, O'zbekiston qoidasi ham ishlayapti.")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
