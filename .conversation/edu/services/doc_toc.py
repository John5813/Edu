"""Rejaga nuqtali chiziq va haqiqiy varaq raqamini qo'yadi.

Ilgari reja shunchaki sarlavhalar ro'yxati edi — o'qituvchi esa har
bo'lim nechanchi varaqdan boshlanishini ko'rishni kutadi:

    Kirish ...................................................... 3
    1. Milliy boylik tushunchasi ................................. 6

Ikki qismdan iborat:

1. `TocPlan.line()` — reja qatorini yozadi. Nuqtalar qo'lda emas,
   Word'ning o'ng tabulyatsiyasi orqali chiziladi: shunda ular chekkaga
   aniq yetib boradi va shrift o'zgarsa ham buzilmaydi.

2. `TocPlan.fill()` — hujjat to'liq yozilib saqlangandan keyin uni
   LibreOffice bilan PDF'ga aylantiradi, har sarlavha qaysi varaqqa
   tushganini topadi va raqamlarni o'sha qatorlarga qo'yadi. Varaq
   raqamini oldindan bilib bo'lmaydi — matn hajmi, rasm va jadvallar
   uni surib yuboradi, shuning uchun hisob oxirida qilinadi.

Agar LibreOffice topilmasa yoki PDF o'qilmasa, nuqtali chiziq ham,
raqam ham olib tashlanadi: reja eskidagidek oddiy ro'yxat bo'lib
qoladi, hujjat esa baribir tayyor bo'ladi.
"""

import asyncio
import glob
import logging
import os
import re
import shutil
import subprocess
import tempfile

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt

logger = logging.getLogger(__name__)

# Sarlavhani PDF matnida qidirganda shuncha belgisi solishtiriladi.
# Uzun sarlavha varaqda ikki qatorga o'ralishi mumkin, shuning uchun
# butun matnni emas, boshini taqqoslaymiz.
_MATCH_CHARS = 42

# LibreOffice katta hujjatni ham shu vaqt ichida aylantirishi kerak.
_CONVERT_TIMEOUT = 300


def _normalize(text: str) -> str:
    """Taqqoslash uchun matnni soddalashtiradi.

    Reja va sarlavha bir xil matn bo'lsa ham, ular bosh harf, apostrof
    turi (ʻ, ', ') va bo'shliqlar bilan farq qiladi — PDF'dan o'qilgan
    matnda esa qator uzilishlari bor. Hammasini bir ko'rinishga
    keltiramiz.
    """
    text = (text or "").lower()
    text = re.sub(r"[‘’ʻʼ′'`´]", "", text)
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return text.strip()


class TocPlan:
    """Reja qatorlarini yig'adi va oxirida raqamlarini to'ldiradi."""

    def __init__(self, width_in: float, size: float = 14.0):
        # Nuqtalar shu nuqtagacha chiziladi — o'ng chekka.
        self.width_in = width_in
        self.size = size
        self._entries = []

    def line(self, doc, text: str, heading: str = "", *, bold: bool = False,
             indent: float = 0.0, spacing: float = 1.5,
             size: float = 0.0, space_after: float = None):
        """Bitta reja qatorini qo'shadi va paragrafni qaytaradi.

        `heading` — hujjat ichidagi sarlavha matni. U reja yozuvidan
        farq qilishi mumkin (masalan "Kirish" va "KIRISH"), shuning
        uchun alohida beriladi; berilmasa, reja matnining o'zi olinadi.
        """
        size = size or self.size
        para = doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.line_spacing = spacing
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if indent:
            fmt.left_indent = Inches(indent)
        if space_after is not None:
            fmt.space_after = Pt(space_after)
        # Tabulyatsiya chap chekinishdan qat'i nazar o'ng chekkada turadi.
        fmt.tab_stops.add_tab_stop(
            Inches(self.width_in), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )

        title_run = para.add_run(text)
        title_run.font.size = Pt(size)
        title_run.font.bold = bold
        title_run.font.name = "Times New Roman"

        tab_run = para.add_run("\t")
        tab_run.font.size = Pt(size)
        tab_run.font.name = "Times New Roman"

        page_run = para.add_run("")
        page_run.font.size = Pt(size)
        page_run.font.bold = bold
        page_run.font.name = "Times New Roman"

        self._entries.append({
            "heading": _normalize(heading or text),
            "tab": tab_run,
            "page": page_run,
        })
        return para

    def __len__(self) -> int:
        return len(self._entries)

    # ────────────────────────────────────────── raqamlarni to'ldirish

    async def fill(self, doc, path: str) -> bool:
        """Hujjatni PDF orqali o'lchab, varaq raqamlarini yozadi.

        `path` — allaqachon saqlangan .docx. Raqamlar qo'yilgach hujjat
        o'sha yo'lga qayta saqlanadi. Muvaffaqiyatli bo'lsa `True`.
        """
        if not self._entries:
            return True

        try:
            pages = await asyncio.to_thread(_page_lines, path)
        except Exception as exc:
            logger.warning("Reja uchun PDF o'lchovi olinmadi: %s", exc)
            pages = []

        numbers = _locate(self._entries, pages) if pages else []
        if not numbers:
            self._strip(doc, path)
            return False

        for entry, number in zip(self._entries, numbers):
            entry["page"].text = str(number)

        await asyncio.to_thread(doc.save, path)
        return True

    def _strip(self, doc, path: str) -> None:
        """O'lchov chiqmasa, nuqtali chiziqni butunlay olib tashlaydi.

        Raqamsiz nuqtalar chekkaga qarab cho'zilib, reja buzilgandek
        ko'rinardi — bunday holda oddiy ro'yxat qolgani ma'qul.
        """
        for entry in self._entries:
            entry["tab"].text = ""
            entry["page"].text = ""
        try:
            doc.save(path)
        except Exception as exc:
            logger.warning("Reja tozalangandan keyin saqlanmadi: %s", exc)


def _page_lines(docx_path: str) -> list:
    """Har varaqdagi matn qatorlarini soddalashtirilgan ko'rinishda qaytaradi."""
    work_dir = tempfile.mkdtemp(prefix="toc_")
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf",
             "--outdir", work_dir, docx_path],
            check=True, timeout=_CONVERT_TIMEOUT, capture_output=True,
        )
        produced = glob.glob(os.path.join(work_dir, "*.pdf"))
        if not produced:
            logger.warning("Reja uchun PDF yaratilmadi: %s", docx_path)
            return []

        try:
            import pymupdf as pdf_reader
        except ImportError:
            import fitz as pdf_reader

        pages = []
        with pdf_reader.open(produced[0]) as document:
            for page in document:
                lines = [_normalize(line) for line in page.get_text().splitlines()]
                pages.append([line for line in lines if line])
        return pages
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired,
            FileNotFoundError) as exc:
        logger.warning("LibreOffice reja uchun ishlamadi: %s", exc)
        return []
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def _starts_with(line: str, heading: str) -> bool:
    """Qator shu sarlavhadan boshlanadimi.

    Sarlavhani matn ichidan emas, qator boshidan qidiramiz: bo'lim nomi
    kirish matnida ham tilga olinishi mumkin, sarlavha esa har doim
    yangi qatordan boshlanadi.
    """
    prefix = heading[:_MATCH_CHARS]
    return bool(prefix) and line.startswith(prefix)


def _matches_on(lines: list, entries: list) -> int:
    """Shu varaqda nechta reja yozuvi uchraydi."""
    return sum(
        1 for entry in entries
        if any(_starts_with(line, entry["heading"]) for line in lines)
    )


def _body_start(entries: list, pages: list) -> int:
    """Hujjat tanasi boshlanadigan varaq indeksi.

    Rejaning o'zida barcha sarlavhalar bor, shuning uchun qidiruvni
    undan keyin boshlash kerak — aks holda hamma bo'lim reja turgan
    varaqda "topilgan" bo'lardi.

    Rejani oxirgi yozuvi bo'yicha topamiz. Reja — hujjatda "Foydalanilgan
    adabiyotlar" birinchi marta uchraydigan joy: ro'yxatning o'zi esa
    hujjat oxirida. Qisqa referatda bir nechta bo'lim bitta varaqqa
    sig'ib qolishi mumkin, shuning uchun eng zich varaqni olish yetarli
    emas edi. Reja ikki varaqqa cho'zilsa ham shu usul ishlaydi: oxirgi
    yozuv ikkinchi varaqda bo'ladi.
    """
    counts = [_matches_on(lines, entries) for lines in pages]
    if not counts or max(counts) < 2:
        return 0

    tail = entries[-1]["heading"]
    for index, lines in enumerate(pages):
        if counts[index] >= 2 and any(_starts_with(line, tail) for line in lines):
            return index + 1

    # Oxirgi yozuv topilmasa (masalan "Ilovalar" hujjatda yo'q) — eng zich
    # varaqdan keyin boshlaymiz.
    return counts.index(max(counts)) + 1


def _locate(entries: list, pages: list) -> list:
    """Har bir yozuv uchun varaq raqamini topadi.

    Sarlavhalar hujjatda reja tartibida keladi, shuning uchun qidiruv
    oldinga siljib boradi — bu matn ichida takrorlangan nomga tushib
    qolmaslikni ta'minlaydi.
    """
    numbers = []
    cursor = _body_start(entries, pages)
    last = cursor + 1
    for entry in entries:
        found = 0
        for index in range(cursor, len(pages)):
            if any(_starts_with(line, entry["heading"]) for line in pages[index]):
                found = index + 1
                cursor = index
                break
        if not found:
            # Topilmasa oldingi raqam qoladi: reja hech bo'lmasa
            # ketma-ket va ishonchli ko'rinadi.
            logger.info("Rejada sarlavha topilmadi: %s", entry["heading"][:60])
            found = last
        numbers.append(found)
        last = found

    return numbers
