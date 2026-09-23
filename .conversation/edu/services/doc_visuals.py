"""Hujjatdagi rasm va jadvallarni raqamlaydi hamda izohlaydi.

O'zbek ilmiy ishlarida har bir rasm va jadval nomlanadi — "1-rasm.",
"2-jadval." — va ostida uning nimani ko'rsatayotgani bir abzatsda
tushuntiriladi. Ilgari bu faqat AI rejalashtirgan diagrammalarda
bajarilardi: qo'shimcha xizmat sifatida qo'yilgan rasm "Rasm. Bo'lim
nomi" deb, qiyoslash jadvali esa umuman nomsiz tushardi.

Raqam HUJJAT bo'yicha ketma-ket borishi kerak, lekin uni har bir
chizuvchiga parametr qilib uzatish o'nlab joyni o'zgartirishni talab
qilardi. Shuning uchun hisoblagich hujjat obyektining o'ziga bog'lanadi:
har hujjat o'zinikini oladi va bir vaqtda yaratilayotgan ikki hujjat
bir-birining raqamini olmaydi.
"""

import weakref
from typing import Optional

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

_LABELS = {
    "uz": {"figure": "{n}-rasm", "table": "{n}-jadval",
           "formula": "Formula {n}", "scheme": "{n}-rasm"},
    "ru": {"figure": "Рисунок {n}", "table": "Таблица {n}",
           "formula": "Формула {n}", "scheme": "Рисунок {n}"},
    "en": {"figure": "Figure {n}", "table": "Table {n}",
           "formula": "Formula {n}", "scheme": "Figure {n}"},
}

# python-docx hujjati lug'at kaliti bo'la olmaydi (u `__eq__` ni qayta
# ta'riflagan, `__hash__` ni esa yo'q qilgan), shuning uchun kalit
# sifatida obyekt identifikatori olinadi. Hujjat yo'qolganda yozuv
# o'zi tozalanadi — ro'yxat cheksiz o'smaydi.
_REGISTRY = {}


class Numbering:
    """Bitta hujjatdagi rasm, jadval va formula raqamlari."""

    def __init__(self, language: str = "uz"):
        self.language = language if language in _LABELS else "uz"
        self.figures = 0
        self.tables = 0
        self.formulas = 0

    # ───────────────────────────────────────────────── raqam va matn

    def _label(self, kind: str, number: int) -> str:
        return _LABELS[self.language][kind].format(n=number)

    def next_figure(self) -> int:
        self.figures += 1
        return self.figures

    def next_table(self) -> int:
        self.tables += 1
        return self.tables

    def next_formula(self) -> int:
        self.formulas += 1
        return self.formulas

    # ─────────────────────────────────────────────────── yozuvchilar

    def figure_caption(self, doc, title: str = "") -> int:
        """Rasm ostidagi nom: "1-rasm. Nomi"."""
        number = self.next_figure()
        self._caption(doc, self._label("figure", number), title,
                      WD_ALIGN_PARAGRAPH.CENTER)
        return number

    def table_caption(self, doc, title: str = "") -> int:
        """Jadval USTIDAGI nom: "1-jadval. Nomi".

        Jadvalning nomi tepasida turadi — o'zbek ishlarida shunday
        rasmiylashtiriladi; rasmniki esa ostida.
        """
        number = self.next_table()
        self._caption(doc, self._label("table", number), title,
                      WD_ALIGN_PARAGRAPH.RIGHT)
        return number

    def _caption(self, doc, label: str, title: str, align) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(2)
        text = f"{label}. {str(title).strip()}" if str(title).strip() else label
        run = paragraph.add_run(text)
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.name = "Times New Roman"

    def note(self, doc, text: str) -> bool:
        """Rasm yoki jadval ostidagi bir abzats izoh."""
        value = " ".join(str(text or "").split())
        if not value:
            return False
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(value)
        run.font.size = Pt(14)
        run.font.name = "Times New Roman"
        return True


def of(doc, language: str = "uz") -> Numbering:
    """Shu hujjatning hisoblagichi — birinchi so'ralganda yaratiladi."""
    key = id(doc)
    counter: Optional[Numbering] = _REGISTRY.get(key)
    if counter is None:
        counter = Numbering(language)
        _REGISTRY[key] = counter
        try:
            weakref.finalize(doc, _REGISTRY.pop, key, None)
        except TypeError:
            # Hujjatga zaif havola qo'yib bo'lmasa, yozuv qo'lda
            # o'chiriladi (`release`).
            pass
    return counter


def release(doc) -> None:
    """Hujjat saqlangandan keyin hisoblagichni bo'shatadi."""
    _REGISTRY.pop(id(doc), None)
