"""HTML slaydli premium taqdimot tizimini tekshiradi.

Eski ikki tizimda joylashuv koddan kelardi: avval AI koordinata aytar,
keyin 24 ta qat'iy qolip bo'lardi. Endi AI butun slaydni HTML/CSS/SVG
qilib chizadi, biz uni brauzerda 1920×1080 da suratga olamiz. Kod
faqat qobiq shartlarini va joylashuv kategoriyalarini beradi.

Shu fayl aynan o'sha kafolatlarni sinaydi:

1. Promptda qobiq shartlari bor, qat'iy varaqma-varaq shablon yo'q.
2. Chala kelgan HTML hujjat tashlab yuboriladi.
3. Reja kelmasa ham slaydlar xilma-xil kategoriyalarda bo'ladi.
4. Brauzer HTML ni haqiqatan 1920×1080 PNG qiladi.
5. PPTX slaydi TAHRIRLANADI: matn — matn qutisi, jadval — jadval,
   blok — shakl. Faqat diagramma rasm bo'lib qoladi.

    python test_html_taqdimot.py
"""

import asyncio
import os
import subprocess
import sys

sys.path.insert(0, ".")
os.environ.setdefault("BOT_TOKEN", "test")

from services.premium_presentation import (  # noqa: E402
    html_extract, html_images, html_render, html_slides, llm_client, themes)

FAILS = []


def check(name, condition, detail=""):
    print(f"  ok   {name}" if condition else f"  XATO {name} — {detail}")
    if not condition:
        FAILS.append(name)


def _page(title="Sinov", extra=""):
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;padding:90px;background:#FFFFFF;
    font-family:'DejaVu Sans','Liberation Sans',Arial,sans-serif}}
    h1{{font-size:64px;color:#10241C}}</style></head>
    <body><h1>{title}</h1>{extra}</body></html>"""


def check_handler_names():
    """Taqdimot oqimida aniqlanmagan o'zgaruvchi qolmaganini tekshiradi.

    Tizim qayta yozilganda eski o'zgaruvchiga qilingan bitta havola
    qolib ketgan edi: slaydlar yaratilib, fayl tayyor bo'lgandan keyin
    "yuborilmoqda" xabarida `NameError` chiqar va mijoz to'lagan puli
    qaytarilar edi. Sintaksis tekshiruvi bunday xatoni ko'rmaydi.
    """
    print("\n0) Oqimda aniqlanmagan nom yo'qligi")
    try:
        from pyflakes import api, reporter
    except ImportError:
        print("  o'tkazildi   pyflakes o'rnatilmagan")
        return

    import io

    files = [
        "bot/handlers/premium_presentation.py",
        "services/premium_presentation/html_render.py",
        "services/premium_presentation/html_slides.py",
        "services/premium_presentation/html_extract.py",
        "services/premium_presentation/pptx_build.py",
    ]
    for path in files:
        out, err = io.StringIO(), io.StringIO()
        api.checkPath(path, reporter.Reporter(out, err))
        undefined = [line for line in out.getvalue().splitlines()
                     if "undefined name" in line]
        check(f"{os.path.basename(path)} da aniqlanmagan nom yo'q",
              not undefined, "; ".join(undefined[:2]))


def check_prompt():
    print("\n1) Qobiq qoidalari")
    theme = themes.get("zumrad")
    rules = html_slides.shell_rules(theme, "uz")

    check("o'lcham aytilgan", "1920px × 1080px" in rules, rules[:80])
    check("tashqi fayl taqiqlangan",
          "Google Fonts" in rules and "YO'Q" in rules)
    check("shrift qat'iy", html_slides.FONT_STACK in rules)
    check("tanlangan rang promptda", theme.accent in rules, theme.accent)
    check("matn kesilishi taqiqlangan", "ellipsis" in rules)
    check("ajratuvchi aytilgan", html_slides.MARKER in rules)
    check("bo'sh joy egasi taqiqlangan", "Lorem ipsum" in rules)

    # Qat'iy shablon bo'lmasligi kerak: kategoriya — yo'nalish, o'rin
    # koordinatasi emas.
    catalogue = html_slides.catalogue_text()
    check("kategoriyalar ro'yxati bor", len(html_slides.CATEGORY_KEYS) >= 12,
          str(len(html_slides.CATEGORY_KEYS)))
    check("kategoriyada koordinata yo'q",
          not any(word in catalogue.lower()
                  for word in ("px", "x:", "y:", "dyuym")), catalogue[:80])
    check("diagramma kategoriyasi bor", "diagramma" in html_slides.CATEGORY_KEYS)
    check("jadval kategoriyasi bor", "jadval" in html_slides.CATEGORY_KEYS)


def check_split():
    print("\n2) Javobni slaydlarga ajratish")
    good = _page("Bir") + f"\n{html_slides.MARKER}\n" + _page("Ikki")
    check("ikkita slayd ajratildi", len(html_slides.split_slides(good)) == 2)

    truncated = _page("Bir") + f"\n{html_slides.MARKER}\n" + "<html><body><h1>Chala"
    parts = html_slides.split_slides(truncated)
    check("chala slayd tashlandi", len(parts) == 1, str(len(parts)))

    fenced = "```html\n" + _page("Bir") + "\n```"
    check("markdown ramkasi olib tashlandi",
          len(html_slides.split_slides(fenced)) == 1)

    thinking = "<think>o'ylayapman</think>" + _page("Bir")
    parts = html_slides.split_slides(thinking)
    check("o'ylash bloki olib tashlandi",
          len(parts) == 1 and "o'ylayapman" not in parts[0])
    check("bo'sh javobdan slayd chiqmaydi", html_slides.split_slides("") == [])


def check_outline():
    print("\n3) Reja va kategoriyalar")

    def fake(system, user, temperature=0.7, max_tokens=16000):
        return {"slides": [{"brief": f"{i}-slayd", "category": "kartalar"}
                           for i in range(1, 9)]}

    original = llm_client._call_openrouter
    try:
        llm_client._call_openrouter = fake
        outline = html_slides.plan_outline("Raqamli iqtisodiyot", 8, "uz")
    finally:
        llm_client._call_openrouter = original

    check("reja to'liq", len(outline) == 8, str(len(outline)))
    check("birinchisi muqova", outline[0]["category"] == "muqova")
    check("oxirgisi yakun", outline[-1]["category"] == "yakun")
    check("ketma-ket takror yo'q",
          all(a["category"] != b["category"]
              for a, b in zip(outline, outline[1:])),
          str([o["category"] for o in outline]))

    # AI javob bermasa ham reja tuzilishi kerak.
    def broken(*a, **k):
        raise RuntimeError("model javob bermadi")

    try:
        llm_client._call_openrouter = broken
        fallback = html_slides.plan_outline("Mavzu", 6, "uz")
    finally:
        llm_client._call_openrouter = original

    check("reja kelmasa ham slaydlar bor", len(fallback) == 6)
    check("zaxira rejada ham xilma-xillik",
          len({o["category"] for o in fallback}) >= 4,
          str([o["category"] for o in fallback]))


def check_writer():
    print("\n4) Slaydlarni bo'laklab yozish")
    calls = []

    def fake(system, user, temperature=0.7, max_tokens=1800):
        calls.append(user)
        pages = [_page(f"Slayd {len(calls)}.{i}") for i in range(3)]
        return f"\n{html_slides.MARKER}\n".join(pages)

    def fake_plan(system, user, temperature=0.7, max_tokens=16000):
        return {"slides": [{"brief": f"{i}", "category": "kartalar"}
                           for i in range(1, 7)]}

    orig_text, orig_json = llm_client._call_openrouter_text, llm_client._call_openrouter
    try:
        llm_client._call_openrouter_text = fake
        llm_client._call_openrouter = fake_plan
        pages = html_slides.write_slides("Mavzu", 6, themes.get("ko'k"), "uz")
    finally:
        llm_client._call_openrouter_text = orig_text
        llm_client._call_openrouter = orig_json

    check("so'ralgan slayd soni chiqdi", len(pages) == 6, str(len(pages)))
    check("bo'laklab so'raldi", len(calls) == 2, f"{len(calls)} ta so'rov")
    check("rejadagi o'rin ko'rsatilgan", "→" in calls[0], calls[0][:60])
    check("ikkinchi bo'lak avvalgisini biladi",
          "takrorlama" in calls[1].lower(), calls[1][:80])


def check_browser_setup():
    print("\n5) Brauzer tayyorligi")
    original_globs = html_render._BROWSER_GLOBS
    original_system = html_render._SYSTEM_BROWSERS
    original_run = subprocess.run
    calls = []

    def fake_run(command, **kwargs):
        calls.append(list(command))

        class Result:
            # Brauzerni yuklab olish tushadi, kutubxonalar esa o'rnatiladi.
            returncode = 0 if "install-deps" in command else 1
            stdout = ""
            stderr = "tarmoq yo'q"

        return Result()

    try:
        # Brauzersiz server holatini yasaymiz.
        html_render._BROWSER_GLOBS = ("yo-q-*/chrome",)
        html_render._SYSTEM_BROWSERS = ()
        html_render._install_done = False
        subprocess.run = fake_run

        check("brauzersiz serverda topilmaydi", not html_render._executable())
        check("available() yolg'on aytmaydi", not html_render.available())

        failure = html_render.install_browser()
        check("o'zi o'rnatishga urinadi", bool(calls), str(calls))
        check("to'g'ri buyruq chaqiriladi",
              calls and calls[0][-3:] == ["install", "chromium"] or
              calls and "install" in calls[0], str(calls[:1]))
        check("xato matni qaytariladi", "tarmoq" in failure, failure)

        # Ikkinchi marta qayta yuklab olishga urinmasin.
        count = len(calls)
        html_render.install_browser()
        check("ikki marta yuklab olmaydi", len(calls) == count, str(len(calls)))

        # Brauzer bor, lekin GTK kutubxonalari yo'q — eng ko'p uchraydigan
        # holat. Kod kutubxonalarni o'rnatib, qayta sinashi kerak.
        calls.clear()
        html_render._deps_done = False
        failure = html_render.install_dependencies()
        check("kutubxonalar o'rnatiladi", bool(calls), str(calls))
        check("avval install-deps sinaladi",
              calls and "install-deps" in calls[0], str(calls[:1]))
        check("kutubxona xatosi qaytmadi", failure == "", failure)
    finally:
        subprocess.run = original_run
        html_render._BROWSER_GLOBS = original_globs
        html_render._SYSTEM_BROWSERS = original_system
        html_render._install_done = False
        html_render._deps_done = False

    check("brauzer topilganda o'rnatish so'ralmaydi",
          html_render.prepare(install=False) == "" if html_render._executable()
          else True)


def check_shot():
    print("\n6) Brauzerda suratga olish")
    if not html_render.available():
        check("brauzer o'rnatilgan", False, html_render._INSTALL_HINT)
        return False

    images = html_render.shoot([_page("Birinchi slayd")], out_dir="temp")
    check("slayd suratga olindi", len(images) == 1, str(len(images)))
    if not images:
        return True
    try:
        from PIL import Image
        with Image.open(images[0]) as picture:
            size = picture.size
            colours = picture.convert("RGB").getcolors(maxcolors=200000) or []
        check("surat 1920×1080", size == (1920, 1080), str(size))
        check("slayd bo'sh emas", len(colours) > 2, str(len(colours)))
    finally:
        for image in images:
            if os.path.exists(image):
                os.remove(image)
    return True


def check_editable():
    print("\n7) Tahrirlanadigan PPTX")
    if not html_render.available():
        check("brauzer o'rnatilgan", False, html_render._INSTALL_HINT)
        return

    theme = themes.get("zumrad")
    font = html_slides.FONT_STACK
    slide = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;padding:84px;background:#FFFFFF;
    font-family:{font};overflow:hidden}}
    h2{{font-size:52px;color:#{theme.heading};font-weight:700}}
    .bar{{width:180px;height:9px;background:#{theme.accent};margin:26px 0 40px}}
    .card{{background:#{theme.accent_soft};border-radius:20px;padding:30px;
    width:520px}}
    .v{{font-size:56px;font-weight:700;color:#{theme.accent}}}
    table{{border-collapse:collapse;width:700px;margin-top:40px}}
    th{{background:#{theme.accent};color:#FFFFFF;font-size:22px;padding:14px}}
    td{{font-size:21px;color:#{theme.body};padding:13px}}
    </style></head><body>
    <h2>Bandlik ko'rsatkichlari<br>ikkinchi qator</h2>
    <div class="bar"></div>
    <div class="card"><div class="v">14,4 mln</div></div>
    <table><tr><th>Hudud</th><th>2026</th></tr>
    <tr><td>Toshkent shahri</td><td>1 260</td></tr></table>
    <svg width="400" height="200"><circle cx="100" cy="100" r="80"
    fill="#{theme.accent}"/></svg></body></html>"""

    path = html_render.render([slide], out_dir="temp", name="sinov")
    try:
        from pptx import Presentation
        from pptx.util import Pt

        presentation = Presentation(path)
        check("slayd yaratildi", len(presentation.slides._sldIdLst) == 1)
        check("slayd 16:9 varaq",
              abs(presentation.slide_width / 914400 - 13.333) < 0.01
              and abs(presentation.slide_height / 914400 - 7.5) < 0.01)

        shapes = list(list(presentation.slides)[0].shapes)
        texts = [s.text_frame.text for s in shapes
                 if s.has_text_frame and s.text_frame.text.strip()]
        check("sarlavha matn bo'lib turibdi",
              any("Bandlik ko'rsatkichlari" in t for t in texts), str(texts[:4]))
        check("qator ko'chishi saqlandi",
              any("\n" in t for t in texts), str(texts[:4]))
        check("ko'rsatkich matni bor", any("14,4 mln" in t for t in texts))

        check("jadval haqiqiy jadval", any(s.has_table for s in shapes))
        table = next(s.table for s in shapes if s.has_table)
        check("jadval kataklari o'qiladi",
              table.cell(0, 0).text == "Hudud"
              and table.cell(1, 0).text == "Toshkent shahri",
              table.cell(0, 0).text)

        pictures = [s for s in shapes if s.shape_type == 13]
        check("diagramma rasm bo'lib qo'yildi", len(pictures) == 1,
              str(len(pictures)))
        check("rasm butun varaqni egallamaydi",
              all(p.width < presentation.slide_width * 0.9 for p in pictures),
              str([p.width for p in pictures]))

        # Bezak bloklari — shakl; matnsiz.
        filled = [s for s in shapes
                  if s.shape_type == 1 and not s.text_frame.text.strip()]
        check("bezak bloklari shakl bo'ldi", len(filled) >= 2, str(len(filled)))

        title = next(s for s in shapes if s.has_text_frame
                     and "Bandlik" in s.text_frame.text)
        run = title.text_frame.paragraphs[0].runs[0]
        check("sarlavha o'lchami saqlandi",
              abs(run.font.size.pt - 26) < 1.5, str(run.font.size.pt))
        check("shrift PowerPointnikiga o'girildi",
              run.font.name == "Arial", str(run.font.name))
        check("rang saqlandi",
              str(run.font.color.rgb) == theme.heading.upper(),
              str(run.font.color.rgb))
        check("slayd foni oq",
              str(list(presentation.slides)[0].background.fill.fore_color.rgb)
              == "FFFFFF")
    finally:
        if os.path.exists(path):
            os.remove(path)


def check_layout_guard():
    """Joylashuv tekshiruvi va tuzatish bosqichi.

    AI HTML ni brauzersiz yozadi, shuning uchun ba'zan matn ustiga
    matn tushadi yoki mazmun yuqoriga to'planib qoladi. Buni faqat
    brauzer ko'radi — shuning uchun slayd chizilishidan oldin
    tekshiriladi va bir marta qayta so'raladi.
    """
    print("\n8) Joylashuv qorovuli")
    if not html_render.available():
        check("brauzer o'rnatilgan", False, html_render._INSTALL_HINT)
        return

    theme = themes.get("ko'k")
    font = html_slides.FONT_STACK

    broken = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;padding:84px;background:#FFFFFF;
    font-family:{font};overflow:hidden}}
    h2{{font-size:48px;color:#{theme.heading}}}
    .line{{position:relative;height:6px;background:#{theme.accent};margin-top:30px}}
    .card{{position:absolute;bottom:100%;margin-bottom:40px;width:320px;
    background:#{theme.accent_soft};padding:24px}}
    .c1{{left:0}} .c2{{left:460px}}
    </style></head><body><h2>Vaqt o'qi</h2><div class="line">
    <div class="card c1"><b>2022</b><div>Birinchi.</div></div>
    <div class="card c2"><b>2023</b><div>Ikkinchi.</div></div></div></body></html>"""

    whole = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;padding:84px;background:#FFFFFF;
    font-family:{font};overflow:hidden;display:flex;flex-direction:column;
    justify-content:space-between}}
    h2{{font-size:48px;color:#{theme.heading}}}
    .row{{display:flex;gap:40px}}
    .card{{flex:1;background:#{theme.accent_soft};padding:30px}}
    .foot{{font-size:20px;color:#{theme.muted}}}
    </style></head><body><h2>To'g'ri joylashgan slayd</h2>
    <div class="row"><div class="card"><b>2022</b><div>Birinchi voqea.</div></div>
    <div class="card"><b>2023</b><div>Ikkinchi voqea.</div></div></div>
    <div class="foot">Manba: statistika qo'mitasi</div></body></html>"""

    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        browser = html_render._launch(playwright)
        try:
            context = browser.new_context(
                viewport={"width": 1920, "height": 1080})
            page = context.new_page()
            page.set_content(broken, wait_until="load")
            problems = html_extract.check_layout(page)
            check("buzuq joylashuv topiladi", bool(problems), str(problems))
            check("matn ustiga matn ko'riladi",
                  any("matn ustiga matn" in item for item in problems),
                  str(problems))
            check("bo'sh pastki qism ko'riladi",
                  any("bo'sh qolgan" in item for item in problems),
                  str(problems))

            page.set_content(whole, wait_until="load")
            check("to'g'ri slaydda shikoyat yo'q",
                  html_extract.check_layout(page) == [],
                  str(html_extract.check_layout(page)))
            context.close()
        finally:
            browser.close()

    # Tuzatish bosqichi chaqiriladimi va natijasi ishlatiladimi.
    asked = []

    def repair(html, problems):
        asked.append(problems)
        return whole

    path = html_render.render([broken], out_dir="temp", name="sinov",
                              repair=repair)
    try:
        check("buzuq slayd tuzatishga yuboriladi", bool(asked), str(asked))
        from pptx import Presentation
        texts = [shape.text_frame.text
                 for shape in list(Presentation(path).slides)[0].shapes
                 if shape.has_text_frame]
        check("tuzatilgan slayd ishlatildi",
              any("To'g'ri joylashgan" in t for t in texts), str(texts[:3]))
    finally:
        if os.path.exists(path):
            os.remove(path)


def check_decoration():
    """Shaffof bezak va doira to'g'ri o'girilishi."""
    print("\n9) Bezak ranglari va shakllari")
    if not html_render.available():
        check("brauzer o'rnatilgan", False, html_render._INSTALL_HINT)
        return

    theme = themes.get("ko'k")
    page_html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;background:#FFFFFF;overflow:hidden;
    position:relative;font-family:{html_slides.FONT_STACK}}}
    .dot{{position:absolute;border-radius:50%;background:#{theme.accent};
    opacity:.08;width:200px;height:200px;top:120px;left:160px}}
    .hidden{{position:absolute;top:-400px;left:40px;width:300px;height:300px;
    background:#{theme.accent}}}
    h1{{font-size:56px;color:#{theme.heading};padding:400px 90px}}
    </style></head><body><div class="dot"></div><div class="hidden"></div>
    <h1>Bezak sinovi</h1></body></html>"""

    path = html_render.render([page_html], out_dir="temp", name="sinov")
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE

        shapes = list(list(Presentation(path).slides)[0].shapes)
        ovals = [s for s in shapes
                 if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.OVAL]
        check("border-radius 50% doira bo'ldi", len(ovals) == 1, str(len(ovals)))
        if ovals:
            colour = str(ovals[0].fill.fore_color.rgb)
            check("shaffof bezak och rangga o'girildi",
                  colour != theme.accent.upper() and colour != "FFFFFF", colour)
        filled = [s for s in shapes if s.shape_type == 1]
        check("slayddan chiqqan blok qo'yilmadi", len(filled) == 1,
              str(len(filled)))
    finally:
        if os.path.exists(path):
            os.remove(path)


def check_photos():
    """Fotosurat oqimi: AI o'rin belgilaydi, Together rasmni chizadi."""
    print("\n10) Fotosuratlar")
    theme = themes.get("zumrad")

    page = ('<html><body>'
            '<img data-prompt="wide photograph of a city skyline" class="photo">'
            '<p>matn</p>'
            '<img data-prompt="close-up of hands writing" class="small" '
            'style="width:400px">'
            '</body></html>')

    check("rasm so'rovlari topiladi",
          len(html_images.requests_in(page)) == 2,
          str(html_images.requests_in(page)))
    check("butun taqdimot bo'yicha sanaladi",
          html_images.count_requests([page, page]) == 4)
    check("so'rovsiz slaydda nol",
          html_images.count_requests(["<html><body>x</body></html>"]) == 0)

    # Together javob bersa — rasm HTML ichiga joylashadi.
    from PIL import Image

    os.makedirs("temp", exist_ok=True)
    made = []

    class Stub:
        async def generate_image(self, prompt, aspect_ratio="16:9"):
            path = os.path.join("temp", f"stub_{len(made)}.png")
            Image.new("RGB", (64, 36), (40, 100, 90)).save(path)
            made.append(prompt)
            return path

    import services.together_service as together_service

    original = together_service.get_together_service
    try:
        together_service.get_together_service = lambda: Stub()
        pages, count = asyncio.run(
            html_images.illustrate([page], theme, "Mavzu"))
    finally:
        together_service.get_together_service = original

    check("ikkala rasm ham chizildi", count == 2, str(count))
    check("rasm HTML ichiga joylashdi",
          pages[0].count('src="data:image/') == 2, str(count))
    check("tavsif Together ga yetib bordi",
          any("skyline" in item for item in made), str(made))

    # Together ishlamasa — rangli blok qoladi, joylashuv buzilmaydi.
    class Broken:
        async def generate_image(self, prompt, aspect_ratio="16:9"):
            raise RuntimeError("kredit yo'q")

    try:
        together_service.get_together_service = lambda: Broken()
        pages, count = asyncio.run(
            html_images.illustrate([page], theme, "Mavzu"))
    finally:
        together_service.get_together_service = original

    check("rasm chiqmasa xato bermaydi", count == 0)
    check("o'rniga rangli blok qoladi",
          pages[0].count("<div ") == 2 and "<img" not in pages[0],
          pages[0][:90])
    check("blok o'z o'lchamini saqlaydi",
          "width:400px" in pages[0], pages[0][-120:])

    # Promptda rasm qoidasi bormi.
    rules = html_slides.shell_rules(theme, "uz")
    check("promptda rasm so'raladi", "data-prompt" in rules)
    check("kamida uchta rasm talab qilinadi", "uchtadan kam bo'lmasin" in rules)
    check("muqovada rasm majburiy", "MUQOVADA albatta" in rules)


def check_icons():
    """Tayyor ikonkalar — eski tizimdagi 142 ta siluet."""
    print("\n11) Ikonkalar")
    from services.premium_presentation import icon_render

    names = icon_render.icon_names()
    check("ikonkalar joyida", len(names) > 100, str(len(names)))
    check("tanish nomlar bor",
          all(name in names for name in ("education", "finance", "research")),
          str(names[:5]))

    rules = html_slides.shell_rules(themes.get("zumrad"), "uz")
    check("promptda ikonka aytilgan", "data-icon" in rules)
    check("promptda ro'yxat berilgan", "education" in rules)

    theme = themes.get("zumrad")
    page = ('<html><body>'
            '<img data-icon="education" class="ikon">'
            '<img data-icon="finance" class="ikon" data-icon-color="FFFFFF">'
            '<img data-icon="yo-q-bunday-ikonka" class="ikon" style="width:60px">'
            '</body></html>')
    filled, count = html_images.apply_icons([page], theme)

    check("ikonkalar qo'yildi", count >= 2, str(count))
    check("HTML ichiga joylashdi",
          filled[0].count('src="data:image/png;base64,') >= 2, str(count))
    check("noma'lum nom slaydni buzmaydi", "<img" not in filled[0]
          or filled[0].count("<img") <= count, filled[0][:80])

    # Rang: aksent va oq — ikki xil fayl bo'lishi kerak.
    education = icon_render.resolve("education")
    check("ikonka fayli topiladi", bool(education), str(education))
    if education:
        painted = icon_render.tinted(education, theme.accent)
        white = icon_render.tinted(education, "FFFFFF")
        check("ikonka bo'yaladi", bool(painted) and painted != white,
              f"{painted} / {white}")
        from PIL import Image

        with Image.open(painted) as image:
            pixels = list(image.convert("RGBA").getdata())
        ink = {p[:3] for p in pixels if p[3] > 200}
        wanted = tuple(int(theme.accent[i:i + 2], 16) for i in (0, 2, 4))
        check("bo'yog'i sxema rangida", ink == {wanted} if ink else False,
              str(list(ink)[:3]))


def check_accent_strip():
    """Bir tomonlama chegara butun ramka bo'lib qolmasin."""
    print("\n12) Aksent chizig'i")
    if not html_render.available():
        check("brauzer o'rnatilgan", False, html_render._INSTALL_HINT)
        return

    theme = themes.get("zumrad")
    page = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><style>
    *{{margin:0;padding:0;box-sizing:border-box}}
    body{{width:1920px;height:1080px;padding:90px;background:#FFFFFF;
    font-family:{html_slides.FONT_STACK};overflow:hidden}}
    .card{{width:600px;height:300px;background:#{theme.accent_soft};
    border-radius:20px;border-top:8px solid #{theme.accent};padding:40px}}
    .boxed{{width:600px;height:200px;margin-top:60px;
    border:4px solid #{theme.accent};padding:30px}}
    </style></head><body>
    <div class="card"><p>Tepasida aksent chizig'i</p></div>
    <div class="boxed"><p>To'liq ramka</p></div></body></html>"""

    path = html_render.render([page], out_dir="temp", name="sinov")
    try:
        from pptx import Presentation

        shapes = [s for s in list(Presentation(path).slides)[0].shapes
                  if s.shape_type == 1]
        # Kartochka + uning tepasidagi tasma + to'liq ramkali blok.
        strips = [s for s in shapes if s.height / 914400 < 0.12]
        check("aksent chizig'i alohida tasma bo'ldi", len(strips) == 1,
              str([round(s.height / 914400, 3) for s in shapes]))
        outlined = [s for s in shapes if s.line.fill.type == 1]
        check("to'liq ramka ramka bo'lib qoldi", len(outlined) == 1,
              str(len(outlined)))
    finally:
        if os.path.exists(path):
            os.remove(path)


def main():
    check_handler_names()
    check_prompt()
    check_split()
    check_outline()
    check_writer()
    check_browser_setup()
    if check_shot():
        check_editable()
        check_layout_guard()
        check_decoration()
    check_photos()
    check_icons()
    if html_render.available():
        check_accent_strip()

    print()
    if FAILS:
        print(f"❌ {len(FAILS)} ta tekshiruv o'tmadi: {', '.join(FAILS[:5])}")
        return 1
    print("✅ HTML slaydlar chiziladi va tahrirlanadigan PPTX ga tushadi.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
