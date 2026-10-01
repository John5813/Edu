"""Bitta "Taqdimot" katalogi: 5 uslub (zamonaviy oqim) + "Chiroyli orqa fonlar" (oddiy oqim).

    python test_taqdimot_uslublari.py
"""
import asyncio, os, re, sys, tempfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from services.premium_presentation import deck_style, deck_styles, html_slides, themes

# ── menyu: "Zamonaviy taqdimot" tugmasi yo'q, hammasi "Taqdimot" da
from bot.keyboards import get_main_keyboard
from translations import get_text
for lang in ("uz", "ru", "en"):
    labels = [b.text for row in get_main_keyboard(lang).keyboard for b in row]
    check(f"{lang}: menyuda zamonaviy taqdimot tugmasi yo'q",
          get_text(lang, "main_menu.premium_presentation") not in labels, labels)
    check(f"{lang}: menyuda Taqdimot tugmasi bor", get_text(lang, "main_menu.presentation") in labels)

from bot.handlers import premium_presentation as pp, documents
check("Taqdimot tugmasi yangi katalogdan boshlanadi",
      all(get_text(l, "main_menu.presentation") in pp.ENTRY_TEXTS for l in ("uz", "ru", "en")))
check("eski zamonaviy tugma nomlari ham qabul qilinadi",
      "✨ Zamonaviy taqdimot" in pp.ENTRY_TEXTS and "⭐ Premium taqdimot" in pp.ENTRY_TEXTS)
check("oddiy oqim kirishi Taqdimot matnini ushlamaydi (ikki marta ishlamasin)",
      not any(k in documents.DOCUMENT_TYPES for k in pp.ENTRY_TEXTS))

# ── har uslub: CSS to'liq, bezaklar haqiqiy elementlar
BODY = ('<section class="slide dark"><div class="body"><h1 class="title big">Mavzu</h1>'
        '<div class="rule"></div></div></section>',
        '<section class="slide"><div class="head"><h2 class="title">Sarlavha</h2><div class="rule"></div></div>'
        '<div class="body"><div class="cols cols-3">'
        '<div class="card line"><div class="card-num">01</div><div class="card-title">A</div><div class="card-note">Izoh.</div></div>'
        '<div class="card line"><div class="card-num">02</div><div class="card-title">B</div><div class="card-note">Izoh.</div></div>'
        '<div class="card line"><div class="card-num">03</div><div class="card-title">C</div><div class="card-note">Izoh.</div></div>'
        '</div></div></section>',
        '<section class="slide"><div class="head"><h2 class="title">Ro\'yxat</h2><div class="rule"></div></div>'
        '<div class="body"><div class="list">'
        '<div class="item"><span class="item-dot"></span><div class="item-text"><b>Bir.</b> Matn.</div></div>'
        '<div class="item"><span class="item-dot"></span><div class="item-text"><b>Ikki.</b> Matn.</div></div>'
        '</div></div></section>')
for key in deck_styles.STYLE_KEYS:
    theme = themes.with_style(themes.get("ko'k"), key)
    css = deck_style.stylesheet(theme)
    left = re.findall(r"#(?:ACCENT|HEADING|BODY|MUTED|SOFT|SOFTER|BAND\w*|INVERT|BACKGROUND|TONE\d|TINT\d|DEEP\d|EDGE|SOFTINK)\b", css)
    check(f"{key}: CSS da almashtirilmagan belgi yo'q", not left, set(left))
    check(f"{key}: uslub CSS ga qo'shilgan", len(css) > len(deck_style.stylesheet(themes.get("ko'k"))))
    pages = html_slides.build_pages(list(BODY), theme)
    text = "".join(pages)
    if key == "jurnal":
        check("jurnal: ro'yxat raqamlari (01, 02)", 'item-num">01' in text and 'item-num">02' in text)
    if key == "kontur":
        check("kontur: burchak belgilari haqiqiy element", text.count('class="kor kor-a"') == 3)
        check("kontur: sarlavha chizig'i haqiqiy element", 'class="kbar"' in text)
    if key == "blok":
        check("blok: muqovada rangli yon tasma", 'class="blok-bar"' in text)
    if key == "qorongu":
        check("qorongu: fon ranglari to'q", theme.background != "FFFFFF" and theme.heading == "FFFFFF")
check("uslub bo'lmasa sxema o'zgarmaydi", themes.with_style(themes.get("ko'k"), "") == themes.get("ko'k"))
check("noma'lum uslub e'tiborga olinmaydi", themes.with_style(themes.get("ko'k"), "nomalum") == themes.get("ko'k"))
check("mavzuga qarab uslub taklifi", deck_styles.suggest("Falsafa asoslari") == "jurnal"
      and deck_styles.suggest("Sun'iy intellekt") == "qorongu")

# ── AI prompti: to'g'ri blok tanlash, bir xillikdan qochish
rules = html_slides.shell_rules(themes.get("ko'k"), "uz")
check("prompt: to'g'ri blok tanlash qoidasi bor", "BLOKNI TO'G'RI TANLANG" in rules and "Kartochkaga qaytaverish" in rules)
check("prompt: eski 'takror mumkin' qoidasi olib tashlangan",
      "xilma-xillik emas" not in rules and "MUMKIN, agar" not in html_slides._user_prompt("M", 1, 3, 10, [], [], 2, "", "", ""))
check("prompt: kartochka soni bo'yicha kod cheklovi yo'q (faqat yo'riqnoma)",
      "30%" not in rules and "foizdan" not in rules)

# ── to'liq quvur: HTML → tahrirlanadigan PPTX (brauzer bo'lsa)
from services.premium_presentation import html_render
if html_render.available():
    from pptx import Presentation
    for key in deck_styles.STYLE_KEYS:
        theme = themes.with_style(themes.get("ko'k"), key)
        pages = html_slides.build_pages(list(BODY), theme)
        out = html_render.render(pages, out_dir=tempfile.mkdtemp(), name=key)
        prs = Presentation(out)
        texts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        check(f"{key}: PPTX {len(pages)} slayd, matn tahrirlanadi", len(prs.slides) == len(pages) and texts >= 8, (len(prs.slides), texts))
else:
    print("  (brauzer yo'q — PPTX sinovi o'tkazib yuborildi)")

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
