"""Mustaqil ish xatolari: bo'sh bo'lim, soxta diagramma, adabiyotlar, reja raqamlari.

    python test_hujjat_xatolari.py
"""
import asyncio, json, os, sys, tempfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from services import ai_service as ais
from services.ai_service import AIService, _chart_is_grounded

def make_ai():
    ai = AIService.__new__(AIService)
    ai.client = MagicMock()
    ai._get_current_model_id = AsyncMock(return_value="model-A")
    return ai

def response(text, finish="stop"):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text), finish_reason=finish)])

# ── 1. Bo'sh javob muvaffaqiyat emas
async def empty_answers():
    ai = make_ai()
    served = []
    async def create(**kw):
        served.append((kw["model"], kw["max_tokens"]))
        if kw["model"] == "model-A": return response("", "stop")
        return response("Haqiqiy matn.")
    ai.client.chat.completions.create = create
    with patch("config.AI_MODEL_FALLBACKS", ["model-B"]):
        out = await ai._make_request(messages=[{"role": "user", "content": "x"}], max_tokens=500)
    check("bo'sh javob — keyingi modelga o'tiladi", out == "Haqiqiy matn." and served[-1][0] == "model-B", served)

    served.clear()
    async def create2(**kw):
        served.append((kw["model"], kw["max_tokens"]))
        if kw["model"] == "model-A" and kw["max_tokens"] <= 500: return response("", "length")
        return response("Uzun matn bo'ldi." if kw["model"] == "model-A" else "")
    ai.client.chat.completions.create = create2
    with patch("config.AI_MODEL_FALLBACKS", ["model-B"]):
        out = await ai._make_request(messages=[{"role": "user", "content": "x"}], max_tokens=500)
    check("token tugab bo'sh qolsa — shu modelga kattaroq limit bilan qayta so'raladi", out == "Uzun matn bo'ldi." and served[0][1] == 500 and served[1][1] > 500, served)

    async def create3(**kw): return response("")
    ai.client.chat.completions.create = create3
    with patch("config.AI_MODEL_FALLBACKS", ["model-B"]):
        try:
            await ai._make_request(messages=[{"role": "user", "content": "x"}], max_tokens=500); ok = False
        except ValueError:
            ok = True
    check("hamma model bo'sh qaytarsa — xato (bo'sh matn jimgina o'tmaydi)", ok)

asyncio.run(empty_answers())

# ── 2. Bo'lim matni: qisqa/bo'sh bo'lsa qayta so'raladi
async def sections():
    ai = make_ai()
    prompts, replies = [], []
    async def fake(messages, max_tokens=0, temperature=0, **kw):
        prompts.append(messages[-1]["content"]); return replies.pop(0)
    ai._make_request = fake
    good = ("Bu to'liq bo'lim matni. " * 60).strip()
    replies[:] = ["Qisqa.", good]
    out = await ai._generate_section_content("Mavzu", "2. Sarlavha bo'limi", 2, 6, "independent_work", "uz")
    check("qisqa javobdan keyin qayta so'raladi va yaxshi matn olinadi", out.startswith("Bu to'liq") and len(prompts) == 2, len(prompts))
    replies[:] = ["", "", ""]; prompts.clear()
    try:
        await ai._generate_section_content("Mavzu", "2. Sarlavha bo'limi", 2, 6, "independent_work", "uz"); ok = False
    except ValueError as exc:
        ok = "yozilmadi" in str(exc)
    check("3 marta ham bo'sh bo'lsa hujjat xato bilan to'xtaydi (bo'sh bo'lim qolmaydi)", ok and len(prompts) == 3, len(prompts))
    check("matn bo'limlari promptida o'ylab topilgan yillarga qarshi qoida bor",
          "kelgusi yillarda" in prompts[0] and "2026-2028" not in prompts[0] and "Prognoz" not in prompts[0] and "Forecast" not in prompts[0], prompts[0][-400:])

asyncio.run(sections())

# ── 3. Diagramma real ma'lumotga tayanishi shart
grounded = {"title": "Aholi soni", "basis": "O'zbekiston Statistika agentligi", "series": [{"name": "mln kishi"}]}
check("manbali diagramma ruxsat", _chart_is_grounded(grounded))
check("manbasiz diagramma rad", not _chart_is_grounded({k: v for k, v in grounded.items() if k != "basis"}))
check("'taxminiy' sarlavhali diagramma rad", not _chart_is_grounded(dict(grounded, title="Yo'nalishlar (taxminiy)")))
check("'daraja' seriyali diagramma rad", not _chart_is_grounded(dict(grounded, series=[{"name": "Rivojlanish darajasi"}])))
check("manbasi 'taxminiy baho' bo'lsa rad", not _chart_is_grounded(dict(grounded, basis="taxminiy baho")))

async def planning():
    ai = make_ai()
    visuals = {"visuals": [
        {"subsection": "1", "kind": "chart", "chart_type": "line", "title": "Qadimgi davr (taxminiy)", "basis": "taxminiy",
         "categories": ["a", "b", "c"], "series": [{"name": "daraja", "values": [1, 2, 3]}], "explanation": "Izoh"},
        {"subsection": "2", "kind": "table", "title": "Davrlar", "headers": ["Davr", "Vakil", "G'oya"],
         "rows": [["Antik", "Aristotel", "Taqlid"]], "explanation": "Izoh"},
        {"subsection": "3", "kind": "chart", "chart_type": "bar", "title": "Nashr soni", "basis": "Milliy kutubxona katalogi",
         "categories": ["2019", "2020", "2021"], "series": [{"name": "Nashr", "values": [4, 5, 6]}], "explanation": "Izoh"},
    ]}
    captured = {}
    async def fake(messages, **kw):
        captured["prompt"] = messages[0]["content"]; return json.dumps(visuals)
    ai._make_request = fake
    plan = await ai.plan_document_visuals("G'arb adabiyoti", [("1", "A"), ("2", "B"), ("3", "C")], "uz")
    kinds = [(p["subsection"], p["kind"]) for p in plan]
    check("o'ylab topilgan diagramma tushib qoladi, jadval va asosli diagramma qoladi", kinds == [("2", "table"), ("3", "chart")], kinds)
    check("prompt: gumanitar mavzuda raqamlashtirish taqiqlangan, jadval matnli", "MAN ETILADI" in captured["prompt"] and "matn" in captured["prompt"] and "basis" in captured["prompt"])

asyncio.run(planning())

# ── 4. Adabiyotlar: mavzuga mos, tilga mos format, alifbo tartibida
async def references():
    ai = make_ai()
    captured = {}
    async def fake(messages, **kw):
        captured["prompt"] = messages[0]["content"]
        return json.dumps({"references": ["Yoqubov U. Kitob. T: Fan, 1990.", "Aristotel. Poetika. Toshkent: Sharq, 2010.",
                                          "Abdurahmonov A. Maqola // Jurnal. 2018. №3. B. 45-52.", "Aristotel. Poetika. Toshkent: Sharq, 2010."]})
    ai._make_request = fake
    refs = await ai._generate_references("G'arb mumtoz adabiyoti", "uz")
    check("adabiyotlar alifbo tartibida, takror yo'q, shahar nomi to'liq",
          [r.split()[0] for r in refs] == ["Abdurahmonov", "Aristotel.", "Yoqubov"] and any("Toshkent: Fan" in r for r in refs), refs)
    pr = captured["prompt"]
    check("prompt: mavzuga mos va faqat mavjud manbalar; huquqiy hujjat faqat huquq mavzusida",
          "DIRECTLY about this topic" in pr and "really exist" in pr and "ONLY when the topic itself is about law" in pr)
    check("prompt: o'zbekcha ro'yxatda inglizcha 'Pages'/'No.' yo'q", "№3" in pr and "B. 45-52" in pr and "Pages X-X" not in pr and "No.X" not in pr)
    check("prompt: yangi yillarni majburlamaydi", "at least three of them from the last five years" not in pr)

asyncio.run(references())

# ── 5. Rasm izohi rasmning o'ziga qarab; vision ishlamasa — rasmni tasvirlamaydigan izoh
async def image_note():
    ai = make_ai()
    img = os.path.join(tempfile.mkdtemp(), "a.png")
    from PIL import Image
    Image.new("RGB", (8, 8), "white").save(img)
    seen = {}
    async def create(**kw):
        seen["kw"] = kw
        return response("Rasmda yonib turgan sham, eski qo'lyozma va patli qalam ko'rinadi. Bu bo'lim mavzusi bilan bog'liq. " * 2)
    ai.client.chat.completions.create = create
    text = await ai.describe_image_file(img, "Bo'lim", "Mavzu", "uz")
    content = seen["kw"]["messages"][0]["content"]
    check("rasm izohi: rasm vision modeliga yuboriladi", any(c.get("type") == "image_url" for c in content) and text.startswith("Rasmda"))
    ai.client.chat.completions.create = AsyncMock(side_effect=RuntimeError("vision yo'q"))
    check("vision ishlamasa bo'sh qaytadi (hujjat to'xtamaydi)", await ai.describe_image_file(img, "B", "M", "uz") == "")
    prompts = []
    async def fake(messages, **kw): prompts.append(messages[0]["content"]); return "Izoh matni uchun yetarli uzun to'liq gap."
    ai._make_request = fake
    await ai.describe_visual("image", "Bo'lim", "Mavzu", "uz")
    check("vision bo'lmaganda ham rasmdagi aniq narsalar o'ylab topilmaydi", "ko'rmayapsiz" in prompts[0] and "tasvirlamang" in prompts[0], prompts[0][-250:])

asyncio.run(image_note())

# ── 6. Reja raqamlari: serverda Times New Roman o'lchamdosh shrift bo'lmasa ham to'g'ri
from docx import Document
from docx.shared import Pt
from services import doc_toc
d = Document()
for i in range(70):
    p = d.add_paragraph(("Bu abzats matni, sahifa hisobini tekshirish uchun yoziladi. " * 8))
    for r in p.runs: r.font.name = "Times New Roman"; r.font.size = Pt(14)
path = os.path.join(tempfile.mkdtemp(), "t.docx"); d.save(path)
with_system = doc_toc._page_lines(path, include_system=True)
bundled_only = doc_toc._page_lines(path, include_system=False)
check("o'rnatilgan Liberation Serif fayllari loyihada bor", all(os.path.exists(os.path.join(doc_toc._FONTS_DIR, f"LiberationSerif-{n}.ttf")) for n in ("Regular", "Bold", "Italic", "BoldItalic")))
check("litsenziya fayli ham bor (SIL OFL)", os.path.exists(os.path.join(doc_toc._FONTS_DIR, "LICENSE-Liberation.txt")))
check("tizim shriftlarisiz (serverda shrift yo'q holat) ham betlar soni bir xil",
      len(bundled_only) == len(with_system) > 3, (len(bundled_only), len(with_system)))
conf = doc_toc._fontconfig(tempfile.mkdtemp())
check("fontconfig Times New Roman -> Liberation Serif almashtiradi", conf and "Times New Roman" in open(conf).read() and "Liberation Serif" in open(conf).read())

# ── 7. Hujjat: bo'sh bo'lim bo'lsa yig'ilmaydi; to'liq hujjatda reja raqamlari haqiqiy betlarga mos
from services import document_service as ds
service = ds.DocumentService.__new__(ds.DocumentService)
service.documents_dir = tempfile.mkdtemp(); service.temp_dir = tempfile.mkdtemp()
service._last_used_icons = set(); service.together = None
body = lambda i: (f"Bu {i}-bo'lim matni, mazmunli va yetarlicha uzun gaplardan iborat. " * 55).strip()
content = {"language": "uz", "author_name": "Ali", "references": ["Ahmad S. Ufq. Toshkent: Fan, 1998."],
           "sections": [{"title": t, "content": body(i)} for i, t in enumerate(
               ["Kirish", "1. Birinchi savol sarlavhasi", "2. Ikkinchi savol sarlavhasi", "3. Uchinchi savol sarlavhasi", "4. To'rtinchi savol", "Xulosa"])]}
async def build(c):
    with patch.object(ds.DocumentService, "_create_independent_work_title_page", AsyncMock()):
        return await service.create_independent_work("Mavzu", c, extras=None)
broken = json.loads(json.dumps(content)); broken["sections"][1]["content"] = ""
try:
    asyncio.run(build(broken)); blocked = False
except ValueError as exc:
    blocked = "Birinchi savol" in str(exc)
check("bo'sh bo'limli hujjat yig'ilmaydi (xato, to'lov olinmaydi)", blocked)

path = asyncio.run(build(content))
doc = Document(path)
toc = {}
for para in doc.paragraphs:
    if "\t" in para.text and para.text.split("\t")[-1].strip().isdigit():
        toc[para.text.split("\t")[0].strip()] = int(para.text.split("\t")[-1])
pages = doc_toc._page_lines(path, include_system=False)      # shrift o'rnatilmagan serverdagidek
def page_of(head):
    key = doc_toc._normalize(head)[:30]
    for number, lines in enumerate(pages, 1):
        if number > 1 and any(l.startswith(key) for l in lines): return number   # 1-bet: reja (titul sahifa sinovda yo'q)
real = {h: page_of(h) for h in toc}
check("reja raqamlari hujjatning haqiqiy betlariga mos", toc and all(toc[h] == real[h] for h in toc), (toc, real))
check("reja qatorlarida nuqtali chiziq va raqam bor (kamida 5 ta qator)", len(toc) >= 5, toc)

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
