"""Saytdagi hujjat xizmatlari boshidan oxirigacha: buyurtma → (katta hujjat — umumiy navbat) → AI → DOCX → yuklab olish.

AI soxta, lekin javob shakli haqiqiy: har so'rovdagi JSON namunasi to'ldirib qaytariladi (namunasi yo'q so'rovga
so'ralgan hajmdagi matn), shuning uchun hujjat yig'uvchining asosiy yo'li ishlaydi, zaxira yo'li emas.
Hujjat yig'uvchi (document_service), navbat, to'lov va fayl — haqiqiy.

    python test_sayt_hujjatlar.py
"""
import asyncio, io, json, os, random, re, sys, tempfile, time, types

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")
FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"), flush=True)
    if not cond: FAILS.append(name)

import database.database as dbmod
from database.database import Database, init_db
from database import web_store
from services import web_jobs, web_kinds  # noqa
import webapp
from aiohttp.test_utils import TestClient, TestServer

TMP = tempfile.mkdtemp()
dbmod.DATABASE_FILE = os.path.join(TMP, "t.db")
web_jobs.RESULTS_DIR = os.path.join(TMP, "results")
SENT = []
class FakeBot:
    async def send_document(self, chat_id, document, **kw): SENT.append(chat_id); return types.SimpleNamespace()
    async def send_message(self, *a, **kw): return types.SimpleNamespace(message_id=1, chat=types.SimpleNamespace(id=1))
webapp.BOT, webapp.BOT_USERNAME = FakeBot(), "Edufayl_bot"

WORDS = ("iqtisodiy rivojlanish jarayonida raqamli texnologiyalar muhim ahamiyat kasb etadi chunki zamonaviy bozor "
         "sharoitida samaradorlik oshadi korxonalar faoliyati takomillashadi va natijada mamlakat taraqqiyoti uchun "
         "mustahkam asos yaratiladi tadqiqot shuni ko'rsatadiki islohotlar izchil davom etmoqda").split()
RNG = random.Random(7)
def prose(n):
    out = []
    while len(out) < n:
        s = [RNG.choice(WORDS) for _ in range(RNG.randint(9, 15))]; s[0] = s[0].capitalize(); s[-1] += "."
        out += s
    return " ".join(out[:n])

LONG_KEYS = re.compile(r"content|text|body|abstract|annotation|introduction|intro|conclusion|discussion|results|method|"
                       r"description|paragraph|summary|relevance|goal|object|subject|novelty|significance|analysis|definition", re.I)

def fill(node, key=""):
    if isinstance(node, dict):
        return {k: fill(v, k) for k, v in node.items()}
    if isinstance(node, list):
        if not node:
            return [prose(8) for _ in range(4)]
        items = [fill(node[0], key) for _ in range(max(len(node), 4 if not isinstance(node[0], dict) else 3))]
        return items
    if isinstance(node, str):
        if LONG_KEYS.search(key or ""):
            return prose(170)
        return prose(6).rstrip(".")
    if isinstance(node, (int, float)) and not isinstance(node, bool):
        return node
    return node

def _has_str(node):
    if isinstance(node, str): return True
    if isinstance(node, dict): return any(_has_str(v) for v in node.values()) or bool(node)
    if isinstance(node, list): return any(_has_str(v) for v in node)
    return False

def _balanced(prompt, i):
    depth, j, instr = 0, i, False
    while j < len(prompt):
        c = prompt[j]
        if c == '"' and prompt[j - 1] != "\\": instr = not instr
        elif not instr and c in "{[": depth += 1
        elif not instr and c in "}]":
            depth -= 1
            if depth == 0: return prompt[i:j + 1]
        j += 1
    return ""

def example(prompt):
    """So'rovdagi JSON namunasi (JSON so'zidan keyingi, ichida matn bor birinchi obyekt/ro'yxat)."""
    starts = [m.end() for m in re.finditer(r"json|JSON", prompt)]
    seen = set()
    for start in starts:
        for i in range(start, min(len(prompt), start + 4000)):
            if prompt[i] not in "{[" or i in seen: continue
            seen.add(i)
            raw = _balanced(prompt, i)
            raw = re.sub(r",\s*\.\.\.\s*", "", raw).replace("…", "")
            raw = re.sub(r"\.\.\.", "", raw)
            raw = re.sub(r",\s*([}\]])", r"\1", raw)
            try:
                node = json.loads(raw)
            except Exception:
                continue
            if _has_str(node):
                return node
    return None

def wanted_words(prompt):
    nums = [int(b or a) for a, b in re.findall(r"(\d{2,4})(?:\s*[-–]\s*(\d{2,4}))?\s*(?:ta\s+)?(?:so['‘’]z|words?|слов)", prompt)]
    return min(max(nums), 3000) if nums else 360

STATS = {"calls": 0, "json": 0}
def answer(messages):
    STATS["calls"] += 1
    prompt = "\n".join(m.get("content", "") for m in messages if isinstance(m.get("content"), str))
    ex = example(prompt) if re.search(r"json", prompt, re.I) else None
    if ex is not None:
        STATS["json"] += 1
        return json.dumps(fill(ex), ensure_ascii=False)
    n = wanted_words(prompt)
    return "\n\n".join(prose(max(40, n // 3)) for _ in range(3))

class _Completions:
    async def create(self, **params):
        await asyncio.sleep(0.01)
        msg = types.SimpleNamespace(content=answer(params["messages"]))
        return types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg, finish_reason="stop")])

async def patch_ai():
    from services import ai_service
    svc = ai_service.get_ai_service()
    svc.client = types.SimpleNamespace(chat=types.SimpleNamespace(completions=_Completions()))
    # boshqa joydagi to'g'ridan-to'g'ri OpenAI mijozlari ham soxta
    for name in dir(ai_service):
        obj = getattr(ai_service, name)
        if hasattr(obj, "client") and obj is not svc:
            try: obj.client = svc.client
            except Exception: pass

async def wait_job(client, job_id, limit=900):
    end = time.time() + limit
    while time.time() < end:
        d = await (await client.get(f"/api/v1/jobs/{job_id}")).json()
        if d["job"]["status"] in ("done", "failed"): return d
        await asyncio.sleep(0.5)
    return d

def docx_info(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    return doc, text

async def main():
    from webapp.server import create_web_app
    from bot.queue_service import get_doc_queue
    await init_db()
    for tid in (100, 101):
        await Database.create_user(tid, f"u{tid}", "Ali Valiyev", "uz")
        await Database.update_user_balance(tid, 5_000_000)
    await patch_ai()
    get_doc_queue().start()
    app = create_web_app()
    clients = {}
    for tid in (100, 101):     # daqiqasiga 6 ta buyurtma cheklovi: turlar ikki mijozga bo'linadi
        c = TestClient(TestServer(app if tid == 100 else create_web_app())); await c.start_server()
        o = {"Origin": f"http://{c.server.host}:{c.server.port}"}
        st = await (await c.post("/api/v1/auth/start", headers=o)).json()
        await web_store.confirm_login(st["token"], tid); await c.get("/api/v1/auth/poll", params={"token": st["token"]})
        clients[tid] = (c, o)
    client, origin = clients[100]
    cat = await (await client.get("/api/v1/catalog")).json()
    kinds = {k["key"]: k for k in cat["kinds"]}
    only = sys.argv[1:] or [k for k in kinds if k != "premium_presentation"]

    for number, key in enumerate(only):
        tid = 100 if number < 4 else 101
        client, origin = clients[tid]
        k = kinds[key]; opt = k["options"]
        topic = "Raqamli iqtisodiyotning rivojlanish yo'llari"
        params = {"topic": topic, "author": "Ali Valiyev", "language": LANGS.get(key, "uz")}
        if opt.get("form") == "thesis":
            params.update(university="Toshkent davlat iqtisodiyot universiteti", faculty="Iqtisodiyot", group="IQ-21")
        else:
            params["size"] = opt["sizes"][0]["key"]
            if opt.get("extras"):
                params["extras"] = [x["key"] for x in opt["extras"] if x["key"] in ("tables", "formulas")]
        t0 = time.time(); STATS.update(calls=0, json=0)
        bal0 = (await Database.get_user(tid)).balance
        r = await client.post("/api/v1/jobs", json={"kind": key, "params": params}, headers=origin)
        data = await r.json()
        if r.status != 200:
            check(f"{key}: buyurtma qabul qilindi", False, data); continue
        price = bal0 - data["balance"]
        job = await wait_job(client, data["job"]["id"])
        ok = job["job"]["status"] == "done"
        check(f"{key}: tayyor ({time.time() - t0:.0f} s, AI so'rovlari {STATS['calls']}, JSON {STATS['json']}, narx {price})", ok, job["job"].get("error"))
        if not ok:
            continue
        r = await client.get(f"/api/v1/jobs/{data['job']['id']}/file")
        body = await r.read()
        cd = r.headers.get("Content-Disposition", "")
        try:
            doc, text = docx_info(body)
        except Exception as exc:
            check(f"{key}: DOCX ochiladi", False, exc); continue
        words = len(text.split())
        check(f"{key}: DOCX ochiladi, mavzu va muallif bor ({words} so'z, {len(doc.paragraphs)} abzats, fayl {cd.split('filename*=')[-1][:60]})",
              topic.split()[0] in text and "Ali Valiyev" in text and words > 300, text[:300])
        check(f"{key}: Telegramga yuborildi", tid in SENT, SENT); SENT.clear()

    total = 0
    for c, _ in clients.values():
        total += len((await (await c.get("/api/v1/jobs")).json())["jobs"])
        await c.close()
    check("«Hujjatlarim» ro'yxatida hammasi", total == len(only), total)
    print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))

LANGS = {"referat": "ru", "thesis": "en"}     # bir nechta tur boshqa tilda
asyncio.run(main())
sys.exit(1 if FAILS else 0)
