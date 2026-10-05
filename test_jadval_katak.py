"""Zamonaviy taqdimot jadvali: katakchalar (to'r) aniq ko'rinadi.

    python test_jadval_katak.py
"""
import os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from pptx import Presentation
from services.premium_presentation import html_extract, html_render, html_slides, pptx_build, themes

BODY = ('<section class="slide"><div class="head"><h2 class="title">Jadval</h2><div class="rule"></div></div>'
        '<div class="body"><table><tr><th>Fan</th><th>Aloqa</th><th>Misol</th></tr>'
        '<tr><td>Fiziologiya</td><td>Organizm</td><td>Qon aylanishi</td></tr>'
        '<tr><td>Anatomiya</td><td>Tuzilma</td><td>Suyaklar</td></tr>'
        '<tr><td>Biokimyo</td><td>Moddalar</td><td>Fermentlar</td></tr></table></div></section>')
DARK = BODY.replace('class="slide"', 'class="slide dark"')

def lum(hexv):
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b

from playwright.sync_api import sync_playwright
with sync_playwright() as pw:
    browser = html_render._launch(pw)
    context = browser.new_context(viewport={"width": 1920, "height": 1080})
    page = context.new_page()
    out = os.environ.get("KATAK_OUT")
    for style in ("toza", "jurnal", "blok", "kontur", "qorongu"):
        for label, body in (("oddiy", BODY), ("to'q", DARK)):
            theme = themes.with_style(themes.get("ko'k"), style)
            html = html_slides.build_pages([body], theme)[0]
            page.set_content(html, wait_until="load")
            layout = html_extract.read_layout(page)
            tables = [b for b in layout["blocks"] if b["kind"] == "table"]
            ok = bool(tables)
            check(f"{style}/{label}: jadval topildi", ok)
            if not ok:
                continue
            block = tables[0]
            grid, width = block.get("gridColor"), block.get("gridWidth")
            check(f"{style}/{label}: katak chizig'i bor", bool(grid) and width >= 1.5, (grid, width))
            bg = (layout.get("background") or "FFFFFF")[:6]
            if grid and bg:
                check(f"{style}/{label}: chiziq fondan ajraladi", abs(lum(grid) - lum(bg)) >= 40, (grid, bg))
            prs = Presentation()
            prs.slide_width, prs.slide_height = int(13.333 * 914400), int(7.5 * 914400)
            pptx_build.add_slide(prs, layout)
            xml = " ".join(sh._element.xml for sh in prs.slides[0].shapes if sh.has_table)
            check(f"{style}/{label}: pptx katakchalarida 4 tomon chizig'i",
                  all(tag in xml for tag in ("<a:lnL", "<a:lnR", "<a:lnT", "<a:lnB")) and "w=\"" in xml)
            if out and style in ("jurnal", "toza") :
                tone = "dark" if label == "to'q" else "light"
                page.screenshot(path=os.path.join(out, f"{style}_{tone}.png"))
    browser.close()

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
