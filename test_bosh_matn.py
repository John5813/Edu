"""Oddiy taqdimot: slaydda bo'sh matn qutisi ("Double-tap to add text") qolmasin.

Sabab: model ba'zan slayd matnini lug'at ({"points": [...]}), lug'atlar ro'yxati yoki `bullets` / `points`
maydonida qaytaradi — kod faqat `content`/`text`/`items` ni o'qirdi, slayd rasm (yoki raqam) bilan, lekin
matnsiz chiqardi. Matnsiz asosiy slaydni qayta yozdirish ham yo'q edi.

    python test_bosh_matn.py
"""
import asyncio
import os
import random
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond:
        FAILS.append(name)


from PIL import Image
from pptx import Presentation

from services import slide_layouts

print("1) Matn har qanday shakldan o'qiladi")
cases = [({"content": {"points": ["A bir.", "B ikki."]}}, "A bir.\nB ikki."),
         ({"content": [{"point": "p1"}, {"point": "p2"}]}, "p1\np2"),
         ({"bullets": ["x.", "y."]}, "x.\ny."),
         ({"points": ["bitta"]}, "bitta"),
         ({"content": "oddiy matn"}, "oddiy matn"),
         ({"content": "", "items": [{"title": "T", "text": "matn"}]}, "matn")]
for slide, want in cases:
    got = slide_layouts.ensure_content(dict(slide))["content"]
    check(f"{list(slide)[0]} → {want!r}", got == want, got)
item = slide_layouts._read_item({"title": "Sarlavha", "point": "Noma'lum kalitdagi matn"})
check("noma'lum kalitdagi band matni yo'qolmaydi", item["text"] == "Noma'lum kalitdagi matn", item)

print("2) Matnsiz asosiy slayd qayta yoziladi")
from services.ai_service import AIService

ai = AIService.__new__(AIService)
asked = []


async def fake_slide(topic, title, language, layout):
    asked.append(title)
    return f"{title} haqida matn."

ai._generate_slide_content = fake_slide
content = {"slides": [{"title": "Muqova", "layout": "cover", "content": ""},
                      {"title": "A-Levels dasturi", "layout": "left_image", "content": ""},
                      {"title": "Lug'at", "layout": "", "content": {"points": ["bor"]}},
                      {"title": "To'liq", "layout": "", "content": "Matn bor."}]}
out = asyncio.run(ai._fill_empty_main_slides(content, "Ta'lim", "uz"))
check("faqat matnsiz asosiy slayd so'raldi", asked == ["A-Levels dasturi"], asked)
check("matn qo'yildi", out["slides"][1]["content"] == "A-Levels dasturi haqida matn.", out["slides"][1])
check("lug'atdagi matn satrga aylandi", out["slides"][2]["content"] == "bor")

print("3) Tayyor PPTX da bo'sh matn qutisi yo'q")
from services.document_service import DocumentService
from services.template_service import TemplateService

tmp = tempfile.mkdtemp()
img = os.path.join(tmp, "src.png")
Image.new("RGB", (800, 600), (90, 140, 200)).save(img)


class FakeImages:
    async def _make(self, *a, **k):
        path = os.path.join(tmp, f"{random.random()}.png")
        shutil.copy(img, path)
        return path

    generate_cover_image = generate_slide_image = generate_panoramic_image = _make


text = ("Buyuk Britaniyada ta'lim tizimi o'ziga xos. Talabalarning 80% qismi A-Levels dasturini tanlaydi. "
        "Bu dastur universitetga kirishning asosiy yo'li.")
shapes = {"plain": {"content": text}, "dict": {"content": {"points": [text]}},
          "bullets": {"bullets": text.split(". ")}, "stat": {"content": text, "stat": {"value": "80%"}},
          "list": {"content": [{"point": s} for s in text.split(". ")]}}
slides = [{"title": "Ta'lim", "content": "", "layout": "cover"}]
for name, extra in shapes.items():
    for layout in ("", "left_image", "stat_hero", "image_left_bullets", "right_image"):
        slides.append({"title": f"{name} {layout}", "layout": layout, **extra})
service = DocumentService()
service.together = FakeImages()
path = asyncio.run(service.create_presentation_with_smart_images("Ta'lim", {"slides": slides}, "Ali", "uz",
                                                                  TemplateService(), "template_1"))
empty = []
for number, slide in enumerate(Presentation(path).slides, 1):
    for shape in slide.shapes:
        if shape.shape_type == 17 and not shape.text_frame.text.strip():       # matn qutisi
            empty.append(number)
check("bo'sh matn qutisi yo'q", not empty, empty)
os.remove(path)
shutil.rmtree(tmp, ignore_errors=True)

print("\nNATIJA:", "HAMMASI O'TDI" if not FAILS else f"{len(FAILS)} ta xato: {FAILS}")
sys.exit(1 if FAILS else 0)
