"""Har bir rasm va jadval nomlanishini va izohlanishini tekshiradi.

O'zbek ilmiy ishlarida har bir rasm va jadval nomlanadi — "1-rasm.",
"2-jadval." — va ostida uning nimani ko'rsatayotgani bir abzatsda
tushuntiriladi. Ilgari bu faqat AI rejalashtirgan diagrammalarda
bajarilardi: qo'shimcha xizmat sifatida qo'yilgan rasm "Rasm. Bo'lim
nomi" deb, qiyoslash jadvali esa umuman nomsiz tushardi. Raqamlar ham
har chizuvchining o'z hisobida borar, hujjatda ikkita "1-rasm" bo'lib
qolardi.

    python test_rasm_jadval_raqami.py
"""

import asyncio
import glob
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, ".")
os.environ.setdefault("BOT_TOKEN", "test")

from docx import Document  # noqa: E402

from services import doc_visuals  # noqa: E402

FAILS = []


def check(name, condition, detail=""):
    print(f"  ok   {name}" if condition else f"  XATO {name} — {detail}")
    if not condition:
        FAILS.append(name)


def check_numbering():
    print("\n1) Hisoblagich")
    doc = Document()
    counter = doc_visuals.of(doc, "uz")
    counter.figure_caption(doc, "Bandlik dinamikasi")
    counter.table_caption(doc, "Hududlar kesimi")
    counter.figure_caption(doc, "Ishsizlik darajasi")
    counter.table_caption(doc, "")
    texts = [p.text for p in doc.paragraphs if p.text.strip()]

    check("rasmlar ketma-ket raqamlanadi",
          texts[0].startswith("1-rasm.") and texts[2].startswith("2-rasm."),
          str(texts))
    check("jadvallar alohida raqamlanadi",
          texts[1].startswith("1-jadval.") and texts[3] == "2-jadval",
          str(texts))
    check("nom matnga ulanadi", "Bandlik dinamikasi" in texts[0], texts[0])

    # Bir hujjatda bitta hisoblagich, ikki hujjatda — ikkita.
    check("bir hujjat bitta hisoblagich", doc_visuals.of(doc) is counter)
    other = Document()
    check("boshqa hujjat o'zinikini oladi",
          doc_visuals.of(other).figures == 0)

    for language, figure, table in (("ru", "Рисунок 1.", "Таблица 1."),
                                    ("en", "Figure 1.", "Table 1.")):
        page = Document()
        numbering = doc_visuals.of(page, language)
        numbering.figure_caption(page, "Nomi")
        numbering.table_caption(page, "Nomi")
        lines = [p.text for p in page.paragraphs if p.text.strip()]
        check(f"«{language}» nomlari to'g'ri",
              lines[0].startswith(figure) and lines[1].startswith(table),
              str(lines))

    print("\n2) Izoh")
    page = Document()
    numbering = doc_visuals.of(page, "uz")
    check("izoh yoziladi", numbering.note(page, "Jadvalda bandlik ko'rsatilgan."))
    check("bo'sh izoh yozilmaydi", not numbering.note(page, "   "))


async def check_document():
    print("\n3) Hujjatda ketma-ketlik")
    from services.document_service import DocumentService

    body = "Mehnat bozori holati yildan yilga o'zgarib bormoqda. " * 40
    content = {
        "title": "Aholi bandligi",
        "language": "uz",
        "plan_style": "oddiy",
        "introduction": "Hammaga ma'lumki, mehnat bozori o'zgardi. " * 16,
        "intro_points": {"goal": "Tahlil qilish.", "tasks": ["Bir", "Ikki"]},
        "sections": [{"title": f"{i}-savol nomi", "content": body}
                     for i in range(1, 5)],
        "visuals": [
            {"subsection": "1", "kind": "chart", "chart_type": "column",
             "title": "Bandlik dinamikasi",
             "categories": ["2024", "2025", "2026"],
             "series": [{"name": "Band aholi", "values": [13.9, 14.1, 14.4]}],
             "explanation": "Diagrammada band aholi soni ko'rsatilgan. "
                            "Ko'rsatkich uch yil davomida o'sgan. "
                            "Buning sababi yangi ish o'rinlarida."},
            {"subsection": "2", "kind": "table", "title": "Hududlar kesimi",
             "headers": ["Hudud", "2025-yil", "2026-yil"],
             "rows": [["Toshkent shahri", "1 200", "1 260"],
                      ["Samarqand viloyati", "980", "1 010"]],
             "explanation": "Jadvalda hududlar bo'yicha bandlik keltirilgan. "
                            "Toshkentda ko'rsatkich eng yuqori. "
                            "Shu bois hududiy dasturlar zarur."},
            {"subsection": "3", "kind": "chart", "chart_type": "pie",
             "title": "Tarmoqlar ulushi",
             "categories": ["Sanoat", "Qishloq xo'jaligi", "Xizmatlar"],
             "series": [{"name": "Ulush", "values": [30, 28, 42]}],
             "explanation": "Diagrammada tarmoqlar ulushi ko'rsatilgan. "
                            "Xizmatlar yetakchi o'rinda. "
                            "Bu tendensiya davom etmoqda."},
            {"subsection": "4", "kind": "table", "title": "Yosh guruhlari",
             "headers": ["Yosh", "Ulush"],
             "rows": [["16-30", "34%"], ["31-50", "48%"]],
             "explanation": "Jadvalda yosh guruhlari berilgan. "
                            "Eng katta ulush o'rta yoshda. "
                            "Yoshlar bandligi alohida e'tibor talab qiladi."},
        ],
        "conclusion": "Xulosa. " * 20,
        "references": ["Karimov B. — Toshkent, 2024."],
    }

    service = DocumentService()
    path = await service.create_course_work(
        "Aholi bandligini ta'minlash", content, "Ortiqov Orif", "uz")
    work = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", work, path], capture_output=True, timeout=300)
    produced = glob.glob(os.path.join(work, "*.pdf"))
    check("PDF yaratildi", bool(produced), work)
    if not produced:
        return

    import pymupdf
    with pymupdf.open(produced[0]) as document:
        text = " ".join(" ".join(page.get_text().split()) for page in document)
    os.remove(path)

    figures = re.findall(r"(\d+)-rasm\.", text)
    tables = re.findall(r"(\d+)-jadval\.", text)
    check("ikkita rasm nomlandi", figures == ["1", "2"], str(figures))
    check("ikkita jadval nomlandi", tables == ["1", "2"], str(tables))
    check("rasm nomi matn bilan",
          "1-rasm. Bandlik dinamikasi" in text, text[:0] or "yo'q")
    check("jadval nomi matn bilan", "2-jadval. Yosh guruhlari" in text)
    check("har birining izohi bor",
          all(sentence in text for sentence in (
              "Buning sababi yangi ish o'rinlarida.",
              "Shu bois hududiy dasturlar zarur.",
              "Bu tendensiya davom etmoqda.",
              "Yoshlar bandligi alohida e'tibor talab qiladi.")))
    check("eski nomsiz ko'rinish qolmadi",
          "Rasm. " not in text and "Sxema. " not in text, text[:0] or "")


async def main():
    check_numbering()
    await check_document()

    print()
    if FAILS:
        print(f"❌ {len(FAILS)} ta tekshiruv o'tmadi: {', '.join(FAILS[:5])}")
        return 1
    print("✅ Har bir rasm va jadval nomlangan va izohlangan.")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
