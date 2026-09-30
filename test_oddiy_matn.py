"""Oddiy taqdimot: uzun matn qirqilmaydi, xulosa bo'sh qolmaydi, rasm cho'zilmaydi.

    python test_oddiy_matn.py
"""
import asyncio, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from services import slide_fit
from services.ai_service import AIService
from services.document_service import DocumentService

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

SENT = ("Bu sinov uchun yozilgan {n}-gap bo'lib, uning mazmuni taqdimot matnining uzunligini "
        "tekshirishga xizmat qiladi va hech qanday real ma'lumotni ifodalamaydi.")
def text_of(sentences): return " ".join(SENT.format(n=i + 1) for i in range(sentences))

short, medium, long_ = text_of(3), text_of(14), text_of(20)
print(f"so'zlar: qisqa={slide_fit.words(short)} o'rta={slide_fit.words(medium)} uzun={slide_fit.words(long_)}")

# ── sig'dirish rejasi
def comfortable(item):
    return slide_fit.fits(item["text"], slide_fit.BODY_W, slide_fit.BODY_H, 18) or bool(
        item["columns"] and all(slide_fit.fits(c, slide_fit.COLUMN_W, slide_fit.BODY_H, 16) for c in item["columns"]))

p = slide_fit.plan_text(short)
check("qisqa matn — bitta blok", len(p) == 1 and p[0]["columns"] is None)

counts, ok_join, ok_fit = [], True, True
for n in range(3, 26):
    text = text_of(n)
    plan = slide_fit.plan_text(text)
    counts.append(len(plan))
    ok_join &= " ".join(i["text"] for i in plan).strip() == text.strip()
    ok_fit &= all(comfortable(i) for i in plan)
check("hech qaysi uzunlikda gap tashlanmaydi", ok_join)
check("har bo'lak o'z slaydiga qulay sig'adi (18 pt bitta blok yoki 16 pt ikki ustun)", ok_fit)
check("matn uzayganda slaydlar soni kamaymaydi", counts == sorted(counts), counts)
check("400 so'zli matn kamida ikki slaydga bo'linadi", len(slide_fit.plan_text(long_)) >= 2)
check("bo'linganda oxirgi bo'lak juda qisqa emas",
      all(slide_fit.words(slide_fit.plan_text(text_of(n))[-1]["text"]) >= 25 for n in range(14, 26)))
check("oraliq uzunlikda ikki ustun ishlatiladi", any(any(i["columns"] for i in slide_fit.plan_text(text_of(n))) for n in range(9, 14)))

# ── prepare: layoutlar
slides = [
    {"title": "Uzun", "content": long_, "layout": "text_with_numbers"},
    {"title": "Rasmli uzun", "content": medium, "layout": "right_image"},
    {"title": "Rasmli qisqa", "content": short, "layout": "right_image"},
    {"title": "Gorizontal uzun", "content": medium, "layout": "horizontal_image"},
    {"title": "Ustunli", "content": "", "layout": "two_column",
     "columns": [{"column_content": long_}, {"column_content": short}]},
    {"title": "Reja", "content": "", "layout": "plan", "plan_items": ["a"]},
]
out = slide_fit.prepare(slides, "uz")
titles = [(s["title"], s["layout"]) for s in out]
check("uzun slayd davomi bilan bo'lindi", any("(davomi)" in t for t, _ in titles), titles)
check("sig'maydigan rasmli slayd matn slaydiga o'tdi",
      ("Rasmli uzun", "text_with_numbers") in titles or any(t.startswith("Rasmli uzun") and l == "text_with_numbers" for t, l in titles), titles)
check("qisqa rasmli slayd o'zgarmadi", ("Rasmli qisqa", "right_image") in titles)
check("uzun ustunli slayd matn slaydiga o'tdi", any(t.startswith("Ustunli") and l == "text_with_numbers" for t, l in titles), titles)
check("reja va boshqa layoutlar tegilmadi", ("Reja", "plan") in titles)

# ── bo'sh xulosa to'ldiriladi
async def fill_tests():
    svc = AIService.__new__(AIService)
    mains = [{"title": f"Asosiy {i}", "content": f"{i}-asosiy slaydning birinchi fikri. Ikkinchi fikr.", "layout": "two_column"} for i in range(1, 5)]
    def deck():
        return {"slides": [{"title": "", "content": "", "layout": "cover"},
                           {"title": "Reja", "content": "", "layout": "plan", "plan_items": []},
                           {"title": "Kirish", "content": "", "layout": "intro"},
                           *[dict(m) for m in mains],
                           {"title": "Xulosa", "content": "", "layout": "conclusion"},
                           {"title": "", "content": "", "layout": "thanks"}]}

    async def good(messages, **kw):
        return '{"intro": "Kirish matni yozildi.", "conclusion": "Xulosa matni yozildi."}'
    svc._make_request = good
    d = await svc._fill_empty_fixed_slides(deck(), "Sinov mavzu", "uz")
    by = {s["layout"]: s for s in d["slides"]}
    check("AI xulosa va kirishni yozdi", by["conclusion"]["content"] == "Xulosa matni yozildi." and by["intro"]["content"] == "Kirish matni yozildi.")
    check("reja bandlari asosiy slaydlardan olindi", len(by["plan"]["plan_items"]) == 4, by["plan"]["plan_items"])

    async def bad(messages, **kw): raise RuntimeError("tarmoq yo'q")
    svc._make_request = bad
    d = await svc._fill_empty_fixed_slides(deck(), "Sinov mavzu", "uz")
    by = {s["layout"]: s for s in d["slides"]}
    check("AI ishlamasa ham xulosa bo'sh emas (slaydlardan yig'iladi)",
          len(by["conclusion"]["content"]) > 60 and "asosiy slaydning birinchi fikri" in by["conclusion"]["content"], by["conclusion"]["content"])
    check("AI ishlamasa kirish ham bo'sh emas", "Sinov mavzu" in by["intro"]["content"], by["intro"]["content"])

    d = deck(); d["slides"][-2]["content"] = "Tayyor xulosa."
    svc._make_request = bad
    d = await svc._fill_empty_fixed_slides(d, "Sinov mavzu", "uz")
    check("tayyor xulosa tegilmaydi", {s["layout"]: s for s in d["slides"]}["conclusion"]["content"] == "Tayyor xulosa.")
asyncio.run(fill_tests())

# ── rasm cho'zilmaydi
tmp = tempfile.mkdtemp()
def make_img(path, size):
    im = Image.new("RGB", size, (200, 220, 240)); d = ImageDraw.Draw(im)
    d.ellipse((size[0] // 2 - 120, size[1] // 2 - 120, size[0] // 2 + 120, size[1] // 2 + 120), fill=(220, 60, 60))
    im.save(path)

class FakeTogether:
    async def generate_slide_image(self, topic, title, lang):
        path = os.path.join(tmp, "s.png"); make_img(path, (1024, 768)); return path
    async def generate_panoramic_image(self, topic, title, lang):
        path = os.path.join(tmp, "p.png"); make_img(path, (1024, 768)); return path
    async def generate_cover_image(self, topic, lang):
        path = os.path.join(tmp, "c.png"); make_img(path, (1024, 768)); return path

async def build():
    svc = DocumentService.__new__(DocumentService)
    svc.documents_dir = tmp; svc.together = FakeTogether(); svc.use_icons = True; svc._last_used_icons = set()
    slides = [{"title": "", "content": "", "layout": "cover"},
              {"title": "Gorizontal", "content": short, "layout": "horizontal_image"},
              {"title": "O'ng rasm", "content": short, "layout": "right_image"},
              {"title": "Uzun", "content": long_, "layout": "text_with_numbers"},
              {"title": "Xulosa", "content": medium, "layout": "conclusion"}]
    return Presentation(await svc.create_presentation_with_smart_images("Sinov", {"slides": slides}, "Talaba", "uz"))

prs = asyncio.run(build())
slides_out = list(prs.slides)
expected = 5 + (len(slide_fit.plan_text(long_)) - 1) + (len(slide_fit.plan_text(medium)) - 1)
check("uzun slaydlar rejadagidek bo'linib, jami slayd soni to'g'ri", len(slides_out) == expected, (len(slides_out), expected))

def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width / 914400 > 3]
for label, index in (("muqova", 0), ("gorizontal", 1), ("o'ng rasm", 2)):
    pic = pictures(slides_out[index])[0]
    box = pic.width / pic.height
    shown = (1024 * (1 - pic.crop_left - pic.crop_right)) / (768 * (1 - pic.crop_top - pic.crop_bottom))
    check(f"{label}: rasm nisbati saqlangan (cho'zilmagan)", abs(box - shown) / box < 0.02, (round(box, 2), round(shown, 2)))

def body_boxes(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.top / 914400 >= 1.9]
sizes = [p.font.size.pt for s in slides_out[3:] for b in body_boxes(s) for p in b.text_frame.paragraphs if p.font.size]
check("uzun matn slaydlarida shrift o'qiladigan (>= 16 pt)", sizes and min(sizes) >= 16, sizes)
check("davomi slaydining sarlavhasi bor", any("(davomi)" in sh.text_frame.text for s in slides_out[3:] for sh in s.shapes if sh.has_text_frame))

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
