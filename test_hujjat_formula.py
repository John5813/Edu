"""Hujjatdagi formulalar: aniq fan mustaqil ishida determinant va ko'paytmalar buzilmaydi.

- `clean_text` ko'paytmani ("2*5", "1*5*9"), modulni ("|A|") va indeksni ("a_{11}") o'chirmaydi;
  markdown ajratish ("**so'z**", "*so'z*") esa avvalgidek olib tashlanadi;
- `$...$` ichidagi formula tozalashga tushmaydi; abzatslar (`keep_paragraphs`) saqlanadi;
- aniq fan mavzusida modelga formulani LaTeX bilan yozish aytiladi (boshqa mavzuda — yo'q);
- Word: qator ichidagi formula belgilarga o'giriladi, matritsa/determinant — jadval (qator va ustunlar);
- snoskalar soni avvalgidek: har betga bittadan.

    python test_hujjat_formula.py
"""
import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from docx import Document
from services import doc_math
from services.ai_service import clean_text, math_rule
from services.document_service import DocumentService

print("1) clean_text matematikani buzmaydi")
out = clean_text("det(A) = (2 * 5) - (3 * 4); det(B) = (1*5*9 + 2*6*7) va |A| hamda a_11 a_22*a_33.")
check("ko'paytma saqlanadi", "(2 * 5) - (3 * 4)" in out and "1*5*9 + 2*6*7" in out, out)
check("modul belgisi saqlanadi", "|A|" in out, out)
check("indeks saqlanadi", "a_11 a_22*a_33" in out, out)
out = clean_text("Bu **muhim** va *asosiy* hamda __yana__ tushuncha.")
check("markdown ajratish olib tashlanadi", out == "Bu muhim va asosiy hamda yana tushuncha.", out)
out = clean_text(r"Formula $\det(A) = a_{11}a_{22} - a_{12}a_{21}$ va $$\begin{vmatrix} a & b \\ c & d \end{vmatrix}$$ tayyor.")
check("$...$ ichiga tegilmaydi", r"$\det(A) = a_{11}a_{22} - a_{12}a_{21}$" in out
      and r"\begin{vmatrix} a & b \\ c & d \end{vmatrix}" in out, out)
check("formuladan tashqaridagi teskari chiziq avvalgidek o'chadi", clean_text(r"oddiy \matn") == "oddiy matn")
check("markdown jadval chiziqlari olib tashlanadi",
      clean_text("| a | b |\n|---|---|\n| 1 | 2 |").split() == ["a", "b", "1", "2"],
      clean_text("| a | b |\n|---|---|\n| 1 | 2 |"))
out = clean_text("Birinchi  abzats\nshu yerda.\n\nIkkinchi abzats.", keep_paragraphs=True)
check("abzatslar saqlanadi", out == "Birinchi abzats shu yerda.\n\nIkkinchi abzats.", repr(out))
check("sukut bo'yicha bitta qator", "\n" not in clean_text("bir\n\nikki"))

print("2) Aniq fan ko'rsatmasi")
check("determinant mavzusi — formula qoidasi", "$" in math_rule("Determinantni xisoblash usullari", "uz"))
check("matritsa mavzusi (ru)", "LaTeX" in math_rule("Матрицы и определители", "ru"))
check("tarix mavzusida yo'q", math_rule("Amir Temur davlati", "uz") == "")

print("3) Formulani belgilarga o'girish")
check("qator ichidagi formula", doc_math.plain(r"C_{ij} = (-1)^{i+j} M_{ij}") == "Cᵢⱼ = (-1)ⁱ⁺ʲ Mᵢⱼ",
      doc_math.plain(r"C_{ij} = (-1)^{i+j} M_{ij}"))
check("ko'paytma nuqta bilan", doc_math.plain(r"2 \cdot 5") == "2 · 5", doc_math.plain(r"2 \cdot 5"))
check("narx formula emas", doc_math.inline("Narx $5 va $10 edi.") == "Narx $5 va $10 edi.")
parts = doc_math.segments(r"\det(B) = \begin{vmatrix} 1 & 2 \\ 3 & 4 \end{vmatrix} = -2")
check("determinant — qator va ustunlar", parts == [("text", "det(B) ="), ("matrix", "vmatrix", [["1", "2"], ["3", "4"]]),
                                                  ("text", "= -2")], parts)
check("abzatsdagi alohida formula ajratiladi",
      [k for k, _ in doc_math.blocks(r"Agar $$A = \begin{pmatrix} a & b \\ c & d \end{pmatrix}$$ bo'lsa, $x^2$.")]
      == ["text", "math", "text"])

print("4) Word hujjati")
text = clean_text(r"""Birinchi abzats: $\det(A) = ad - bc$ formulasi.

$$\det(B) = \begin{vmatrix} 1 & 2 & 3 \\ 4 & 5 & 6 \\ 7 & 8 & 9 \end{vmatrix} = 0$$

Ikkinchi abzats: ko'paytma 2*5 va $C_{ij} = (-1)^{i+j} M_{ij}$.""", keep_paragraphs=True)
svc = DocumentService.__new__(DocumentService)
doc = Document()
nxt = svc._add_body_with_footnotes(doc, text, ["Manba"], 1)
body = "\n".join(p.text for p in doc.paragraphs)
check("hujjatda $ va \\ qolmaydi", "$" not in body and "\\" not in body, body)
check("qator ichidagi formula belgilarda", "det(A) = ad - bc" in body and "Cᵢⱼ = (-1)ⁱ⁺ʲ Mᵢⱼ" in body, body)
check("ko'paytma saqlangan", "2*5" in body, body)
check("determinant jadval: 3 qator, ichida 1..9", len(doc.tables) == 1 and len(doc.tables[0].rows) == 3
      and [c.text for c in doc.tables[0].rows[1].cells][1:4] == ["4", "5", "6"],
      [[c.text for c in r.cells] for r in doc.tables[0].rows] if doc.tables else None)
check("abzatslar alohida", sum(1 for p in doc.paragraphs if p.text.strip()) == 2,
      [p.text for p in doc.paragraphs])
check("snoska qo'yildi (qisqa matn — bitta)", nxt == 2, nxt)

long_text = "\n\n".join("Gap " + str(i) + " " + "so'z " * 120 + "tugadi." for i in range(6))
doc = Document()
nxt = svc._add_body_with_footnotes(doc, long_text, ["A", "B"], 1)
check("uzun matnda snoskalar har betga bittadan", nxt - 1 == max(1, round(len(long_text) / 2550)), nxt - 1)

print()
if FAILS:
    print(f"XATOLAR: {len(FAILS)}"); sys.exit(1)
print("Hammasi joyida")
