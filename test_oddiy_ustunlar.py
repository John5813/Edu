"""Oddiy taqdimot: ustunsiz kelgan ikki/uch ustunli slayd yarim bo'sh chiqmaydi.

Model yetishmagan slaydlarni qo'shimcha yozganda `columns` bermasdi va butun
matn birinchi ustunga tushib, qolgani bo'sh qolardi (15-slayddan keyin).

    python test_oddiy_ustunlar.py
"""
import asyncio, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

from pptx import Presentation
from services import slide_layouts
# Bu test eski shablonlarning chizilishini tekshiradi: shablon tanlash o'chirilgan
# (tanlash alohida test_oddiy_shablonlar.py da tekshiriladi).
slide_layouts.assign = lambda slides, *a, **k: slides
from services.document_service import DocumentService, resolve_columns, split_into_columns

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

BODY = ("Birinchi gap mahsulot hayoti davri haqida. Ikkinchi gap bozor tendensiyalarini yoritadi. "
        "Uchinchi gap tashqi axborot oqimlarini tushuntiradi. To'rtinchi gap xulosani beradi.")

two = split_into_columns(BODY, 2)
check("matn 2 ta ustunga gap chegarasida bo'linadi", len(two) == 2 and all(c.endswith(".") for c in two), two)
check("3 ta ustunga bo'linadi", len(split_into_columns(BODY, 3)) == 3)
check("bitta gap bo'linmaydi", split_into_columns("Faqat bitta gap.", 2) == ["Faqat bitta gap."])
check("bo'sh matn — ustun yo'q", split_into_columns("", 2) == [])
check("bo'sh ustun tashlanib matn qayta taqsimlanadi",
      len(resolve_columns({"columns": [{"column_content": BODY}, {"column_content": ""}]}, 2)) == 2)
check("to'g'ri berilgan ustunlar o'zgarmaydi",
      [c["column_content"] for c in resolve_columns(
          {"columns": [{"column_content": "A. B."}, {"column_content": "C. D."}]}, 2)] == ["A. B.", "C. D."])

async def build():
    svc = DocumentService.__new__(DocumentService)
    svc.documents_dir = tempfile.mkdtemp(); svc.together = None
    svc.use_icons = True; svc._last_used_icons = set()
    slides = [
        {"title": "", "content": "", "layout": "cover"},
        {"title": "Ikki ustun", "content": BODY, "layout": "two_column"},
        {"title": "Uch ustun", "content": BODY, "layout": "three_column"},
        {"title": "Bitta gap", "content": "Faqat bitta qisqa gap.", "layout": "two_column"},
        {"title": "Bo'sh ustunli", "content": BODY, "layout": "two_column",
         "columns": [{"column_content": "Faqat birinchi."}, {"column_content": ""}]},
    ]
    return Presentation(await svc.create_presentation_with_smart_images(
        "Sinov", {"slides": slides}, "Talaba", "uz"))

def bodies(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
            and sh.top / 914400 >= 1.9]

prs = asyncio.run(build())
slides = list(prs.slides)
b2, b3, b1, be = (bodies(slides[i]) for i in (1, 2, 3, 4))
check("ikki ustunli: ikkala ustunda matn bor", len(b2) == 2 and all(len(b.text_frame.text) > 40 for b in b2),
      [len(b.text_frame.text) for b in b2])
check("uch ustunli: uchala ustunda matn bor", len(b3) == 3, len(b3))
check("bitta gap: bitta quti, butun kenglikda", len(b1) == 1 and b1[0].width / 914400 > 11, [b.width for b in b1])
check("bo'sh ustun yarim bo'sh slayd qoldirmaydi", all(b.width / 914400 > 5 for b in be) and
      (len(be) == 2 or be[0].width / 914400 > 11), [(b.left, b.width) for b in be])

# ── Ikonkalar shablon rangida va doira ichida, matn chetiga tekislangan
import io
import numpy as np
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE
from services.template_service import TemplateService

async def build_with_template(tid):
    svc = DocumentService.__new__(DocumentService)
    svc.documents_dir = tempfile.mkdtemp(); svc.together = None
    svc.use_icons = True; svc._last_used_icons = set()
    slides = [{"title": "", "content": "", "layout": "cover"},
              {"title": "Ikki ustun", "content": "", "layout": "two_column",
               "columns": [{"column_content": "Bozor tendensiyalari va raqobatchilar talablari muhimdir. Mahsulot maqsadi aniq."},
                           {"column_content": "Ma'lumotlar bazasi tahlili mahsulotni rivojlantirishga yordam beradi."}]}]
    return Presentation(await svc.create_presentation_with_smart_images(
        "Sinov", {"slides": slides}, "Talaba", "uz",
        template_service=TemplateService(), template_id=tid))

ts = TemplateService()
for tid in ("template_10", "template_5"):
    slide = list(asyncio.run(build_with_template(tid)).slides)[1]
    accent = tuple(ts.get_readable_colors(tid)["title"])
    circles = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    icons = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE
             and sh.width / 914400 < 1.5]
    check(f"{tid}: har ikonka doira ichida", len(circles) == 2 and len(icons) == 2, (len(circles), len(icons)))
    colours = []
    for icon in icons:
        arr = np.asarray(Image.open(io.BytesIO(icon.image.blob)).convert("RGBA"))
        colours.append(tuple(int(x) for x in np.median(arr[arr[..., 3] > 200][:, :3], axis=0)))
    check(f"{tid}: ikonka shablon rangida (ko'k emas)", all(c == accent for c in colours), (colours, accent))
    text_left = [sh.left for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                 and sh.top / 914400 >= 1.9]
    check(f"{tid}: doira ustun matni bilan bir chapdan", min(abs(c.left - (t + 91440)) for c in circles for t in text_left) < 20000,
          [(c.left, t) for c in circles for t in text_left][:2])


print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
