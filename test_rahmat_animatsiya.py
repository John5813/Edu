""""Rahmat" animatsiyasi: taqdimot yuborilgach mijoz oxiriga harakatlanuvchi personajli sahifa qo'shadi.

- katalogdagi har animatsiyaning fayllari bor (manba, namuna video, muqova);
- sahifa qo'shiladi: GIF personaj panel rangida (chegara ko'rinmaydi), ovoz — "audio" aloqasi bilan,
  slayd ochilganda o'zi boshlanadi, yozuvlar navbat bilan paydo bo'ladi;
- funksiya sukut bo'yicha O'CHIQ, admin "🎛 Funksiyalar boshqaruvi" dan yoqadi;
- tugma faqat yoqilganda chiqadi; mini oyna: ro'yxat, fayllar, egasi bo'lmagan so'rov rad etiladi,
  tanlov taqdimotni bot orqali yuboradi, har tanlov ASL taqdimotdan yig'iladi.

    python test_rahmat_animatsiya.py
"""
import asyncio
import hashlib
import hmac
import json
import os
import shutil
import sys
import tempfile
import time
from urllib.parse import urlencode

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond:
        FAILS.append(name)


from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from services.premium_presentation import themes
from services.thanks_anim import catalog
from services.thanks_anim import slide as thanks_slide

TMP = tempfile.mkdtemp()
theme = themes.for_deck("Moliyaviy savodxonlik", "toza", "kam")

print("1) Katalog")
keys = [a.key for a in catalog.ANIMATIONS]
check("nomlar takrorlanmaydi", len(keys) == len(set(keys)))
missing = [a.key for a in catalog.ANIMATIONS if not a.ready]
check("har animatsiyaning fayllari bor", not missing, missing)
check("muqova rasmlari bor", all(os.path.exists(a.path("poster.jpg")) for a in catalog.ANIMATIONS))
check("noma'lum nom — None", catalog.get("yoq") is None and catalog.get("../x") is None)
check("nom tilga qarab", catalog.get("pul_mushuk").name("ru") == "Кот считает деньги")

print("2) Sahifa qo'shiladi")
prs = Presentation()
prs.slide_width, prs.slide_height = 12192000, 6858000
for _ in range(3):
    prs.slides.add_slide(prs.slide_layouts[6])
base = os.path.join(TMP, "base.pptx")
prs.save(base)
anim = catalog.get("anime_qiz")
out = thanks_slide.add_thanks_slide(base, os.path.join(TMP, "out.pptx"), anim, theme, "uz", "Javlonbek")
deck = Presentation(out)
check("bitta sahifa qo'shildi", len(deck.slides) == 4)
last = deck.slides[-1]
texts = " ".join(sh.text_frame.text for sh in last.shapes if sh.has_text_frame)
check("rahmat yozuvi va muallif", "E'tiboringiz uchun rahmat!" in texts and "Javlonbek" in texts, texts)
pics = [sh for sh in last.shapes if sh.shape_type == 13]
gif = [p for p in pics if p.image.content_type == "image/gif"]
check("personaj — GIF", len(gif) == 1)
audio = [p for p in pics if p._element.find(".//" + qn("a:audioFile")) is not None]
check("ovoz bor va slayddan tashqarida", len(audio) == 1 and audio[0].left < 0)
check("GIF slayd ichida", gif and gif[0].left > 0 and gif[0].top >= 0
      and gif[0].left + gif[0].width <= deck.slide_width and gif[0].top + gif[0].height <= deck.slide_height)

# GIF foni panel rangi bilan aynan bir xil — chegara ko'rinmaydi
import io
with Image.open(io.BytesIO(gif[0].image.blob)) as im:
    frames = im.n_frames
    bg = max(im.convert("RGB").getcolors(1 << 16))[1]
panels = [sh for sh in last.shapes if sh.shape_type == 1 and sh.fill.type == 1
          and abs(sh.width - deck.slide_width * 860 / 1920) < 20000]
panel_rgb = tuple(panels[0].fill.fore_color.rgb) if panels else None
check("GIF harakatlanadi (ko'p kadr)", frames > 20, frames)
check("GIF foni = panel rangi", panel_rgb == bg, (panel_rgb, bg))

timing = last._element.find(qn("p:timing"))
xml = timing is not None and __import__("lxml.etree", fromlist=["x"]).tostring(timing).decode()
check("slayd ochilganda o'zi boshlanadi", xml and 'evt="onBegin"' in xml and 'nodeType="clickEffect"' not in xml)
check("ovoz o'zi boshlanadi", xml and 'cmd="playFrom(0.0)"' in xml and "<p:audio>" in xml)
check("yozuvlar navbat bilan paydo bo'ladi", xml and xml.count('presetClass="entr"') == 5 and xml.count("<p:bldP") == 5,
      xml and xml.count('presetClass="entr"'))
import zipfile
with zipfile.ZipFile(out) as z:
    rels = z.read("ppt/slides/_rels/slide4.xml.rels").decode()
    first = z.read("ppt/slides/_rels/slide1.xml.rels").decode()
check("ovoz aloqasi — audio (video emas)", "relationships/audio" in rels and "relationships/video" not in rels)
check("boshqa sahifalarga tegilmadi", "audio" not in first)

print("3) Tillar")
check("rus", thanks_slide.texts("ru")[1] == "Спасибо за внимание!")
check("kirill", thanks_slide.texts("uz-cyrl")[1].startswith("Эътиборингиз"))
check("noma'lum til — ingliz", thanks_slide.texts("de")[1].startswith("Thank you"))
out_ru = thanks_slide.add_thanks_slide(base, os.path.join(TMP, "ru.pptx"), catalog.get("pul_mushuk"), theme, "ru", "")
ru_texts = " ".join(sh.text_frame.text for sh in Presentation(out_ru).slides[-1].shapes if sh.has_text_frame)
check("muallifsiz ham chiziladi", "Спасибо" in ru_texts and "Автор" not in ru_texts, ru_texts)

print("4) Funksiya sukut bo'yicha o'chiq, admin yoqadi")
import database.database as dbmod
from database.database import Database, init_db

dbmod.DATABASE_FILE = os.path.join(TMP, "t.db")
from services import thanks_anim


async def db_part():
    await init_db()
    db = Database()
    off = await db.get_feature_status(thanks_anim.FEATURE, default=False)
    other = await db.get_feature_status("boshqa_funksiya")
    from bot.handlers.admin import _feature_keyboard

    kb = await _feature_keyboard(db)
    row = [b for r in kb.inline_keyboard for b in r if "Rahmat animatsiyasi" in b.text]
    await db.set_feature_status(thanks_anim.FEATURE, True)
    on = await db.get_feature_status(thanks_anim.FEATURE, default=False)
    kb2 = await _feature_keyboard(db)
    row2 = [b for r in kb2.inline_keyboard for b in r if "Rahmat animatsiyasi" in b.text]
    return off, other, row, on, row2


off, other, row, on, row2 = asyncio.run(db_part())
check("sukut — o'chiq", off is False)
check("boshqa funksiyalar sukuti o'zgarmadi (yoqilgan)", other is True)
check("admin panelida tugma bor", row and "O'chirilgan" in row[0].text and row[0].callback_data == "toggle_thanks_anim_on",
      row and row[0].text)
check("yoqilgandan keyin", on is True and row2 and "Yoqilgan" in row2[0].text
      and row2[0].callback_data == "toggle_thanks_anim_off")

print("5) Tugma va mini oyna")
import webapp

webapp.WEBAPP_DOMAIN = "edufayl.uz"
thanks_anim.STORE = os.path.join(TMP, "store")
kb = thanks_anim.offer(base, topic="Moliyaviy savodxonlik", style="toza", volume="kam", language="uz",
                       author="Javlonbek", filename="Taqdimot_Moliya.pptx", user_id=42, chat_id=42, user_lang="uz")
btn = kb.inline_keyboard[0][0] if kb else None
check("tugma mini oynani ochadi", btn and btn.web_app and btn.web_app.url.startswith("https://edufayl.uz/anim?token="),
      btn and btn.web_app)
token = btn.web_app.url.split("token=")[1].split("&")[0]
info = webapp.DOC_TOKENS.get(token, {})
check("asl taqdimot nusxasi saqlandi", info.get("kind") == "thanks_anim" and os.path.exists(info.get("file_path", "")))


def init_data(user_id):
    fields = {"auth_date": str(int(time.time())), "user": json.dumps({"id": user_id}), "query_id": "q"}
    check_str = "\n".join(f"{k}={fields[k]}" for k in sorted(fields))
    secret = hmac.new(b"WebAppData", os.environ["BOT_TOKEN"].encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check_str.encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


class FakeBot:
    def __init__(self):
        self.sent = []

    async def send_document(self, chat_id, document, caption=None, **kw):
        self.sent.append((chat_id, document.path, document.filename, caption,
                          len(Presentation(document.path).slides)))


async def web_part():
    from aiohttp.test_utils import TestClient, TestServer

    from webapp.server import create_web_app

    webapp.BOT = FakeBot()
    client = TestClient(TestServer(create_web_app()))
    await client.start_server()
    try:
        page = await client.get(f"/anim?token={token}")
        lst = await (await client.get("/api/anim/list?lang=ru")).json()
        media = await client.get("/anim-media/anime_qiz/preview.mp4")
        bad = await client.get("/anim-media/anime_qiz/packed.mp4")
        trav = await client.get("/anim-media/..%2F..%2Fconfig/preview.mp4")
        noauth = await client.post(f"/api/anim/{token}", json={"key": "anime_qiz"})
        stranger = await client.post(f"/api/anim/{token}", json={"key": "anime_qiz"},
                                     headers={"X-Telegram-Init-Data": init_data(7)})
        unknown = await client.post(f"/api/anim/{token}", json={"key": "yoq"},
                                    headers={"X-Telegram-Init-Data": init_data(42)})
        ok1 = await client.post(f"/api/anim/{token}", json={"key": "anime_qiz"},
                                headers={"X-Telegram-Init-Data": init_data(42)})
        ok2 = await client.post(f"/api/anim/{token}", json={"key": "pul_mushuk"},
                                headers={"X-Telegram-Init-Data": init_data(42)})
        return (page.status, lst, media.status, bad.status, trav.status, noauth.status, stranger.status,
                unknown.status, ok1.status, ok2.status)
    finally:
        await client.close()


page, lst, media, bad, trav, noauth, stranger, unknown, ok1, ok2 = asyncio.run(web_part())
check("sahifa ochiladi", page == 200)
items = lst.get("items", [])
check("ro'yxatda hamma animatsiya, tilda", len(items) == len(catalog.ANIMATIONS)
      and any(i["name"] == "Кот считает деньги" for i in items), len(items))
check("namuna video beriladi", media == 200)
check("manba va boshqa fayllar berilmaydi", bad == 404 and trav == 404, (bad, trav))
check("initData siz — rad", noauth == 401)
check("boshqa odam — rad", stranger == 403)
check("noma'lum animatsiya — rad", unknown == 400)
check("tanlov ishlaydi", ok1 == 200 and ok2 == 200, (ok1, ok2))
sent = webapp.BOT.sent
check("bot ikki marta yubordi, asl nomi bilan", len(sent) == 2 and all(s[2] == "Taqdimot_Moliya.pptx" for s in sent), sent)
check("har tanlov asl taqdimotdan (ustma-ust emas)", [s[4] for s in sent] == [4, 4], [s[4] for s in sent])
check("vaqtinchalik fayllar o'chirildi", not [f for f in os.listdir(thanks_anim.STORE) if "_" in f],
      os.listdir(thanks_anim.STORE))

shutil.rmtree(TMP, ignore_errors=True)
print("\nNATIJA:", "HAMMASI O'TDI" if not FAILS else f"{len(FAILS)} ta xato: {FAILS}")
sys.exit(1 if FAILS else 0)
