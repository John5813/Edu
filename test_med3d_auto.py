"""3D Anatomiya: mavzu → AI reja → atlasdan 3D modellar → PowerPoint. Tarmoqsiz (atlas va AI soxta).

- lug'at: ikkala tana, ayol/erkakka xos organlar, o'zbekcha nomlar, `inside` bog'lanishlari;
- reja tekshiruvi: anatomiya emas → NotAnatomy; noma'lum organ, ichma-ich organlar, teri, tana almashishi;
- yashirin organ: kamera klapanga yaqinlashsa yurak chiqariladi;
- sahnalar: mayda organga yaqinlashganda yangi sahna, 3D ramka MAX_FRAME dan oshmaydi;
- PPTX: slaydlar soni, Morph, bir xil GLB bitta faylda, yorliqlar;
- bot: tugma → ogohlantirish (Ha/Yo'q) → "Yo'q" — bu siz uchun emas, "Ha" — mavzu; manba so'ralmaydi.

    python test_med3d_auto.py
"""
import asyncio, copy, re, os, sys, tempfile, zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import numpy as np
from services.med3d import auto
from services.med3d.deck import MAX_FRAME
from services.med3d.lugat import BY_ID

LEX = auto.lexicon()
MALE, FEMALE = LEX["bodies"]["male"], LEX["bodies"]["female"]

print("1) Lug'at")
check("ikkala tana bor, organlar ko'p", len(MALE) > 60 and len(FEMALE) > 60, (len(MALE), len(FEMALE)))
check("ayolga xos organlar faqat ayolda", all(o in FEMALE and o not in MALE for o in ("bachadon", "tuxumdonlar", "yoldosh")))
check("erkakka xos organlar faqat erkakda", all(o in MALE and o not in FEMALE for o in ("moyaklar", "prostata")))
check("oshqozon faqat erkakda, o'pka faqat ayolda", "oshqozon" in MALE and "oshqozon" not in FEMALE
      and "opkalar" in FEMALE and "opkalar" not in MALE)
check("har organda qism, chegara va namuna", all(v["parts"] and len(v["samples"]) > 5 and v["tris"] > 0
                                               for b in (MALE, FEMALE) for v in b.values()))
check("o'zbekcha nomlar bor", all(BY_ID[o]["uz"] for b in (MALE, FEMALE) for o in b))
check("AI jadvali faqat mavjud organlar", "bachadon | Bachadon | Uterus | female" in auto.organ_table()
      and "ong_opka" in auto.organ_table())

print("2) Reja tekshiruvi")
try:
    auto.normalize({"anatomy": False, "reason": "Tarixga oid", "slides": []}, 5, "uz")
    check("anatomiya emas → NotAnatomy", False)
except auto.NotAnatomy as e:
    check("anatomiya emas → NotAnatomy (sabab bilan)", "Tarix" in str(e))
raw = {"anatomy": True, "body": "male", "names": {"bachadon": "Uterus"}, "slides": [
    {"title": "Ko'rinish", "text": "Matn", "show": ["bachadon", "tuxumdonlar", "nimadir"], "rot": [999, 0, 0],
     "labels": ["bachadon", "oshqozon"]},
    {"title": "Bachadon", "text": "Matn", "show": ["bachadon", "teri"], "focus": ["bachadon"], "rot": [0, 30, 0]},
    {"title": "Tuxumdon", "text": "Matn", "show": ["tuxumdonlar"], "rot": [0, -30, 0], "exit": "qayoqqa"},
]}
p = auto.normalize(raw, 3, "uz")
check("tana organlarga qarab almashtirildi (erkak → ayol)", p["body"] == "female", p["body"])
check("noma'lum organ tashlandi, yorliq faqat ko'rinayotganlardan", p["slides"][0]["show"] == ["bachadon", "tuxumdonlar"]
      and p["slides"][0]["labels"] == ["bachadon"], p["slides"][0])
check("teri boshqa organ bilan birga bo'lmaydi", p["slides"][1]["show"] == ["bachadon"])
check("burchak chegaralangan, exit tuzatilgan", p["slides"][0]["rot"][0] == 45 and p["slides"][2]["exit"] == "up")
turns = [{"rot": [0, y, 0]} for y in (0, 30, -60, 180, 90, -120)]
auto.smooth_turns(turns)
ys = [t["rot"][1] for t in turns]
check("burilish silliq: qo'shni slaydlar orasida ≤ 45°", all(abs(b - a) <= auto.TURN_MAX + 1e-6 for a, b in zip(ys, ys[1:])), ys)
check("o'zbekcha nom lug'atdan (AI nomi emas)", p["names"]["bachadon"] == "Bachadon")
nested = auto.normalize({"anatomy": True, "body": "male", "slides": [
    {"title": "a", "text": "b", "show": ["skelet", "bosh_suyagi", "yurak"], "rot": [0, 0, 0]}] * 3}, 3, "en")
check("ichma-ich organlar: skelet bilan bosh suyagi birga emas", nested["slides"][0]["show"] == ["skelet", "yurak"],
      nested["slides"][0]["show"])
check("boshqa tilda nom AI dan yoki inglizcha", nested["names"]["yurak"] == "Heart")
try:
    auto.normalize({"anatomy": True, "slides": [{"show": ["yurak"]}]}, 8, "uz")
    check("slaydlar kam bo'lsa — xato (qayta urinish)", False)
except ValueError:
    check("slaydlar kam bo'lsa — xato (qayta urinish)", True)

joint = auto.normalize({"anatomy": True, "body": "male", "slides": [
    {"title": "a", "text": "b", "show": ["chakka_jag_bogimi", "tishlar"], "rot": [0, 0, 0]},
    {"title": "a", "text": "b", "show": ["pastki_jag", "tishlar", "medial_qanotsimon"], "rot": [0, 0, 0]},
    {"title": "a", "text": "b", "show": ["tizza_boylamlari"], "rot": [0, 0, 0]}]}, 3, "uz")
js = joint["slides"]
check("bo'g'im slaydiga suyagi qo'shiladi (havoda suzmaydi)", "bosh_suyagi" in js[0]["show"]
      and js[0]["support"] == ["bosh_suyagi"], js[0]["show"])
check("tishlar faqat pastki jag' bilan — olinadi (yuqori qatori havoda qolardi)", "tishlar" not in js[1]["show"],
      js[1]["show"])
check("tizza boylamlari — son va boldir suyaklari bilan", {"son_suyagi", "boldir_suyaklari"} <= set(js[2]["show"]),
      js[2]["show"])
from services.med3d import director as _dir
from services.med3d.camera import sides as _sides
_male = auto.lexicon()["bodies"]["male"]
check("juft mushak ikki bo'lak (o'ng/chap), yurak — bitta", len(_sides(_male["chaynov_mushagi"]["samples"])) == 2
      and len(_sides(_male["yurak"]["samples"])) == 1)
_view = _dir.best_view(_male, {"show": ["bosh_suyagi", "tishlar", "chaynov_mushagi"], "focus": ["chaynov_mushagi"],
                               "rot": [0, 0, 0]})
check("chaynov mushagi oldidan emas, yon tomondan ko'rsatiladi", abs(_view[1]) >= 30, _view)

sick_plan = auto.normalize({"anatomy": True, "kind": "disease", "side": "right", "body": "male", "slides": [
    {"title": "a", "text": "b", "stage": "anatomy", "show": ["kurak", "yelka_suyagi", "qirra_usti_mushagi"]},
    {"title": "a", "text": "b", "stage": "disease", "show": ["kurak", "yelka_suyagi", "qirra_usti_mushagi"],
     "lesion": ["qirra_usti_mushagi", "yurak"], "picture": "torn supraspinatus tendon"},
    {"title": "a", "text": "b", "stage": "treatment", "show": ["kurak", "qirra_usti_mushagi"], "picture": "x"}]},
    3, "uz")
ss = sick_plan["slides"]
check("kasallik: bosqichlar va zararlangan organ (ko'rinishda bo'lmagani olinmaydi)",
      sick_plan["kind"] == "disease" and [x["stage"] for x in ss] == ["anatomy", "disease", "treatment"]
      and ss[1]["lesion"] == ["qirra_usti_mushagi"] and "qirra_usti_mushagi" in ss[1]["labels"], ss)
check("rasm faqat zararlangan organli slaydda", bool(ss[1]["picture"]) and not ss[2]["picture"])
_sc = auto.normalize({"anatomy": True, "kind": "disease", "slides": [
    {"title": "a", "text": "b", "stage": "treatment", "show": ["yurak"], "scene": "a patient with a doctor"},
    {"title": "a", "text": "b", "stage": "disease", "show": ["yurak"], "scene": "x"},
    {"title": "a", "text": "b", "stage": "treatment", "show": ["yurak"], "scene": "y"},
    {"title": "a", "text": "b", "stage": "treatment", "show": ["yurak"], "scene": "z"}]}, 4, "uz")["slides"]
check("davolash tasviri faqat davolash slaydlarida, ko'pi bilan 2 ta",
      [bool(x["scene"]) for x in _sc] == [True, False, True, False], [x["scene"] for x in _sc])
check("bitta tomon (o'ng yelka)", sick_plan["side"] == "right")
_lat = auto.normalize({"anatomy": True, "kind": "disease", "slides": [
    {"title": "a", "text": "b", "show": ["yurak"], "latin": "Infarctus myocardii"}] * 3}, 3, "uz")
check("kasallikning lotincha nomi saqlanadi", _lat["slides"][0]["latin"] == "Infarctus myocardii")
_pics = {"slides": [{"picture": "x"}, {"picture": ""}]}
_saved = os.environ.pop("TOGETHER_API_KEY", None)
check("rasm xizmati kalitsiz — tayyor emas (3D boshlanmaydi)", not auto.pictures_ready())
if _saved:
    os.environ["TOGETHER_API_KEY"] = _saved
_orig_plan, _orig_ready = auto.plan, auto.pictures_ready
auto.plan = lambda *a, **k: copy.deepcopy(sick_plan)
auto.pictures_ready = lambda: False
try:
    asyncio.run(auto.build_deck("Rotator manjet", 3))
    check("rasm xizmati ishlamasa — 3D boshlanmaydi (PicturesUnavailable)", False)
except auto.PicturesUnavailable:
    check("rasm xizmati ishlamasa — 3D boshlanmaydi (PicturesUnavailable)", True)
finally:
    auto.plan, auto.pictures_ready = _orig_plan, _orig_ready
check("kasallik rasmlari uchun kuchli modellar zanjiri", auto.picture_models()[0].startswith("ByteDance-Seed/Seedream-4"))
plain = auto.normalize({"anatomy": True, "slides": [{"title": "a", "text": "b", "show": ["yurak"],
                                                      "lesion": ["yurak"], "picture": "x"}] * 3}, 3, "uz")
check("anatomiya mavzusida bosqich, qizil o'choq va rasm yo'q",
      all(not x["stage"] and not x["lesion"] and not x["picture"] for x in plain["slides"]))
from services.med3d import manba as _manba
_one = auto.one_side(auto.lexicon()["bodies"]["male"], "right", _manba.atlas("male"))
check("o'ng tomon: kurak suyagining faqat o'ng nusxasi (x < 0)", max(_one["kurak"]["max"][0], 0) == 0
      and all(p[0] < 0 for p in _one["kurak"]["samples"]), _one["kurak"]["max"])

print("3) Yashirin organ va sahnalar")
slides = [{"title": "", "text": "", "show": ["yurak", "yurak_klapanlari"], "focus": ["yurak_klapanlari"],
           "labels": ["yurak", "yurak_klapanlari"], "rot": [0, 0, 0], "exit": "up"}]
auto.reveal_hidden(MALE, slides)
check("klapanga yaqinlashilsa yurak chiqariladi", slides[0]["show"] == ["yurak_klapanlari"]
      and slides[0]["labels"] == ["yurak_klapanlari"], slides[0])
plan = {"body": "male", "names": {}, "slides": [
    {"title": "Hazm", "text": "t", "show": ["jigar", "oshqozon", "ingichka_ichak", "yogon_ichak"], "focus": None,
     "labels": ["jigar", "oshqozon"], "rot": [0, 0, 0], "exit": "up"},
    {"title": "Jigar", "text": "t", "show": ["jigar", "ot_pufagi"], "focus": None, "labels": ["jigar"],
     "rot": [0, -30, 0], "exit": "up"},
    {"title": "O't", "text": "t", "show": ["ot_pufagi", "ot_yollari"], "focus": None, "labels": ["ot_pufagi"],
     "rot": [0, 30, 0], "exit": "up"},
    {"title": "Yana", "text": "t", "show": ["jigar", "oshqozon", "ingichka_ichak", "yogon_ichak"], "focus": None,
     "labels": [], "rot": [0, 40, 0], "exit": "up"}]}
runs = auto.split_runs(MALE, plan["slides"])
check("mayda organ (o't pufagi) uchun alohida sahna", len(runs) >= 2 and [2] in [m for m, _, _ in runs] or
      any(2 in m and 0 not in m for m, _, _ in runs), [m for m, _, _ in runs])

heart = [{"show": ["skelet", "yurak", "aorta_yoyi", "opka_arteriyasi", "yuqori_kovak_vena"]},
         {"show": ["yurak", "yurak_klapanlari", "buyraklar"]}]
near = auto.neighbours(MALE, "yurak", heart)
check("yurakka tegib turgan tomirlar unga soya tushiradi (bitta yorug'lik)",
      {"aorta_yoyi", "opka_arteriyasi", "yuqori_kovak_vena"} <= set(near), near)
check("skelet, ichidagi klapan va uzoqdagi buyrak soya tushirmaydi",
      not {"skelet", "yurak_klapanlari", "buyraklar"} & set(near), near)

from services.med3d.camera import Camera
rng = np.random.default_rng(1)
ball = rng.normal(size=(300, 3)) * 0.05                                  # orqadagi kichik organ
wall = np.c_[rng.uniform(-0.4, 0.4, (3000, 2)), np.full(3000, 0.3)]     # oldidagi katta organ (z — kameraga)
cam = Camera([-0.5, -0.5, -0.5], [0.5, 0.5, 0.5])
check("katta organ ortida to'liq yopilgan organ — ko'rinmaydi (yorliq qo'yilmaydi)",
      cam.visible_share(ball, (0, 0, 0), wall) < 0.05, cam.visible_share(ball, (0, 0, 0), wall))
check("yonida turgan organ — ko'rinadi", cam.visible_share(ball + [0.7, 0, 0], (0, 0, 0), wall) > 0.9)
one = {"show": ["devor", "shar"], "focus": ["shar"], "labels": ["devor", "shar"], "rot": [0, 0, 0]}
auto.reveal_focus(cam, [one], {"devor": wall.tolist(), "shar": ball.tolist()})
check("fokusdagi organni yopgan organ shu slayddan olinadi", one["show"] == ["shar"] and one["labels"] == ["shar"], one)


class FakeAtlas:
    """Atlas o'rniga: har qism — organ chegarasidagi kichik kub (tarmoq va 130 MB atlas kerak emas)."""
    def __init__(self, body):
        self.by_id, self.where = {}, {}
        for organ, info in body.items():
            for pid in info["parts"]:
                self.by_id[pid] = {"id": pid}
                self.where.setdefault(pid, (np.array(info["min"]), np.array(info["max"])))

    def _one(self, part):
        lo, hi = self.where[part["id"]]
        c = np.array([[x, y, z] for x in (0, 1) for y in (0, 1) for z in (0, 1)], dtype=np.float32)
        pos = (lo + c * (hi - lo)).astype(np.float32)
        nrm = (c - 0.5).astype(np.float32)
        idx = np.array([0, 1, 3, 0, 3, 2, 4, 6, 7, 4, 7, 5, 0, 4, 5, 0, 5, 1, 2, 3, 7, 2, 7, 6, 0, 2, 6, 0, 6, 4,
                        1, 5, 7, 1, 7, 3], dtype=np.uint32)
        return pos, nrm, idx


print("4) PPTX")
out_dir = tempfile.mkdtemp()
path = auto.make_deck("Hazm a'zolari", plan, out_dir, atlas=FakeAtlas(MALE))
z = zipfile.ZipFile(path)
names = z.namelist()
slide_files = sum(1 for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
autos = ["advTm" in z.read(f"ppt/slides/slide{i}.xml").decode() for i in range(1, slide_files + 1)]
check("4 ta asosiy slayd (+ o'zi o'tadigan yordamchilar)", autos.count(False) == 4 and not autos[0], autos)
check("organlar to'plami o'zgarganda yordamchi slayd bor", any(autos), autos)
from pptx import Presentation
prs = Presentation(path)
sizes, models = [], []
from lxml import etree
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main"}
for i in range(1, slide_files + 1):
    root = etree.fromstring(z.read(f"ppt/slides/slide{i}.xml"))
    for gf in root.iterfind(".//p:graphicFrame", NS):
        sizes.append(int(gf.find("p:xfrm/a:ext", NS).get("cx")) / 914400)
        models.append(gf.find(".//p:cNvPr", NS).get("name"))
check("3D ramkalar MAX_FRAME dan oshmaydi", max(sizes) <= MAX_FRAME + 1e-3, max(sizes))
check("model nomlari Morph uchun (!!m3d_...)", all(m.startswith("!!m3d_") for m in models), models[:4])
glbs = [n for n in names if n.endswith(".glb")]
check("bir xil GLB bitta faylda (fayl soni < model soni)", 0 < len(glbs) < len(models), (len(glbs), len(models)))
morphs = [bool(etree.fromstring(z.read(f"ppt/slides/slide{i}.xml")).findall(".//p159:morph", NS))
          for i in range(1, slide_files + 1)]
check("Morph 2-slayddan boshlab", morphs == [False] + [True] * (slide_files - 1), morphs)
# Haqiqiy 3D: slayddagi organlar bitta modelda (PowerPoint alohida modellarni ustma-ust qatlam qilib chizadi).
first = [g.find(".//p:cNvPr", NS).get("name") for g in etree.fromstring(z.read("ppt/slides/slide1.xml")).iterfind(
    ".//p:graphicFrame", NS) if 0 <= int(g.find("p:xfrm/a:off", NS).get("x")) < 13.4 * 914400]
check("1-slayd: 4 organ — bitta 3D model", len(first) == 1, first)
labels = [s.text_frame.text for s in prs.slides[0].shapes if s.name.startswith("lbl_") and s.has_text_frame
          and s.text_frame.text]
joined = " ".join(labels).lower()
check("yorliqlar: o'zbekcha nom + lotincha nom (uslubdan qat'i nazar)", all(w in joined for w in ("jigar", "hepar", "oshqozon", "gaster")), labels)
check("manba (CC BY) izohda va fayl xususiyatlarida", "BodyParts3D" in prs.slides[0].notes_slide.notes_text_frame.text
      and "BodyParts3D" in (prs.core_properties.comments or ""))

# Joyida almashish: fade — organlar uchib kelmaydi; auto — yordamchi slayd o'zi o'tadi va raqamga kirmaydi.
from services.med3d import deck as med3d_deck
deck_plan, scenes = auto.build_scenes("male", plan, FakeAtlas(MALE))
two = [dict(deck_plan["slides"][0], scene=deck_plan["slides"][0]["scene"]) for _ in range(3)]
two[1].update(fade=True, auto=True, frame_of=0, labels=[])
p2 = med3d_deck.build({"scene": deck_plan["scene"], "slides": two}, os.path.join(out_dir, "fade.pptx"), scenes=scenes)
z2 = zipfile.ZipFile(p2)
xml = [z2.read(f"ppt/slides/slide{i}.xml").decode() for i in (1, 2, 3)]
pages = [[sh.text_frame.text for sh in s.shapes if sh.name == "!!page"][0] for s in Presentation(p2).slides]
check("yordamchi slayd avtomatik o'tadi (advTm), boshqalari yo'q", 'advTm="0"' in xml[1]
      and "advTm" not in xml[2], [("advTm" in x) for x in xml])
check("yordamchi raqamga kirmaydi", pages == ["01 / 02", "02 / 02", "02 / 02"], pages)
offs = [int(m) for m in re.findall(r'<p:graphicFrame.*?<a:off x="(-?\d+)"', xml[1], re.S)]
check("fade slaydida slayd tashqarisidagi (uchib keladigan) 3D model nusxasi yo'q",
      offs and all(-914400 < x < 13.4 * 914400 for x in offs), offs)


print("4b) Uslublar, rejissyor, savol slaydi")
from services.med3d import design, director
check("12 ta dizayn uslubi", len(design.STYLES) >= 12, list(design.STYLES))
big = {k: v for k, v in plan.items() if k not in ("directed", "style")}
big["slides"] = [dict(s, title=f"{s['title']} {k}") for k in range(2) for s in plan["slides"]]
big["slides"] = [copy.deepcopy(s) for s in big["slides"]]
big["slides"][0]["labels"] = list(big["slides"][0]["show"])     # savol uchun kamida 3 ta yorliqli umumiy ko'rinish
path_q = auto.make_deck("Hazm a'zolari katta", big, tempfile.mkdtemp(), atlas=FakeAtlas(MALE), style="atlas")
zq = zipfile.ZipFile(path_q)
nq = sum(1 for n in zq.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
xmls = [zq.read(f"ppt/slides/slide{i}.xml").decode() for i in range(1, nq + 1)]
quiz = [x for x in xmls if "!!answers" in x]
check("8+ slaydda xulosadan oldin savol slaydi (javoblar kartasi)", len(quiz) == 1, nq)
check("javoblar bosilganda chiqadi (clickEffect)", quiz and 'nodeType="clickEffect"' in quiz[0])
digits = re.findall(r'name="lbl_[^"]+".*?<a:t>(\d+)</a:t>', quiz[0], re.S) if quiz else []
check("savol slaydida nomlar o'rnida raqamlar", len(digits) >= 3, digits)
moves = big.get("directed") and [s.get("layout") for s in big["slides"]]
check("kompozitsiyalar ketma-ket takrorlanmaydi", all(a != b for a, b in zip(moves, moves[1:])), moves)
seen_moves = director.direct(MALE, copy.deepcopy(big["slides"]), "boshqa mavzu")
check("kamera harakatlari ketma-ket takrorlanmaydi", all(a != b or a == "tinch" for a, b in
                                                        zip(seen_moves[1:], seen_moves[2:])), seen_moves)
check("uslub ochiq fonda: atlas (krem fon)", "F6F1E7" in xmls[0], xmls[0][:200])
from services.med3d import deck as med3d_deck_q
check("sifat nazorati: yorliq va model matnga tegmaydi, yorliqlar ustma-ust emas",
      not [x for _, x in med3d_deck_q.LAST_QUALITY if x], [x for _, x in med3d_deck_q.LAST_QUALITY if x][:2])

print("5) Bot: tugma → ogohlantirish → Ha / Yo'q")
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage
from bot.handlers import premium_presentation as pp
import config
config.ADMIN_IDS = [1]
from bot.states import PremiumPresentationStates as PS


class Chat:
    def __init__(self):
        self.next_id, self.sent, self.deleted, self.edited = 100, [], [], []
        self.bot = SimpleNamespace(delete_message=self.delete)

    async def delete(self, chat_id, message_id):
        self.deleted.append(message_id)

    def message(self, text=None):
        chat, msg = self, MagicMock()
        msg.text, msg.bot = text, self.bot
        msg.chat = SimpleNamespace(id=1)
        msg.from_user = SimpleNamespace(id=1, first_name="Ali")

        async def answer(body, **kw):
            chat.next_id += 1
            chat.sent.append((chat.next_id, body, kw.get("reply_markup")))
            return SimpleNamespace(message_id=chat.next_id)

        async def edit_text(body, **kw):
            chat.edited.append(body)
        async def answer_photo(photo, caption=None, **kw):
            return await answer(caption, **kw)
        msg.answer, msg.edit_text, msg.delete, msg.answer_photo = answer, edit_text, AsyncMock(), answer_photo
        return msg


def buttons(markup):
    return [b for row in markup.inline_keyboard for b in row]


def callback(chat, data, user=1):
    cb = MagicMock(); cb.data = data; cb.answer = AsyncMock()
    cb.message = chat.message(); cb.bot = chat.bot
    cb.from_user = SimpleNamespace(id=user)
    return cb


def make_db():
    db = MagicMock()
    user = SimpleNamespace(language="uz", first_name="Ali", balance=100000, id=1)
    db.get_user = AsyncMock(return_value=user)
    db.get_active_channels = AsyncMock(return_value=[])
    db.get_feature_status = AsyncMock(return_value=False)
    return db


async def bot_flow():
    chat, db = Chat(), make_db()
    state = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=1, user_id=1))
    await pp.premium_presentation_start(chat.message("🌟 Taqdimot"), state, db)
    kinds = [b for b in buttons(chat.sent[-1][2]) if (b.callback_data or "").startswith("prem_ppt_kind:")]
    check("3D Anatomiya tugmasi birinchi", kinds[0].callback_data == "prem_ppt_kind:med3d" and "3D Anatomiya" in kinds[0].text,
          kinds[0].text)

    await pp.premium_ppt_kind_selected(callback(chat, "prem_ppt_kind:med3d"), state, db)
    text, markup = chat.sent[-1][1], chat.sent[-1][2]
    cbs = [b.callback_data for b in buttons(markup)]
    check("bosilganda birinchi ogohlantirish: faqat tana tuzilishi", "tana tuzilishi" in text and "kompyuter" in text, text[:120])
    check("Ha va Yo'q tugmalari", cbs == ["prem_med3d:ha", "prem_med3d:yoq"], cbs)
    check("tur hali tanlanmagan (javob kutiladi)", (await state.get_data()).get("kind") != "med3d")

    await pp.premium_med3d_answer(callback(chat, "prem_med3d:yoq"), state, db)
    check("Yo'q → \"bu siz uchun emas\"", any("siz uchun emas" in t for t in chat.edited), chat.edited)
    check("Yo'q → turlar qayta ko'rsatiladi", await state.get_state() == PS.waiting_for_kind.state
          and any((b.callback_data or "").startswith("prem_ppt_kind:") for b in buttons(chat.sent[-1][2])))

    await pp.premium_ppt_kind_selected(callback(chat, "prem_ppt_kind:med3d"), state, db)
    await pp.premium_med3d_answer(callback(chat, "prem_med3d:ha"), state, db)
    data = await state.get_data()
    check("Ha → tur 3D Anatomiya, mavzu so'raladi", data.get("kind") == "med3d"
          and await state.get_state() in (PS.waiting_for_topic.state, PS.waiting_for_topic_text.state),
          (data.get("kind"), await state.get_state()))
    await pp.premium_ppt_got_topic(chat.message("Yurakning tuzilishi"), state, db)
    await pp.premium_ppt_skip_name(callback(chat, "prem_ppt_skip_name"), state, db)
    await pp.premium_ppt_skip_preferences(callback(chat, "prem_ppt_skip_preferences"), state, db)
    check("istaklardan keyin manba so'ralmaydi — dizayn uslubi so'raladi",
          await state.get_state() == PS.waiting_for_style.state, await state.get_state())
    style_cbs = [b.callback_data for b in buttons(chat.sent[-1][2]) if (b.callback_data or "").startswith("prem_m3d_style:")]
    check("12 ta uslub va «Avtomatik» tugmasi", len(style_cbs) == 13 and "prem_m3d_style:auto" in style_cbs, style_cbs)
    await pp.premium_ppt_previous(callback(chat, "prem_ppt_prev"), state, db)
    check("uslubdan orqaga — istaklar", await state.get_state() == PS.waiting_for_preferences.state,
          await state.get_state())
    await pp.premium_ppt_skip_preferences(callback(chat, "prem_ppt_skip_preferences"), state, db)
    await pp.premium_m3d_style(callback(chat, "prem_m3d_style:kobalt"), state, db)
    check("uslub tanlandi → slaydlar soni", await state.get_state() == PS.waiting_for_count.state
          and (await state.get_data()).get("m3d_style") == "kobalt", await state.get_state())
    counts = [b.text for b in buttons(chat.sent[-1][2]) if (b.callback_data or "").startswith("prem_ppt_count:")]
    check("slaydlar soni 5–15, narx 3D Anatomiyaniki (10 slayd — 25 000)",
          len(counts) == len(auto.SLIDE_OPTIONS) and counts[0].startswith("5 ta slayd")
          and any("10 ta slayd" in c and "25" in c for c in counts), counts)
    check("rasm manbasi va yozuv animatsiyasi taklif qilinmaydi",
          pp._photo_source({"kind": "med3d"}) is None and not await pp._anim_offer({"kind": "med3d"}))
    check("orqaga: slaydlar sonidan — uslubga, uslubdan — istaklarga",
          pp._previous_step({"kind": "med3d", "step": "count"}) == "m3dstyle"
          and pp._previous_step({"kind": "med3d", "step": "m3dstyle"}) == "prefs")
    summary, _ = pp._summary(dict(await state.get_data(), slide_count=10), "uz") if hasattr(pp, "_summary") else ("", None)
    check("xulosada tanlangan uslub", not summary or "Kobalt" in summary, summary[:300])

    other = Chat()
    state2 = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=2, user_id=2))
    cb = callback(other, "prem_ppt_kind:med3d", user=2)
    await pp.premium_ppt_kind_selected(cb, state2, db)
    check("3D Anatomiya hammaga ochiq (admin bo'lmaganga ham ogohlantirish chiqadi)",
          not (cb.answer.await_args and cb.answer.await_args.kwargs.get("show_alert")))
    os.environ["MED3D_OCHIQ"] = "0"
    state3 = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=3, user_id=3))
    cb3 = callback(Chat(), "prem_ppt_kind:med3d", user=3)
    await pp.premium_ppt_kind_selected(cb3, state3, db)
    os.environ.pop("MED3D_OCHIQ", None)
    check("MED3D_OCHIQ=0 — admin bo'lmaganga «tez orada»", cb3.answer.await_args is not None
          and cb3.answer.await_args.kwargs.get("show_alert"))
    from services import uz_script as _uzs
    check("3D fayl nomida ko'rish sharti — mijoz tanlagan tilda",
          "PowerPoint 2019+" in pp.med3d_filename("Yurak: tuzilishi", "uz") and "telefonda" in pp.med3d_filename("x", "uz")
          and "на телефоне" in pp.med3d_filename("x", "ru") and "Slide Show" in pp.med3d_filename("x", "en")
          and "телефонда кўринмайди" in pp.med3d_filename("x", _uzs.UZ_CYRILLIC_LANG)
          and ":" not in pp.med3d_filename("Yurak: tuzilishi", "uz") and pp.med3d_filename("x", "uz").endswith(".pptx"))
    check("3D yo'riqnoma videosi bor (taqdimot bilan yuboriladi)", os.path.exists(pp.MED3D_GUIDE_VIDEO)
          and os.path.getsize(pp.MED3D_GUIDE_VIDEO) < 10_000_000)
    _msg = MagicMock()
    _msg.answer_video = AsyncMock(return_value=SimpleNamespace(video=SimpleNamespace(file_id="VID1")))
    _msg.answer = AsyncMock()
    await pp._send_med3d_guide(_msg, "uz")
    await pp._send_med3d_guide(_msg, "ru")
    _calls = _msg.answer_video.await_args_list
    check("yo'riqnoma: video izohi — ko'rish sharti, ikkinchi marta file_id bilan (qayta yuklanmaydi)",
          len(_calls) == 2 and _calls[1].kwargs["video"] == "VID1" and "PowerPoint" in _calls[0].kwargs["caption"]
          and not _msg.answer.await_args_list)
    pp._med3d_guide_id = None
    state4 = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=4, user_id=4))
    await state4.update_data(topic="Qon tahlili", kind="med3d", slide_count=10, m3d_style="kobalt")
    await pp.premium_med3d_alternative(callback(Chat(), "prem_m3d_alt:ha", user=4), state4, db)
    alt = await state4.get_data()
    check("3D qilib bo'lmaydigan mavzu — «oddiy taqdimot» tugmasi zamonaviy turga o'tkazadi, mavzu saqlanadi",
          alt.get("kind") == "modern" and alt.get("topic") == "Qon tahlili"
          and await state4.get_state() == pp.PremiumPresentationStates.waiting_for_style.state, alt)


asyncio.run(bot_flow())

check("3D Anatomiya narxi: 10 slayd — 25 000 so'm (zamonaviydan 5 baravar)", auto.price_for(10) == 25000,
      auto.price_for(10))
check("3D Anatomiya narxi 3D Pro'dan baland va 500 ga yaxlit",
      all(auto.price_for(n) % 500 == 0 for n in auto.SLIDE_OPTIONS))

print()
print("HAMMASI O'TDI" if not FAILS else f"XATOLAR: {len(FAILS)} — {FAILS}")
sys.exit(1 if FAILS else 0)
