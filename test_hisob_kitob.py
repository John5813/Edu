"""Hisob-kitob taqdimotlari: kod hisoblaydi, diagramma o'qlari, hisob mavzu turi.

    python test_hisob_kitob.py
"""
import os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from services.premium_presentation import (deck_calc as calc, deck_charts, deck_shape,
                                           deck_style, html_slides, llm_client, themes)
TH = themes.get("ko'k")

# ── 6. KOD hisoblaydi
close = lambda a, b: abs(a - b) < 1e-6
check("daraja: 8.1*(1+0.009)^10", close(calc.evaluate("P0*(1+r)**t", {"P0": 8.1, "r": 0.009, "t": 10}), 8.859244369700455))
check("^ daraja va % belgisi", close(calc.evaluate("1000*(1+5%)^3"), 1157.625))
check("ikki baravar oshish davri (ln)", close(calc.evaluate("ln(2)/ln(1+0.9%)"), 77.36240945156692))
check("tug'ilish koeffitsiyenti", close(calc.evaluate("(140000/7800000)*1000"), 17.94871794871795))
check("o'zgaruvchilar bir-biriga tayanadi", calc.parse_vars("a=2;b=a*3")["b"] == 6)
for bad in ("__import__('os').system('x')", "open('f')", "(1).__class__", "2**99999",
            "1/0", "x+1", "[1,2]", "lambda: 1", "print(1)", "a.b"):
    try:
        calc.evaluate(bad); ok = False
    except calc.CalcError:
        ok = True
    check(f"xavfli/yaroqsiz ifoda rad etiladi: {bad[:22]}", ok)

out = calc.apply('<div class="calc" data-kind="line" data-range="0:10:5" data-vars="P0=8.1;r=0.009" '
                 'data-labels="2025,2030,2035" data-series="Asosiy: P0*(1+r)**t|Tez: P0*(1+2*r)**t" '
                 'data-unit="mlrd" data-xlabel="Yil"></div>')
check("calc diagramma: chart bloki, hisoblangan qiymatlar", 'class="chart"' in out and "8.1,8.471,8.859" in out, out)
check("calc diagramma: yorliq va X o'qi sarlavhasi saqlanadi", 'data-labels="2025,2030,2035"' in out and 'data-xlabel="Yil"' in out)
out = calc.apply('<span data-calc="(T/A)*1000" data-vars="T=140;A=7800" data-fmt="1">?</span>')
check("data-calc: raqam o'zi hisoblanadi", ">17,9<" in out and "data-calc" not in out, out)
out = calc.apply('<div class="kpi-value" data-calc="ln(2)/ln(1+0.9%)" data-fmt="0" data-suffix=" yil">?</div>')
check("data-calc: qo'shimcha matn (yil)", "77 yil" in out, out)
out = calc.apply('<div class="calc" data-kind="bar" data-vars="T=140;O=60;A=7800" '
                 'data-series="Tug\'ilish: (T/A)*1000|O\'lim: (O/A)*1000" data-unit="promille"></div>')
check("calc ustunli: har ustun o'z formulasidan", 'data-series="17.949,7.692"' in out and "Tug&#x27;ilish,O&#x27;lim" in out, out)
check("noto'g'ri formulali diagramma slaydni buzmaydi (tashlanadi)",
      calc.apply('<p>a</p><div class="calc" data-series="Bad: zzz+1" data-range="0:3"></div><p>b</p>') == "<p>a</p><p>b</p>")
check("noto'g'ri data-calc: model yozgan matn qoladi",
      ">?<" in calc.apply('<span data-calc="zzz+">?</span>'))
check("nuqta soni cheklangan", calc.apply('<div class="calc" data-kind="line" data-range="0:1000" data-series="a: t"></div>') == "")

# ── 4. Diagramma: X/Y o'qlar, bitta birlik
body = deck_charts.draw(calc.apply('<div class="calc" data-kind="line" data-range="0:10:2" data-vars="P0=8.1;r=0.009" '
                        'data-labels="2025,2027,2029,2031,2033,2035" data-series="Past: P0*(1+r)**t" '
                        'data-unit="mlrd" data-xlabel="Yil"></div>'), TH)
check("chiziqli: Y o'qi shkalasi (gorizontal to'r)", body.count('stroke-opacity="0.28"') >= 4, body.count('stroke-opacity="0.28"'))
check("chiziqli: tor oraliqda nol bazadan emas (yassi chiziq yo'q)", ">7,5<" in body or ">8<" in body and ">0<" not in body)
check("chiziqli: X o'qi sarlavhasi", ">Yil<" in body)
bar = deck_charts.draw('<div class="chart" data-kind="bar" data-labels="A,B" data-series="Q: 12,18" data-unit="u"></div>', TH)
check("ustunli: Y o'qi bor va noldan boshlanadi", ">0<" in bar and bar.count('stroke-opacity="0.28"') >= 3)
two = deck_charts.draw('<div class="chart" data-kind="line" data-labels="1,2,3" '
                       'data-series="Aholi (mlrd): 7.9,8,8.1|Urbanizatsiya (%): 57,58,59"></div>', TH)
check("ikki xil birlik: ikkita alohida diagramma", two.count("<svg") == 2 and "chart-stack" in two, two.count("<svg"))
check("ikki xil birlik: birliklar qavsdan olinadi", ">mlrd<" in two and ">%<" in two)
same = deck_charts.draw('<div class="chart" data-kind="line" data-labels="1,2,3" data-series="A: 10,11,12|B: 9,10,13"></div>', TH)
check("o'lchami yaqin qatorlar bitta diagrammada", same.count("<svg") == 1)
donut = deck_charts.draw('<div class="chart" data-kind="donut" data-labels="a,b,c" data-series="U: 25,65,10"></div>', TH)
check("halqa diagramma chiziladi", donut.count("<path") == 3 and "65%" in donut)
many = deck_charts.draw(calc.apply('<div class="calc" data-kind="line" data-range="0:50:2" data-vars="P=6;r=0.01" '
                                   'data-series="a: P*(1+r)**t"></div>'), TH)
check("ko'p nuqtada yozuvlar siyraklashadi", many.count("<text") < 40, many.count("<text"))
apos = deck_charts.draw(calc.apply('<div class="calc" data-kind="bar" data-vars="a=1;b=2" data-series="Tug\'ilish: a|O\'lim: b"></div>'), TH)
check("apostrofli yorliq ikki marta escape qilinmaydi", "&amp;#x27;" not in apos and "&amp;" not in apos.replace("&amp;#","X"), apos[:80])
ticks = deck_charts._nice_ticks(0, 18)
check("chiroyli bo'linmalar (0, 5, 10, 15, 20)", ticks == [0, 5, 10, 15, 20], ticks)

# ── 1. Hisob mavzu turi
for topic in ("Global Demografik hisob kitoblar va proagnozlarni hisoblash formulalari",
              "Moliyaviy prognoz formulalari", "Расчёт процентов", "Compound interest calculation"):
    check(f"hisob oilasi: {topic[:34]}", deck_shape.of(topic) == "hisob", deck_shape.of(topic))
check("hisob bo'lmagan mavzu o'zgarmaydi", deck_shape.of("Falsafa asoslari") != "hisob")
guide = deck_shape.guidance("hisob")
check("hisob yo'riqnomasi: formula + ishlangan misol + diagramma + kod hisoblaydi",
      all(w in guide for w in ("formula", "`misol`", "DIAGRAMMA", "`calc`")), guide[:80])
check("oila nomlari ro'yxatida hisob bor", "hisob" in deck_shape.names())

captured = {}
orig = llm_client._call_openrouter
llm_client._call_openrouter = lambda system, prompt, **k: captured.setdefault("p", prompt) and {"slides": []}
try:
    html_slides.plan_outline("Demografik hisob-kitoblar va prognoz formulalari", 10, "uz")
    plan_hisob = captured.pop("p")
    html_slides.plan_outline("Falsafa asoslari", 10, "uz")
    plan_other = captured.pop("p")
finally:
    llm_client._call_openrouter = orig
check("reja so'rovida hisob mavzusi uchun formula/misol/diagramma talabi", "HISOB-KITOB mavzusi" in plan_hisob)
check("boshqa mavzuda bu talab yo'q", "HISOB-KITOB mavzusi" not in plan_other)

# ── 2, 3, 5. Prompt
rules = html_slides.shell_rules(TH, "uz")
check("prompt: hisoblangan va statistik raqam ajratilgan", "a) HISOBLANGAN raqam" in rules and "b) STATISTIK FAKT" in rules)
check("prompt: hisoblangan raqam RUXSAT, kod hisoblaydi", "RUXSAT" in rules and "calc" in rules)
check("prompt: manba yili (bugungi va kelgusi) yozilmaydi", "bugungi va kelgusi yillar" in rules)
check("prompt: raqam to'qish taqiqi saqlangan", "RAQAMNI O'YLAB TOPMANG" in rules)
blocks = deck_style.BLOCKS
check("namunalar: line, bar va donut", all(f'data-kind="{k}"' in blocks for k in ("line", "bar", "donut")))
check("namunalar: calc diagramma va data-calc", 'class="calc"' in blocks and 'data-calc=' in blocks)
check("namunalar: bitta birlik qoidasi", "BITTA birlikda" in blocks)

# ── 5. Manba yili
g = html_slides.guard_source_years
check("manba yili: bugungi yil 'taxminiy' ga almashadi", g("(BMT, 2026)") == "(BMT, taxminiy)", g("(BMT, 2026)"))
check("manba yili: o'tgan yil o'zgarmaydi", g("(BMT, 2020)") == "(BMT, 2020)")
check("manba yili: ruscha/inglizcha", g("(ООН, 2027)", "ru") == "(ООН, оценка)" and g("(UN, 2027)", "en") == "(UN, estimate)")
check("manba yili: oddiy qavs o'zgarmaydi", g("(2026 yil)") == "(2026 yil)" and g("o'sish (2026)") == "o'sish (2026)")

# ── to'liq quvur
page = html_slides.build_pages([
    '<section class="slide"><div class="head"><h2 class="title">T</h2></div><div class="body">'
    '<div class="calc" data-kind="donut" data-series="A: 25|B: 75"></div>'
    '<div class="misol-answer">Javob: <span data-calc="(T/A)*1000" data-vars="T=140;A=7800" data-fmt="1">?</span> p</div>'
    '<p class="note">(BMT, 2026)</p></div></section>'], TH, "uz")[0]
check("quvur: calc diagramma SVG ga aylanadi", "<svg" in page and 'class="calc"' not in page.split("<!--manba")[0])
check("quvur: javob hisoblanadi, manba yili himoyalanadi", "17,9" in page and "BMT, taxminiy" in page)
check("quvur: model yozgan asl slayd (calc) tuzatish uchun saqlanadi",
      'class="calc"' in html_slides.source_of(page))

# ── Formula: model qo'sh teskari chiziq yozsa ham, so'zli pastki indeks ham o'qiladi
from services.premium_presentation import deck_math, html_images
f = deck_math.formula(r"CDR = \\frac{D}{P} \\times 1000")
check("qo'sh teskari chiziqli kasr to'g'ri chiqadi", 'class="frac"' in f and "×" in f and "\\" not in f, f)
f = deck_math.formula("A_{o'sish} = (T - O) + M")
check("so'zli pastki indeks o'qiladigan matn bo'lib qoladi", "A_(o'sish)" in f, f)
check("raqamli indeks Unicode bo'lib qoladi", "x₁" in deck_math.formula("x_1") or "₁" in deck_math.formula("x_1"))

# ── Diagramma ranglari bir-biridan ajraladi
for key in ("ko'k", "yashil", "to'q sariq", "qizil"):
    colours = deck_charts.palette(themes.get(key))
    check(f"diagramma ranglari turli: {key}", len(set(colours)) == len(colours) >= 5
          and all(deck_charts._hue_gap(a, b) >= deck_charts._MIN_HUE_GAP
                  for i, a in enumerate(colours) for b in colours[i + 1:]), colours)
check("birinchi rang sxemaning asosiy rangi", deck_charts.palette(themes.get("ko'k"))[0] == themes.get("ko'k").chart[0])
night = deck_charts.palette(themes.with_style(themes.get("ko'k"), "qorongu"))
check("qorong'u uslubda ranglar ochroq va turli", len(set(night)) >= 5 and night[1] != "F59E0B", night)

# ── Mayda matn: brauzer topadi, tuzatish qisqartirishga ruxsat beradi
check("mayda matn qoidasi tekshiruvda bor", "juda mayda" in __import__("services.premium_presentation.html_extract", fromlist=["x"])._CHECK_SCRIPT)
before = '<section class="slide"><div class="body"><p>' + "so'z " * 40 + "</p></div></section>"
after = '<section class="slide"><div class="body"><p>' + "so'z " * 8 + "</p></div></section>"
check("oddiy tuzatishda matn 80% yo'qolsa rad", html_slides._rewritten(before, after) != "")
check("mayda matn tuzatishida qisqartirish qabul", html_slides._rewritten(before, after, True) == "")

# ── O'ylab topilgan iqtibos manbasi
g = html_slides.guard_quote_sources('<p class="quote-by">— BMT Aholi Jamg\'armasi (UNFPA), 2026-yil hisobotidan</p>')
check("kelgusi yilli 'hisobotidan' iqtibos manbasi olib tashlanadi", "2026" not in g and "UNFPA" in g, g)
g = html_slides.guard_quote_sources('<p class="quote-by">— Amir Temur</p>')
check("haqiqiy muallif tegilmaydi", "Amir Temur" in g)
check("prompt: iqtibos faqat haqiqiy", "IQTIBOS faqat HAQIQIY" in html_slides.shell_rules(TH, "uz"))

# ── Rasm: oddiy realistik foto, diagramma va yozuvsiz
pp = html_images.photo_prompt("an infographic diagram of a population pyramid with labels, wide shot")
check("rasm tavsifi: diagramma/infografika so'zlari yo'q", not re.search(r"infographic diagram|labels,", pp.split("no ")[0]), pp)
check("rasm tavsifi: realistik foto va yozuvsiz", pp.startswith("realistic natural photograph") and "no text" in pp and "no diagram" in pp, pp)
check("prompt: rasm faqat oddiy foto", "FAQAT oddiy, realistik" in deck_style.BLOCKS)

from services.premium_presentation import html_render
if html_render.available():
    from pptx import Presentation
    pages = html_slides.build_pages([
        '<section class="slide"><div class="head"><h2 class="title">Prognoz</h2><div class="rule"></div></div><div class="body">'
        '<div class="calc" data-kind="line" data-range="0:10:2" data-vars="P0=8.1;r=0.009" '
        'data-labels="2025,2027,2029,2031,2033,2035" data-series="Past: P0*(1+r)**t|Yuqori: P0*(1+2*r)**t" '
        'data-unit="mlrd" data-xlabel="Yil"></div></div></section>',
        '<section class="slide"><div class="head"><h2 class="title">Ko\'rsatkich</h2><div class="rule"></div></div><div class="body">'
        '<div class="cols cols-2"><div class="kpi"><div class="kpi-value" data-calc="(140000/7800000)*1000" data-fmt="1">?</div>'
        '<div class="kpi-label">YTK</div></div><div class="kpi"><div class="kpi-value" data-calc="ln(2)/ln(1+0.9%)" data-fmt="0">?</div>'
        '<div class="kpi-label">Ikki baravar</div></div></div></div></section>'], TH, "uz")
    prs = Presentation(html_render.render(pages, out_dir=tempfile.mkdtemp(), name="hisob"))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
    check("PPTX: 2 slayd, hisoblangan raqamlar matn bo'lib tahrirlanadi", len(prs.slides) == 2 and "17,9" in text and "77" in text, text[:120])
    check("PPTX: diagramma rasm sifatida tushgan", pics >= 1, pics)
else:
    print("  (brauzer yo'q — PPTX sinovi o'tkazib yuborildi)")

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
