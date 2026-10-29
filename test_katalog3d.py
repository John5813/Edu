"""Umumiy 3D katalog (katalog3d): ro'yxat butunligi, litsenziyalar, qidiruv va slaydga joylash.

- katalog.json: id lar takrorlanmaydi, faqat CC0 / CC BY 4.0, CC BY da muallif va attribution bor;
- har modelning preview rasmi bor, tanlov.jsonl bilan bir xil;
- nomlarda taqiqlangan narsa (o'yinchoq, anime, qurol...) yo'q;
- qidir(): o'zbekcha so'z bo'yicha topadi, fan filtri ishlaydi;
- joylash(): slaydda am3d model va ostida litsenziya yozuvi; model yo'q bo'lsa FileNotFoundError.

    python test_katalog3d.py
"""
import io, json, os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import numpy as np
from pptx import Presentation
from pptx.util import Inches

TMP = tempfile.mkdtemp()
os.environ["KATALOG3D_DIR"] = TMP
from services import katalog3d
from services.med3d import glb

print("1) Ro'yxat butunligi")
data = json.load(open(katalog3d.KATALOG_JSON, encoding="utf-8"))
mods = data["modellar"]
ids = [m["id"] for m in mods]
check("2000 dan ko'p model", len(mods) > 2000, len(mods))
check("soni maydoni to'g'ri", data["soni"] == len(mods))
check("id lar takrorlanmaydi", len(ids) == len(set(ids)))
check("faqat CC0 / CC BY 4.0", {m["litsenziya"] for m in mods} <= {"CC0 1.0", "CC BY 4.0"},
      {m["litsenziya"] for m in mods})
ccby = [m for m in mods if m["litsenziya"] == "CC BY 4.0"]
check("CC BY da muallif bor", all(m["muallif"] for m in ccby), [m["id"] for m in ccby if not m["muallif"]][:5])
check("CC BY attribution da muallif va litsenziya",
      all(m["muallif"] in m["attribution"] and "CC BY 4.0" in m["attribution"] for m in ccby))
check("har modelda attribution_full", all(m["attribution_full"] for m in mods))
check("fayl yo'li <fan>/<id>/model.glb", all(m["fayl"] == f"{m['fan'].replace(chr(39), '')}/{m['id']}/model.glb"
                                             for m in mods))
check("o'zbekcha nom bo'sh emas", all(m["nom_uz"].strip() for m in mods))
check("har modelning preview rasmi bor", all(katalog3d.preview_yoli(i) for i in ids),
      [i for i in ids if not katalog3d.preview_yoli(i)][:5])
tanlov = [json.loads(l) for l in open(os.path.join(HERE, "data", "katalog3d", "tanlov.jsonl"), encoding="utf-8")]
check("tanlov.jsonl katalog bilan bir xil", sorted(t["uid"] for t in tanlov) == sorted(ids))
check("tanlovda manba URL yoki HF yo'li", all(t.get("url") or t.get("yol") for t in tanlov))
TAQIQ = ("toy", "doll", "anime", "cartoon", "rubik", "lowpoly", "low poly", "stylized", "weapon", "gun",
         "rifle", "sword", "fnaf", "pokemon", "minecraft", "o'yinchoq")
TAQIQ_RE = re.compile(r"(?<![a-z'])(" + "|".join(re.escape(w) for w in TAQIQ) + r")(?![a-z])")
bad = [m["nom"] for m in mods if TAQIQ_RE.search((m["nom"] + " " + m["nom_uz"]).lower())]
check("nomlarda taqiqlangan so'z yo'q", not bad, bad[:5])
check("bir nechta fan bor", len({m["fan"] for m in mods}) >= 10)

print("2) Qidiruv")
r = katalog3d.qidir("bosh suyagi")
check("'bosh suyagi' topiladi", r and all("bosh suyag" in m["nom_uz"].lower() for m in r[:5]), [m["nom_uz"] for m in r[:5]])
r = katalog3d.qidir("sher", fan="san'at")
check("fan filtri", r and all(m["fan"] == "san'at" for m in r))
check("bo'sh so'rov — bo'sh ro'yxat", katalog3d.qidir("  ") == [])
check("apostrofsiz ham topadi", katalog3d.qidir("qal'a") and katalog3d.qidir("qala"))

print("3) Slaydga joylash")
mid = next(m["id"] for m in ccby)
m = katalog3d.katalog()[mid]
check("yuklanmagan model — None", katalog3d.model_yoli(mid) is None)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
try:
    katalog3d.joylash(slide, mid, 1, 1, 4)
    check("yuklanmagan model — FileNotFoundError", False)
except FileNotFoundError:
    check("yuklanmagan model — FileNotFoundError", True)
pos = np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0]], dtype=np.float32)
nrm = np.array([[0, 0, 1]] * 3, dtype=np.float32)
path = os.path.join(TMP, m["fayl"]); os.makedirs(os.path.dirname(path))
open(path, "wb").write(glb.write_glb([glb.Primitive(pos, nrm, np.array([0, 1, 2]), (0.8, 0.8, 0.8, 1))]))
katalog3d.joylash(slide, mid, 1, 1, 4)
buf = io.BytesIO(); prs.save(buf)
xml = slide.shapes._spTree.xml
check("am3d:model3d bor", "model3d" in xml)
check("ostida attribution yozuvi", any(s.has_text_frame and s.text_frame.text == m["attribution"] for s in slide.shapes))
check("GLB taqdimotga qo'shilgan", any(p.partname.endswith(".glb") for p in prs.part.package.iter_parts()))

print()
print("HAMMASI OK" if not FAILS else f"{len(FAILS)} ta XATO: {FAILS}")
sys.exit(1 if FAILS else 0)
