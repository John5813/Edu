"""Yozuv animatsiyasi: sahifa ochilganda mazmun o'zi, navbat bilan paydo bo'ladi (Infografik va Zamonaviy).

- har sahifada avtomatik boshlanadigan (bosishsiz) kirish animatsiyalari, sarlavha birinchi;
- fon, katta panel, chetdan chiqib turgan bezak, jadval/diagramma animatsiyasiz qoladi;
- punktlar chizilish tartibida birin-ketin, bir sahifa 2,5 soniyadan oshmaydi;
- XML PowerPoint sxemasiga mos (sxema bo'lsa tekshiriladi), animatsiyasi bor sahifaga tegilmaydi;
- funksiya sukut bo'yicha o'chiq; yoqilganda mijoz xulosada "✨ Yozuv animatsiyasi" ni tanlaydi.

    python test_yozuv_animatsiya.py
"""
import asyncio
import os
import random
import re
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond:
        FAILS.append(name)


from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from services.premium_presentation import slide_anim, themes
from services.premium_presentation.designer import kit, styles

TMP = tempfile.mkdtemp()
theme = themes.for_deck("Turizm", "", "kam")


def ctx(slide):
    return styles.Ctx(slide=slide, pal=styles.Palette.of(theme), rng=random.Random(1), topic_icon="globe")


def items(n):
    return [{"title": f"Yo'nalish {i + 1}", "note": f"Qisqa izoh {i + 1}", "value": f"{i + 1}0%"} for i in range(n)]


print("1) Animatsiya qo'shiladi")
prs = Presentation()
prs.slide_width, prs.slide_height = 12192000, 6858000
spec = {"title": "Sarlavha", "lead": "Bosh gap", "text": "Matn", "items": items(4), "credit": "",
        "points": ["birinchi", "ikkinchi"], "sides": [{"name": "A", "items": ["a"]}, {"name": "B", "items": ["b"]}]}
for name in ("cover_x", "leaves", "arrows", "finale"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    styles.BY_NAME[name].draw(ctx(s), dict(spec))
# oddiy sahifa: fon, chetdan chiqqan bezak, sarlavha, 3 karta, diagramma
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = kit.rect(s, 0, 0, 1920, 1080)
edge = kit.oval(s, 1900, 1060, 200)
title = kit.text(s, 100, 60, 1200, 100, [[("Sarlavha", 48, "111111", True)]])
cards = [kit.rect(s, 100 + i * 500, 400, 400, 300) for i in range(3)]
chart_data = CategoryChartData()
chart_data.categories = ["A", "B"]
chart_data.add_series("S", (1, 2))
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0), Emu(2000000), Emu(1000000), chart_data)
# allaqachon animatsiyasi bor sahifa (masalan "Rahmat")
done = prs.slides.add_slide(prs.slide_layouts[6])
kit.rect(done, 100, 100, 200, 200)
done._element.append(etree.fromstring('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                                      '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot"/></p:par>'
                                      '</p:tnLst></p:timing>'))
path = os.path.join(TMP, "d.pptx")
prs.save(path)

count = slide_anim.apply(path)
deck = Presentation(path)
check("6 tadan 5 sahifaga qo'shildi (animatsiyalisiga tegilmadi)", count == 5, count)


def timing_of(slide):
    t = slide._element.find(qn("p:timing"))
    return etree.tostring(t).decode() if t is not None else ""


def steps(xml):
    """[(spid, delay)] — kirish animatsiyalari tartibi."""
    return [(int(spid), int(delay)) for delay, spid in
            re.findall(r'presetClass="entr"[^>]*><p:stCondLst><p:cond delay="(\d+)"/>.*?spid="(\d+)"', xml)]


ok_auto, ok_long = True, True
for i, slide in enumerate(list(deck.slides)[:5], 1):
    xml = timing_of(slide)
    if 'evt="onBegin"' not in xml or "clickEffect" in xml or "afterEffect" in xml:
        ok_auto = False
    st = steps(xml)
    if not st or max(d for _, d in st) + 600 > 2500:
        ok_long = False
check("o'zi boshlanadi (bosish shart emas)", ok_auto)
check("har sahifa 2,5 soniyada ochilib bo'ladi", ok_long)
check("animatsiyalisi o'zgarmadi", timing_of(deck.slides[5]).count("cTn") == 1)

plain = deck.slides[4]
st = steps(timing_of(plain))
ids = [spid for spid, _ in st]
check("fon animatsiyasiz", bg.shape_id not in ids)
check("chetdan chiqqan bezak animatsiyasiz", edge.shape_id not in ids)
check("diagramma animatsiyasiz", chart.shape_id not in ids)
check("sarlavha birinchi, kechikishsiz", st and st[0] == (title.shape_id, 0), st[:2])
card_delays = [d for spid, d in st if spid in {c.shape_id for c in cards}]
check("kartalar chizilish tartibida birin-ketin", len(card_delays) == 3 and card_delays == sorted(card_delays)
      and len(set(card_delays)) == 3, card_delays)

leaves = steps(timing_of(deck.slides[1]))
check("barglar sahifasida ko'p qism navbat bilan chiqadi", len(leaves) >= 8 and len({d for _, d in leaves}) >= 6,
      len(leaves))

again = slide_anim.apply(path)
check("qayta qo'llash ikki marta qo'shmaydi", again == 0)

schema_file = None
for root in ("/mnt/skills/public/pptx/scripts/office/schemas/ISO-IEC29500-4_2016",):
    if os.path.exists(os.path.join(root, "pml.xsd")):
        schema_file = os.path.join(root, "pml.xsd")
if schema_file:
    here = os.getcwd()
    os.chdir(os.path.dirname(schema_file))
    schema = etree.XMLSchema(etree.parse("pml.xsd"))
    os.chdir(here)
    bad = [i for i, sl in enumerate(deck.slides, 1) if not schema.validate(sl._element.getroottree())]
    check("PowerPoint sxemasiga mos", not bad, (bad, str(schema.error_log.last_error)[:200]))
else:
    print("  --   sxema topilmadi, tekshiruv o'tkazib yuborildi")

print("2) Mijoz tanlovi va admin")
from bot.handlers import premium_presentation as pp

base = {"topic": "Turizm", "slide_count": 10, "presentation_language": "uz", "volume": "kam", "style": "toza"}
_, kb_off = pp._summary({**base, "text_anim_offer": False}, "uz")
check("admin o'chirgan — tanlov yo'q", not any(b.callback_data == "prem_ppt_anim" for r in kb_off.inline_keyboard for b in r))
text_on, kb_on = pp._summary({**base, "text_anim_offer": True}, "uz")
btn = [b for r in kb_on.inline_keyboard for b in r if b.callback_data == "prem_ppt_anim"]
check("yoqilgan — tanlov bor, sukut 'Bor'", btn and btn[0].text == "✨ Yozuv animatsiyasi: ✅ Bor" and "✅ Bor" in text_on,
      btn and btn[0].text)
check("tanlansa — PowerPoint slayd-shou haqida aytiladi", "Slayd-shou" in text_on)
text_no, kb_no = pp._summary({**base, "text_anim_offer": True, "text_anim": False}, "uz")
btn = [b for r in kb_no.inline_keyboard for b in r if b.callback_data == "prem_ppt_anim"]
check("mijoz o'chirsa — 'Yo'q'", btn and btn[0].text.endswith("❌ Yo'q"))
check("o'chirsa — eslatma yo'q", "Slayd-shou" not in text_no)
check("generatsiyada: tanlov hisobga olinadi",
      pp._text_anim_choice({"text_anim_offer": True}) is True
      and pp._text_anim_choice({"text_anim_offer": True, "text_anim": False}) is False
      and pp._text_anim_choice({"text_anim": True}) is None)

import database.database as dbmod
from database.database import Database, init_db

dbmod.DATABASE_FILE = os.path.join(TMP, "t.db")


async def admin_part():
    await init_db()
    db = Database()
    off = await db.get_feature_status(slide_anim.FEATURE, default=False)
    from bot.handlers.admin import _feature_keyboard

    row = [b for r in (await _feature_keyboard(db)).inline_keyboard for b in r if "Yozuv animatsiyasi" in b.text]
    await db.set_feature_status(slide_anim.FEATURE, True)
    row2 = [b for r in (await _feature_keyboard(db)).inline_keyboard for b in r if "Yozuv animatsiyasi" in b.text]
    return off, row, row2


off, row, row2 = asyncio.run(admin_part())
check("sukut — o'chiq", off is False)
check("admin panelida", row and row[0].callback_data == "toggle_text_anim_on" and "O'chirilgan" in row[0].text)
check("yoqilgandan keyin", row2 and row2[0].callback_data == "toggle_text_anim_off" and "Yoqilgan" in row2[0].text)

print("\nNATIJA:", "HAMMASI O'TDI" if not FAILS else f"{len(FAILS)} ta xato: {FAILS}")
sys.exit(1 if FAILS else 0)
