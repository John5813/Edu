"""Taqdimotda matn ustma-ust tushmasligi: muqova sarlavhasi, abzatslar, diagramma yorliqlari.

    python test_taqdimot_ustma_ust.py
"""
import os, sys, tempfile, zipfile, re, asyncio

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from services.premium_presentation import deck_charts, pptx_build, chart_data, html_slides as hs, html_render, themes

# 1. Diagramma yorliqlari o'raladi
lines = deck_charts._wrap_lines("Qadimgi Yunoniston (polislar)", 22, 150)
check("uzun yorliq bir necha qatorga bo'linadi", 1 < len(lines) <= 3, lines)
check("bo'lingan yorliq so'zlari yo'qolmaydi", "".join(lines).replace("…", "").startswith("Qadimgi"), lines)
wrapped, extra = deck_charts._x_labels(["A", "Qadimgi Yunoniston (polislar)"], {0, 1}, 22, 150)
check("past chekka uzun yorliqqa joy ajratadi", extra > 0 and 1 in wrapped, (wrapped, extra))
wrapped2, extra2 = deck_charts._x_labels(["2019", "2020"], {0, 1}, 22, 150)
check("qisqa yorliqqa qo'shimcha joy kerak emas", extra2 == 0, extra2)

# 2. Abzatslar bitta quti, muqova sarlavhasi pastdan tayanadi
def block(y, h, text, **kw):
    b = {"kind": "text", "x": 130, "y": y, "w": 945, "h": h, "text": text, "size": 31,
         "lineHeight": 49.6, "weight": 400, "color": "112233", "family": "Georgia", "lines": 5, "flow": 1}
    b.update(kw); return b

layout = {"background": "FFFFFF", "blocks": [
    block(410, 248, "Birinchi abzats matni."), block(678, 248, "Ikkinchi abzats matni."),
    {"kind": "text", "x": 130, "y": 100, "w": 800, "h": 60, "text": "Yakka", "size": 30, "lines": 1}]}
prs = pptx_build.new_presentation()
pptx_build.add_slide(prs, layout)
boxes = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
check("ikki abzats bitta quti bo'ldi (yakka matn alohida)", len(boxes) == 2, len(boxes))
par = [s for s in boxes if "abzats" in s.text_frame.text][0]
check("quti ikkala abzatsni qamraydi", len(par.text_frame.paragraphs) == 2 and par.top < 420 * 6350 and par.height >= 500 * 6350,
      (len(par.text_frame.paragraphs), par.top, par.height))
check("abzatslar orasida brauzerdagi oraliq saqlangan", par.text_frame.paragraphs[1].space_before.pt > 5,
      par.text_frame.paragraphs[1].space_before)

layout2 = {"background": "FFFFFF", "blocks": [
    {"kind": "text", "x": 130, "y": 349, "w": 1660, "h": 225, "text": "Uzun muqova sarlavhasi", "size": 104,
     "lineHeight": 112, "lines": 2, "valign": "bottom"}]}
prs2 = pptx_build.new_presentation()
pptx_build.add_slide(prs2, layout2)
tb = [s for s in prs2.slides[0].shapes if s.has_text_frame][0]
check("muqova sarlavhasi pastdan tayanadi", tb.text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM)

# 3. Haqiqiy chiqarishda: ustunga raqam, muqova sarlavhasiga 'bottom'
theme = themes.with_style(themes.get("zumrad"), "toza")
cover = ('<section class="slide dark"><div class="body"><h1 class="title big">Kiberxavfsizlik va Shaxsiy '
         'Ma\'lumotlar Himoyasi</h1><div class="rule"></div><p class="lead">Izoh.</p></div></section>')
slide = ('<section class="slide"><div class="head"><h2 class="title">Sarlavha</h2><div class="rule"></div></div>'
         '<div class="body"><div class="split"><div class="par-col"><p class="par">Birinchi abzats uzun matn '
         'bilan, ikkinchi qatorga o\'tadigan darajada uzun bo\'lsin.</p><p class="par">Ikkinchi abzats ham '
         'shunday uzun matn bilan yoziladi.</p></div><div class="rasm" data-prompt="x"><p class="rasm-matn">'
         'Qo\'shimcha.</p></div></div></div></section>')
pages = hs.build_pages([cover, slide], theme, "uz")
out = tempfile.mkdtemp()
path = html_render.render(pages, out_dir=out, name="t")
prs3 = Presentation(path)
def frames(slide_):
    return [s for s in slide_.shapes if s.has_text_frame and s.text_frame.text.strip()]
title = [s for s in frames(prs3.slides[0]) if "Kiberxavfsizlik" in s.text_frame.text]
check("haqiqiy muqovada sarlavha pastdan tayanadi", title and title[0].text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM)
pars = [s for s in frames(prs3.slides[1]) if "abzats" in s.text_frame.text]
check("haqiqiy slaydda abzatslar bitta qutida", len(pars) == 1 and len(pars[0].text_frame.paragraphs) == 2,
      [len(p.text_frame.paragraphs) for p in pars])

# 4. Diagramma promptida qisqa yorliq ko'rsatmasi bor
check("diagramma promptida qisqa yorliq ko'rsatmasi bor", "1-3 so'z" in open(
    "services/premium_presentation/prompts/uz.py", encoding="utf8").read())

print("\nXATO:" if FAILS else "\nHAMMASI YAXSHI", FAILS or "")
sys.exit(1 if FAILS else 0)
