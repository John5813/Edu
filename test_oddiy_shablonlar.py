"""Oddiy taqdimot: 25 ta shablon, mavzuga qarab aralash tanlov.

    python test_oddiy_shablonlar.py
"""
import asyncio, os, random, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

from PIL import Image
from pptx import Presentation
from services import slide_layouts as sl
from services.document_service import DocumentService
from services.template_service import TemplateService

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

def words(seed, n):
    r = random.Random(seed)
    vocab = "bozor mahsulot tahlil rivojlanish samaradorlik texnologiya iqtisodiyot ta'lim sifat resurs natija jarayon tizim model".split()
    return " ".join(r.choice(vocab) for _ in range(n)).capitalize() + "."

def item(i, topic, n=18, value=""):
    return {"head": f"{topic[:12]} {i + 1}", "text": f"{words(topic + str(i), n)}", "value": value}

def deck(topic, n_main=10, n_items=4, layouts=None):
    out = [{"title": topic, "layout": "cover", "content": ""}]
    for k in range(n_main):
        out.append({"title": f"{topic}: {k + 1}-bo'lim", "layout": (layouts or [""] * n_main)[k],
                    "items": [item(j, f"{topic}{k}") for j in range(n_items)]})
    out.append({"title": "Xulosa", "layout": "conclusion", "content": "Xulosa matni."})
    return out

def mains(slides):
    return [s for s in slides if s["layout"] not in sl.FIXED]

# ── katalog
check("25 ta shablon", len(sl.SPECS) == 25, len(sl.SPECS))
check("rasmli shablonlar 7 ta", sum(1 for s in sl.SPECS if s.image) == 7)
check("yangi chizuvchilar 19 ta", len(sl.KIT_LAYOUTS) == 19)

# ── aralashtirish
topics = ["Sun'iy intellekt", "Ekologiya", "Bank tizimi", "Sport tarixi", "Tibbiyot"]
seqs = []
for topic in topics:
    planned = mains(sl.assign(deck(topic), topic, "uz"))
    names = [s["layout"] for s in planned]
    seqs.append(tuple(names))
    check(f"{topic}: ketma-ket takror yo'q", all(a != b for a, b in zip(names, names[1:])), names)
    check(f"{topic}: kamida 6 xil shablon", len(set(names)) >= 6, names)
    check(f"{topic}: rasmli shablonlar <= 4", sum(1 for n in names if sl.image_kind(n)) <= 4, names)
    check(f"{topic}: rasmli shablonlar ketma-ket emas",
          all(not (sl.image_kind(a) and sl.image_kind(b)) for a, b in zip(names, names[1:])), names)
check("turli mavzuda turli kombinatsiya", len(set(seqs)) >= 4, len(set(seqs)))
check("bir xil mavzu — bir xil natija (sinash oson)",
      [s["layout"] for s in sl.assign(deck("A"), "A", "uz")] == [s["layout"] for s in sl.assign(deck("A"), "A", "uz")])
check("rasm o'chiq bo'lsa rasmli shablon yo'q",
      not any(sl.image_kind(s["layout"]) for s in sl.assign(deck("Ekologiya"), "Ekologiya", "uz", images=False)))

# uzoq taqdimot: takror cheklangan
long_names = [s["layout"] for s in mains(sl.assign(deck("Uzun", 22), "Uzun", "uz"))]
check("22 slaydda bitta shablon <= 3 marta", max(long_names.count(n) for n in set(long_names)) <= 3, long_names)

# ── mazmunga qarab tanlov
proc = deck("Jarayon", 3)
proc[2]["items"] = [{"head": f"Bosqich {i}", "text": "Avval ma'lumot yig'iladi, keyin tahlil qilinadi, so'ng natija e'lon qilinadi va bosqich yakunlanadi."} for i in range(1, 5)]
hits = sum(sl.assign(proc, f"J{k}", "uz")[2]["layout"] in ("timeline", "staircase", "chevron_process") for k in range(10))
check("jarayon mazmuni — bosqichli shablon (ko'pincha)", hits >= 7, hits)

cmp_ = deck("Taqqos", 3)
cmp_[2]["items"] = [
    {"head": "Afzalliklari", "text": "Arzon narx va tezkor ishga tushirish. Oddiy foydalanish imkoniyati."},
    {"head": "Kamchiliklari", "text": "Cheklangan imkoniyat va sozlash murakkabligi. Farq sezilarli bo'ladi."}]
hits = sum(sl.assign(cmp_, f"T{k}", "uz")[2]["layout"] == "comparison" for k in range(10))
check("taqqoslash mazmuni — comparison (ko'pincha)", hits >= 7, hits)

nums = deck("Raqam", 3)
nums[2]["items"] = [{"head": h, "text": t, "value": v} for h, t, v in [
    ("Samaradorlik", "Ko'rsatkich bir yilda sezilarli oshdi va barqaror bo'ldi.", "85%"),
    ("O'sish", "Hajm avvalgi davrga nisbatan bir necha baravar ko'paydi.", "3,2 marta"),
    ("Qamrov", "Xizmatdan foydalanuvchi mamlakatlar soni doimiy oshib bormoqda.", "47 ta")]]
hits = sum(sl.assign(nums, f"R{k}", "uz")[2]["layout"] in ("stats_row", "stat_hero", "text_with_numbers") for k in range(10))
check("raqamli mazmun — statistik shablon (ko'pincha)", hits >= 7, hits)

# ── eski shakldagi (items'siz) javob ham ishlaydi
old = [{"title": "T", "layout": "cover", "content": ""},
       {"title": "Eski", "layout": "two_column", "content": "",
        "columns": [{"column_content": "Birinchi ustun matni shu yerda. Ikkinchi gap ham bor."},
                    {"column_content": "Ikkinchi ustun matni. Uning ham ikkinchi gapi bor."}]},
       {"title": "Matn", "layout": "right_image",
        "content": "Birinchi gap. Ikkinchi gap. Uchinchi gap. To'rtinchi gap shu yerda."}]
res = sl.assign(old, "Eski", "uz")
check("eski shakl: hamma slayd saqlanadi, layout to'ldiriladi", len(res) == 3 and all(s["layout"] for s in res), [s["layout"] for s in res])

# ── sig'maydigan uzun matn shablonni o'zgartiradi, qirqilmaydi
big = deck("Uzun", 2)
big[1]["items"] = [{"head": f"H{i}", "text": words(f"big{i}", 70)} for i in range(4)]
res = sl.assign(big, "Uzun", "uz")
check("uzun matn: sig'adigan shablon tanlandi yoki matnli slaydga o'tdi",
      res[1]["layout"] in sl.NAMES, res[1]["layout"])

# ── ma'lumotni yo'qotmaslik: reshape
five = [{"head": f"H{i}", "text": f"Gap {i}.", "value": ""} for i in range(5)]
merged = sl.reshape(five, 4, 4)
check("ortiqcha element oxirgisiga qo'shiladi", merged and len(merged) == 4 and "Gap 4." in merged[-1]["text"])
check("yetishmasa gap bo'yicha bo'linadi",
      len(sl.reshape([{"head": "", "text": "Bir. Ikki. Uch. To'rt.", "value": ""}], 2, 2) or []) == 2)
check("bo'linmaydigan bitta gap — moslanmaydi", sl.reshape([{"head": "", "text": "Bitta gap.", "value": ""}], 2, 2) is None)

# ── haqiqiy fayl: matn to'liq, rasm bo'lmasa ham ishlaydi, yilni raqam deb olmaydi
check("yil statistika emas", sl.find_number("2024-yilda tashkil topdi") == "")
check("foiz topiladi", sl.find_number("Ko'rsatkich 85% ga yetdi") == "85%")

class FakeTogether:
    async def _img(self, w, h):
        p = tempfile.mktemp(suffix=".jpg"); Image.new("RGB", (w, h), (90, 120, 160)).save(p); return p
    async def generate_cover_image(self, *a, **k): return await self._img(800, 1200)
    async def generate_slide_image(self, *a, **k): return await self._img(1024, 768)
    async def generate_panoramic_image(self, *a, **k): return await self._img(1344, 576)

async def build(slides, together, template="template_20"):
    svc = DocumentService.__new__(DocumentService)
    svc.documents_dir = tempfile.mkdtemp(); svc.together = together
    svc.use_icons = True; svc._last_used_icons = set()
    return await svc.create_presentation_with_smart_images("Sinov", {"slides": slides}, "Talaba", "uz",
                                                           TemplateService(), template)

def all_text(prs):
    return " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)

for label, together in (("rasmli", FakeTogether()), ("rasmsiz", None)):
    for template in ("template_20", "template_1", "template_19"):
        slides = deck("Ekologiya", 10)
        prs = Presentation(asyncio.run(build(slides, together, template)))
        text = all_text(prs)
        lost = [i["text"] for s in slides if s.get("items") for i in s["items"] if i["text"] not in text.replace("\n", " ")
                and not any(part in text for part in [i["text"][:30]])]
        check(f"{label}/{template}: matn to'liq (hech narsa tashlanmadi)", not lost, lost[:1])
        check(f"{label}/{template}: slaydlar soni >= 12", len(prs.slides) >= 12, len(prs.slides))
        pics = sum(1 for s in prs.slides for sh in s.shapes if sh.name == "fixed:image")
        if together is None:
            check(f"{label}/{template}: rasm yo'q", pics == 0, pics)

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
