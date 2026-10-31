"""3D taqdimot (katalog bo'yicha, har qanday mavzu): qidiruv, AI reja, slaydlar va bot oqimi.

- qidiruv.score: nom va teg bo'yicha ball, qo'shimchaga chidamli, diskda yo'q va og'ir modellar olinmaydi;
- qidiruv.top: AI tanlovi (asosiy/yordamchi), AI ishlamasa — nom bo'yicha aniq moslik;
- reja.normalize: muqova/yakun, noto'g'ri kompozitsiya tuzatiladi, uchlikda 3 model, raqamlarda ≥2 raqam,
  ketma-ket bir xil kompozitsiya yo'q, ro'yxatli slaydlar ko'pi bilan 2 ta, yangi kompozitsiyalar (paragraf,
  xronologiya, taqqoslash, kartalar, bosqichlar, fikr, savol) va ularning maydonlari, matn yo'qolmaydi;
- yukla: serverda yo'q model kerak bo'lganda o'zi yuklanadi (bir marta), xato bo'lsa taklif qilinmaydi;
- deck.build: har slaydda haqiqiy 3D ob'ekt (rasm emas), Morph, qismlarga yorliq-chiziq yo'q,
  sahnadan chiqib-kiruvchi nusxalar slayddan tashqarida, CC BY mualliflari izohda;
- bot: 3D tugmasi → mavzu → katalogda topildi → ism → … → slaydlar soni (3D taqdimot narxlari);
  hech narsa topilmasa — boshqa turlar taklif qilinadi.

    python test_kat3d.py
"""
import asyncio, io, os, sys, tempfile, time, zipfile
from types import SimpleNamespace
from unittest import mock
from unittest.mock import AsyncMock, MagicMock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")
TMP = tempfile.mkdtemp()
os.environ["KATALOG3D_DIR"] = TMP
os.environ["KAT3D_REVIEW"] = "0"

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import numpy as np
from pptx import Presentation
from pptx.util import Emu

from services import katalog3d
from services.kat3d import ai, deck, qidiruv, reja
import services.kat3d as kat3d
from services.med3d import glb

CAT = katalog3d.katalog()
NEFERTITI = "ce5b14926e494558ab584375a8d63ca7"          # Nefertiti byusti (CC BY)
SKARABEY = "c8a568f1c1a746d4b2939b2f9a46f3d9"           # Ulkan skarabey haykali (CC BY)
SKARABEY2 = "b4cd7baabeea42e192c54eb02ca72c1a"
SMILODON = "si_19a0e367-a669-499b-b8c6-6809e5c4a249"     # CC0
IDS = [NEFERTITI, SKARABEY, SKARABEY2, SMILODON]
check("test modellari katalogda bor", all(i in CAT for i in IDS), [i for i in IDS if i not in CAT])


def put_model(mid):
    """Kichik (bitta uchburchak) GLB — diskda "yuklangan" model."""
    pos = np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0]], dtype=np.float32)
    nrm = np.array([[0, 0, 1]] * 3, dtype=np.float32)
    path = os.path.join(TMP, CAT[mid]["fayl"])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    open(path, "wb").write(glb.write_glb([glb.Primitive(pos, nrm, np.array([0, 1, 2]), (0.8, 0.7, 0.5, 1))]))


print("1) Qidiruv — ball")
have = frozenset(IDS)
r = qidiruv.score("Qadimgi Misr", ["Nefertiti", "scarab", "Egypt"], ["tarix"], available=have)
ids = [m for _, m in r]
check("Nefertiti va skarabeylar topildi", {NEFERTITI, SKARABEY, SKARABEY2} <= set(ids), ids)
check("mavzuga aloqasiz model olinmadi", SMILODON not in ids, ids)
check("nomdagi moslik birinchi", ids[0] in (NEFERTITI, SKARABEY, SKARABEY2), r[:3])
check("qo'shimchaga chidamli: «skarabeylar» ~ «skarabey»",
      SKARABEY in [m for _, m in qidiruv.score("skarabeylar", [], available=have)])
check("diskda yo'q model taklif qilinmaydi", qidiruv.score("Nefertiti", [], available=frozenset()) == [])
with mock.patch.object(qidiruv, "MAX_MB", 1.0):
    check("og'ir model (Telegram 50 MB) olinmaydi", qidiruv.score("Nefertiti", [], available=have) == [])
check("bo'sh so'rov — bo'sh ro'yxat", qidiruv.score("va", [], available=have) == [])

print("2) Qidiruv — AI tanlovi")
def fake_ai(expand, select):
    def call(system, user, *a, **kw):
        if system == qidiruv.EXPAND_SYSTEM:
            if isinstance(expand, Exception):
                raise expand
            return expand
        if isinstance(select, Exception):
            raise select
        return select(user)
    return call

def codes_of(user, *mids):
    """select() so'rovidagi c-kodlar: model nomi bo'yicha."""
    out = []
    for line in user.splitlines():
        for mid in mids:
            if line.startswith("c") and f"| {CAT[mid]['nom_uz']} |" in line:
                out.append(line.split(" | ")[0])
    return out

with mock.patch("services.kat3d.yukla.mumkin", return_value=have), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(
         {"en": ["Nefertiti", "scarab"], "uz": ["Misr"], "fan": ["tarix"]},
         lambda u: {"asosiy": codes_of(u, NEFERTITI), "yordamchi": codes_of(u, SKARABEY, SKARABEY2)})):
    t = qidiruv.top("Qadimgi Misr")
check("asosiy model bor — taqdimot bo'ladi", t["ok"] and t["checked"], t)
check("asosiy birinchi, keyin yordamchilar", t["ids"][:1] == [NEFERTITI] and set(t["ids"]) == {NEFERTITI, SKARABEY, SKARABEY2},
      t["ids"])
with mock.patch("services.kat3d.yukla.mumkin", return_value=have), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(
         {"en": ["scarab"], "uz": [], "fan": []}, lambda u: {"asosiy": [], "yordamchi": codes_of(u, SKARABEY)})):
    t = qidiruv.top("Qo'ng'izlar")
check("faqat 1 yordamchi — yetarli emas", not t["ok"] and t["ids"] == [], t)
with mock.patch("services.kat3d.yukla.mumkin", return_value=have), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(RuntimeError("402"), RuntimeError("402"))):
    t = qidiruv.top("Nefertiti")
check("AI ishlamasa — nom bo'yicha aniq moslik", t["ok"] and t["ids"][:1] == [NEFERTITI] and not t["checked"], t)
with mock.patch("services.kat3d.yukla.mumkin", return_value=have), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(RuntimeError("402"), RuntimeError("402"))):
    t = qidiruv.top("Kvant kompyuterlari")
check("hech narsa topilmasa — ok emas", not t["ok"] and not t["ids"], t)
GLOBE = next(i for i in CAT if i.startswith("8a3f6e66955e"))           # Yer globusi (ramziy)
have2 = have | {GLOBE}
with mock.patch("services.kat3d.yukla.mumkin", return_value=have2), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(
         {"en": ["Nefertiti", "scarab", "globe"], "uz": [], "fan": []},
         lambda u: {"asosiy": codes_of(u, NEFERTITI), "yordamchi": codes_of(u, SKARABEY, SKARABEY2),
                    "ramziy": codes_of(u, GLOBE), "coverage": 90})):
    t = qidiruv.top("Qadimgi Misr")
check("qamrov yuqori — to'liq 3D (tugmalarsiz)", t["level"] == "full" and t["ids"][0] == NEFERTITI, t)
with mock.patch("services.kat3d.yukla.mumkin", return_value=have2), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(
         {"en": ["globe", "scarab"], "uz": [], "fan": []},
         lambda u: {"asosiy": [], "yordamchi": codes_of(u, SKARABEY), "ramziy": codes_of(u, GLOBE), "coverage": 30})):
    t = qidiruv.top("Jahon iqtisodiyoti va Misr")
check("qisman mos — uch yo'l, ramziy modellar ham tanlovda", t["level"] == "partial" and t["ok"]
      and set(t["choices"]) == {SKARABEY, GLOBE}, t)
with mock.patch("services.kat3d.yukla.mumkin", return_value=have2), \
     mock.patch.object(ai, "call_json", side_effect=fake_ai(
         {"en": ["globe"], "uz": [], "fan": []}, lambda u: {"asosiy": [], "yordamchi": [], "ramziy": [], "coverage": 5})):
    t = qidiruv.top("Kvant maydon nazariyasi")
check("hech narsa mos emas — none", t["level"] == "none" and not t["ok"] and not t["choices"], t)
pic = qidiruv.collage([NEFERTITI, SKARABEY])
check("topilgan modellar kollaji (JPEG)", pic and pic[:2] == b"\xff\xd8")
check("tanlash uchun raqamlangan kollaj", (qidiruv.collage([NEFERTITI, SKARABEY], numbered=True) or b"")[:2] == b"\xff\xd8")

print("3) Reja — normalize")
codes = {f"m{i}": m for i, m in enumerate(IDS)}
long = "Juda uzun sarlavha " * 10
raw = {"style": "yoq_uslub", "slides": [
    {"layout": "chap", "title": long, "points": ["a"], "models": ["m0"]},
    {"layout": "chap", "title": "Ikkinchi", "points": ["• birinchi fikr", "", "ikkinchi"], "models": ["m0"]},
    {"layout": "chap", "title": "Uchinchi", "points": ["x"], "models": ["m1"]},
    {"layout": "uchlik", "title": "Uchlik", "items": [{"model": "m1", "name": "A", "text": "a"}], "models": []},
    {"layout": "raqamlar", "title": "Raqamlar", "stats": [{"value": "146 m", "label": "balandligi"}], "models": ["m2"]},
    {"layout": "galaktika", "title": "Noma'lum", "points": ["y"], "models": ["m9"]},
    {"layout": "chap", "title": "Oxirgi", "points": ["z"], "models": ["m0", "m1", "m2"]},
    {"layout": "chap", "title": "Ortiqcha", "points": ["z"], "models": ["m0"]}]}
p = reja.normalize(raw, 7, IDS, codes)
lays = [s["layout"] for s in p["slides"]]
check("aniq 7 slayd (ortiqchasi tashlanadi)", len(lays) == 7, lays)
check("1-slayd muqova, oxirgisi yakun", lays[0] == "muqova" and lays[-1] == "yakun", lays)
check("noma'lum kompozitsiya — oddiy matnli", lays[5] in ("chap", "ong", "matn", "qahramon"), lays)
check("ro'yxatli slaydlar ko'pi bilan 2 ta", sum(l in ("chap", "ong") for l in lays[1:-1]) <= 2, lays)
nom = next(s for s in p["slides"] if s["title"] == "Noma'lum")
check("ortiqcha ro'yxat boshqa dizaynga o'tdi, matni saqlandi",
      nom["layout"] not in ("chap", "ong") and "y" in " ".join(nom["points"] + nom["paragraphs"] + [nom["subtitle"]]),
      nom)
check("ketma-ket bir xil kompozitsiya yo'q",
      all(a != b for a, b in zip(lays, lays[1:]) if a in ("chap", "ong", "qahramon")), lays)
check("uchlikda aniq 3 xil model", lays[3] == "uchlik" and len(set(p["slides"][3]["models"])) == 3,
      p["slides"][3])
check("bitta raqamli «raqamlar» — oddiy slaydga aylandi", lays[4] != "raqamlar", lays)
check("sarlavha qisqartirildi", len(p["slides"][0]["title"]) <= reja.LIMITS["title"] and p["slides"][0]["title"].endswith("…"))
check("bo'sh va belgili punktlar tozalandi", p["slides"][1]["points"] == ["birinchi fikr", "ikkinchi"], p["slides"][1]["points"])
check("katalogda yo'q kod (m9) olinmadi", all(m in IDS for s in p["slides"] for m in s["models"]))
check("har slaydda kamida bitta model", all(s["models"] for s in p["slides"]))
check("yakunda bir nechta model", len(p["slides"][-1]["models"]) >= 2, p["slides"][-1]["models"])
check("noma'lum uslub — bo'sh (keyin fan bo'yicha)", p["style"] == "")
limit = (CAT[NEFERTITI]["hajm"] + CAT[SMILODON]["hajm"]) / 1e6 + 0.05     # skarabeylar sig'maydi
check("umumiy hajm chegarasi: muhimlari qoladi", reja._fit_size(IDS, limit) == [NEFERTITI, SMILODON],
      reja._fit_size(IDS, limit))
check("fan bo'yicha uslub: tarix — muzey vitrinasi", reja._style_for([NEFERTITI]) == "vitrina")

print("3b) Har varoq har xil: yangi kompozitsiyalar")
S = "Bu to'liq ma'lumotli gap bo'lib, mavzuning muhim jihatini tushuntiradi."
raw = {"slides": [
    {"layout": "muqova", "title": "Misr", "subtitle": "Kirish", "models": ["m0", "m1"]},
    {"layout": "matn", "title": "Paragraf", "paragraphs": [S + " " + S, S], "models": ["m0"]},
    {"layout": "xronologiya", "title": "Vaqt", "events": [{"date": "1912", "text": S}, {"date": "1922", "text": S},
                                                         {"date": "1950", "text": S}], "models": ["m1"]},
    {"layout": "xronologiya", "title": "Kam voqea", "events": [{"date": "1912", "text": "Birinchi voqea."},
                                                              {"date": "", "text": "Sanasiz."}], "models": ["m1"]},
    {"layout": "taqqoslash", "title": "Ikki tomon", "columns": [{"name": "A", "points": [S, S]},
                                                                {"name": "B", "points": [S]}], "models": ["m2"]},
    {"layout": "kartalar", "title": "Kartalar", "cards": [{"title": f"K{k}", "text": S} for k in range(4)], "models": ["m3"]},
    {"layout": "bosqichlar", "title": "Bosqichlar", "steps": [{"title": f"B{k}", "text": S} for k in range(4)],
     "models": ["m0"]},
    {"layout": "fikr", "title": "Fikr", "quote": "Buyuk fikr.", "source": "", "text": S, "models": ["m1"]},
    {"layout": "savol", "title": "Savol", "question": "Bilasizmi?", "answer": S, "models": ["m2"]},
    {"layout": "kartalar", "title": "Yana kartalar", "cards": [{"title": "K", "text": S}] * 3, "models": ["m3"]},
    {"layout": "yakun", "title": "Xulosa", "points": [S, S], "models": ["m0", "m1"]}]}
p = reja.normalize(raw, 11, IDS, codes)
lays = [s["layout"] for s in p["slides"]]
check("yangi kompozitsiyalar saqlandi", lays[1:3] == ["matn", "xronologiya"] and
      {"taqqoslash", "kartalar", "bosqichlar", "fikr", "savol"} <= set(lays), lays)
check("2 ta voqeali xronologiya — oddiy matnga, voqealari yo'qolmadi",
      lays[3] != "xronologiya" and "Birinchi voqea" in " ".join(p["slides"][3]["points"] + p["slides"][3]["paragraphs"]),
      p["slides"][3])
check("ketma-ket hech qachon bir xil dizayn yo'q", all(a != b for a, b in zip(lays, lays[1:])), lays)
check("taqqoslashda aniq 2 ustun", len(p["slides"][4]["columns"]) == 2)
check("har slaydda 1 model (uchlik/yakundan tashqari)", all(len(s["models"]) == 1 for s in p["slides"][1:-1]))
long_par = "Juda uzun paragraf. " * 60
p2 = reja.normalize({"slides": [{"layout": "muqova", "title": "M"}, {"layout": "matn", "paragraphs": [long_par]},
                                {"layout": "yakun", "points": ["a"]}]}, 3, IDS, codes)
check("paragraf chegarasi", len(p2["slides"][1]["paragraphs"][0]) <= reja.LIMITS["para"])

print("4) Slaydlar — haqiqiy 3D va Morph")
for mid in IDS:
    put_model(mid)
katalog3d.mavjud(ttl=0)
plan = {"topic": "Qadimgi Misr", "language": "uz", "style": "vitrina", "slides": [
    {"layout": "muqova", "kicker": "Tarix", "title": "Qadimgi Misr", "subtitle": "Nil bo'yidagi sivilizatsiya",
     "points": [], "stats": [], "items": [], "models": [NEFERTITI, SKARABEY], "notes": "Kirish."},
    {"layout": "chap", "kicker": "Malika", "title": "Nefertiti", "subtitle": "", "points": ["Exnatonning xotini",
     "1912-yili topilgan"], "stats": [], "items": [], "models": [NEFERTITI], "notes": ""},
    {"layout": "raqamlar", "kicker": "Raqamlar", "title": "Xeops piramidasi", "subtitle": "", "points": [],
     "stats": [{"value": "146 m", "label": "balandligi"}, {"value": "2,3 mln", "label": "tosh blok"}],
     "items": [], "models": [SKARABEY], "notes": ""},
    {"layout": "uchlik", "kicker": "Ramzlar", "title": "Muqaddas jonzotlar", "subtitle": "", "points": [], "stats": [],
     "items": [{"model": SKARABEY, "name": "Skarabey", "text": "quyosh"},
               {"model": SKARABEY2, "name": "Tumor", "text": "himoya"},
               {"model": SMILODON, "name": "Mushuksimon", "text": "kuch"}],
     "models": [SKARABEY, SKARABEY2, SMILODON], "notes": ""},
    {"layout": "qahramon", "kicker": "Haykal", "title": "Byust", "subtitle": "", "points": ["Tutmos ustaxonasi"],
     "stats": [], "items": [], "models": [NEFERTITI], "notes": ""},
    {"layout": "yakun", "kicker": "Xulosa", "title": "Meros", "subtitle": "", "points": ["Yozuv va me'morchilik"],
     "stats": [], "items": [], "models": [NEFERTITI, SKARABEY, SMILODON], "notes": ""}]}
out = os.path.join(TMP, "kat3d.pptx")
deck.build(plan, out)
prs = Presentation(out)
W, H = prs.slide_width, prs.slide_height
check("6 ta slayd", len(prs.slides) == 6, len(prs.slides))
z = zipfile.ZipFile(out)
glbs = [n for n in z.namelist() if n.endswith(".glb")]
check("har model taqdimotda bir marta (GLB)", len(glbs) == 4, glbs)
xmls = [z.read(f"ppt/slides/slide{i}.xml").decode() for i in range(1, 7)]
check("har slaydda haqiqiy 3D ob'ekt (am3d:model3d)", all("model3d" in x for x in xmls))
check("Morph o'tishi (2-slayddan)", all("morph" in x for x in xmls[1:]), [("morph" in x) for x in xmls])


def models_on(slide):
    """(nom, markaz_x, markaz_y) — "!!m_" bilan boshlanuvchi 3D ob'ektlar (graphicFrame emas, AlternateContent)."""
    out = []
    for el in slide.shapes._spTree.iter():
        if el.tag.endswith("}cNvPr") and (el.get("name") or "").startswith("!!m_"):
            frame = el.getparent().getparent()
            off = next((e for e in frame.iter() if e.tag.endswith("}off")), None)
            ext = next((e for e in frame.iter() if e.tag.endswith("}ext") and e.get("cx")), None)
            if off is not None and ext is not None and frame.tag.endswith("}graphicFrame"):
                out.append((el.get("name"), int(off.get("x")) + int(ext.get("cx")) // 2,
                            int(off.get("y")) + int(ext.get("cy")) // 2))
    return out


inside = lambda x, y: 0 <= x <= W and 0 <= y <= H
names = [[n for n, x, y in models_on(s) if inside(x, y)] for s in prs.slides]
check("slayddagi modellar rejadagidek", [len(n) for n in names] == [2, 1, 1, 3, 1, 3], [len(n) for n in names])
check("bir model bir slaydda bir xil nom bilan (Morph uni topadi)",
      all(len(n) == len(set(n)) for n in names), names)
m_nef = "!!m_" + NEFERTITI[:24]
check("Nefertiti muqovadan keyingi slaydga uchib o'tadi (bir xil nom)", m_nef in names[0] and m_nef in names[1])
ghosts = [[n for n, x, y in models_on(s) if not inside(x, y)] for s in prs.slides]
check("chiqayotgan model keyingi slaydda sahna tashqarisida kutadi",
      "!!m_" + SKARABEY[:24] in ghosts[1], ghosts[1])
check("kirayotgan model oldingi slaydda sahna tashqarisidan uchib keladi",
      "!!m_" + SKARABEY2[:24] in ghosts[2], ghosts[2])
turns = [x.count('presetID="37"') for x in xmls]
check("sayyora bo'lmagan modellar o'z joyida aylanmaydi", turns == [0] * 6, turns)
plan_spin = dict(plan, spin=[NEFERTITI])
deck.build(plan_spin, os.path.join(TMP, "spin.pptx"))
zs = zipfile.ZipFile(os.path.join(TMP, "spin.pptx"))
turns = [zs.read(f"ppt/slides/slide{i}.xml").decode().count('presetID="37"') for i in range(1, 7)]
nef = ["!!m_" + NEFERTITI[:24] in n for n in names]
check("rejadagi «spin» modeli o'z joyida sekin aylanadi (Turntable)", turns == [int(x) for x in nef], (turns, nef))
x_spin = zs.read("ppt/slides/slide2.xml").decode()
check("aylanish 3D ko'rinish burchagi bo'yicha, cheksiz takrorlanadi",
      '3d.view.rotation.y' in x_spin and 'repeatCount="indefinite"' in x_spin)
EARTH = "5f9c35be31a047928eace8b415a8ee3a"
check("sayyoralar aylanadi, haykal va mashinalar yo'q",
      deck.is_planet(EARTH) and not deck.is_planet(NEFERTITI) and not deck.is_planet("28c2d309449848aea6b8b35e25d4acb6"))
lines = [sh for s in prs.slides for sh in s.shapes if sh.shape_type is not None and "CONNECTOR" in str(sh.shape_type)]
check("qismlarga yorliq-chiziq yo'q (anatomiya emas)", not lines, len(lines))
texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
check("muqovada sarlavha", "Qadimgi Misr" in texts, texts)
check("CC BY muallifi slaydda ko'rsatilgan", any(CAT[NEFERTITI]["attribution"] in sh.text_frame.text
                                                  for sh in prs.slides[1].shapes if sh.has_text_frame))
last_notes = prs.slides[-1].notes_slide.notes_text_frame.text
check("oxirgi slayd izohida hamma CC BY mualliflari",
      all(CAT[m]["muallif"] in last_notes for m in IDS if CAT[m]["litsenziya"] == "CC BY 4.0"), last_notes[:300])
check("fayl telefonda ham ochiladi: fallback rasm shaffof PNG", any(n.endswith(".png") for n in z.namelist()))

print("4b) Yangi kompozitsiyalar chiziladi")
p = reja.normalize(raw, 11, IDS, codes)
p.update(style="atlas", topic="Misr", language="uz")
over = dict(p["slides"][1], layout="chap", points=["Juda uzun punkt matni, u to'liq gap bo'lib davom etadi " * 8] * 5,
            subtitle=S, notes="")
p["slides"].insert(2, over)
out2 = os.path.join(TMP, "xil.pptx")
deck.build(p, out2)
prs2 = Presentation(out2)
shape_names = [{sh.name for sh in sl.shapes} for sl in prs2.slides]
def has(prefix):
    return any(any(n.startswith(prefix) for n in names) for names in shape_names)
for prefix, what in (("para_", "paragraf"), ("dot_", "vaqt chizig'i"), ("col_", "taqqoslash ustunlari"),
                     ("card_", "kartalar"), ("num_", "bosqich raqamlari"), ("qmark_", "iqtibos belgisi"),
                     ("badge_", "savol nishoni")):
    check(f"{what} chizildi", has(prefix), prefix)
check("har slaydda sarlavha (Morph uchun)", all("!!title" in n for n in shape_names), [sorted(n)[:3] for n in shape_names])
notes2 = prs2.slides[2].notes_slide.notes_text_frame.text if prs2.slides[2].has_notes_slide else ""
body = [sh for sh in prs2.slides[2].shapes if sh.name.startswith("body_")]
kept = len(body[0].text_frame.paragraphs) if body else 0
check("sig'magan punktlar o'chmadi — ma'ruzachi izohiga ko'chdi", 0 < kept < 5 and notes2.count("•") == 5 - kept,
      (kept, notes2[:120]))

print("4c) Aralash taqdimot: tanlangan 3D + diagramma + rasm")
mix_raw = {"slides": [
    {"layout": "muqova", "title": "Inflyatsiya", "models": []},
    {"layout": "matn", "title": "Nima?", "paragraphs": [S + " " + S], "models": []},
    {"layout": "diagramma", "title": "Turlari", "chart": {"type": "column", "labels": ["A", "B", "C"], "values": ["3%", 25, "200"], "unit": "%"},
     "subtitle": S, "points": [S], "models": []},
    {"layout": "diagramma", "title": "Buzuq grafik", "chart": {"labels": ["A"], "values": [1]}, "points": [S, S], "models": []},
    {"layout": "rasm", "title": "Hayotda", "photo": "supermarket prices", "points": [S, S], "models": []},
    {"layout": "fikr", "title": "Muvozanat", "quote": "Narx — tarozi.", "text": S, "models": ["m1"]},
    {"layout": "qahramon", "title": "Modelsiz qahramon", "subtitle": S, "points": [S], "models": []},
    {"layout": "yakun", "title": "Xulosa", "points": [S, S], "models": []}]}
mix_ids = [NEFERTITI, SKARABEY, SMILODON]
mix_codes = {f"m{i}": m for i, m in enumerate(mix_ids)}
p = reja.normalize(mix_raw, 8, mix_ids, mix_codes, mix=True)
lays = [s["layout"] for s in p["slides"]]
used_models = {m for s in p["slides"] for m in s["models"]}
check("aralash: har tanlangan model kamida bir slaydda", used_models == set(mix_ids), used_models)
check("aralash: ko'p slayd 3D'siz", sum(1 for s in p["slides"] if not s["models"]) >= 3, [len(s["models"]) for s in p["slides"]])
check("diagramma: qiymatlar songa aylandi", p["slides"][2]["layout"] == "diagramma" and p["slides"][2]["chart"]["values"] == [3.0, 25.0, 200.0],
      p["slides"][2])
check("bitta qiymatli grafik — matnli slaydga", lays[3] != "diagramma", lays)
check("rasm slaydida qidiruv so'rovi", p["slides"][4]["layout"] == "rasm" and p["slides"][4]["photo"] == "supermarket prices")
check("modelsiz «qahramon» — oddiy matnli slayd", p["slides"][6]["layout"] != "qahramon" or p["slides"][6]["models"], p["slides"][6])
full_p = reja.normalize(mix_raw, 8, mix_ids, mix_codes)
check("to'liq rejimda diagramma/rasm yo'q, har slaydda model",
      not {"diagramma", "rasm"} & {s["layout"] for s in full_p["slides"]} and all(s["models"] for s in full_p["slides"]))
from PIL import Image as _Img
photo_path = os.path.join(TMP, "foto.jpg")
_Img.new("RGB", (1200, 700), (180, 120, 60)).save(photo_path)
for s_ in p["slides"]:
    if s_["layout"] == "rasm":
        s_["photo_path"] = photo_path
p.update(style="tungi", topic="Inflyatsiya", language="uz")
out3 = os.path.join(TMP, "mix.pptx")
deck.build(p, out3)
prs3 = Presentation(out3)
z3 = zipfile.ZipFile(out3)
check("diagramma — PowerPoint'ning o'z grafigi (tahrirlanadi)", any(n.startswith("ppt/charts/chart") for n in z3.namelist()))
check("rasm slaydida foto", any(sh.name.startswith("photo_") for sl in prs3.slides for sh in sl.shapes))
model_slides = sum(1 for sl in prs3.slides if any(inside(x, y) for _, x, y in models_on(sl)))   # arvohlar ekrandan tashqarida
check("3D faqat tanlangan joylarda", model_slides == sum(1 for s_ in p["slides"] if s_["models"]), model_slides)
check("aralash narx 40% arzon", kat3d.price_for(10, "mix") == round(kat3d.price_for(10) * 0.6 / 100) * 100 < kat3d.price_for(10))

print("5) build_deck — diskda model yo'q bo'lsa")
try:
    asyncio.run(kat3d.build_deck("Mavzu", 5, ["yoq_model"]))
    check("modelsiz buyurtma — NoModels", False)
except kat3d.NoModels:
    check("modelsiz buyurtma — NoModels", True)

print("5b) Avtomatik yuklash (3D Anatomiya atlasi kabi)")
from services.kat3d import yukla
QUYOSH = "9ef1c68fbb944147bcfcc891d3912645"
calls = []

def fake_download(url, path, tries=4):
    calls.append(url)
    if "xato" in url:
        raise OSError("tarmoq")
    open(path, "wb").write(b"raw")

class FakeKonv:
    p = None                                          # haqiqiy node jarayoni yo'q — qorovul kuzatmaydi
    closed = 0

    def close(self):
        FakeKonv.closed += 1

    def __call__(self, job):
        pos = np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0]], dtype=np.float32)
        data = glb.write_glb([glb.Primitive(pos, np.array([[0, 0, 1]] * 3, dtype=np.float32), np.array([0, 1, 2]),
                                            (1, 1, 1, 1))])
        open(job["out"], "wb").write(data)
        return {"ok": True, "size": len(data), "tris": 1}

src = {QUYOSH: {"uid": QUYOSH, "manba": "objaverse", "url": "https://example/quyosh.glb", "nom": "Sun",
                "muallif": "X", "sahifa": "https://example", "litsenziya": "CC BY 4.0"},
       SMILODON: {"uid": SMILODON, "manba": "smithsonian", "url": "https://example/xato.glb", "nom": "Smilodon",
                  "litsenziya": "CC0 1.0"}}
os.remove(os.path.join(TMP, CAT[SMILODON]["fayl"]))
with mock.patch.object(yukla, "ishlaydi", return_value=True), mock.patch.object(yukla, "manbalar", return_value=src), \
     mock.patch.object(yukla, "yuklab_ol", side_effect=fake_download), \
     mock.patch.object(yukla, "_konvertor", return_value=FakeKonv()):
    check("serverda yo'q, lekin manbasi bor model taklif qilinadi", QUYOSH in yukla.mumkin())
    yukla.oldindan([QUYOSH]); yukla.oldindan([QUYOSH])
    got = yukla.ensure([QUYOSH, SMILODON, NEFERTITI], timeout=30)
    check("kerak bo'lganda o'zi yuklandi (diskdagi bilan birga)", got == [QUYOSH, NEFERTITI], got)
    check("bir model bir marta yuklanadi", calls.count("https://example/quyosh.glb") == 1, calls)
    qdir = os.path.dirname(os.path.join(TMP, CAT[QUYOSH]["fayl"]))
    check("model.glb, LICENSE.txt, meta.json", sorted(os.listdir(qdir)) == ["LICENSE.txt", "meta.json", "model.glb"],
          os.listdir(qdir))
    check("litsenziyada muallif", "CC BY 4.0" in open(os.path.join(qdir, "LICENSE.txt"), encoding="utf-8").read())
    check("xom fayllar o'chirildi", os.listdir(os.path.join(TMP, "_xom")) == [], os.listdir(os.path.join(TMP, "_xom")))
    check("yuklanmagan model endi taklif qilinmaydi", SMILODON not in yukla.mumkin())
    check("diskdagilar ro'yxati yangilandi", QUYOSH in katalog3d.mavjud())
check("har model alohida jarayonda — ish tugashi bilan yopiladi (xotira qaytadi)", FakeKonv.closed >= 1, FakeKonv.closed)

print("5c) Xotira: 2 GB li server qotib qolmasin")
OGIR = dict(src[QUYOSH], uchburchak=6938773, hajm=251_000_000, tekstura_max=8192)
YENGIL = dict(src[QUYOSH], uchburchak=500_000, hajm=80_000_000, tekstura_max=8192)
check("og'ir model ~1,6 GB, oddiy ~350 MB deb baholanadi",
      1500 <= yukla.ehtiyoj_mb(OGIR) <= 2100 and 300 <= yukla.ehtiyoj_mb(YENGIL) <= 450,
      (yukla.ehtiyoj_mb(OGIR), yukla.ehtiyoj_mb(YENGIL)))
os.remove(os.path.join(TMP, CAT[QUYOSH]["fayl"])); katalog3d.mavjud(ttl=0)
with mock.patch.object(yukla, "ishlaydi", return_value=True), mock.patch.object(yukla, "byudjet_mb", return_value=880), \
     mock.patch.object(yukla, "manbalar", return_value={QUYOSH: OGIR}), \
     mock.patch.object(yukla, "yuklab_ol", side_effect=fake_download) as dl:
    check("2 GB serverda sig'maydigan og'ir model taklif qilinmaydi", QUYOSH not in yukla.mumkin())
    check("... va yuklanmaydi ham", yukla.oldindan([QUYOSH]) == 0 and yukla.ensure([QUYOSH], timeout=5) == []
          and not dl.called)
with mock.patch.object(yukla, "ishlaydi", return_value=True), mock.patch.object(yukla, "byudjet_mb", return_value=880), \
     mock.patch.object(yukla, "manbalar", return_value={QUYOSH: YENGIL}):
    check("oddiy model 2 GB serverda ham taklif qilinadi", QUYOSH in yukla.mumkin())
with mock.patch.object(yukla, "ishlaydi", return_value=True), \
     mock.patch.object(yukla, "_start", side_effect=lambda mid, pool: mid) as start:
    yukla.oldindan([f"m{i}" for i in range(20)])
    check("mavzu yozilganda fonda ko'pi bilan 8 ta model", start.call_count == yukla.OLDINDAN_MAX == 8, start.call_count)

import subprocess as _sp
class SlowKonv:
    def __init__(self):
        self.p = _sp.Popen([sys.executable, "-c", "import time; time.sleep(60)"])
    def __call__(self, job):
        self.p.wait()                                 # node o'ldirilsa readline bo'sh qaytadi
        raise RuntimeError("konvertor to'xtab qoldi")
    def close(self):
        self.p.kill()
slow = []
with mock.patch.object(yukla, "_konvertor", side_effect=lambda: slow.append(SlowKonv()) or slow[-1]), \
     mock.patch.object(yukla, "_meminfo", return_value={"MemAvailable": 120, "MemTotal": 1970}):
    t0 = time.time()
    try:
        yukla._aylantir("x.glb", "y.glb", 400)
        check("bo'sh xotira tugasa qorovul o'girishni to'xtatadi", False)
    except yukla.XotiraBand as exc:
        check("bo'sh xotira tugasa qorovul o'girishni to'xtatadi", time.time() - t0 < 5 and slow[0].p.poll() is not None,
              str(exc))
with mock.patch.object(yukla, "_meminfo", return_value={"MemAvailable": 300}):
    t0 = time.time()
    check("xotira bo'shamasa o'girish boshlanmaydi (kutadi, keyin voz kechadi)",
          yukla._kut_xotira(400, timeout=0.1) is False and time.time() - t0 < 5)
with mock.patch.object(yukla, "_meminfo", return_value={"MemTotal": 1970}):
    check("ruxsat — server xotirasining 45% i (2 GB da ~880 MB)", 850 <= yukla.byudjet_mb() <= 900, yukla.byudjet_mb())
put_model(QUYOSH); katalog3d.mavjud(ttl=0)
with mock.patch.object(yukla, "ishlaydi", return_value=False):
    check("node/paketlar yo'q — faqat diskdagilar", yukla.mumkin() == katalog3d.mavjud())
put_model(SMILODON)
katalog3d.mavjud(ttl=0)

print("6) Bot oqimi")
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage
from bot.handlers import premium_presentation as pp
from bot.states import PremiumPresentationStates as PS
from services.med3d import auto as med3d_auto


class Chat:
    def __init__(self):
        self.next_id, self.sent, self.photos = 100, [], []
        self.bot = SimpleNamespace(delete_message=AsyncMock())

    def message(self, text=None):
        chat, msg = self, MagicMock()
        msg.text, msg.bot = text, self.bot
        msg.chat = SimpleNamespace(id=1)
        msg.from_user = SimpleNamespace(id=1, first_name="Ali")

        async def answer(body, **kw):
            chat.next_id += 1
            chat.sent.append([chat.next_id, body, kw.get("reply_markup")])
            note = MagicMock(message_id=chat.next_id)

            async def edit(new, **kw2):
                chat.sent.append([note.message_id, new, kw2.get("reply_markup")])
            note.edit_text = edit
            return note

        async def photo(file, **kw):
            chat.photos.append(file)
        msg.answer, msg.answer_photo = answer, photo
        msg.edit_text, msg.delete = AsyncMock(), AsyncMock()
        return msg


def callback(chat, data):
    cb = MagicMock(); cb.data = data; cb.answer = AsyncMock()
    cb.message = chat.message(); cb.bot = chat.bot
    cb.from_user = SimpleNamespace(id=1)
    return cb


def make_db():
    db = MagicMock()
    db.get_user = AsyncMock(return_value=SimpleNamespace(language="uz", first_name="Ali", balance=10**6, id=1))
    db.get_active_channels = AsyncMock(return_value=[])
    db.get_feature_status = AsyncMock(return_value=False)
    return db


def buttons(markup):
    return [b for row in markup.inline_keyboard for b in row] if markup else []


prefetched = []


async def flow(found, kat=None):
    chat, db = Chat(), make_db()
    state = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=1, user_id=1))
    await pp.premium_presentation_start(chat.message(pp.MED3D_ENTRY_TEXTS[0]), state, db)
    kat = kat or ({"ok": True, "ids": [NEFERTITI, SKARABEY], "asosiy": [NEFERTITI], "checked": True} if found
                  else {"ok": False, "ids": [], "asosiy": [], "checked": True})
    with mock.patch.object(med3d_auto, "check_topic", return_value={"ok": False, "checked": True}), \
         mock.patch.object(qidiruv, "top", return_value=kat), \
         mock.patch("services.kat3d.yukla.oldindan", side_effect=prefetched.extend):
        await pp.premium_ppt_got_topic(chat.message("Qadimgi Misr"), state, db)
    return chat, db, state


async def bot_flow():
    chat, db, state = await flow(found=True)
    data = await state.get_data()
    check("katalogda topildi — tur «3D taqdimot», modellar buyurtmada",
          data.get("kind") == pp.KIND_KAT3D and data.get("kat3d_ids") == [NEFERTITI, SKARABEY], data)
    found_msg = [s[1] for s in chat.sent if CAT[NEFERTITI]["nom_uz"] in str(s[1])]
    check("mijozga topilgan modellar nomi ko'rsatildi", found_msg, [s[1] for s in chat.sent])
    check("va rasmi (kollaj) yuborildi", len(chat.photos) == 1)
    check("topilgan modellar to'lovdan oldin fonda yuklana boshladi", prefetched == [NEFERTITI, SKARABEY], prefetched)
    check("keyin darhol ism so'raladi (promokodsiz)", await state.get_state() == PS.waiting_for_client_name.state,
          await state.get_state())
    await pp.premium_ppt_skip_name(callback(chat, "prem_ppt_skip_name"), state, db)
    await pp.premium_ppt_skip_preferences(callback(chat, "prem_ppt_skip_preferences"), state, db)
    check("manba va uslub so'ralmaydi — darhol slaydlar soni", await state.get_state() == PS.waiting_for_count.state,
          await state.get_state())
    counts = [b.callback_data for b in buttons(chat.sent[-1][2]) if b.callback_data.startswith("prem_ppt_count:")]
    check("3D taqdimot slaydlar soni varianti", counts == [f"prem_ppt_count:{n}" for n in kat3d.SLIDE_OPTIONS], counts)
    await pp.premium_ppt_got_count(callback(chat, "prem_ppt_count:8"), state, db)
    summary = chat.sent[-1][1]
    check("xulosada 3D modellar soni va narx", "3D modellar" in summary and "<b>2</b>" in summary
          and pp._price_of(await state.get_data(), 8) == kat3d.price_for(8), summary)

    chat, db, state = await flow(found=False)
    last = chat.sent[-1]
    check("mos 3D topilmadi — boshqa turlar taklif qilinadi",
          "prem_m3d_alt:kind:modern" in [b.callback_data for b in buttons(last[2])]
          and await state.get_state() != PS.waiting_for_client_name.state, last[1])
    await pp.premium_med3d_alternative(callback(chat, "prem_m3d_alt:kind:modern"), state, db)
    data = await state.get_data()
    check("boshqa tur tanlandi — 3D belgisi tozalandi, mavzu saqlandi",
          data.get("kind") == "modern" and not data.get("standalone_med3d") and data.get("topic") == "Qadimgi Misr",
          data)
    await pp.premium_ppt_previous(callback(chat, "prem_ppt_prev"), state, db)
    check("orqaga — oddiy mavzu so'rovi (3D matni emas)",
          await state.get_state() == PS.waiting_for_topic_text.state and "3D taqdimot" not in chat.sent[-1][1],
          (await state.get_state(), chat.sent[-1][1][:80]))

async def partial_flow():
    choices = [NEFERTITI, SKARABEY, SKARABEY2, SMILODON, GLOBE, QUYOSH]
    part = {"level": "partial", "ok": True, "ids": choices[:4], "choices": choices, "asosiy": [], "checked": True}
    chat, db, state = await flow(True, part)
    last = chat.sent[-1]
    cbs = [b.callback_data for b in buttons(last[2])]
    check("qisman mos — uch tugma: to'liq / aralash / 3D'siz", cbs == ["prem_k3m:full", "prem_k3m:mix", "prem_k3m:plain"], cbs)
    check("aralash 40% arzon ekani aytildi", "40%" in last[1], last[1][:120])
    check("keyingi qadamga o'tilmadi — mijoz tanlaydi", await state.get_state() != PS.waiting_for_client_name.state)
    cb = callback(chat, "prem_k3m:plain")
    await pp.premium_kat3d_mode(cb, state, db)
    plain = [b.callback_data for b in buttons(cb.message.edit_text.call_args.kwargs["reply_markup"])]
    check("3D'siz — infografik va zamonaviy tugmalari", plain[:2] == ["prem_m3d_alt:kind:info", "prem_m3d_alt:kind:modern"], plain)
    await pp.premium_kat3d_mode(callback(chat, "prem_k3m:full"), state, db)
    data = await state.get_data()
    check("to'liq 3D — modellar avtomatik tanlandi", data["kat3d_mode"] == "full" and data["kat3d_ids"] == choices[:4]
          and await state.get_state() == PS.waiting_for_client_name.state, data)

    chat, db, state = await flow(True, part)
    await pp.premium_kat3d_mode(callback(chat, "prem_k3m:mix"), state, db)
    check("aralash — ism so'raladi", await state.get_state() == PS.waiting_for_client_name.state)
    await pp.premium_ppt_skip_name(callback(chat, "prem_ppt_skip_name"), state, db)
    await pp.premium_ppt_skip_preferences(callback(chat, "prem_ppt_skip_preferences"), state, db)
    labels = [b.text for b in buttons(chat.sent[-1][2]) if b.callback_data.startswith("prem_ppt_count:")]
    check("slaydlar soni narxlari 40% arzon", f"{kat3d.price_for(5, 'mix'):,}" in labels[0], labels[:2])
    photos = len(chat.photos)
    await pp.premium_ppt_got_count(callback(chat, "prem_ppt_count:5"), state, db)
    text, markup = chat.sent[-1][1], chat.sent[-1][2]
    picks = [b for b in buttons(markup) if b.callback_data.startswith("prem_k3p:") and b.callback_data != "prem_k3p:done"]
    check("model tanlash: rasm va 2 ustunli tugmalar", len(chat.photos) == photos + 1 and len(picks) == len(choices)
          and all(len(r) <= 2 for r in markup.inline_keyboard[:-2]) and await state.get_state() == PS.waiting_for_kat3d_pick.state,
          (len(picks), [len(r) for r in markup.inline_keyboard]))
    cb = callback(chat, "prem_k3p:done")
    await pp.premium_kat3d_pick(cb, state, db)
    check("hech narsa tanlanmasa — davom etmaydi", cb.answer.call_args.kwargs.get("show_alert")
          and await state.get_state() == PS.waiting_for_kat3d_pick.state)
    for k in range(5):
        await pp.premium_kat3d_pick(callback(chat, f"prem_k3p:{k}"), state, db)
    cb = callback(chat, "prem_k3p:5")
    await pp.premium_kat3d_pick(cb, state, db)
    data = await state.get_data()
    check("slaydlar sonidan (5) ko'p tanlab bo'lmaydi", len(data["kat3d_ids"]) == 5 and cb.answer.call_args.kwargs.get("show_alert"),
          data["kat3d_ids"])
    await pp.premium_kat3d_pick(callback(chat, "prem_k3p:0"), state, db)
    check("bosilsa tanlov bekor bo'ladi", NEFERTITI not in (await state.get_data())["kat3d_ids"])
    await pp.premium_kat3d_pick(callback(chat, "prem_k3p:done"), state, db)
    summary = chat.sent[-1][1]
    check("xulosa: aralash 3D va 40% arzon narx", "Aralash 3D" in summary
          and pp._price_of(await state.get_data(), 5) == kat3d.price_for(5, "mix"), summary)
    await pp.premium_ppt_previous(callback(chat, "prem_ppt_prev"), state, db)
    check("xulosadan orqaga — model tanlashga", await state.get_state() == PS.waiting_for_kat3d_pick.state)
    await pp.premium_ppt_previous(callback(chat, "prem_ppt_prev"), state, db)
    check("tanlashdan orqaga — slaydlar soniga", await state.get_state() == PS.waiting_for_count.state)

asyncio.run(bot_flow())
asyncio.run(partial_flow())

print()
print("HAMMASI OK" if not FAILS else f"{len(FAILS)} ta XATO: {FAILS}")
sys.exit(1 if FAILS else 0)
