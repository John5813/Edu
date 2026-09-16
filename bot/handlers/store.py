"""Saytdagi tayyor ishni botda sotib olish.

Sayt faqat ko'rgazma rasmlarini ko'rsatadi; faylning o'zi yopiq Telegram
kanalida turadi. "Sotib olish" tugmasi mijozni shu yerga — botga olib
keladi, chunki to'lov tizimi (balans, Stars, to'ldirish) allaqachon shu
yerda ishlaydi.
"""

import asyncio
import logging
import os
import re
import shutil

from aiogram import F, Router
from aiogram.filters import Command, CommandObject, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.types import CallbackQuery, InlineKeyboardButton, InlineKeyboardMarkup, Message
from aiogram.utils.keyboard import InlineKeyboardBuilder

import webapp
from bot import checkout as pay, uploads
from bot.states import StorePublishStates, StoreStates
from config import ADMIN_IDS, STORE_VAULT_CHAT_ID
from database.database import Database
from translations import get_text

router = Router()
logger = logging.getLogger(__name__)

CHECKOUT = pay.Checkout(service="store", back_callback="store_cancel")

_CODE_RE = re.compile(r"^[A-Z0-9]{8}$")


def _t(lang: str, uz: str, ru: str, en: str) -> str:
    return {"uz": uz, "ru": ru, "en": en}.get(lang, uz)


def _card(item: dict, lang: str) -> str:
    """Sotib olishdan oldingi ma'lumot oynasi."""
    price = f"{item['price']:,}".replace(",", " ")
    lines = {
        "uz": [f"📊 <b>{item['title']}</b>", ""],
        "ru": [f"📊 <b>{item['title']}</b>", ""],
        "en": [f"📊 <b>{item['title']}</b>", ""],
    }[lang if lang in ("uz", "ru", "en") else "uz"]
    if item.get("description"):
        lines.append(item["description"])
        lines.append("")
    if item.get("slide_count"):
        count = item["slide_count"]
        docx = (item.get("file_type") or "pptx") == "docx"
        lines.append(_t(lang,
                        f"📄 {'Varaqlar' if docx else 'Slaydlar'}: {count} ta",
                        f"📄 {'Страниц' if docx else 'Слайдов'}: {count}",
                        f"📄 {'Pages' if docx else 'Slides'}: {count}"))
    # Ichki kod mijozga ko'rsatilmaydi — u faqat admin uchun.
    lines.append("")
    lines.append(_t(lang, f"💰 Narxi: <b>{price} so'm</b>",
                   f"💰 Цена: <b>{price} сум</b>",
                   f"💰 Price: <b>{price} so'm</b>"))
    return "\n".join(lines)


async def _deliver(message: Message, item: dict, lang: str) -> bool:
    """Faylni ombordan mijozga yuboradi."""
    try:
        await message.answer_document(
            document=item["file_id"],
            caption=f"📊 {item['title']}",
            parse_mode=None,
        )
        return True
    except Exception as exc:
        logger.error("Do'kon fayli yuborilmadi (%s): %s", item["public_code"], exc)
        await message.answer(_t(
            lang,
            "❌ Faylni yuborishda xatolik bo'ldi. Administratorga murojaat qiling.",
            "❌ Не удалось отправить файл. Обратитесь к администратору.",
            "❌ Could not send the file. Please contact the administrator.",
        ))
        return False


async def _pending(state: FSMContext) -> dict | None:
    """FSM'dagi buyurtmani katalogdagi joriy holat bilan qayta o'qiydi."""
    data = await state.get_data()
    code = data.get("store_code")
    if not code:
        return None
    return await Database.get_store_item(code)


async def _sell(message: Message, item: dict, lang: str, user_id: int,
                state: FSMContext, db: Database, charge: bool) -> None:
    """To'lovni yechib, faylni yetkazadi va sotuvni yozadi."""
    if charge:
        await db.update_user_balance(user_id, -item["price"])
    if not await _deliver(message, item, lang):
        if charge:  # yetkazilmadi — pul qaytariladi
            await db.update_user_balance(user_id, item["price"])
        return
    await Database.record_store_sale(item["id"], user_id, item["price"])
    pay.forget(user_id)
    await state.clear()
    await message.answer(_t(
        lang,
        "✅ Xaridingiz uchun rahmat! Ish har doim shu yerdan qayta yuklab olinadi.",
        "✅ Спасибо за покупку! Работу всегда можно скачать здесь повторно.",
        "✅ Thanks for your purchase! You can download this work again any time.",
    ))


# ── Saytdan kelish ─────────────────────────────────────────────────────

@router.message(CommandStart(deep_link=True, magic=F.args.startswith("buy_")))
async def store_open(message: Message, command: CommandObject, state: FSMContext,
                     db: Database):
    await state.clear()
    code = (command.args or "")[4:].strip().upper()

    user = await db.get_user(message.from_user.id)
    if not user:
        user = await db.create_user(
            telegram_id=message.from_user.id,
            username=message.from_user.username,
            first_name=message.from_user.first_name,
            language="uz",
        )
    lang = user.language if user else "uz"

    item = await Database.get_store_item(code) if _CODE_RE.match(code) else None
    if not item:
        await message.answer(_t(
            lang,
            "❌ Bu ish topilmadi — u katalogdan olib tashlangan bo'lishi mumkin.",
            "❌ Работа не найдена — возможно, она удалена из каталога.",
            "❌ This work was not found — it may have been removed from the catalogue.",
        ))
        return

    # Ilgari sotib olingan bo'lsa qayta to'lov so'ralmaydi.
    if await Database.has_bought_store_item(message.from_user.id, item["id"]):
        await message.answer(_t(
            lang,
            "✅ Bu ishni siz allaqachon sotib olgansiz — mana u:",
            "✅ Вы уже покупали эту работу — вот она:",
            "✅ You have already bought this work — here it is:",
        ))
        await _deliver(message, item, lang)
        return

    await state.update_data(store_code=item["public_code"], price=item["price"])
    await state.set_state(StoreStates.waiting_for_payment)
    await message.answer(
        _card(item, lang),
        parse_mode="HTML",
        reply_markup=pay.payment_keyboard(CHECKOUT, lang, item["price"]),
    )


# ── To'lov ─────────────────────────────────────────────────────────────

@router.callback_query(F.data == CHECKOUT.pay_balance,
                       StoreStates.waiting_for_payment)
async def store_pay_balance(callback: CallbackQuery, state: FSMContext, db: Database):
    await callback.answer()
    user = await db.get_user(callback.from_user.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await _expired(callback.message, state, lang)
        return

    balance = user.balance if user else 0
    if balance < item["price"]:
        await pay.send_shortfall(callback.message, CHECKOUT, lang, item["price"], balance)
        return

    try:
        await callback.message.edit_reply_markup(reply_markup=None)
    except Exception:
        pass
    await _sell(callback.message, item, lang, callback.from_user.id, state, db, charge=True)


@router.callback_query(F.data == CHECKOUT.pay_other,
                       StoreStates.waiting_for_payment)
async def store_pay_other(callback: CallbackQuery, state: FSMContext, db: Database):
    await callback.answer()
    user = await db.get_user(callback.from_user.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await _expired(callback.message, state, lang)
        return
    try:
        await callback.message.edit_reply_markup(
            reply_markup=pay.other_methods_keyboard(CHECKOUT, lang, item["price"]))
    except Exception:
        pass


@router.callback_query(F.data == CHECKOUT.pay_back,
                       StoreStates.waiting_for_payment)
async def store_pay_back(callback: CallbackQuery, state: FSMContext, db: Database):
    await callback.answer()
    user = await db.get_user(callback.from_user.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await _expired(callback.message, state, lang)
        return
    try:
        await callback.message.edit_reply_markup(
            reply_markup=pay.payment_keyboard(CHECKOUT, lang, item["price"]))
    except Exception:
        pass


@router.callback_query(F.data == CHECKOUT.recheck,
                       StoreStates.waiting_for_payment)
async def store_recheck(callback: CallbackQuery, state: FSMContext, db: Database):
    """Balans to'ldirilgandan keyin buyurtmani yo'qotmasdan davom etish."""
    user = await db.get_user(callback.from_user.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await callback.answer()
        await _expired(callback.message, state, lang)
        return

    balance = user.balance if user else 0
    if balance < item["price"]:
        await callback.answer(
            get_text(lang, "pay_still_short", balance=balance, price=item["price"]),
            show_alert=True,
        )
        return

    await callback.answer()
    await _sell(callback.message, item, lang, callback.from_user.id, state, db, charge=True)


@router.callback_query(F.data == CHECKOUT.pay_stars,
                       StoreStates.waiting_for_payment)
async def store_pay_stars(callback: CallbackQuery, state: FSMContext, db: Database):
    await callback.answer()
    user = await db.get_user(callback.from_user.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await _expired(callback.message, state, lang)
        return

    sent = await pay.send_invoice(
        callback.message, CHECKOUT, lang, item["price"],
        title=item["title"],
        description=_t(lang, "Tayyor ish — do'kondan xarid",
                       "Готовая работа — покупка из каталога",
                       "Ready-made work — catalogue purchase"),
    )
    if sent:
        try:
            await callback.message.edit_reply_markup(reply_markup=None)
        except Exception:
            pass


@router.message(StoreStates.waiting_for_payment, F.successful_payment)
async def store_successful_stars(message: Message, state: FSMContext, db: Database):
    payment = message.successful_payment
    if not payment or not CHECKOUT.owns_payload(payment.invoice_payload):
        return

    user = await db.get_user(message.chat.id)
    lang = user.language if user else "uz"
    item = await _pending(state)
    if not item:
        await _expired(message, state, lang)
        return
    # Stars to'lovi Telegram tomonida yechilgan — balansdan olinmaydi.
    await _sell(message, item, lang, message.chat.id, state, db, charge=False)


@router.callback_query(F.data == "store_cancel")
async def store_cancel(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    pay.forget(callback.from_user.id)
    await state.clear()
    try:
        await callback.message.edit_reply_markup(reply_markup=None)
    except Exception:
        pass


async def _expired(message: Message, state: FSMContext, lang: str) -> None:
    await state.clear()
    await message.answer(_t(
        lang,
        "⏳ Buyurtma eskirdi yoki ish katalogdan olib tashlandi. Saytdan qaytadan tanlang.",
        "⏳ Заказ устарел или работа удалена из каталога. Выберите заново на сайте.",
        "⏳ The order expired or the work was removed. Please pick it again on the site.",
    ))


# ── Admin: katalogga qo'yish ───────────────────────────────────────────

SKIP = "-"
# Do'kon taqdimot ham, Word hujjati ham qabul qiladi; PDF esa tozalab
# bo'lmaydi — undagi matn qatlami tahrirlanmaydi.
STORE_FORMATS = (".pptx", ".docx")


def _is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


@router.message(Command("nashr"))
async def publish_start(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    if not STORE_VAULT_CHAT_ID:
        await message.answer(
            "❌ <code>STORE_VAULT_CHAT_ID</code> sozlanmagan.\n\n"
            "Yopiq kanal oching, botni u yerga administrator qiling va kanal "
            "ID'sini (<code>-100...</code>) shu o'zgaruvchiga yozing.",
            parse_mode="HTML",
        )
        return
    # Slaydlarni rasmga aylantirish LibreOffice'ga tayanadi. Buni oldindan
    # aytmasak, admin faylni yuborib, savollarga javob berib bo'lgach
    # xatoga uchrardi.
    if not shutil.which("soffice"):
        await message.answer(
            "❌ LibreOffice o'rnatilmagan — varaqlardan ko'rgazma rasmi olib "
            "bo'lmaydi.\n\nServerda:\n"
            "<code>apt install libreoffice-impress libreoffice-writer</code>\n\n"
            "Impress taqdimot uchun, Writer esa Word hujjatlari uchun kerak.",
            parse_mode="HTML",
        )
        return
    await state.clear()
    await state.set_state(StorePublishStates.waiting_for_file)
    await message.answer(
        "📤 <b>Katalogga qo'yish</b>\n\n"
        "Tayyor <b>.pptx</b> faylni yuboring.\n"
        "Bekor qilish uchun /bekor.",
        parse_mode="HTML",
    )


@router.message(Command("bekor"), StateFilter(StorePublishStates))
async def publish_cancel(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    data = await state.get_data()
    path = data.get("pub_path")
    if path:
        try:
            os.remove(path)
        except OSError:
            pass
    await state.clear()
    await message.answer("Bekor qilindi.")


@router.message(StorePublishStates.waiting_for_file, F.document)
async def publish_got_file(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    upload = await uploads.receive(message, "uz", accept=STORE_FORMATS,
                                   prefix="store_pub", extract=False)
    if upload is None:
        return

    await state.update_data(pub_path=upload.path)
    await state.set_state(StorePublishStates.waiting_for_work_type)
    await message.answer("🗃 Bu qanday ish?", reply_markup=_work_type_keyboard())


def _work_type_keyboard() -> InlineKeyboardMarkup:
    from config import STORE_WORK_LABELS

    keyboard = InlineKeyboardBuilder()
    for key, label in STORE_WORK_LABELS.items():
        keyboard.add(InlineKeyboardButton(text=label, callback_data=f"pubtype:{key}"))
    keyboard.adjust(2)
    return keyboard.as_markup()


@router.callback_query(F.data.startswith("pubtype:"),
                       StorePublishStates.waiting_for_work_type)
async def publish_got_work_type(callback: CallbackQuery, state: FSMContext):
    if not _is_admin(callback.from_user.id):
        return
    from config import STORE_WORK_LABELS, work_label

    key = callback.data.split(":", 1)[1]
    if key not in STORE_WORK_LABELS:
        await callback.answer("Noma'lum tur", show_alert=True)
        return

    await callback.answer()
    await state.update_data(pub_work_type=key)
    await state.set_state(StorePublishStates.waiting_for_customer)
    try:
        await callback.message.edit_text(f"🗃 Ish turi: <b>{work_label(key)}</b>",
                                         parse_mode="HTML")
    except Exception:
        pass
    await callback.message.answer(
        "👤 Ish kimga tayyorlangan edi? <b>Ism-familiyasini</b> yozing — "
        "u fayldan tozalanadi.\n\n"
        f"Ism noma'lum bo'lsa <code>{SKIP}</code> yuboring.",
        parse_mode="HTML",
    )


@router.message(StorePublishStates.waiting_for_customer, F.text)
async def publish_got_customer(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    name = (message.text or "").strip()
    await state.update_data(pub_customer="" if name == SKIP else name)
    await state.set_state(StorePublishStates.waiting_for_title)
    await message.answer("📝 Saytda ko'rinadigan <b>sarlavha</b>ni yozing.",
                         parse_mode="HTML")


@router.message(StorePublishStates.waiting_for_title, F.text)
async def publish_got_title(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    title = (message.text or "").strip()
    if len(title) < 3:
        await message.answer("Sarlavha juda qisqa. Qaytadan yozing.")
        return
    from config import store_price

    await state.update_data(pub_title=title[:300])
    await state.set_state(StorePublishStates.waiting_for_price)
    data = await state.get_data()
    suggested = store_price(data.get("pub_work_type", ""))
    await message.answer(
        f"💰 <b>Narxini</b> so'mda yozing.\n\n"
        f"Bu tur uchun odatdagisi: <b>{suggested:,}</b> so'm — "
        f"rozi bo'lsangiz <code>{SKIP}</code> yuboring.".replace(",", " "),
        parse_mode="HTML")


def _category_prompt(title: str) -> str:
    from services.store_taxonomy import classify

    guess = classify(title)
    if guess:
        return (f"🗂 <b>Fan</b>: mavzuga qarab <b>{guess}</b> deb aniqlandi.\n\n"
                f"Rozi bo'lsangiz <code>{SKIP}</code> yuboring, "
                f"yoki boshqa nom yozing.")
    return ("🗂 <b>Fan</b> nomini yozing (masalan: <code>iqtisodiyot</code>).\n"
            f"Kerak bo'lmasa <code>{SKIP}</code>.")


@router.message(StorePublishStates.waiting_for_price, F.text)
async def publish_got_price(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    from config import store_price

    raw = (message.text or "").replace(" ", "").replace(",", "")
    if raw == SKIP:
        data = await state.get_data()
        price = store_price(data.get("pub_work_type", ""))
    elif raw.isdigit() and 0 < int(raw) <= 10_000_000:
        price = int(raw)
    else:
        await message.answer("Narx faqat raqamlardan iborat bo'lsin. Qaytadan yozing.")
        return
    await state.update_data(pub_price=price)
    await state.set_state(StorePublishStates.waiting_for_category)
    data = await state.get_data()
    await message.answer(_category_prompt(data.get("pub_title", "")), parse_mode="HTML")


@router.message(StorePublishStates.waiting_for_category, F.text)
async def publish_got_category(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    category = (message.text or "").strip().lower()
    await state.update_data(pub_category="" if category == SKIP else category[:50])
    await state.set_state(StorePublishStates.waiting_for_description)
    await message.answer(
        "🧾 Qisqa <b>tavsif</b> yozing — u saytda sarlavha ostida turadi.\n"
        f"Kerak bo'lmasa <code>{SKIP}</code>.",
        parse_mode="HTML",
    )


@router.message(StorePublishStates.waiting_for_description, F.text)
async def publish_got_description(message: Message, state: FSMContext):
    if not _is_admin(message.from_user.id):
        return
    from services.store_publisher import publish_work

    description = (message.text or "").strip()
    data = await state.get_data()
    path = data.get("pub_path", "")
    await state.clear()

    status = await message.answer("⏳ Tozalanmoqda va katalogga qo'yilmoqda...")
    try:
        result = await publish_work(
            message.bot, path,
            title=data.get("pub_title", ""),
            price=int(data.get("pub_price", 0)),
            customer_name=data.get("pub_customer", ""),
            description="" if description == SKIP else description[:1000],
            category=data.get("pub_category", ""),
            work_type=data.get("pub_work_type", ""),
        )
    except Exception as exc:
        logger.exception("Katalogga qo'yilmadi: %s", exc)
        await status.edit_text(f"❌ Katalogga qo'yilmadi: {exc}")
        return
    finally:
        try:
            os.remove(path)
        except OSError:
            pass

    code = result["public_code"]
    unit = "varaq" if result["file_type"] == "docx" else "slayd"
    await status.edit_text(
        f"✅ <b>Katalogga qo'shildi</b>\n\n"
        f"🔖 Kod: <code>{code}</code>\n"
        f"📄 {result['slide_count']} {unit}, {result['preview_count']} ta ko'rgazma rasmi\n"
        f"🌐 {webapp.public_url('/shop/' + code)}\n\n"
        f"Olib tashlash: <code>/nashr_ochir {code}</code>",
        parse_mode="HTML",
    )


def _conversion_works(extension: str) -> bool:
    """Shu turdagi faylni rasmga aylantirib bo'ladimi.

    Paket bor-yo'qligini tekshirish o'rniga haqiqiy o'tkazish sinaladi:
    LibreOffice filtrlari alohida paketlarda keladi va `soffice` mavjud
    bo'la turib .docx yoki .pptx ni ocholmasligi mumkin.
    """
    import tempfile

    from services.premium_presentation.qa import discard_images, pptx_to_images

    folder = tempfile.mkdtemp()
    path = os.path.join(folder, f"sinov{extension}")
    try:
        if extension == ".pptx":
            from pptx import Presentation
            from pptx.util import Inches

            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)) \
                .text_frame.text = "sinov"
            deck.save(path)
        else:
            from docx import Document

            document = Document()
            document.add_paragraph("sinov")
            document.save(path)

        images = pptx_to_images(path)
        discard_images(images)
        return bool(images)
    except Exception as exc:
        logger.warning("Konvertatsiya sinovi xato berdi (%s): %s", extension, exc)
        return False
    finally:
        shutil.rmtree(folder, ignore_errors=True)


@router.message(Command("dokon"))
async def store_status(message: Message):
    """Do'kon ishlashiga kerak bo'lgan hamma narsani tekshiradi.

    Avtomatik nashr mijozga yetkazishni buzmaslik uchun xatolarni yutadi,
    shuning uchun nosozlik faqat shu yerda ko'rinadi.
    """
    if not _is_admin(message.from_user.id):
        return

    from config import STORE_AUTO_PUBLISH, STORE_PREVIEW_DIR
    from services import store_publisher

    status = await message.answer("⏳ Tekshirilmoqda...")
    lines = ["🏪 <b>Do'kon holati</b>\n"]

    auto = "yoqilgan" if STORE_AUTO_PUBLISH else "o'chirilgan"
    lines.append(f"{'✅' if STORE_AUTO_PUBLISH else '⏸'} Avtomatik nashr: {auto}")

    if STORE_VAULT_CHAT_ID:
        try:
            probe = await message.bot.send_message(
                STORE_VAULT_CHAT_ID, "Do'kon tekshiruvi — bu xabar o'chiriladi.")
            try:
                await message.bot.delete_message(STORE_VAULT_CHAT_ID, probe.message_id)
            except Exception:
                pass
            lines.append(f"✅ Ombor kanali: <code>{STORE_VAULT_CHAT_ID}</code> — yozish mumkin")
        except Exception as exc:
            lines.append(f"❌ Ombor kanali: <code>{STORE_VAULT_CHAT_ID}</code>\n"
                         f"    {type(exc).__name__}: {str(exc)[:160]}\n"
                         f"    <i>Bot o'sha kanalda administrator ekanini tekshiring.</i>")
    else:
        lines.append("❌ Ombor kanali: <code>STORE_VAULT_CHAT_ID</code> sozlanmagan")

    loop = asyncio.get_running_loop()
    for extension, label, package in ((".pptx", "Taqdimot", "libreoffice-impress"),
                                      (".docx", "Word hujjati", "libreoffice-writer")):
        ok = await loop.run_in_executor(None, _conversion_works, extension)
        lines.append(f"{'✅' if ok else '❌'} {label} rasmga aylantirish"
                     + ("" if ok else f"\n    <i>Kerak: apt install {package}</i>"))

    lines.append(f"🌐 Sayt: {webapp.public_url('/shop')}")

    try:
        total = (await Database.list_store_items(limit=1))["total"]
        lines.append(f"📦 Katalogda: {total} ta ish")
    except Exception as exc:
        lines.append(f"❌ Katalog o'qilmadi: {exc}")

    try:
        folders = len(os.listdir(STORE_PREVIEW_DIR))
        lines.append(f"🖼 Ko'rgazma papkalari: {folders} ta")
    except OSError:
        pass

    if store_publisher.LAST_ERROR:
        last = store_publisher.LAST_ERROR
        lines.append(f"\n⚠️ <b>Oxirgi nashr xatosi</b> ({last.get('at', '')})\n"
                     f"    {last.get('title', '')} — {last.get('work_type', '')}\n"
                     f"    <code>{last.get('error', '')}</code>")

    await status.edit_text("\n".join(lines), parse_mode="HTML",
                           disable_web_page_preview=True)


@router.message(Command("nashr_ochir"))
async def publish_remove(message: Message, command: CommandObject):
    if not _is_admin(message.from_user.id):
        return
    code = (command.args or "").strip().upper()
    if not _CODE_RE.match(code):
        await message.answer("Foydalanish: <code>/nashr_ochir KOD</code>", parse_mode="HTML")
        return
    if await Database.set_store_item_active(code, False):
        from services.store_publisher import discard_previews
        discard_previews(code)
        await message.answer(f"✅ {code} katalogdan olib tashlandi.")
    else:
        await message.answer(f"❌ {code} topilmadi.")
