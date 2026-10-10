"""3D Pro taqdimotni PowerPoint'ning haqiqiy shakllari bilan yig'adi: Morph o'tishi, yorliq-chiziqlar, animatsiya.

Morph qanday ishlaydi: har slaydda bir xil nomli shakllar bor — `!!panel` (katta rangli fon shakli),
`!!stripe` (shaffof qiya chiziq), `!!obj` (3D obyekt), `!!title`, `!!page`. PowerPoint "!!" bilan
boshlangan bir xil nomli shakllarni qo'shni slaydlarda "bitta narsa" deb biladi va ularni silliq
siljitadi, kattalashtiradi, shaklini o'zgartiradi: panel butun varaqdan doiraga aylanadi, obyekt markazdan
chapga o'tib kattalashadi. Yangi shakllar (yorliqlar, punktlar) esa o'tishdan keyin navbat bilan chiqadi.

Morph PowerPoint 2019/2021/365 da ishlaydi; eski versiya va telefon ilovalari o'rniga oddiy "Fade"
o'tishini ko'rsatadi (`mc:Fallback`) — hech narsa buzilmaydi.
"""
import logging
import math
import os
import uuid
from typing import Dict, List, Optional, Tuple

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .images import Picture, edge_points

log = logging.getLogger(__name__)

W, H = 13.333, 7.5
FONT = "Segoe UI"
MORPH_MS = 1400

# Ranglar: bg — ochiq fon, ink — fondagi matn, panel — katta rangli shakl, on — panel ustidagi matn,
# pill — yorliq foni, pill_ink — yorliq matni, accent — ochiq fondagi urg'u.
PALETTES: Dict[str, Dict[str, str]] = {
    "indigo": dict(bg="F1F3FB", ink="1D2142", muted="5B6180", panel="4B5BDB", on="FFFFFF", on_muted="DDE1FF",
                   pill="FFFFFF", pill_ink="3E4DC9", accent="4B5BDB", line="FFFFFF"),
    "emerald": dict(bg="EEF7F2", ink="12302A", muted="4E6B62", panel="17885F", on="FFFFFF", on_muted="D3F1E4",
                    pill="FFFFFF", pill_ink="137451", accent="17885F", line="FFFFFF"),
    "violet": dict(bg="F4F1FD", ink="241A44", muted="625A80", panel="6A4BE0", on="FFFFFF", on_muted="E4DDFF",
                   pill="FFFFFF", pill_ink="5A3CD0", accent="6A4BE0", line="FFFFFF"),
    "crimson": dict(bg="FBF1F2", ink="3A1219", muted="7A5960", panel="BE3149", on="FFFFFF", on_muted="FFDCE2",
                    pill="FFFFFF", pill_ink="A8253B", accent="BE3149", line="FFFFFF"),
    "teal": dict(bg="EEF7F8", ink="10313A", muted="4F6B72", panel="0F8494", on="FFFFFF", on_muted="CDEFF3",
                 pill="FFFFFF", pill_ink="0B6D7A", accent="0F8494", line="FFFFFF"),
    "amber": dict(bg="FBF5EC", ink="33230E", muted="7A6448", panel="B35F0B", on="FFFFFF", on_muted="FFE6C7",
                  pill="FFFFFF", pill_ink="9A4F06", accent="B35F0B", line="FFFFFF"),
    "graphite": dict(bg="14161C", ink="EEF0F6", muted="A3A9BA", panel="2A3040", on="FFFFFF", on_muted="C9CEDC",
                     pill="F2B33D", pill_ink="15171D", accent="F2B33D", line="F2B33D"),
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# ───────────────────────────────────────────────────────────── yordamchilar

def _set_alpha(shape, alpha: float) -> None:
    """Shakl to'ldirish rangining shaffofligi (alpha 0..1)."""
    fill = shape._element.spPr.find(qn("a:solidFill"))
    if fill is None or not len(fill):
        return
    color = fill[0]
    for old in color.findall(qn("a:alpha")):
        color.remove(old)
    node = etree.SubElement(color, qn("a:alpha"))
    node.set("val", str(int(max(0.0, min(1.0, alpha)) * 100000)))


def box(slide, name: str, x, y, w, h, color: str, *, round_: float = 0.0, alpha: float = 1.0,
        rotation: float = 0.0, line: Optional[str] = None, line_pt: float = 0.0):
    """Yumaloq to'rtburchak (round_=0.5 va w=h — doira). Hamma panel bir xil shakl turi — Morph silliq."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.adjustments[0] = max(0.0, min(0.5, round_))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    if alpha < 1:
        _set_alpha(shape, alpha)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_pt or 1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if rotation:
        shape.rotation = rotation
    if shape.has_text_frame:
        shape.text_frame.text = ""
    return shape


def fit_size(text: str, w: float, h: float, start: float, smallest: float, line: float = 1.18) -> float:
    """Matn qutiga sig'adigan shrift o'lchami (taxminiy: o'rtacha harf eni ≈ 0.53 em)."""
    size = start
    paragraphs = [p for p in str(text or "").split("\n")] or [""]
    while size > smallest:
        per_line = max(1, int(w * 72 / (size * 0.53)))
        lines = sum(max(1, math.ceil(len(p) / per_line)) for p in paragraphs)
        if lines * size * line / 72 <= h:
            return size
        size -= 1
    return smallest


def text(slide, name: str, x, y, w, h, value: str, *, size: float, color: str, bold: bool = False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, smallest: Optional[float] = None,
         line: float = 1.12, caps: bool = False, spacing: Optional[float] = None):
    value = str(value or "")
    if caps:
        value = value.upper()
    if smallest:
        size = fit_size(value, w, h, size, smallest, line + 0.06)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    for index, part in enumerate(value.split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        run = para.add_run()
        run.text = part
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = _rgb(color)
        if spacing is not None:
            run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return shape


def pill(slide, name: str, x, y, label: str, pal: Dict[str, str], *, size: float = 11.5,
         right: bool = False, fill: Optional[str] = None, ink: Optional[str] = None) -> Tuple[object, float]:
    """Yumaloq yorliq ("TYPES OF BONES" kabi). Eni matnga qarab. (shakl, eni) qaytaradi."""
    label = (label or "").upper()
    width = min(3.6, max(0.9, len(label) * size * 0.66 / 72 + 0.42))
    height = 0.36 * size / 11.5
    left = x - width if right else x
    shape = box(slide, name, left, y, width, height, fill or pal["pill"], round_=0.5)
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Inches(0.12)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = _rgb(ink or pal["pill_ink"])
    run._r.get_or_add_rPr().set("spc", "60")
    return shape, width


def place(slide, picture: Optional[Picture], name: str, x, y, w, h, *, valign: str = "middle",
          halign: str = "center") -> Optional[Tuple[float, float, float, float]]:
    """Rasmni qutiga nisbatini saqlab joylaydi. Qaytaradi: (x, y, w, h) dyuymda yoki None."""
    if picture is None:
        return None
    ratio = picture.width / max(1, picture.height)
    pw, ph = (w, w / ratio) if w / h < ratio else (h * ratio, h)
    px = x + {"left": 0, "right": w - pw}.get(halign, (w - pw) / 2)
    py = y + {"top": 0, "bottom": h - ph}.get(valign, (h - ph) / 2)
    shape = slide.shapes.add_picture(picture.path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    shape.name = name
    if not picture.transparent:
        # Fon olib tashlanmagan rasm — oq kartochka sifatida, yumaloq burchak bilan.
        geom = shape._element.spPr.find(qn("a:prstGeom"))
        if geom is not None:
            geom.set("prst", "roundRect")
    return px, py, pw, ph


# ───────────────────────────────────────────────────────────── Morph va animatsiya

_MORPH = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
    '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="{ms}">'
    '<p159:morph option="byObject"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'spd="med"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')


def add_morph(slide, ms: int = MORPH_MS) -> None:
    node = etree.fromstring(_MORPH.format(ms=ms))
    anchor = slide._element.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = slide._element.find(qn("p:cSld"))
    anchor.addnext(node)


def add_entrances(slide, steps: List[Tuple[object, str, int]]) -> None:
    """Shakllar o'tishdan keyin o'zi navbat bilan chiqadi: [(shakl, "fade"|"float"|"zoom", kechikish ms)]."""
    from services.premium_presentation import slide_anim

    if not steps:
        return
    plan, builds = [], []
    for shape, kind, delay in steps:
        plan.append((shape.shape_id, kind, delay))
        if etree.QName(shape._element).localname == "sp" and shape._element.find(qn("p:txBody")) is not None:
            filled = shape.has_text_frame and shape.text_frame.text.strip()
            builds.append(f'<p:bldP spid="{shape.shape_id}" grpId="0"{"" if filled else " animBg=\"1\""}/>')
    timing = etree.fromstring(slide_anim._timing(plan, builds))
    ext = slide._element.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(timing)
    else:
        slide._element.append(timing)


# ───────────────────────────────────────────────────────────── slaydlar

class Ctx:
    def __init__(self, pal: Dict[str, str], pictures: Dict[str, Picture], anchors: Dict[int, list],
                 author: str, year: str, total: int):
        self.pal, self.pictures, self.anchors = pal, pictures, anchors
        self.author, self.year, self.total = author, year, total


def _picture(ctx: Ctx, slide_data: Dict) -> Optional[Picture]:
    return ctx.pictures.get((slide_data.get("object") or {}).get("id", ""))


def _page(slide, ctx: Ctx, index: int, on_panel: bool) -> None:
    color = ctx.pal["on_muted"] if on_panel else ctx.pal["muted"]
    text(slide, "!!page", W - 1.55, H - 0.55, 1.1, 0.3, f"{index + 1:02d} / {ctx.total:02d}", size=10,
         color=color, align=PP_ALIGN.RIGHT, spacing=1)


def _title(slide, ctx: Ctx, value: str, x, y, w, h, *, on_panel: bool, size: float = 30, smallest: float = 22):
    return text(slide, "!!title", x, y, w, h, value, size=size, smallest=smallest, bold=True,
                color=ctx.pal["on"] if on_panel else ctx.pal["ink"], anchor=MSO_ANCHOR.TOP, line=1.05)


def _stripe(slide, ctx: Ctx, x: float, *, on_panel: bool):
    color = ctx.pal["on"] if on_panel else ctx.pal["panel"]
    return box(slide, "!!stripe", x, -1.4, 2.1, H + 2.8, color, alpha=0.09 if on_panel else 0.06, rotation=14)


def slide_cover(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    box(slide, "!!panel", 0, 0, W, H, pal["panel"])
    _stripe(slide, ctx, 6.2, on_panel=True)
    place(slide, _picture(ctx, s), "!!obj", 0.35, 0.45, 6.2, 6.7)
    _title(slide, ctx, s["title"], 6.95, 1.6, 5.85, 2.5, on_panel=True, size=46, smallest=30)
    box(slide, "!!dot", 6.98, 4.32, 0.9, 0.08, pal["pill"], round_=0.5)
    sub = text(slide, "cover_sub", 6.95, 4.6, 5.7, 1.3, s.get("subtitle", ""), size=18, smallest=14,
               color=pal["on_muted"], line=1.15)
    meta = " · ".join(p for p in (ctx.author, ctx.year) if p)
    steps = [(sub, "float", 200)]
    if meta:
        steps.append((text(slide, "cover_meta", 6.95, 6.25, 5.7, 0.4, meta, size=13, color=pal["on_muted"]),
                      "fade", 500))
    return steps


def _callout_slots(anchors: List[Tuple[float, float]], heights: List[float], top: float,
                   bottom: float) -> List[float]:
    """Yorliqlar balandligi: iloji boricha o'z nuqtasi ro'parasida, bir-birini bosmasin, varaqdan chiqmasin."""
    order = sorted(range(len(anchors)), key=lambda i: anchors[i][1])
    gap = 0.28
    ys = [0.0] * len(anchors)
    cursor = top
    for i in order:
        ys[i] = max(cursor, anchors[i][1] - 0.22)
        cursor = ys[i] + heights[i] + gap
    overflow = cursor - gap - bottom
    if overflow > 0:
        # Pastdan sig'madi: hammasini yuqoriga suramiz (tepa chegaragacha), keyin oraliqni qisqartiramiz.
        shift = min(overflow, min(ys[i] for i in order) - top)
        for i in order:
            ys[i] -= shift
        cursor = top
        for i in order:
            ys[i] = max(cursor, min(ys[i], bottom - heights[i]))
            cursor = ys[i] + heights[i] + 0.12
    return ys


def slide_callouts(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    box(slide, "!!panel", 0, 0, W, H, pal["panel"])
    _stripe(slide, ctx, 8.7, on_panel=True)
    picture = _picture(ctx, s)
    placed = place(slide, picture, "!!obj", 4.3, 1.4, 4.75, 5.85, valign="bottom")
    _title(slide, ctx, s["title"], 0.6, 0.42, 8.6, 0.95, on_panel=True)
    callouts = s.get("callouts") or []
    if not callouts:
        return []
    ox, oy, ow, oh = placed or (4.3, 1.4, 4.75, 5.85)

    # Nuqtalar: vision javobi, bo'lmasa obyekt chetidan.
    found = ctx.anchors.get(index) or [None] * len(callouts)
    n = len(callouts)
    default_sides = ["left" if i % 2 == 0 else "right" for i in range(n)]
    rows = [0.25 + 0.5 * (i // 2) / max(1, (n - 1) // 2) if n > 2 else 0.35 + 0.3 * (i % 2) for i in range(n)]
    fallback = edge_points(picture, default_sides, rows) if picture else [(0.3 if d == "left" else 0.7, r)
                                                                         for d, r in zip(default_sides, rows)]
    rel = [found[i] if i < len(found) and found[i] else fallback[i] for i in range(n)]
    anchors = [(ox + ow * x, oy + oh * y) for x, y in rel]
    center = ox + ow / 2
    sides = ["left" if ax < center else "right" for ax, _ in anchors]
    # Bir tomonga 3 tadan ortiq tushmasin.
    for side, other in (("left", "right"), ("right", "left")):
        same = [i for i in range(n) if sides[i] == side]
        while len(same) > 3 or (len(same) == n and n >= 3):
            move = max(same, key=lambda i: anchors[i][0]) if side == "left" else min(same, key=lambda i: anchors[i][0])
            sides[move] = other
            same.remove(move)

    text_w = 3.45
    heights = []
    sizes = []
    for c in callouts:
        size = fit_size(c["text"], text_w, 1.0, 14, 11, 1.2)
        lines = max(1, math.ceil(len(c["text"]) / max(1, int(text_w * 72 / (size * 0.53)))))
        sizes.append(size)
        heights.append(0.4 + 0.1 + lines * size * 1.2 / 72)

    steps = []
    delay = 150
    for side in ("left", "right"):
        group = [i for i in range(n) if sides[i] == side]
        if not group:
            continue
        ys = _callout_slots([anchors[i] for i in group], [heights[i] for i in group], 1.55, H - 0.55)
        for i, top in zip(group, ys):
            c = callouts[i]
            ax, ay = anchors[i]
            if side == "left":
                p, pw = pill(slide, f"call_pill_{i}", 0.6, top, c["label"], pal, size=12.5)
                start = (0.6 + pw + 0.04, top + 0.2)
                body = text(slide, f"call_text_{i}", 0.6, top + 0.5, text_w, heights[i] - 0.44, c["text"],
                            size=sizes[i], color=pal["on_muted"], line=1.15)
            else:
                p, pw = pill(slide, f"call_pill_{i}", W - 0.6, top, c["label"], pal, size=12.5, right=True)
                start = (W - 0.6 - pw - 0.04, top + 0.2)
                body = text(slide, f"call_text_{i}", W - 0.6 - text_w, top + 0.5, text_w, heights[i] - 0.44,
                            c["text"], size=sizes[i], color=pal["on_muted"], line=1.15, align=PP_ALIGN.RIGHT)
            line = slide.shapes.add_connector(MSO_CONNECTOR.CURVE, Inches(start[0]), Inches(start[1]),
                                              Inches(ax), Inches(ay))
            line.name = f"call_line_{i}"
            line.line.color.rgb = _rgb(pal["line"])
            line.line.width = Pt(1.5)
            dot = box(slide, f"call_dot_{i}", ax - 0.08, ay - 0.08, 0.16, 0.16, pal["pill"], round_=0.5,
                      line=pal["panel"], line_pt=2)
            steps += [(p, "float", delay), (line, "fade", delay + 120), (dot, "zoom", delay + 220),
                      (body, "float", delay + 160)]
            delay += 330
    return steps


def slide_points(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    picture = _picture(ctx, s)
    has_pic = picture is not None
    box(slide, "!!panel", -0.3, 0, 5.75 if has_pic else 0.55, H, pal["panel"])
    _stripe(slide, ctx, 3.7 if has_pic else -2.0, on_panel=True)
    place(slide, picture, "!!obj", 0.25, 0.55, 4.95, 6.5)
    x0 = 6.0 if has_pic else 1.3
    width = W - x0 - 0.7
    _title(slide, ctx, s["title"], x0, 0.55, width, 1.0, on_panel=False)
    box(slide, "!!dot", x0, 1.55, 0.9, 0.07, pal["accent"], round_=0.5)
    points = s.get("points") or []
    steps = []
    count = max(1, len(points))
    gap = (H - 2.05 - 0.6) / count
    for i, item in enumerate(points):
        top = 2.0 + i * gap
        badge = box(slide, f"pt_badge_{i}", x0, top, 0.46, 0.46, pal["accent"], round_=0.22)
        frame = badge.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = str(i + 1)
        run.font.name, run.font.size, run.font.bold = FONT, Pt(15), True
        # Urg'u rangi panel bilan bir xil bo'lsa — oq raqam; boshqacha (graphite: sariq) — to'q raqam.
        run.font.color.rgb = _rgb(pal["on"] if pal["accent"] == pal["panel"] else pal["pill_ink"])
        shapes = [badge]
        y = top
        if item.get("head"):
            shapes.append(text(slide, f"pt_head_{i}", x0 + 0.7, y - 0.02, width - 0.7, 0.42, item["head"],
                               size=17, smallest=13, bold=True, color=pal["ink"]))
            y += 0.44
        shapes.append(text(slide, f"pt_text_{i}", x0 + 0.7, y, width - 0.7, gap - (y - top) - 0.12,
                           item["text"], size=14, smallest=11, color=pal["muted"], line=1.15))
        steps += [(shapes[0], "zoom", 150 + i * 280)] + [(sh, "float", 230 + i * 280) for sh in shapes[1:]]
    return steps


def _cell_borders(cell, color: str, width_pt: float = 0.75) -> None:
    """Jadval katagining ingichka och chegaralari (sukutdagi qora chiziq o'rniga)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    for index, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(int(width_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tc_pr.insert(index, ln)


def slide_table(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    picture = _picture(ctx, s)
    box(slide, "!!panel", 8.9, -1.6, 9.6, 9.6, pal["panel"], round_=0.5)
    _stripe(slide, ctx, -0.6, on_panel=False)
    place(slide, picture, "!!obj", 0.3, 0.9, 4.55, 6.2)
    x0 = 5.2 if picture else 0.8
    width = W - x0 - 0.75
    _title(slide, ctx, s["title"], x0, 0.5, width, 1.0, on_panel=False)
    box(slide, "!!dot", x0, 1.5, 0.9, 0.07, pal["accent"], round_=0.5)
    rows = s.get("rows") or []
    columns = s.get("columns") or ["", ""]
    row_h = min(0.82, (H - 2.0 - 0.7) / (len(rows) + 1))
    shape = slide.shapes.add_table(len(rows) + 1, 2, Inches(x0), Inches(1.85), Inches(width),
                                   Inches(row_h * (len(rows) + 1)))
    shape.name = "table"
    table = shape.table
    table.columns[0].width = Inches(width * 0.36)
    table.columns[1].width = Inches(width * 0.64)
    tbl_pr = shape._element.find(".//" + qn("a:tblPr"))
    if tbl_pr is not None:                        # PowerPoint'ning tayyor jadval uslubi o'chadi
        for attr in ("firstRow", "bandRow"):
            tbl_pr.set(attr, "0")
        style = tbl_pr.find(qn("a:tableStyleId"))
        if style is not None:
            tbl_pr.remove(style)
    data = [columns] + rows
    for r, values in enumerate(data):
        table.rows[r].height = Inches(row_h)
        for c, value in enumerate(values[:2]):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = _rgb(pal["panel"])
            else:
                cell.fill.fore_color.rgb = _rgb("FFFFFF" if r % 2 else "F3F4F8")
            cell.margin_left = cell.margin_right = Inches(0.16)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _cell_borders(cell, "E2E5EE" if r else pal["panel"])
            frame = cell.text_frame
            frame.word_wrap = True
            para = frame.paragraphs[0]
            run = para.add_run()
            run.text = str(value or "")
            run.font.name = FONT
            run.font.size = Pt(fit_size(run.text, width * (0.36 if c == 0 else 0.64) - 0.32, row_h - 0.1,
                                        15 if r else 15, 10.5))
            run.font.bold = r == 0 or c == 0
            # Jadval qatorlari har palitrada oq/och kulrang — matni to'q (to'q "graphite" fonda ham o'qiladi).
            run.font.color.rgb = _rgb(pal["on"] if r == 0 else ("1D2142" if c == 0 else "4B5068"))
    return [(shape, "fade", 250)]


def slide_stats(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    picture = _picture(ctx, s)
    box(slide, "!!panel", 0, 4.5, W, 3.0, pal["panel"])
    _stripe(slide, ctx, 9.4, on_panel=False)
    place(slide, picture, "!!obj", 8.35, 0.3, 4.7, 6.9, valign="bottom")
    width = 7.3 if picture else W - 1.4
    _title(slide, ctx, s["title"], 0.7, 0.55, width, 1.0, on_panel=False)
    steps = []
    if s.get("text"):
        steps.append((text(slide, "stats_text", 0.7, 1.65, width, 2.5, s["text"], size=18, smallest=13,
                           color=pal["muted"], line=1.2), "float", 150))
    stats = s.get("stats") or []
    cell = (8.0 if picture else W - 1.4) / max(1, len(stats))
    for i, item in enumerate(stats):
        x = 0.7 + i * cell
        value = (item["value"] + (" " + item["unit"] if item.get("unit") else "")).strip()
        steps.append((text(slide, f"st_val_{i}", x, 4.75, cell - 0.25, 0.95, value, size=44, smallest=26,
                           bold=True, color=pal["on"], line=1.0), "zoom", 300 + i * 260))
        steps.append((text(slide, f"st_lab_{i}", x, 5.72, cell - 0.25, 0.4, item.get("label", ""), size=15,
                           smallest=12, bold=True, color=pal["on"]), "float", 400 + i * 260))
        if item.get("text"):
            steps.append((text(slide, f"st_txt_{i}", x, 6.12, cell - 0.3, 1.0, item["text"], size=12,
                               smallest=10, color=pal["on_muted"], line=1.15), "float", 480 + i * 260))
    return steps


def slide_focus(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    picture = _picture(ctx, s)
    box(slide, "!!panel", -1.7, -1.0, 9.5, 9.5, pal["panel"], round_=0.5)
    _stripe(slide, ctx, 8.1, on_panel=False)
    place(slide, picture, "!!obj", -0.5, 0.2, 8.1, 7.3)
    x0 = 8.15 if picture else 1.0
    width = W - x0 - 0.6
    _title(slide, ctx, s["title"], x0, 1.1, width, 1.5, on_panel=False, size=32)
    steps = []
    if s.get("lead"):
        steps.append((text(slide, "focus_lead", x0, 2.75, width, 1.2, s["lead"], size=19, smallest=14, bold=True,
                           color=pal["accent"], line=1.15), "float", 200))
    if s.get("text"):
        steps.append((text(slide, "focus_text", x0, 4.0, width, 2.7, s["text"], size=15, smallest=11.5,
                           color=pal["muted"], line=1.25), "float", 420))
    return steps


def slide_timeline(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    picture = _picture(ctx, s)
    box(slide, "!!panel", 0, 0, W, 3.35, pal["panel"])
    _stripe(slide, ctx, 6.3, on_panel=True)
    place(slide, picture, "!!obj", 9.05, 0.25, 3.9, 4.35, valign="bottom")
    _title(slide, ctx, s["title"], 0.7, 0.75, 8.0 if picture else W - 1.4, 1.6, on_panel=True, size=32)
    steps_data = s.get("steps") or []
    n = max(1, len(steps_data))
    left, right = 0.8, W - 0.8
    axis_y = 4.75
    axis = box(slide, "tl_axis", left, axis_y - 0.025, right - left, 0.05, pal["accent"], alpha=0.35)
    steps = [(axis, "fade", 100)]
    span = (right - left) / n
    for i, item in enumerate(steps_data):
        x = left + span * i + 0.05
        dot = box(slide, f"tl_dot_{i}", x, axis_y - 0.15, 0.3, 0.3, pal["bg"], round_=0.5, line=pal["accent"],
                  line_pt=3)
        when = text(slide, f"tl_when_{i}", x, axis_y + 0.3, span - 0.3, 0.5, item.get("when", ""), size=19,
                    smallest=13, bold=True, color=pal["accent"])
        body = text(slide, f"tl_text_{i}", x, axis_y + 0.82, span - 0.3, 1.75, item.get("text", ""), size=13,
                    smallest=10.5, color=pal["muted"], line=1.18)
        steps += [(dot, "zoom", 200 + i * 260), (when, "float", 260 + i * 260), (body, "float", 320 + i * 260)]
    return steps


def slide_conclusion(slide, s: Dict, ctx: Ctx, index: int) -> List:
    pal = ctx.pal
    box(slide, "!!panel", 0, 0, W, H, pal["panel"])
    _stripe(slide, ctx, 7.5, on_panel=True)
    picture = _picture(ctx, s)
    place(slide, picture, "!!obj", 8.7, 1.3, 4.3, 5.8)
    width = 7.4 if picture else W - 1.6
    _title(slide, ctx, s["title"], 0.8, 0.75, width, 1.3, on_panel=True, size=40, smallest=28)
    box(slide, "!!dot", 0.82, 2.08, 0.9, 0.08, pal["pill"], round_=0.5)
    steps = []
    for i, point in enumerate([p for p in s.get("points") or [] if p]):
        top = 2.55 + i * 1.35
        check = box(slide, f"cn_check_{i}", 0.8, top + 0.02, 0.42, 0.42, pal["pill"], round_=0.5)
        frame = check.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = "✓"
        run.font.name, run.font.size, run.font.bold = FONT, Pt(14), True
        run.font.color.rgb = _rgb(pal["pill_ink"])
        body = text(slide, f"cn_text_{i}", 1.45, top, width - 0.7, 1.2, point, size=19, smallest=14,
                    color=pal["on"], line=1.15)
        steps += [(check, "zoom", 200 + i * 300), (body, "float", 260 + i * 300)]
    return steps


RENDER = {"cover": slide_cover, "callouts": slide_callouts, "points": slide_points, "table": slide_table,
          "stats": slide_stats, "focus": slide_focus, "timeline": slide_timeline,
          "conclusion": slide_conclusion}
# Sahifa raqami turgan pastki o'ng burchak rangli panel ustida bo'ladigan joylashuvlar.
ON_PANEL = {"cover", "callouts", "conclusion", "table", "stats"}


def build(deck: Dict, pictures: Dict[str, Picture], anchors: Dict[int, list], *, author: str = "",
          year: str = "", out_path: Optional[str] = None, morph: bool = True) -> str:
    """PPTX ni yig'adi va yo'lini qaytaradi."""
    pal = PALETTES.get(deck.get("palette") or "", PALETTES["indigo"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    slides = deck["slides"]
    ctx = Ctx(pal, pictures, anchors, author, year, len(slides))
    for index, data in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(pal["bg"])
        render = RENDER.get(data["layout"], slide_points)
        try:
            steps = render(slide, data, ctx, index)
        except Exception:
            log.exception("%d-slayd chizilmadi (%s)", index + 1, data["layout"])
            raise
        if index:
            _page(slide, ctx, index, data["layout"] in ON_PANEL)
        if morph and index:
            add_morph(slide)
        add_entrances(slide, steps)
    out_path = out_path or os.path.join("temp", f"pro3d_{uuid.uuid4().hex[:10]}.pptx")
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    prs.save(out_path)
    return out_path
