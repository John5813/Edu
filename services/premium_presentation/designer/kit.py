"""Vektor dizaynerning chizish asboblari: 1920x1080 px koordinatada PowerPoint'ning haqiqiy shakllari.

Hammasi tahrirlanadi: freeform va tayyor geometriya, chiziqli/radial gradient, soya, shaffoflik, rasmni
shakl ichiga kesish, bo'yalgan ikonka. HTML orqali o'tmaydi — shuning uchun to'rtburchakdan boshqa har
qanday shakl, haqiqiy soya va shaffoflik chiqadi.
"""
import io
import math
import os
import tempfile
from typing import List, Optional, Sequence, Tuple

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W_PX, H_PX = 1920, 1080
EMU_W = 12192000
PX = EMU_W / W_PX

# Telefon va WPS da ham bor shriftlar: sarlavha qalin Arial (jurnal uslubida Times New Roman).
SANS = "Arial"
SERIF = "Times New Roman"

LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE, BOTTOM = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM
SHAPE = MSO_SHAPE


def px(v: float) -> Emu:
    return Emu(int(round(v * PX)))


# ───────────────────────────────────────────── ranglar

def hex_rgb(value: str) -> Tuple[int, int, int]:
    value = (value or "000000").lstrip("#")
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def mix(a: str, b: str, t: float) -> str:
    """a dan b ga t ulushda aralashma (t=0 — a, t=1 — b)."""
    (r1, g1, b1), (r2, g2, b2) = hex_rgb(a), hex_rgb(b)
    return "%02X%02X%02X" % (round(r1 + (r2 - r1) * t), round(g1 + (g2 - g1) * t), round(b1 + (b2 - b1) * t))


def luminance(value: str) -> float:
    r, g, b = (c / 255 for c in hex_rgb(value))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


# ───────────────────────────────────────────── to'ldirish, chiziq, effekt

def _color(parent, hex_: str, alpha: float = 100):
    c = etree.SubElement(parent, qn("a:srgbClr"))
    c.set("val", hex_.lstrip("#").upper())
    if alpha < 100:
        a = etree.SubElement(c, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    return c


def _set_fill(shape, el) -> None:
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:blipFill", "a:pattFill"):
        for old in spPr.findall(qn(tag)):
            spPr.remove(old)
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    geom.addnext(el)


def fill(shape, hex_: str, alpha: float = 100):
    el = etree.Element(qn("a:solidFill"))
    _color(el, hex_, alpha)
    _set_fill(shape, el)
    return shape


def nofill(shape):
    _set_fill(shape, etree.Element(qn("a:noFill")))
    return shape


def grad(shape, stops: Sequence, angle: float = 90, radial: bool = False, focus=(50, 50)):
    """stops: [(pos 0-100, rang[, shaffoflik 0-100])]. angle — chiziqli yo'nalish (0 — chapdan o'ngga)."""
    g = etree.Element(qn("a:gradFill"))
    g.set("rotWithShape", "1")
    lst = etree.SubElement(g, qn("a:gsLst"))
    for stop in stops:
        gs = etree.SubElement(lst, qn("a:gs"))
        gs.set("pos", str(int(stop[0] * 1000)))
        _color(gs, stop[1], stop[2] if len(stop) > 2 else 100)
    if radial:
        p = etree.SubElement(g, qn("a:path"))
        p.set("path", "circle")
        r = etree.SubElement(p, qn("a:fillToRect"))
        fx, fy = focus
        r.set("l", str(fx * 1000)); r.set("t", str(fy * 1000))
        r.set("r", str((100 - fx) * 1000)); r.set("b", str((100 - fy) * 1000))
    else:
        lin = etree.SubElement(g, qn("a:lin"))
        lin.set("ang", str(int((angle % 360) * 60000)))
        lin.set("scaled", "0")
    _set_fill(shape, g)
    return shape


def line(shape, hex_: Optional[str] = None, width: float = 2, alpha: float = 100, dash: str = ""):
    if hex_ is None:
        shape.line.fill.background()
        return shape
    shape.line.width = px(width)
    shape.line.color.rgb = RGBColor.from_string(hex_.upper())
    ln = shape._element.spPr.find(qn("a:ln"))
    if alpha < 100:
        clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", dash)
    return shape


def _effects(shape):
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    return eff


def shadow(shape, blur=30, dist=10, angle=90, alpha=40, color="000000"):
    s = etree.SubElement(_effects(shape), qn("a:outerShdw"))
    s.set("blurRad", str(int(blur * PX))); s.set("dist", str(int(dist * PX)))
    s.set("dir", str(int((angle % 360) * 60000))); s.set("algn", "ctr"); s.set("rotWithShape", "0")
    _color(s, color, alpha)
    return shape


def inner_shadow(shape, blur=20, dist=6, angle=90, alpha=45, color="000000"):
    s = etree.SubElement(_effects(shape), qn("a:innerShdw"))
    s.set("blurRad", str(int(blur * PX))); s.set("dist", str(int(dist * PX)))
    s.set("dir", str(int((angle % 360) * 60000)))
    _color(s, color, alpha)
    return shape


def glow(shape, size=20, color="FFFFFF", alpha=40):
    g = etree.SubElement(_effects(shape), qn("a:glow"))
    g.set("rad", str(int(size * PX)))
    _color(g, color, alpha)
    return shape


# ───────────────────────────────────────────── shakllar

def box(s, x, y, w, h, kind=MSO_SHAPE.RECTANGLE, rot: float = 0, adj: Sequence[float] = ()):
    sh = s.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    line(sh)
    if rot:
        sh.rotation = rot
    for i, v in enumerate(adj):
        sh.adjustments[i] = v
    return sh


def rect(s, x, y, w, h, **k):
    return box(s, x, y, w, h, MSO_SHAPE.RECTANGLE, **k)


def oval(s, cx, cy, r, ry=None):
    ry = ry or r
    return box(s, cx - r, cy - ry, 2 * r, 2 * ry, MSO_SHAPE.OVAL)


def poly(s, pts: List[Tuple[float, float]], close: bool = True):
    fb = s.shapes.build_freeform(px(pts[0][0]), px(pts[0][1]), scale=1.0)
    fb.add_line_segments([(px(x), px(y)) for x, y in pts[1:]], close=close)
    sh = fb.convert_to_shape()
    line(sh)
    return sh


def arc_pts(cx, cy, r, a0, a1, steps: int = 48):
    """Yoy nuqtalari (gradus, 0 — o'ngga, soat strelkasiga teskari; y pastga)."""
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / steps)),
             cy - r * math.sin(math.radians(a0 + (a1 - a0) * i / steps))) for i in range(steps + 1)]


def ring_sector(s, cx, cy, r0, r1, a0, a1):
    return poly(s, arc_pts(cx, cy, r1, a0, a1) + arc_pts(cx, cy, r0, a1, a0))


def rotate(pts, angle, ox=0.0, oy=0.0):
    c, sn = math.cos(math.radians(angle)), math.sin(math.radians(angle))
    return [(ox + x * c - y * sn, oy + x * sn + y * c) for x, y in pts]


def ribbon(points, width):
    """Markaziy chiziq bo'ylab qalinligi `width` bo'lgan lenta (ko'pburchak)."""
    left, right = [], []
    for i, (x, y) in enumerate(points):
        a = points[max(i - 1, 0)]
        b = points[min(i + 1, len(points) - 1)]
        dx, dy = b[0] - a[0], b[1] - a[1]
        n = math.hypot(dx, dy) or 1
        nx, ny = -dy / n * width / 2, dx / n * width / 2
        left.append((x + nx, y + ny))
        right.append((x - nx, y - ny))
    return left + right[::-1]


# ───────────────────────────────────────────── matn

def _estimate(paras, w: float, scale: float) -> float:
    """Matn bloki balandligini taxmin qiladi (px): Arial harf kengligi ~0,52 em, katta harf ~0,66 em."""
    total = 0.0
    for para, spacing in paras:
        size = max(r[1] for r in para) * scale
        chars = 0.0
        for run in para:
            t = run[0]
            upper = sum(1 for ch in t if ch.isupper())
            em = 0.56 if (len(run) > 3 and run[3]) else 0.52
            chars += (len(t) + upper * 0.25) * em * run[1] * scale / size
        per_line = max(w / (size * 1.0), 1)
        lines = max(1, math.ceil(chars * 1.08 / per_line)) if chars else 1
        total += lines * size * spacing + 0.0
    return total


def _widest_word(paras) -> float:
    """Eng uzun so'zning taxminiy kengligi (px, telefonning kengroq shrifti uchun zaxira bilan)."""
    widest = 0.0
    for para in paras:
        for run in para:
            em = 0.62 if (len(run) > 3 and run[3]) else 0.56
            for word in run[0].split():
                upper = sum(1 for ch in word if ch.isupper())
                widest = max(widest, (len(word) + upper * 0.25) * em * run[1])
    return widest


BOOST = 4          # px (= 2 pt): matn va yorliqlarga qo'shimcha o'lcham
MAX_GROW = 1.3     # joy bo'lsa matn shuncha martagacha kattalashadi


def text(s, x, y, w, h, paras, align=LEFT, anchor=TOP, gap: float = 0, line_spacing: float = 1.15,
         fit: bool = True, min_scale: float = 0.55, grow: bool = False):
    """paras: [[(matn, px o'lcham, rang, qalin, shrift, harf oralig'i, kursiv), ...], ...].

    `fit` — matn qutiga sig'maguncha hamma o'lcham birdek kichraytiriladi (taxminiy o'lchov bilan):
    AI yozgan matn uzunligi oldindan aniq emas, quti esa dizaynda qat'iy.
    """
    paras = [[r for r in para if r and r[0]] for para in paras]
    paras = [p for p in paras if p]
    if not paras:
        return None
    # Oddiy matn (sarlavha emas) +2 pt: mijozlar telefonda shriftni mayda deb topdi.
    paras = [[(r[0], r[1] + BOOST if r[1] <= 44 else r[1]) + tuple(r[2:]) for r in para] for para in paras]
    scale = 1.0
    if fit:
        measured = [(p, line_spacing * 1.18) for p in paras]
        widest = _widest_word(paras)
        # Matn kam bo'lsa — kattalashadi (joy bo'sh qolmasin), ko'p bo'lsa — kichrayadi.
        if grow and max(r[1] for para in paras for r in para) <= 48:   # sarlavhalar kattalashmaydi
            scale = MAX_GROW
        while scale > min_scale:
            need = _estimate(measured, w, scale) + gap * scale * (len(paras) - 1)
            # Eng uzun so'z ham bir qatorga sig'sin — aks holda PowerPoint uni o'rtasidan bo'ladi.
            if need <= h and widest * scale <= w:
                break
            scale -= 0.04
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    # "Matnga qarab kattalash" o'chiq: WPS uni ko'rib matnni o'ramasdan qutini o'ngga cho'zardi
    # va izoh slayd chetidan chiqib ketardi.
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if gap and i < len(paras) - 1:
            p.space_after = Pt(gap * scale / 2)
        for run in para:
            t, size, hex_ = run[0], run[1], run[2]
            bold = run[3] if len(run) > 3 else False
            font = run[4] if len(run) > 4 else SANS
            spc = run[5] if len(run) > 5 else 0
            italic = run[6] if len(run) > 6 else False
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(max(size * scale, 9) / 2)
            f.bold = bool(bold)
            f.italic = bool(italic)
            f.name = font
            f.color.rgb = RGBColor.from_string(hex_.upper())
            if spc:
                r._r.get_or_add_rPr().set("spc", str(int(spc)))
    return tb


# ───────────────────────────────────────────── rasm va ikonka

class Photo:
    """Slayd rasmi (bayt). `crop(w, h)` — markazdan kesilgan JPEG fayl yo'li (vaqtinchalik)."""

    def __init__(self, raw: bytes):
        from PIL import Image

        self.image = Image.open(io.BytesIO(raw)).convert("RGB")
        self._files: List[str] = []

    def crop(self, w: float, h: float) -> str:
        im = self.image
        sw, sh = im.size
        target = w / h
        if sw / sh > target:
            nw = int(sh * target)
            im = im.crop(((sw - nw) // 2, 0, (sw - nw) // 2 + nw, sh))
        else:
            nh = int(sw / target)
            im = im.crop((0, (sh - nh) // 2, sw, (sh - nh) // 2 + nh))
        fd, path = tempfile.mkstemp(suffix=".jpg", prefix="dz_")
        os.close(fd)
        im.save(path, quality=88)
        self._files.append(path)
        return path

    def cleanup(self) -> None:
        for path in self._files:
            try:
                os.remove(path)
            except OSError:
                pass
        self._files.clear()


def photo(s, pic: Photo, x, y, w, h, geom: str = "", region=None):
    """Rasm. geom — shakl ('hexagon' ...). region=(X, Y, RW, RH): rasm shu katta maydonni egallaydi deb
    hisoblanib, (x, y, w, h) bo'lagi kesib olinadi (mozaika uchun)."""
    if region:
        rx, ry, rw, rh = region
        path = pic.crop(rw, rh)
    else:
        path = pic.crop(w, h)
    shape = s.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    if region:
        shape.crop_left = (x - rx) / rw
        shape.crop_top = (y - ry) / rh
        shape.crop_right = (rx + rw - x - w) / rw
        shape.crop_bottom = (ry + rh - y - h) / rh
    if geom:
        shape._element.spPr.find(qn("a:prstGeom")).set("prst", geom)
    return shape


def icon(s, name: str, cx, cy, size, color: str):
    from .. import icon_render

    path = icon_render.resolve(name) or os.path.join(icon_render.ICONS_DIR, "idea.png")
    tinted = icon_render.tinted(path, "#" + color) or path
    return s.shapes.add_picture(tinted, px(cx - size / 2), px(cy - size / 2), px(size), px(size))
