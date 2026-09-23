"""HTML slaydni brauzerda suratga oladi va PPTX ga yig'adi.

Slaydning ko'rinishi endi HTML/CSS bilan belgilanadi, ya'ni uni
python-pptx bilan qayta chizish shart emas: brauzer nima ko'rsatsa,
PowerPointda ham aynan o'sha turadi. Har slayd 1920×1080 PNG bo'lib,
13.333×7.5 dyuymli varaqni to'liq egallaydi.

Brauzer bitta marta ishga tushiriladi va slaydlar navbat bilan
olinadi: server kichik (1 vCPU / 2 GB), bir vaqtda ikkita Chromium
oynasi ochilsa xotira yetmaydi.
"""

import glob
import logging
import os
import re
import tempfile
from typing import List

log = logging.getLogger("html_render")

SLIDE_W_PX = 1920
SLIDE_H_PX = 1080
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

# Chromium root ostida ishlaganda sandbox bilan darhol yopiladi.
_LAUNCH_ARGS = [
    "--no-sandbox",
    "--disable-setuid-sandbox",
    "--disable-dev-shm-usage",
    "--disable-gpu",
    "--disable-software-rasterizer",
    "--hide-scrollbars",
]

# Sahifa tashqi fayl so'ramaydi, shuning uchun "networkidle" ni kutish
# shart emas — u ba'zan bekorga sekundlar yeydi.
_WAIT_UNTIL = "load"
_TIMEOUT_MS = 30000


# Playwright brauzerni O'Z versiyasi bo'yicha qidiradi: kutubxona
# yangilansa, u yangi raqamli papkani kutadi va serverdagi eski brauzerni
# ko'rmaydi ("Executable doesn't exist at .../chromium_headless_shell-1243").
# Shuning uchun brauzer diskdan o'zimiz ham qidiramiz.
#
# Playwright brauzerlarni shu yerda saqlaydi:
#   $PLAYWRIGHT_BROWSERS_PATH yoki ~/.cache/ms-playwright
# ichida chromium-<raqam>/ va chromium_headless_shell-<raqam>/ papkalari.
_BROWSER_GLOBS = (
    "chromium-*/chrome-linux/chrome",
    "chromium_headless_shell-*/chrome-headless-shell-linux64/chrome-headless-shell",
)

# Tizimga o'rnatilgan brauzerlar — Playwright papkasi umuman bo'lmasa.
_SYSTEM_BROWSERS = (
    "/opt/pw-browsers/chromium",
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
    "/usr/bin/google-chrome",
    "/usr/bin/google-chrome-stable",
)

_INSTALL_HINT = ("Serverda brauzer topilmadi. Bir marta shuni bajaring:\n"
                 "  venv/bin/playwright install --with-deps chromium")


def _browser_roots() -> List[str]:
    roots = []
    configured = os.getenv("PLAYWRIGHT_BROWSERS_PATH", "").strip()
    if configured and configured != "0":
        roots.append(configured)
    roots.append(os.path.expanduser("~/.cache/ms-playwright"))
    # Bot boshqa foydalanuvchi ostida ishga tushirilgan bo'lishi mumkin,
    # brauzer esa root ostida o'rnatilgan bo'ladi.
    roots.append("/root/.cache/ms-playwright")
    return [root for root in roots if os.path.isdir(root)]


def _executable() -> str:
    """Diskdagi eng yangi Chromium — Playwright topolmaganda ishlatiladi."""
    override = os.getenv("PREMIUM_CHROMIUM_PATH", "").strip()
    if override and os.path.exists(override):
        return override

    found = []
    for root in _browser_roots():
        for pattern in _BROWSER_GLOBS:
            found.extend(glob.glob(os.path.join(root, pattern)))
    if found:
        # Papka nomidagi raqam — build raqami; eng kattasi eng yangisi.
        def build_number(path: str) -> int:
            match = re.search(r"-(\d+)[/\\]", path)
            return int(match.group(1)) if match else 0

        return max(found, key=build_number)

    for path in _SYSTEM_BROWSERS:
        if os.path.exists(path):
            return path
    return ""


def available() -> bool:
    """Playwright ham, brauzer ham shu serverda bormi.

    Ilgari bu faqat kutubxona import bo'lishini tekshirardi va brauzer
    yo'qligi mijoz to'lovdan keyin bilinardi.
    """
    try:
        from playwright.sync_api import sync_playwright  # noqa: F401
    except ImportError:
        return False
    return bool(_executable())


def _launch(playwright):
    """Brauzerni ishga tushiradi — uch usulni ketma-ket sinab.

    1. Playwright o'zi bilgan brauzer (odatdagi holat).
    2. To'liq Chromium: yangi Playwright `headless=True` uchun alohida
       "headless shell" ni kutadi, serverda esa ko'pincha faqat to'liq
       Chromium o'rnatilgan bo'ladi.
    3. Diskdan topilgan brauzer yo'li.
    """
    attempts = []
    try:
        return playwright.chromium.launch(headless=True, args=_LAUNCH_ARGS)
    except Exception as exc:
        attempts.append(str(exc).splitlines()[0][:160])

    try:
        return playwright.chromium.launch(
            headless=True, args=_LAUNCH_ARGS, channel="chromium")
    except Exception as exc:
        attempts.append(str(exc).splitlines()[0][:160])

    path = _executable()
    if not path:
        raise RuntimeError(f"{_INSTALL_HINT}\n\n" + "\n".join(attempts))

    log.warning("Playwright brauzerni topmadi (%s) — %s ishlatiladi",
                attempts[0], path)
    return playwright.chromium.launch(
        headless=True, args=_LAUNCH_ARGS, executable_path=path)


def shoot(html_slides: List[str], out_dir: str = "temp") -> List[str]:
    """Har HTML hujjatni PNG qilib saqlaydi va yo'llarini qaytaradi.

    Bitta slayd chizilmasa, qolganlari baribir chiqadi — taqdimot
    bitta xato tufayli butunlay yo'qolmaydi.
    """
    from playwright.sync_api import sync_playwright

    os.makedirs(out_dir, exist_ok=True)
    paths: List[str] = []

    with sync_playwright() as playwright:
        browser = _launch(playwright)
        try:
            context = browser.new_context(
                viewport={"width": SLIDE_W_PX, "height": SLIDE_H_PX},
                device_scale_factor=1,
            )
            try:
                page = context.new_page()
                for index, html in enumerate(html_slides, 1):
                    path = os.path.join(
                        out_dir, f"slide_{os.getpid()}_{index:02d}.png")
                    try:
                        page.set_content(html, wait_until=_WAIT_UNTIL,
                                         timeout=_TIMEOUT_MS)
                        page.screenshot(path=path, type="png",
                                        clip={"x": 0, "y": 0,
                                              "width": SLIDE_W_PX,
                                              "height": SLIDE_H_PX})
                        paths.append(path)
                    except Exception as exc:
                        log.error("%d-slayd chizilmadi: %s", index, exc)
            finally:
                context.close()
        finally:
            browser.close()

    log.info("Slaydlar suratga olindi: %d/%d", len(paths), len(html_slides))
    return paths


def build_pptx(image_paths: List[str], out_dir: str = "temp",
               name: str = "taqdimot") -> str:
    """PNG larni to'liq varaqni egallagan slaydlarga aylantiradi."""
    from pptx import Presentation
    from pptx.util import Inches

    if not image_paths:
        raise RuntimeError("Birorta slayd suratga olinmadi")

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W_IN)
    presentation.slide_height = Inches(SLIDE_H_IN)
    blank = presentation.slide_layouts[6]

    for path in image_paths:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            path, left=0, top=0,
            width=presentation.slide_width,
            height=presentation.slide_height)

    os.makedirs(out_dir, exist_ok=True)
    handle, out_path = tempfile.mkstemp(
        prefix=f"{name}_", suffix=".pptx", dir=out_dir)
    os.close(handle)
    presentation.save(out_path)
    return out_path


def render(html_slides: List[str], out_dir: str = "temp",
           name: str = "taqdimot") -> str:
    """HTML → PNG → PPTX. Vaqtinchalik rasmlar o'chiriladi."""
    images = shoot(html_slides, out_dir)
    try:
        return build_pptx(images, out_dir, name)
    finally:
        for path in images:
            try:
                os.remove(path)
            except OSError:
                pass
