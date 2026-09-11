import logging
import os
import uuid

from pptx import Presentation

from services.project_work import palettes, variety

from . import config
from .image_client import generate_image
from .layouts import SLIDE_H, SLIDE_W, render_canvas
from .models import Brief, Slide, VisualElement

log = logging.getLogger("renderer")


def _soft(hex_colour: str, amount: float = 0.86) -> str:
    """Rangni oqqa yaqinlashtiradi — matn ostidagi yumshoq fon uchun."""
    raw = (hex_colour or "").lstrip("#")
    if len(raw) != 6:
        return "F1F3F6"
    try:
        channels = [int(raw[i:i + 2], 16) for i in (0, 2, 4)]
    except ValueError:
        return "F1F3F6"
    return "".join(f"{int(round(c + (255 - c) * amount)):02X}" for c in channels)


def _shorten(text: str, limit: int) -> str:
    clean = " ".join((text or "").split())
    if len(clean) <= limit:
        return clean
    cut = clean[:limit].rsplit(" ", 1)[0]
    return cut.rstrip(",;:") + "…"


def _replace_with_message(element, slide: Slide, brief: Brief) -> None:
    """Yaratilmagan rasm o'rniga slaydning asosiy fikrini qo'yadi.

    Ilgari bu yerga shunchaki to'q rangli to'rtburchak chizilardi. Rasm
    generatsiyasi ishlamay qolganda slaydning yarmini ma'nosiz rangli
    maydon egallab olar va taqdimot "ortiqcha narsalar" bilan to'lib
    ketardi. Endi o'sha joy slaydning asosiy xabarini ko'taradi: bu
    dizayn nuqtai nazaridan ataylab qo'yilgandek ko'rinadi.
    """
    fill = _soft(element.fill or (brief.theme.accent if brief.theme else "2A78D6"))
    element.type = "rect"
    element.radius = True
    element.fill = fill
    element.prompt = None

    width = element.w or 5.0
    height = element.h or 4.0
    message = _shorten(slide.key_text or slide.title or "", 240)
    if not message or width < 1.8 or height < 1.0:
        return

    pad_x = min(0.45, width * 0.12)
    pad_y = min(0.45, height * 0.12)
    inner_w = width - 2 * pad_x
    inner_h = height - 2 * pad_y

    # Shrift qutiga qarab: uzun matn kichik quti ichiga sig'sin.
    per_line = max(int(inner_w / 0.11), 12)
    lines = max(1, -(-len(message) // per_line))
    size = 18.0 if lines * 0.34 <= inner_h else max(13.0, inner_h / max(lines, 1) / 0.34 * 18.0)
    size = max(13.0, min(size, 20.0))

    slide.canvas.elements.append(VisualElement(
        type="text",
        x=element.x + pad_x,
        y=element.y + pad_y,
        w=inner_w,
        h=inner_h,
        text=message,
        size=size,
        italic=True,
        align="left",
        color=palettes.on_fill(fill),
    ))


def build_presentation(brief: Brief, user_id: int | None = None) -> str:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]  # to'liq bo'sh layout

    # Butun taqdimot bo'ylab bitta ikonka ikki marta ishlatilmasin.
    used_icons: set[str] = set()

    # Rang sxemasi taqdimot boshida bir marta tanlanadi: PowerPoint'ning
    # standart ranglari har taqdimotni bir xil qilib qo'yardi.
    palette = variety.choose_palette((brief.topic, brief.theme.primary), user_id)

    wanted = failed = 0

    for s in brief.slides:
        slide = prs.slides.add_slide(blank_layout)

        # image elementlari uchun rasm generatsiyasi
        image_paths: dict[int, str] = {}
        for el in list(s.canvas.elements):
            if el.type == "image" and el.prompt:
                wanted += 1
                path = generate_image(el.prompt)
                if path:
                    image_paths[id(el)] = path
                else:
                    # Rasmni jimgina tashlab ketish slaydni yarim bo'sh qoldiradi
                    # va aynan shu sababdan taqdimotlar rasmsiz chiqib ketgan edi.
                    log.error(
                        "Slayd %s: rasm yaratilmadi (TOGETHER_API_KEY bormi?) — "
                        "o'rniga asosiy fikr paneli qo'yiladi", s.index
                    )
                    failed += 1
                    _replace_with_message(el, s, brief)

        render_canvas(slide, s, image_paths, used_icons, palette)

    if wanted and failed == wanted:
        # Bitta ham rasm chiqmasa sabab deyarli har doim bitta bo'ladi:
        # kalit yo'q, krediti tugagan yoki model hisobda mavjud emas.
        # Har slayd uchun alohida yozuv bu xulosani ko'mib yuborardi.
        log.error(
            "HECH BIR RASM YARATILMADI (%s ta so'rovning hammasi muvaffaqiyatsiz). "
            "TOGETHER_API_KEY, hisobdagi kredit va PREMIUM_TOGETHER_IMAGE_MODEL "
            "(%s) ni tekshiring — yuqoridagi HTTP javob kodi sababni ko'rsatadi.",
            wanted, config.TOGETHER_IMAGE_MODEL,
        )
    elif failed:
        log.warning("Rasmlar: %s tadan %s tasi yaratilmadi", wanted, failed)

    os.makedirs(config.WORK_DIR, exist_ok=True)
    out_path = os.path.join(config.WORK_DIR, f"ppt_{uuid.uuid4().hex[:10]}.pptx")
    prs.save(out_path)
    log.info("PPTX saqlandi: %s", out_path)
    return out_path
