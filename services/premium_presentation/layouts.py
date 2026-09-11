"""
Kanvas renderer — AI tomonidan tasvirlangan elementlarni python-pptx orqali chizadi.
"""
import logging
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE

log = logging.getLogger("layouts")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}

# python-pptx bu turlarni native qo'llaydi, ya'ni diagramma PowerPoint'da
# tahrirlanadigan bo'lib qoladi — rasm emas.
CHART_TYPE_MAP = {
    "bar":     XL_CHART_TYPE.BAR_CLUSTERED,
    "column":  XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line":    XL_CHART_TYPE.LINE_MARKERS,
    "area":    XL_CHART_TYPE.AREA,
    "pie":     XL_CHART_TYPE.PIE,
    "donut":   XL_CHART_TYPE.DOUGHNUT,
    "radar":   XL_CHART_TYPE.RADAR_MARKERS,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}

# Bitta qiymatlar qatorini rang bo'yicha ajratadigan turlar: bu yerda rang
# har bo'lakka, qolganlarida esa butun qatorga beriladi.
_PER_POINT = {"pie", "donut"}


def _hex(hex_str: str) -> RGBColor:
    h = (hex_str or "000000").lstrip("#").strip()
    if len(h) == 3:
        h = h[0]*2 + h[1]*2 + h[2]*2
    if len(h) != 6:
        return RGBColor(0, 0, 0)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _clamp(val, lo, hi):
    return max(lo, min(hi, val))


def render_canvas(slide, s, image_paths: dict, used_icons: set | None = None,
                  palette=None):
    used_icons = used_icons if used_icons is not None else set()
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _hex(s.canvas.background)

    for el in s.canvas.elements:
        try:
            if el.type == "rect":
                _draw_rect(slide, el)
            elif el.type == "text":
                _draw_text(slide, el)
            elif el.type == "circle":
                _draw_circle(slide, el)
            elif el.type == "image":
                img_path = image_paths.get(id(el))
                if img_path:
                    _draw_image(slide, el, img_path)
            elif el.type == "chart":
                _draw_chart(slide, el, palette)
            elif el.type == "kpi":
                _draw_kpi(slide, el, s)
            elif el.type == "icon":
                _draw_icon(slide, el, used_icons)
            elif el.type == "infographic":
                # Pipeline uni ibtidoiy elementlarga yoyishi kerak edi; bu yerga
                # yetib kelgani — yoyish o'tkazib yuborilganini bildiradi.
                log.warning("Yoyilmagan infografika (slayd %s) chizilmadi", s.index)
        except Exception as exc:
            log.warning("Element chizishda xato (%s, slayd %s): %s", el.type, s.index, exc)


# ─────────────────────────────────────────────────────────── rect

def _draw_rect(slide, el):
    w = _clamp(el.w or 1.0, 0.05, 13.333)
    h = _clamp(el.h or 1.0, 0.05, 7.5)
    x = _clamp(el.x, -0.5, 13.333)
    y = _clamp(el.y, -0.5, 7.5)

    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if el.radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if el.fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _hex(el.fill)
    else:
        shp.fill.background()
    shp.line.fill.background()


# ─────────────────────────────────────────────────────────── kpi

def _draw_kpi(slide, el, s):
    """Ko'rsatkich kartochkasi: katta raqam va uning ostida izoh.

    Bitta raqam diagramma talab qilmaydi — u shunchaki katta yozilishi kerak.
    Kartochka native shakllardan quriladi, ya'ni PowerPoint'da tahrirlanadi.
    """
    w = _clamp(el.w or 3.0, 1.2, 13.333)
    h = _clamp(el.h or 1.9, 0.9, 7.5)
    x = _clamp(el.x, 0.0, 13.333 - w)
    y = _clamp(el.y, 0.0, 7.5 - h)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    card.shadow.inherit = False
    card.fill.solid()
    card.fill.fore_color.rgb = _hex(el.fill or "F4F6F9")
    card.line.fill.background()

    box = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.10),
                                   Inches(w - 0.24), Inches(h - 0.20))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    number = frame.paragraphs[0]
    number.alignment = PP_ALIGN.CENTER
    run = number.add_run()
    run.text = str(el.value or "")
    run.font.size = Pt(_clamp((el.size or 34), 20, 54))
    run.font.bold = True
    run.font.name = el.font or "Calibri"
    run.font.color.rgb = _hex(el.color or "1B2A4A")

    if el.label:
        caption = frame.add_paragraph()
        caption.alignment = PP_ALIGN.CENTER
        crun = caption.add_run()
        crun.text = str(el.label)
        crun.font.size = Pt(12)
        crun.font.name = el.font or "Calibri"
        crun.font.color.rgb = _hex("52514E")


# ─────────────────────────────────────────────────────────── text

def _draw_text(slide, el):
    w = _clamp(el.w or 5.0, 0.5, 13.333)
    h = _clamp(el.h or 1.0, 0.2, 7.5)
    x = _clamp(el.x, 0.0, 13.0)
    y = _clamp(el.y, 0.0, 7.3)

    # Minimal font o'lchami: sarlavha bo'lmasa kamida 13pt
    font_size = el.size or 14
    if not el.bold and font_size < 13:
        font_size = 13

    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    # Ko'p qatorli matnni \n bo'yicha ajratib, har birini alohida paragraf qilamiz
    lines = (el.text or "").split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = ALIGN_MAP.get(el.align or "left", PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(_clamp(font_size, 8, 72))
        run.font.bold = el.bold
        run.font.italic = el.italic
        run.font.name = el.font or "Calibri"
        if el.color:
            run.font.color.rgb = _hex(el.color)


# ─────────────────────────────────────────────────────────── circle

def _draw_circle(slide, el):
    # d berilsa — aylana; w/h berilsa — ellips (masalan cycle preseti halqasi).
    if el.d:
        w = h = _clamp(el.d, 0.1, 7.0)
    else:
        w = _clamp(el.w or 1.0, 0.1, 13.333)
        h = _clamp(el.h or 1.0, 0.1, 7.5)
    x = _clamp(el.x, 0.0, 13.333 - w)
    y = _clamp(el.y, 0.0, 7.5 - h)

    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if el.fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _hex(el.fill)
    else:
        shp.fill.background()
    if el.line:
        shp.line.color.rgb = _hex(el.line)
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()


# ─────────────────────────────────────────────────────────── icon

def _draw_icon(slide, el, used_icons: set):
    """Lokal ikonkani berilgan rangda chizadi (kerak bo'lsa fon shakli bilan).

    Rasm generatsiyasi emas: fayl `assets/icons/` dan olinadi va Pillow orqali
    bo'yaladi. Narxi nol, kutish yo'q, uslub butun taqdimotda bir xil.
    """
    from . import icon_render

    size = _clamp(el.w or el.d or 0.6, 0.15, 3.0)
    x = _clamp(el.x, 0.0, 13.333 - size)
    y = _clamp(el.y, 0.0, 7.5 - size)

    if el.shape != "none":
        shape_type = MSO_SHAPE.OVAL if el.shape == "circle" else MSO_SHAPE.ROUNDED_RECTANGLE
        pad = size * 0.3
        holder = slide.shapes.add_shape(
            shape_type, Inches(x - pad), Inches(y - pad),
            Inches(size + 2 * pad), Inches(size + 2 * pad),
        )
        holder.shadow.inherit = False
        holder.fill.solid()
        holder.fill.fore_color.rgb = _hex(el.fill or "2A78D6")
        holder.line.fill.background()

    icon_path = icon_render.resolve(el.icon, el.text or "", used=used_icons)
    if not icon_path:
        log.warning("Ikonka topilmadi: %s / %s", el.icon, (el.text or "")[:40])
        return

    tinted = icon_render.tinted(icon_path, el.color or "FFFFFF")
    if tinted:
        slide.shapes.add_picture(tinted, Inches(x), Inches(y), Inches(size), Inches(size))


# ─────────────────────────────────────────────────────────── image

def _draw_image(slide, el, img_path: str):
    w = _clamp(el.w or 5.0, 0.5, 13.333)
    h = _clamp(el.h or 4.0, 0.5, 7.5)
    x = _clamp(el.x, 0.0, 13.0)
    y = _clamp(el.y, 0.0, 7.0)
    slide.shapes.add_picture(img_path, Inches(x), Inches(y), Inches(w), Inches(h))


# ─────────────────────────────────────────────────────────── chart

def _xy_data(categories, series_list):
    """Kategoriyalar son bo'lsa, ularni x qiymatlari sifatida ishlatadi."""
    x_values = []
    for category in categories:
        try:
            x_values.append(float(str(category).replace(" ", "").replace(",", ".")))
        except (TypeError, ValueError):
            return None
    if len(x_values) < 2:
        return None

    data = XyChartData()
    added = False
    for entry in series_list:
        if not isinstance(entry, dict):
            continue
        values = entry.get("values") or []
        series = data.add_series(entry.get("name", "Series"))
        for x, y in zip(x_values, values):
            try:
                series.add_data_point(x, float(y))
                added = True
            except (TypeError, ValueError):
                continue
    return data if added else None


def _colour_chart(chart, chart_type: str, palette) -> None:
    """Diagramma ranglarini taqdimot palitrasiga moslaydi.

    PowerPoint'ning standart rang sxemasi har taqdimotda bir xil ko'k-to'q
    sariq bo'lib chiqardi. Bu yerda rang tekshiruvdan o'tgan palitradan
    olinadi, shuning uchun har taqdimot boshqacha ko'rinadi.
    """
    if palette is None:
        return
    try:
        plot = chart.plots[0]
        if chart_type in _PER_POINT:
            # Bir qator, ko'p bo'lak: rang har bo'lakka beriladi.
            points = list(plot.series[0].points)
            for index, point in enumerate(points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _hex(
                    palette.categorical[index % len(palette.categorical)])
            return
        for index, series in enumerate(plot.series):
            colour = palette.categorical[index % len(palette.categorical)]
            if chart_type in ("line", "scatter", "radar"):
                series.format.line.color.rgb = _hex(colour)
                series.format.line.width = Pt(2.25)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = _hex(colour)
    except Exception as exc:
        log.warning("Diagramma rangi qo'llanmadi: %s", exc)

def _draw_chart(slide, el, palette=None):
    """python-pptx native chart + caption (izoh matni) chizadi."""
    w = _clamp(el.w or 7.0, 2.0, 13.0)
    h = _clamp(el.h or 4.0, 1.5, 6.5)   # caption uchun joy qoldiramiz
    x = _clamp(el.x, 0.0, 11.0)
    y = _clamp(el.y, 0.0, 6.0)

    chart_data = CategoryChartData()

    categories = el.categories or ["A", "B", "C"]
    chart_data.categories = categories

    series_list = el.series or [{"name": "Ma'lumot", "values": [1] * len(categories)}]
    for s in series_list:
        if not isinstance(s, dict):
            continue
        name = s.get("name", "Series")
        values = s.get("values", [0] * len(categories))
        if len(values) < len(categories):
            values = list(values) + [0] * (len(categories) - len(values))
        chart_data.add_series(name, tuple(values[:len(categories)]))

    requested = el.chart_type or "column"
    if requested == "scatter":
        # XY_SCATTER kategoriya emas, (x, y) juftlarini talab qiladi. Kategoriyalar
        # son bo'lmasa scatter chizib bo'lmaydi — bunday holatda chiziqqa o'tamiz,
        # aks holda diagramma umuman chizilmay qolardi.
        xy_data = _xy_data(categories, series_list)
        if xy_data is not None:
            chart_data = xy_data
        else:
            log.info("Scatter uchun kategoriyalar son emas — chiziqqa o'tildi")
            requested = "line"

    chart_type = CHART_TYPE_MAP.get(requested, XL_CHART_TYPE.COLUMN_CLUSTERED)

    try:
        chart_frame = slide.shapes.add_chart(
            chart_type, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
        )
        chart = chart_frame.chart

        if el.chart_title:
            chart.has_title = True
            chart.chart_title.text_frame.text = el.chart_title
        else:
            chart.has_title = False

        if el.chart_type in ("pie", "donut"):
            plot = chart.plots[0]
            plot.has_data_labels = True

        if len(series_list) <= 1:
            chart.has_legend = False

        _colour_chart(chart, requested, palette)

        # ── Caption: diagramma ostida izoh matni ──────────────────
        caption_text = el.caption or el.chart_title or ""
        if caption_text:
            cap_y = _clamp(y + h + 0.08, 0.0, 7.3)
            # caption slayd ichida bo'lishini tekshir
            if cap_y + 0.35 <= 7.5:
                cap_tb = slide.shapes.add_textbox(
                    Inches(x), Inches(cap_y), Inches(w), Inches(0.35)
                )
                cap_tf = cap_tb.text_frame
                cap_tf.word_wrap = True
                cap_p = cap_tf.paragraphs[0]
                cap_p.alignment = PP_ALIGN.CENTER
                cap_run = cap_p.add_run()
                cap_run.text = f"📊 {caption_text}"
                cap_run.font.size = Pt(11)
                cap_run.font.italic = True
                cap_run.font.name = "Calibri"
                cap_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    except Exception as exc:
        log.error("Chart chizishda xato: %s", exc)
