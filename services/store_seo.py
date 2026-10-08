"""Do'kondagi ishlar qidiruv tizimlarida (Google, Google Rasmlar, Yandex) topilishi uchun.

Bu yerda:
  - manzil uchun lotin harfli nom (slug): /shop/AB12CD34/iqtisodiyot-asoslari-kurs-ishi
  - ish turlarining ko'plik nomi va boshqa atalishi (taqdimot = prezentatsiya, slayd)
  - mavzuning boshqa yozuvdagi ko'rinishi (lotin ↔ kirill): odamlar ikkalasida ham izlaydi
  - fayldan reja (bo'lim/slayd sarlavhalari) va qisqa parcha olish: sahifada haqiqiy matn bo'lsin
  - eski ishlar uchun shu ma'lumotni fonda to'ldirish (fayl ombordan yuklanadi)
  - IndexNow: yangi sahifa haqida Yandex va Bing'ga darhol xabar berish
"""
import asyncio
import hashlib
import json
import logging
import os
import re
import tempfile
import unicodedata
from typing import Dict, List, Optional

logger = logging.getLogger(__name__)

# ───────────────────────────────────────────────────────── nomlar

# Ish turining ko'plik nomi — katalog sahifasining sarlavhasi uchun ("Tayyor kurs ishlari").
PLURALS = {
    "taqdimot": "taqdimotlar",
    "premium_taqdimot": "zamonaviy taqdimotlar",
    "referat": "referatlar",
    "mustaqil_ish": "mustaqil ishlar",
    "kurs_ishi": "kurs ishlari",
    "loyiha_ishi": "loyiha ishlari",
    "maqola": "maqolalar",
    "tezis": "tezislar",
    "mahsus_ishlanma": "maxsus ishlanmalar",
    "bitiruv_ishi": "bitiruv ishlari",
    "diplom_ishi": "diplom ishlari",
    "dissertatsiya": "dissertatsiyalar",
}

# Odamlar bir narsani turlicha ataydi: sahifa matnida shular ham bo'lsa, har qaysi so'z bilan topiladi.
SYNONYMS = {
    "taqdimot": "taqdimot, prezentatsiya, slayd (PowerPoint, PPTX)",
    "premium_taqdimot": "taqdimot, prezentatsiya, slayd (PowerPoint, PPTX)",
    "referat": "referat (Word, DOCX)",
    "mustaqil_ish": "mustaqil ish, SRS (Word, DOCX)",
    "kurs_ishi": "kurs ishi, kurs loyihasi (Word, DOCX)",
    "loyiha_ishi": "loyiha ishi (Word, DOCX)",
    "maqola": "ilmiy maqola (Word, DOCX)",
    "tezis": "tezis, konferensiya tezisi (Word, DOCX)",
    "bitiruv_ishi": "bitiruv malakaviy ishi, BMI (Word, DOCX)",
    "diplom_ishi": "diplom ishi (Word, DOCX)",
    "dissertatsiya": "dissertatsiya (Word, DOCX)",
}

_DASH_RE = re.compile(r"[^a-z0-9]+")


def _ascii(text: str) -> str:
    from services.uz_script import to_latin

    text = to_latin(str(text or ""))   # kirill (o'zbek va rus) → lotin
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[ʻʼ‘’'`´]", "", text.lower())
    return _DASH_RE.sub("-", text).strip("-")


def slugify(text: str, limit: int = 80) -> str:
    """Manzilga yoziladigan nom: faqat lotin harf, raqam va chiziqcha."""
    slug = _ascii(text)
    if len(slug) > limit:
        slug = slug[:limit].rsplit("-", 1)[0]
    return slug.strip("-")


def type_slug(work_type: str) -> str:
    return (work_type or "").replace("_", "-")


def type_from_slug(slug: str) -> str:
    return (slug or "").replace("-", "_")


def item_slug(title: str, work_type: str = "") -> str:
    """Ish manzilidagi nom: mavzu + ish turi ("iqtisodiyot-asoslari-kurs-ishi")."""
    from config import work_label

    base = slugify(title, 70) or "ish"
    kind = slugify(work_label(work_type))
    if kind and f"-{kind}-" not in f"-{base}-":
        base = f"{base}-{kind}"
    return base


def other_script(title: str) -> str:
    """Mavzuning boshqa yozuvdagi ko'rinishi (lotin → kirill yoki aksincha); farq qilmasa bo'sh."""
    from services.uz_script import has_cyrillic, to_cyrillic, to_latin

    title = (title or "").strip()
    if not title:
        return ""
    if not _UZBEK_ONLY.search(title) and _RUSSIAN.search(title):    # ruscha mavzuning "lotincha"si foydasiz
        return ""
    if not has_cyrillic(title) and _FOREIGN_LATIN.search(title.lower()):   # inglizcha mavzu — kirillchasi yo'q
        return ""
    other = to_latin(title) if has_cyrillic(title) else to_cyrillic(title)
    return other if other.strip() and other != title else ""


# O'zbek lotinida w va (ch dan tashqari) c yo'q.
_FOREIGN_LATIN = re.compile(r"w|c(?!h)|\b(the|and|of|in|for)\b")
_UZBEK_ONLY = re.compile(r"[ўқғҳЎҚҒҲ]")
_RUSSIAN = re.compile(r"[ыщэъЫЩЭЪ]|\w(ии|ой|ая|ые|ого|ение|ость|ский|ская|ское)\b", re.I)


def plural(work_type: str) -> str:
    from config import work_label

    return PLURALS.get(work_type) or (work_label(work_type) or "ishlar").lower()


# ───────────────────────────────────────────────────────── fayldan reja va parcha

_GENERIC = {
    "reja", "mundarija", "kirish", "xulosa", "xulosalar", "xulosa va takliflar", "foydalanilgan adabiyotlar",
    "foydalanilgan adabiyotlar royxati", "adabiyotlar", "adabiyotlar royxati", "etiboringiz uchun rahmat",
    "rahmat", "savollar", "ilovalar", "ilova", "taqdimot rejasi", "mavzu", "план", "содержание",
    "введение", "заключение", "спасибо за внимание", "литература", "список литературы", "contents",
    "introduction", "conclusion", "references", "thank you", "жоспар", "кіріспе", "қорытынды",
}
_TOC_HEAD = {"reja", "mundarija", "содержание", "план", "оглавление", "contents", "жоспар", "мазмұны"}
_COVER_WORDS = ("bajardi", "tekshirdi", "qabul qildi", "tayyorladi", "vazirligi", "universiteti",
                "instituti", "fakulteti", "kafedrasi", "guruh", "talaba", "выполнил", "проверил")
_WORK_NAMES = {"taqdimot", "referat", "mustaqil ish", "kurs ishi", "loyiha ishi", "maqola", "tezis",
               "bitiruv ishi", "diplom ishi", "dissertatsiya", "bitiruv malakaviy ishi", "курсовая работа",
               "реферат", "презентация"}
_NUMBERED = re.compile(r"^\s*(\d+(\.\d+)*[.)]?\s|[IVXLC]+[.)]?\s|\d+\s*-?\s*(bob|bo.?lim)\b)", re.I)
_PAGE_TAIL = re.compile(r"[\s.·…_]*\t?\s*\d{1,3}\s*$")
_MAX_OUTLINE = 20
_EXCERPT = 420


def _key(text: str) -> str:
    from services.store_taxonomy import normalize

    return re.sub(r"[^\w ]+", "", normalize(text)).strip()


def _clean_line(text: str) -> str:
    text = re.sub(r"\s+", " ", (text or "").replace("\t", " ")).strip()
    return text.strip(" .:;—-")


def _useful(title: str, topic_key: str) -> bool:
    key = _key(title)
    if len(key) < 3 or key in _GENERIC or key == topic_key:
        return False
    if key.startswith(("taqdimot rejasi", "reja ", "mundarija ", "план ")):
        return False
    if any(word in key for word in _COVER_WORDS) or key in _WORK_NAMES:
        return False
    # "K = 35,2%", "Ko'rsatkich 2" — jadval va infografika yozuvlari, sarlavha emas.
    if re.search(r"[=%]", title) or re.fullmatch(r"\D{1,20}\s\d+", title.strip()):
        return False
    letters = sum(ch.isalpha() for ch in title)
    return letters >= 4 and letters >= len(title) * 0.5


def _cut(text: str, limit: int = _EXCERPT) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) <= limit:
        return text
    cut = text[:limit].rsplit(" ", 1)[0].rstrip(",;:—- ")
    return cut + "…"


def _unique(lines: List[str], topic: str) -> List[str]:
    topic_key = _key(topic)
    seen, out = set(), []
    for line in lines:
        line = _clean_line(line)
        key = _key(line)
        if not line or key in seen or not _useful(line, topic_key):
            continue
        seen.add(key)
        out.append(line[:160])
        if len(out) >= _MAX_OUTLINE:
            break
    return out


def _extract_pptx(path: str, topic: str) -> Dict:
    from pptx import Presentation

    from services.store_publisher import _iter_shapes

    titles, excerpt = [], ""
    for index, slide in enumerate(Presentation(path).slides):
        if index == 0:          # muqova: mavzu va (tozalangan) muallif joyi
            continue
        boxes = []
        for shape in _iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                boxes.append((shape.top or 0, shape.left or 0, shape.text_frame))
        if not boxes:
            continue
        boxes.sort(key=lambda box: (box[0], box[1]))
        head = boxes[0][2].text.strip().split("\n")[0]
        if len(head) <= 160:
            titles.append(head)
        if not excerpt:
            for _, _, frame in boxes[1:]:
                for para in frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if len(text) >= 100:
                        excerpt = text
                        break
                if excerpt:
                    break
    return {"outline": _unique(titles, topic), "excerpt": _cut(excerpt)}


def _bold(paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold for run in runs)


def _extract_docx(path: str, topic: str) -> Dict:
    from docx import Document

    paragraphs = [p for p in Document(path).paragraphs if p.text.strip()]
    keys = [_key(p.text) for p in paragraphs]

    # 1) Reja (mundarija) — hujjatning o'zi yozgan bo'limlar ro'yxati.
    toc, start = [], None
    for index, key in enumerate(keys):
        if key in _TOC_HEAD:
            start = index + 1
            break
    body_from = 0
    if start is not None:
        index = start
        while index < len(paragraphs) and len(toc) < 40:
            text = paragraphs[index].text
            if len(text) > 220:
                break
            line = _PAGE_TAIL.sub("", text) if "\t" in text or re.search(r"\.{3,}", text) else text
            # Reja tugab, matnning o'zi boshlandi: birinchi band yana uchradi.
            if toc and _key(line) == _key(toc[0]):
                break
            toc.append(line)
            index += 1
        body_from = index

    # 2) Rejasiz hujjat: qalin, qisqa qatorlar (sarlavhalar).
    headings = []
    if len(_unique(toc, topic)) < 2:
        for paragraph in paragraphs[body_from:]:
            text = paragraph.text.strip()
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if len(text) <= 150 and ("heading" in style or "заголов" in style or _bold(paragraph)):
                headings.append(text)
        # Raqamlangan bo'limlar ko'p bo'lsa — faqat ular (qalin yozilgan boshqa qatorlar sarlavha emas).
        numbered = [line for line in headings if _NUMBERED.match(line) or "bob" in _key(line)]
        toc = numbered if len(numbered) >= 3 else headings

    # Parcha: kirishdan keyingi birinchi mazmunli xatboshi (asosiy qism boshidan).
    excerpt, seen_heading = "", False
    for paragraph, key in zip(paragraphs[body_from:], keys[body_from:]):
        text = paragraph.text.strip()
        if len(text) <= 150 and _bold(paragraph) and key not in _GENERIC:
            seen_heading = True
            continue
        if seen_heading and len(text) >= 160 and not text.endswith(":"):
            excerpt = text
            break
    if not excerpt:
        excerpt = next((p.text for p in paragraphs[body_from:] if len(p.text) >= 160), "")
    return {"outline": _unique(toc, topic), "excerpt": _cut(excerpt)}


def extract(path: str, topic: str = "") -> Dict:
    """Fayldan reja (sarlavhalar) va qisqa parcha: {"outline": [...], "excerpt": "..."}."""
    try:
        if path.lower().endswith(".pptx"):
            return _extract_pptx(path, topic)
        if path.lower().endswith(".docx"):
            return _extract_docx(path, topic)
    except Exception as exc:
        logger.warning("Ish matni o'qilmadi (%s): %s", os.path.basename(path), exc)
    return {"outline": [], "excerpt": ""}


def outline_of(row: Dict) -> List[str]:
    try:
        value = json.loads(row.get("outline") or "[]")
    except (TypeError, ValueError):
        return []
    return [str(line) for line in value if str(line).strip()] if isinstance(value, list) else []


# ───────────────────────────────────────────────────────── eski ishlarni to'ldirish

BACKFILL_PAUSE = 2.0        # Telegram'ni bezovta qilmaslik uchun har fayl orasida


async def backfill(bot, *, pause: float = BACKFILL_PAUSE, limit: int = 0) -> int:
    """Rejasi hali olinmagan ishlarni ombordan yuklab, reja va parchani yozadi.

    Har ish bir marta urinib ko'riladi (xato bo'lsa ham belgilanadi), shuning uchun
    bot qayta ishga tushganda boshidan boshlamaydi.
    """
    from database.database import Database

    done = 0
    while True:
        rows = await Database.store_items_without_seo(20)
        if not rows:
            break
        for row in rows:
            data, state = {"outline": [], "excerpt": ""}, 2
            suffix = ".docx" if (row.get("file_type") or "pptx") == "docx" else ".pptx"
            handle, path = tempfile.mkstemp(prefix="seo_", suffix=suffix)
            os.close(handle)
            try:
                await bot.download(row["file_id"], destination=path)
                data = await asyncio.to_thread(extract, path, row.get("title") or "")
                state = 1
            except Exception as exc:
                logger.info("SEO: %s faylini olib bo'lmadi: %s", row["public_code"], exc)
            finally:
                try:
                    os.remove(path)
                except OSError:
                    pass
            await Database.set_store_seo(row["public_code"], data["outline"], data["excerpt"], state)
            done += 1
            if limit and done >= limit:
                return done
            await asyncio.sleep(pause)
    if done:
        logger.info("SEO: %s ta ishning rejasi to'ldirildi", done)
    return done


# ───────────────────────────────────────────────────────── IndexNow (Yandex, Bing)

INDEXNOW_ENDPOINT = "https://yandex.com/indexnow"
INDEXNOW_EXTRA = ("https://api.indexnow.org/indexnow",)


def indexnow_key() -> str:
    from config import BOT_TOKEN

    return hashlib.sha256(("edu-indexnow|" + (BOT_TOKEN or "")).encode()).hexdigest()[:32]


def _public_host() -> Optional[str]:
    """Haqiqiy domen (IP, localhost va Replit sinov domeni emas) — aks holda xabar berilmaydi."""
    import webapp

    url = webapp.public_url("")
    host = re.sub(r"^https?://", "", url).split("/")[0].split(":")[0]
    if not url.startswith("https://") or not host or re.fullmatch(r"[\d.]+", host):
        return None
    if host.startswith("localhost") or host.endswith((".replit.dev", ".repl.co")):
        return None
    return host


async def ping(paths: List[str]) -> bool:
    """Yangi/yangilangan sahifalarni IndexNow orqali bildiradi. Hech qachon xato ko'tarmaydi."""
    import aiohttp

    host = _public_host()
    if not host or not paths:
        return False
    key = indexnow_key()
    body = {"host": host, "key": key, "keyLocation": f"https://{host}/{key}.txt",
            "urlList": [f"https://{host}{path}" for path in paths][:10000]}
    ok = False
    try:
        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=20)) as session:
            for endpoint in (INDEXNOW_ENDPOINT,) + INDEXNOW_EXTRA:
                try:
                    async with session.post(endpoint, json=body) as response:
                        ok = ok or response.status in (200, 202)
                        if response.status not in (200, 202):
                            logger.info("IndexNow %s: %s", endpoint, response.status)
                except Exception as exc:
                    logger.info("IndexNow %s: %s", endpoint, exc)
    except Exception as exc:
        logger.info("IndexNow bajarilmadi: %s", exc)
    return ok


def item_path(code: str, title: str, work_type: str = "") -> str:
    return f"/shop/{code}/{item_slug(title, work_type)}"


async def announce_all() -> None:
    """Katalogdagi hamma sahifani bir marta bildiradi (ro'yxat o'zgarganda yana)."""
    from database.database import Database

    if not _public_host():
        return
    rows = await Database.all_store_codes()
    paths = ["/shop"] + [item_path(r["code"], r.get("title") or "", r.get("work_type") or "") for r in rows]
    stamp = hashlib.sha1("\n".join(paths).encode()).hexdigest()[:16]
    if await Database.get_bot_setting("indexnow_stamp") == stamp:
        return
    if await ping(paths):
        await Database.set_bot_setting("indexnow_stamp", stamp)


async def startup(bot) -> None:
    """Bot ishga tushgach fonda: eski ishlarni to'ldirish, keyin hammasini bildirish."""
    await asyncio.sleep(30)
    try:
        await backfill(bot)
    except Exception as exc:
        logger.warning("SEO to'ldirish to'xtadi: %s", exc)
    try:
        await announce_all()
    except Exception as exc:
        logger.info("IndexNow e'loni bajarilmadi: %s", exc)
