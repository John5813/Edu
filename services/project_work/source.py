"""Mijoz bergan manbani matnga aylantirish.

Loyiha ishi qo'llanmadan, taqdimotdan, muassasa saytidan yoki mijozning o'z
tushuntirishidan kelib chiqishi mumkin. Bu modul har qaysisini bitta narsaga
— matnga — keltiradi, keyin mazmun generatori o'shanga tayanadi.
"""

import logging
import os
from dataclasses import dataclass

logger = logging.getLogger(__name__)

KIND_AI = "ai"          # AI o'zi yozadi, tashqi manba yo'q
KIND_TEXT = "text"      # mijoz matn bilan tushuntirgan
KIND_FILE = "file"      # DOCX / PDF / PPTX
KIND_URL = "url"        # bir yoki bir nechta sayt

# Bir manbadan olinadigan maksimal matn. Undan ortig'i baribir promptga
# sig'maydi va sifatga hissa qo'shmaydi.
_MAX_SOURCE_CHARS = 60_000


@dataclass
class SourceMaterial:
    kind: str = KIND_AI
    text: str = ""
    label: str = ""

    @property
    def has_content(self) -> bool:
        return bool(self.text.strip())


def _trim(text: str) -> str:
    text = (text or "").strip()
    return text[:_MAX_SOURCE_CHARS]


def from_instructions(text: str) -> SourceMaterial:
    return SourceMaterial(kind=KIND_TEXT, text=_trim(text), label="")


async def from_file(local_path: str, file_name: str) -> SourceMaterial:
    """DOCX, PDF yoki PPTX fayldan matn ajratadi."""
    lowered = file_name.lower()
    if lowered.endswith(".pptx"):
        text = _read_pptx(local_path)
    elif lowered.endswith(".pdf"):
        text = await _read_pdf(local_path)
    elif lowered.endswith(".docx"):
        text = _read_docx(local_path)
    else:
        raise ValueError(f"qo'llab-quvvatlanmaydigan fayl turi: {file_name}")

    if len(text.split()) < 20:
        raise ValueError("fayldan yetarli matn ajratilmadi")
    return SourceMaterial(kind=KIND_FILE, text=_trim(text), label=file_name)


def _read_docx(path: str) -> str:
    from docx import Document

    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pptx(path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        if lines:
            parts.append(f"[{number}-slayd] " + "\n".join(lines))
    return "\n\n".join(parts)


async def _read_pdf(path: str) -> str:
    from services.book_translate_service import auto_convert_pdf_to_docx

    docx_path = await auto_convert_pdf_to_docx(path)
    try:
        return _read_docx(docx_path)
    finally:
        try:
            os.remove(docx_path)
        except OSError:
            pass


async def from_urls(urls: list) -> SourceMaterial:
    """Bir yoki bir nechta sayt manzilidan matn yig'adi."""
    from services.url_book_service import fetch_multiple_urls

    result = await fetch_multiple_urls(urls)
    text = result.get("combined_content") or ""
    if result.get("total_word_count", 0) < 20 or not text.strip():
        failures = [r.get("error") or "fetch_failed" for r in result.get("results", []) if not r.get("ok")]
        raise ValueError(failures[0] if failures else "saytdan yetarli matn olinmadi")

    reached = [r["url"] for r in result.get("results", []) if r.get("ok")]
    labels = ", ".join(_domain(url) for url in reached) or ", ".join(_domain(u) for u in urls)
    return SourceMaterial(kind=KIND_URL, text=_trim(text), label=labels)


def _domain(url: str) -> str:
    from urllib.parse import urlparse

    try:
        return urlparse(url).netloc or url
    except Exception:
        return url
