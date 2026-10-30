"""Umumiy 3D katalog (tarix, biologiya, texnika, san'at...) — taqdimot slaydiga tayyor `.glb` modellar.

Ro'yxat `data/katalog3d/katalog.json` da (git'da): har yozuvda o'zbekcha nom, fan, kalit so'zlar,
litsenziya va muallif. Modellarning o'zi (~12,5 GB) git'ga qo'shilmaydi: serverga kerak bo'lganda o'zi
yuklanadi (`services.kat3d.yukla`, 3D Anatomiya atlasi kabi) — `data/katalog3d_modellar/` ga (joyni
`KATALOG3D_DIR` o'zgartiradi). Hammasini birdan yuklash: `tools/katalog3d/yukla.sh` yoki `KATALOG3D_OLDINDAN=1`.

Manbalar: Objaverse/Sketchfab (faqat CC0 va CC BY 4.0) va Smithsonian Open Access (CC0).
CC BY modellarda muallifni ko'rsatish majburiy — `joylash()` yozuvni model ostiga o'zi qo'yadi,
to'liq matn `attribution_full` da ("Manbalar" slaydi uchun).

Yangi katalog (masalan NASA) qo'shish: `data/katalog3d/` ga xuddi shu tuzilishdagi `<nom>.json` qo'yiladi
({"modellar": [...]}, `fayl` — `KATALOG3D_DIR` ga nisbatan yo'l, rasmi `preview/<id>.jpg`). Hammasi bitta
ro'yxatga birlashadi: qidiruv va 3D taqdimot ularni o'zi ko'radi.
"""
import glob
import json
import os
import re
import threading
import time
from functools import lru_cache
from typing import Dict, List, Optional

from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KATALOG_DIR = os.path.join(ROOT, "data", "katalog3d")
KATALOG_JSON = os.path.join(KATALOG_DIR, "katalog.json")
PREVIEW_DIR = os.path.join(KATALOG_DIR, "preview")
MASK_DIR = os.path.join(KATALOG_DIR, "mask")          # preview ning shaffoflik niqobi (fonsiz model rasmi uchun)
MODELS_DIR = os.getenv("KATALOG3D_DIR") or os.path.join(ROOT, "data", "katalog3d_modellar")


@lru_cache(maxsize=1)
def katalog() -> Dict[str, dict]:
    """Hamma katalog bitta lug'atda: avval asosiy `katalog.json`, keyin qolgan `*.json` lar (alifbo tartibida)."""
    files = [KATALOG_JSON] + sorted(f for f in glob.glob(os.path.join(KATALOG_DIR, "*.json")) if f != KATALOG_JSON)
    out: Dict[str, dict] = {}
    for path in files:
        if not os.path.isfile(path):
            continue
        with open(path, encoding="utf-8") as fh:
            for m in json.load(fh).get("modellar", []):
                out.setdefault(m["id"], m)
    return out


_mavjud_lock = threading.Lock()
_mavjud_cache: Dict[str, object] = {"at": 0.0, "ids": frozenset()}


def mavjud(ttl: float = 600.0) -> frozenset:
    """Serverda haqiqatan yuklangan modellar (id lar). Diskda yo'q model mijozga taklif qilinmaydi.
    10 daqiqada bir yangilanadi — yuklash davom etayotgan serverda ham yangi modellar o'zi qo'shiladi."""
    with _mavjud_lock:
        if time.time() - float(_mavjud_cache["at"]) < ttl:
            return _mavjud_cache["ids"]                      # type: ignore[return-value]
        ids = frozenset(i for i, m in katalog().items() if os.path.isfile(os.path.join(MODELS_DIR, m["fayl"])))
        _mavjud_cache.update(at=time.time(), ids=ids)
        return ids


def _norm(s: str) -> str:
    return re.sub(r"[ʻʼ'`‘’]", "", s.lower())


def qidir(soz: str, fan: Optional[str] = None, limit: int = 10) -> List[dict]:
    """O'zbekcha (yoki inglizcha) so'z bo'yicha qidirish; nomdagi moslik kalit so'zdagidan ustun."""
    words = [w for w in _norm(soz).split() if len(w) > 1]
    if not words:
        return []
    found = []
    for m in katalog().values():
        if fan and m["fan"] != fan:
            continue
        nom = _norm(m["nom_uz"] + " " + m["nom"])
        kalit = _norm(" ".join(m.get("kalit", [])))
        ball = sum(3 if w in nom else 1 if w in kalit else 0 for w in words)
        if ball:
            found.append((ball, m.get("sifat", 0), m))
    found.sort(key=lambda t: (t[0], t[1]), reverse=True)
    return [m for _, _, m in found[:limit]]


def model_yoli(model_id: str) -> Optional[str]:
    """Diskdagi `.glb` yo'li; model hali yuklanmagan bo'lsa None."""
    m = katalog().get(model_id)
    if not m:
        return None
    path = os.path.join(MODELS_DIR, m["fayl"])
    return path if os.path.isfile(path) else None


def preview_yoli(model_id: str) -> Optional[str]:
    path = os.path.join(PREVIEW_DIR, model_id + ".jpg")
    return path if os.path.isfile(path) else None


def maska_yoli(model_id: str) -> Optional[str]:
    path = os.path.join(MASK_DIR, model_id + ".png")
    return path if os.path.isfile(path) else None


def joylash(slide, model_id: str, x: float, y: float, size: float, parts=None, rot=(0, 0, 0),
            lights: str = "studio", izoh: bool = True, izoh_pt: int = 9):
    """Slaydga katalog modelini qo'yadi (x, y, size — dyuymda) va ostiga muallif/litsenziya yozuvini.

    Yozuv CC BY uchun majburiy, CC0 uchun ham qo'yiladi. Model diskda bo'lmasa FileNotFoundError.
    """
    from services.med3d.pptx3d import add_model3d

    path = model_yoli(model_id)
    if not path:
        raise FileNotFoundError(f"katalog3d: model yuklanmagan: {model_id}")
    m = katalog()[model_id]
    with open(path, "rb") as fh:
        node = add_model3d(slide, fh.read(), m["nom_uz"], x, y, size, rot, parts, lights)
    if izoh:
        box = slide.shapes.add_textbox(Inches(x), Inches(y + size), Inches(size), Inches(0.3))
        tf = box.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = m["attribution"]
        run.font.size = Pt(izoh_pt)
    return node
