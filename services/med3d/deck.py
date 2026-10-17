"""Tibbiy 3D taqdimotni rejadan yig'adi: chapda matn, o'ngda 3D model, yorliq chiziqlari, Morph o'tishi.

Reja (keyinchalik AI tuzadi):

    {"scene": "skull",
     "slides": [
        {"title": "Bosh suyagi", "text": "...", "show": ["cranium", "mandible"], "rot": [0, 0, 0],
         "labels": ["frontal", "parietal_r", "mandible"]},
        {"title": "Pastki jag'", "text": "...", "show": ["mandible"], "rot": [70, 0, 0],
         "labels": ["incisors_lower", "molars_lower_l"], "exit": "up"}]}

- `show` — slaydda ko'rinadigan guruhlar; oldingi slaydda bo'lib, bunda yo'q guruh `exit` tomonga
  (sukut bo'yicha yuqoriga) slayddan chiqib ketadi — Morph uni silliq olib chiqadi.
- `focus` — kamera qaysi guruhlarga moslanadi (sukut — `show`); `rot` — gradus (X, Y, Z).
- Bir slayddagi hamma guruh bir xil ramka va aylanishda — qismlar bir-biriga aniq yopishadi.
"""
import logging
import os
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from services.pro3d.build import add_entrances, add_morph

from . import ASSETS_DIR, design, load_scene
from .camera import Camera
from .pptx3d import add_model3d

log = logging.getLogger(__name__)

W, H = 13.333, 7.5
MORPH_MS = 2200
HELPER_MS = 1300
PAL = dict(bg="0B1220", ink="F3F6FB", muted="9FB0C8", accent="38BDF8", pill="14213A", pill_ink="E8F3FF",
           line="38BDF8")
TEXT_BOX = (0.65, 0.75, 4.35)                  # x, y, eni — chapdagi matn ustuni
MODEL_REGION = (7.05, 0.6, 4.2, 6.3)           # yorliqlar bo'lsa: ikki yonida yorliq ustunlari
MODEL_REGION_WIDE = (5.4, 0.45, 7.5, 6.6)      # yorliqsiz slayd: model kattaroq
FILL = 0.9
# 3D ramka tomoni (dyuym) — bundan katta ramkani PowerPoint juda katta tasvir sifatida chizadi
# (sekin va xira bo'lishi mumkin). Mayda organ uchun alohida, kichikroq sahna ishlatiladi.
MAX_FRAME = 26.0
LABEL_H, LABEL_GAP, LABEL_MAX_W = 0.36, 0.5, 3.0
LABEL_H2, LABEL_GAP2 = 0.54, 0.66              # ikki qatorli yorliq: o'zbekcha nom + lotincha nom
HIDDEN_SHARE = 0.08                            # organ yuzasining shundan kami ko'rinsa — yorliq qo'yilmaydi
ZOOM_LIMIT = 1.6                               # fokusga yaqinlashish: butun ko'rinishdan ko'pi bilan shuncha katta
LABEL_X_MIN, LABEL_X_MAX = 5.15, W - 0.2
EXITS = {"up": (0, -1), "down": (0, 1), "left": (-1, 0), "right": (1, 0)}
# Chiqib ketishdagi qo'shimcha burilish (X, Y, Z gradus) va masshtab. Har burilish < 180° — Morph qaysi
# yo'nalishda aylantirishidan qat'i nazar aylanish yaqqol ko'rinadi; Z (ekran tekisligida) — buralish (spiral).
SPINS = {"down": ((55, -150, 165), 0.5), "up": ((-70, 120, -40), 0.7),
         "left": ((25, -165, -35), 0.8), "right": ((25, 165, 35), 0.8)}


def _wrap(angle: float) -> float:
    return angle - 360.0 if angle > 359.0 else (angle + 360.0 if angle < -359.0 else angle)


class Frame:
    """Bir slayddagi 3D ramka: markaz (dyuym), tomoni, aylanish."""

    def __init__(self, cx: float, cy: float, size: float, rot: Sequence[float]):
        self.cx, self.cy, self.size, self.rot = cx, cy, size, tuple(float(a) for a in rot)

    @property
    def left(self) -> float:
        return self.cx - self.size / 2

    @property
    def top(self) -> float:
        return self.cy - self.size / 2

    def to_slide(self, u: float, v: float) -> Tuple[float, float]:
        return self.cx + u * self.size, self.cy + v * self.size

    def moved_out(self, direction: str) -> "Frame":
        """Slayd tashqarisidagi joy — uchib chiqish/kirish uchun. Model to'g'ri chiziqda sirpanib ketmaydi: Morph
        burilishni ham oraliqlaydi, shuning uchun tashqaridagi nusxa boshqa burchakda va kichikroq turadi —
        pastga ketayotgan organ o'z o'qi atrofida buralib (spiral) uzoqlashadi, yon tomonga ketayotgani orqasiga
        aylanib, qiyshayib chiqadi. Kirishda xuddi shu harakat teskari: aylanib kelib, o'z holatiga yoziladi."""
        dx, dy = EXITS.get(direction, EXITS["up"])
        spin, shrink = SPINS.get(direction, SPINS["up"])
        size = self.size * shrink
        cx = self.cx + dx * (W + size) * 0.75 if dx else self.cx
        cy = -size / 2 - 0.4 if dy < 0 else (H + size / 2 + 0.4 if dy > 0 else self.cy)
        rot = tuple(_wrap(a + d) for a, d in zip(self.rot, spin))
        return Frame(cx, cy, size, rot)


def fit_frame(cam: Camera, points, rot, region, fill: float = FILL) -> Frame:
    u0, v0, u1, v1 = cam.extent(points, rot)
    rx, ry, rw, rh = region
    size = min(rw * fill / max(u1 - u0, 1e-6), rh * fill / max(v1 - v0, 1e-6), MAX_FRAME)
    return Frame(rx + rw / 2 - size * (u0 + u1) / 2, ry + rh / 2 - size * (v0 + v1) / 2, size, rot)


TEXT_EDGE = 5.15            # chapdagi matn ustuni shu yerda tugaydi — 3D model undan chapga o'tmasin


def keep_clear(cam: Camera, frame: Frame, focus_points, all_points, area=TEXT_EDGE) -> Frame:
    """Model matnga tegmasin: `area` — model chiqib turishi mumkin bo'lgan to'rtburchak (x0, y0, x1, y1).

    Matn chapda bo'lsa — model o'ngga, o'ngda bo'lsa — chapga, pastda bo'lsa — tepaga suriladi (fokusdagi organ
    slayddan chiqib ketmaguncha), bo'lmasa — fokus atrofida kichraytiriladi. Matn yo'q tomondan slayd chetidan
    chiqib turish mumkin (kesilgan "yaqin plan" ko'rinishi). Eski chaqiruv: `area` — chap chegara (son).
    """
    ax0, ay0, ax1, ay1 = (area, -9.0, 99.0, 99.0) if isinstance(area, (int, float)) else area
    for _ in range(14):
        u0, v0, u1, v1 = cam.extent(all_points, frame.rot)
        x0, y0 = frame.to_slide(u0, v0)
        x1, y1 = frame.to_slide(u1, v1)
        dx = ax0 - x0 if x0 < ax0 - 1e-3 else (ax1 - x1 if x1 > ax1 + 1e-3 else 0.0)
        dy = ay0 - y0 if y0 < ay0 - 1e-3 else (ay1 - y1 if y1 > ay1 + 1e-3 else 0.0)
        if not dx and not dy:
            return frame
        f0, g0, f1, g1 = cam.extent(focus_points, frame.rot)
        fx0, fy0 = frame.to_slide(f0, g0)
        fx1, fy1 = frame.to_slide(f1, g1)
        # Model joyga sig'adi (faqat bir chetidan chiqqan) — suriladi; ikkala chetidan chiqsa surish bir tomonni
        # tuzatib, ikkinchisini battar qiladi — kichraytiriladi.
        fits = x1 - x0 <= ax1 - ax0 + 1e-3 and y1 - y0 <= ay1 - ay0 + 1e-3
        if fits and fx0 + dx >= -0.1 and fx1 + dx <= W - 0.25 and fy0 + dy >= -0.2 and fy1 + dy <= H + 0.2:
            frame = Frame(frame.cx + dx, frame.cy + dy, frame.size, frame.rot)
            continue
        fu, fv = (f0 + f1) / 2, (g0 + g1) / 2
        fx, fy = frame.cx + fu * frame.size, frame.cy + fv * frame.size
        size = frame.size * 0.85
        frame = Frame(fx - fu * size, fy - fv * size, size, frame.rot)
    return frame


def _too_big(cam: Camera, frame: Frame, points) -> bool:
    """Kamera boshqa organga yaqinlashganda bu organ slayddan bir yarim barobar katta bo'lib ketadimi.

    Masalan, klapanlarga yaqinlashganda yonidagi aorta butun ekranni to'sib qo'yadi — bunday organ shu
    slayddan chiqariladi (Morph uni uchirib yuboradi), keyingi umumiy ko'rinishda qaytadi.
    """
    u0, v0, u1, v1 = cam.extent(points, frame.rot)
    return (u1 - u0) * frame.size > 1.5 * W or (v1 - v0) * frame.size > 1.5 * H


def _spread(ys: List[float], top: float, bottom: float, gap: float) -> List[float]:
    """Yorliqlarni anchor balandligiga yaqin, lekin ustma-ust tushmaydigan qilib joylaydi."""
    out = [min(max(y, top), bottom) for y in ys]
    for i in range(1, len(out)):
        out[i] = max(out[i], out[i - 1] + gap)
    if out and out[-1] > bottom:
        out[-1] = bottom
        for i in range(len(out) - 2, -1, -1):
            out[i] = min(out[i], out[i + 1] - gap)
    return out


def _near(seg, point, tol: float = 0.13) -> bool:
    """Chiziq (seg) boshqa yorliq nuqtasi ustidan o'tadimi."""
    (x1, y1), (x2, y2), (px, py) = seg[0], seg[1], point
    dx, dy = x2 - x1, y2 - y1
    t = max(0.0, min(1.0, ((px - x1) * dx + (py - y1) * dy) / (dx * dx + dy * dy or 1e-9)))
    return (x1 + t * dx - px) ** 2 + (y1 + t * dy - py) ** 2 < tol * tol


def _cross(a, b) -> bool:
    """Ikki chiziq kesishadimi."""
    def orient(p, q, r):
        return (q[0] - p[0]) * (r[1] - p[1]) - (q[1] - p[1]) * (r[0] - p[0])
    (p1, p2), (q1, q2) = a, b
    return (orient(p1, p2, q1) * orient(p1, p2, q2) < 0) and (orient(q1, q2, p1) * orient(q1, q2, p2) < 0)


def _untangle(edges: List[float], anchors: List[Tuple[float, float]], ys: List[float]) -> List[float]:
    """Chiziq qo'shni nuqta ustidan o'tsa, ikki yorliq o'rnini almashtiradi (kesishishlar kamaysa)."""
    def conflicts(order):
        segs = [((e, y), a) for e, y, a in zip(edges, order, anchors)]
        n = len(segs)
        over = sum(_near(segs[i], anchors[j]) for i in range(n) for j in range(n) if i != j)
        return over + sum(_cross(segs[i], segs[j]) for i in range(n) for j in range(i + 1, n))

    ys = list(ys)
    for _ in range(len(ys) * 2):
        best, swap = conflicts(ys), None
        for i in range(len(ys) - 1):
            trial = ys[:i] + [ys[i + 1], ys[i]] + ys[i + 2:]
            score = conflicts(trial)
            if score < best:
                best, swap = score, trial
        if swap is None:
            break
        ys = swap
    return ys


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# ───────────────────────────────────────────── yorliqlar

def _label_size(st, value: str, latin: str) -> Tuple[float, float, int, int]:
    """Yorliq (eni, bo'yi, nom shrifti, lotin shrifti)."""
    size = 11 if len(value) <= 24 else 10
    small = 9 if len(latin) <= 28 else 8
    name_w = len(value) * size * 0.56 / 72 + 0.36
    latin_w = len(latin) * small * 0.52 / 72 + 0.36 if latin else 0
    if st.label == "pill":
        width = min(LABEL_MAX_W, max(1.0, name_w, latin_w))
        return width, (LABEL_H2 if latin else LABEL_H), size, small
    # tag / plain: nom (kartochka yoki matn) va ostida lotincha nom
    width = min(LABEL_MAX_W, max(1.0, name_w, latin_w))
    return width, (0.34 + (0.24 if latin else 0)), size, small


def _draw_marker(slide, st, key: str, item: dict, dx: float = 0.0) -> List:
    """Savol slaydi: nom o'rniga raqamli doira (javoblar bosilganda chiqadi)."""
    d = 0.42
    x = item["x"] + item["w"] - d if item["side"] == "left" else item["x"]
    circle = design.shape(slide, f"lbl_{key}", MSO_SHAPE.OVAL, x + dx, item["y"] - d / 2, d, d, st.accent,
                          line=st.pill_line if not st.light else st.bg, line_pt=1.5)
    frame = circle.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = str(item["number"])
    run.font.name, run.font.size, run.font.bold = st.head_font, Pt(15), True
    run.font.color.rgb = _rgb(st.bg if not st.light else "FFFFFF")
    return [circle]


def _draw_label(slide, st, key: str, item: dict, dx: float = 0.0) -> List:
    """Bitta yorliq (nom va lotincha nom). dx — siljitish (oldingi slaydda slayd tashqarisidagi nusxa uchun).

    Qaytaradi: Morph bilan uchib kiradigan shakllar (kartochka, lotincha matn)."""
    if item.get("number"):
        return _draw_marker(slide, st, key, item, dx)
    x, y, w, h = item["x"] + dx, item["y"], item["w"], item["h"]
    size, small, value, latin = item["size"], item["small"], item["uz"], item["la"]
    align = PP_ALIGN.RIGHT if item["side"] == "left" else PP_ALIGN.LEFT
    shapes = []
    if st.label == "pill":
        pill = design.shape(slide, f"lbl_{key}", MSO_SHAPE.ROUNDED_RECTANGLE, x, y - h / 2, w, h, st.pill,
                            round_=st.pill_round if not latin else 0.35, line=st.pill_line, line_pt=0.75)
        frame = pill.text_frame
        frame.margin_left = frame.margin_right = Inches(0.1)
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.word_wrap = False
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = value
        run.font.name, run.font.size, run.font.bold = st.body_font, Pt(size), True
        run.font.color.rgb = _rgb(st.pill_ink)
        if latin:
            second = frame.add_paragraph()
            second.alignment = PP_ALIGN.CENTER
            lat = second.add_run()
            lat.text = latin
            lat.font.name, lat.font.size, lat.font.italic = st.body_font, Pt(small), True
            lat.font.color.rgb = _rgb(st.latin if not st.light else st.body)
        shapes.append(pill)
        return shapes
    top = y - h / 2
    if st.label == "tag":
        tw = min(w, len(value) * size * 0.6 / 72 + 0.4)
        tx = x + w - tw if item["side"] == "left" else x
        tag = design.shape(slide, f"lbl_{key}", MSO_SHAPE.ROUNDED_RECTANGLE, tx, top, tw, 0.32, st.pill,
                           round_=0.5)
        frame = tag.text_frame
        frame.margin_left = frame.margin_right = Inches(0.08)
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.word_wrap = False
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = value.upper()
        run.font.name, run.font.size, run.font.bold = st.head_font, Pt(size - 1), True
        run.font.color.rgb = _rgb(st.pill_ink)
        shapes.append(tag)
    else:                                    # plain: matn va ostida ingichka chiziq
        shapes.append(design.textbox(slide, f"lbl_{key}", x, top, w, 0.3, value, size=size + 1, color=st.ink,
                                     font=st.head_font, bold=True, align=align, anchor=MSO_ANCHOR.BOTTOM))
        shapes.append(design.shape(slide, f"lbl_{key}_rule", MSO_SHAPE.RECTANGLE, x, top + 0.31, w, 0.012,
                                   st.line))
    if latin:
        shapes.append(design.textbox(slide, f"lbl_{key}_la", x, top + 0.34, w, 0.24, latin, size=small,
                                     color=st.latin, font=st.body_font, italic=True, align=align))
    return shapes


def _connector(slide, st, key: str, start, end) -> List:
    kind = {"curve": MSO_CONNECTOR.CURVE, "elbow": MSO_CONNECTOR.ELBOW}.get(st.connector, MSO_CONNECTOR.STRAIGHT)
    line = slide.shapes.add_connector(kind, Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]))
    line.name = f"lbl_{key}_line"
    line.line.color.rgb = _rgb(st.line)
    line.line.width = Pt(1.25)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(end[0] - 0.06), Inches(end[1] - 0.06), Inches(0.12),
                                 Inches(0.12))
    dot.name = f"lbl_{key}_dot"
    dot.fill.solid()
    dot.fill.fore_color.rgb = _rgb(st.accent if not st.light else st.line)
    dot.line.color.rgb = _rgb(st.bg if st.light else "FFFFFF")
    dot.line.width = Pt(1)
    dot.shadow.inherit = False
    return [dot, line]


def _labels(cam: Camera, scene: dict, frame: Frame, keys: Sequence[str], visible: Sequence[str],
            model_box: Tuple[float, float, float, float], st, lay, offsets: Dict[str, Tuple[float, float]]
            ) -> List[Tuple[str, dict]]:
    """Yorliqlar joyi: [(kalit, {x, y, w, h, side, ax, ay, ...})]. Chizilmaydi — chizish `_draw_label` da."""
    known = scene["labels"]
    items = []
    x_lo, x_hi = lay.labels_x
    for key in keys:
        info = known.get(key)
        if not info or info["group"] not in visible:
            log.warning("med3d: yorliq tashlab ketildi: %s", key)
            continue
        ox, oy = offsets.get(info["group"], (0.0, 0.0))
        # Birlashgan modelda (bir nechta organ — bitta GLB) to'suvchilar yorliqning o'zida beriladi.
        others = info.get("occluders") or [p for g in visible if g != info["group"] for p in scene["groups"][g]["samples"]]
        radius = 0.12 / frame.size
        if others and cam.visible_share(info["samples"], frame.rot, others, radius) < HIDDEN_SHARE:
            log.info("med3d: organ boshqasi orqasida yopilgan, yorliq qo'yilmadi: %s", key)
            continue
        # Uzun organ (qizilo'ngach, aorta) qisman slayddan chiqib turishi mumkin: nuqta slayd ichidagi qismidan.
        pts = np.asarray(info["samples"], float)
        uv, _ = cam.project(pts, frame.rot)
        sx, sy = frame.cx + ox + uv[:, 0] * frame.size, frame.cy + oy + uv[:, 1] * frame.size
        inside = (sx >= x_lo) & (sx <= x_hi) & (sy >= 0.15) & (sy <= H - 0.15)
        if inside.sum() >= 3:
            pts = pts[inside]
        u, v = cam.anchor(pts, frame.rot, others, radius)
        ax, ay = frame.cx + ox + u * frame.size, frame.cy + oy + v * frame.size
        if not (x_lo <= ax <= x_hi and 0.15 <= ay <= H - 0.15):
            log.info("med3d: yorliq nuqtasi slayddan tashqarida, qo'yilmadi: %s", key)
            continue
        items.append((key, info["uz"], ax, ay, info.get("la", "")))
    if not items:
        return []
    left, top, right, bottom = model_box
    mid = (left + right) / 2
    y_lo, y_hi = lay.labels_y
    out = []
    for side in ("left", "right"):
        col = sorted((it for it in items if (it[2] < mid) == (side == "left")), key=lambda it: it[3])
        if not col:
            continue
        sizes = [_label_size(st, it[1], it[4]) for it in col]
        tallest = max(h for _, h, _, _ in sizes)
        gap = tallest + 0.14
        ys = _spread([it[3] for it in col], y_lo + tallest / 2, y_hi - tallest / 2, gap)   # yorliq to'liq chegarada
        widths = [w for w, _, _, _ in sizes]
        if side == "left":
            edges = [max(x_lo + wd, min(left - 0.3, it[2] - 0.5)) for it, wd in zip(col, widths)]
        else:
            edges = [min(x_hi - wd, max(right + 0.3, it[2] + 0.5)) for it, wd in zip(col, widths)]
        ys = _untangle(edges, [(it[2], it[3]) for it in col], ys)
        for (key, uz, ax, ay, la), ly, edge, (w, h, size, small) in zip(col, ys, edges, sizes):
            x = edge - w if side == "left" else edge
            start = (edge, ly) if st.label != "plain" else (edge, ly - h / 2 + 0.32)
            if st.label == "tag":
                start = (edge, ly - h / 2 + 0.16)
            out.append((key, dict(uz=uz, la=la, x=x, y=ly, w=w, h=h, size=size, small=small, side=side,
                                  start=start, ax=ax, ay=ay)))
    return out


# ───────────────────────────────────────────── matn ustuni

def _text_column(slide, data: dict, number: int, total: int, source: str, st, lay) -> List:
    tx, ty, tw, th = lay.title
    bx, by, bw, bh = lay.body
    if lay.panel:
        px, py, pw, ph = lay.panel
    else:                                     # boshqa kompozitsiyada panel slayd tashqarisida kutib turadi
        px, py, pw, ph = -6.0, 0.45, 5.05, 6.6
    design.shape(slide, "!!panel", MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pw, ph, st.panel, alpha=st.panel_alpha,
                 round_=0.06)
    ax, ay = lay.accent
    design.shape(slide, "!!accent", MSO_SHAPE.ROUNDED_RECTANGLE, ax, ay, 0.9, 0.07, st.accent, round_=0.5)
    title = design.textbox(slide, "!!title", tx, ty, tw, th, data.get("title", ""), size=lay.title_size,
                           smallest=20, color=st.ink, font=st.head_font, bold=True, line=1.04, caps=st.title_caps)
    body = design.textbox(slide, "!!body", bx, by, bw, bh, data.get("text", ""), size=lay.body_size,
                          smallest=11, color=st.body, font=st.body_font, line=1.25)
    fx, fy = (W - 6.9, 0.22) if lay.key == "pastki" else (0.65, H - 0.62)
    design.textbox(slide, "!!page", fx, fy, 1.2, 0.3, f"{number + 1:02d} / {total:02d}", size=10, color=st.muted,
                   font=st.body_font)
    design.textbox(slide, "!!source", fx + 1.3, fy, 5.4, 0.3, f"3D: {source}", size=8, color=st.muted,
                   font=st.body_font)
    return [title, body]


_CNVPR_NAME = etree.XPath("./*/p:cNvPr/@name",
                          namespaces={"p": "http://schemas.openxmlformats.org/presentationml/2006/main"})


def _answers_timing(slide, auto_steps, click_steps) -> None:
    """Animatsiya: `auto_steps` o'tishdan keyin o'zi, `click_steps` — bosilganda (savol slaydidagi javoblar)."""
    from services.premium_presentation import slide_anim

    def group(steps, cid, first_click):
        xml = []
        for n, (shape, kind, delay) in enumerate(steps):
            part, used = slide_anim._effect(cid, shape.shape_id, kind, delay)
            if first_click and n == 0:
                part = part.replace('nodeType="withEffect"', 'nodeType="clickEffect"', 1)
            xml.append(part)
            cid += used
        return "".join(xml), cid

    auto_xml, cid = group(auto_steps, 7, False)
    click_xml, cid = group(click_steps, cid + 2, True)
    pars = []
    if auto_xml:
        pars.append('<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
                    '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
                    f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                    f'{auto_xml}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    if click_xml:
        pars.append('<p:par><p:cTn id="5" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
                    '<p:childTnLst><p:par><p:cTn id="6" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                    f'<p:childTnLst>{click_xml}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    builds = "".join(f'<p:bldP spid="{sh.shape_id}" grpId="0"/>' for sh, _, _ in auto_steps + click_steps
                     if etree.QName(sh._element).localname == "sp" and sh.has_text_frame and sh.text_frame.text.strip())
    xml = (f'<p:timing {slide_anim.P_NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
           f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
           f'nodeType="mainSeq"><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
           f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           f'{"<p:bldLst>" + builds + "</p:bldLst>" if builds else ""}</p:timing>')
    node = etree.fromstring(xml)
    ext = slide._element.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(node)
    else:
        slide._element.append(node)


def _answer_card(slide, st, lay, items: List[Tuple[str, dict]]):
    """Savol slaydidagi javoblar kartasi (bosilganda chiqadi): 1 — Jigar (Hepar) ..."""
    bx, by, bw, bh = lay.body
    lines = [f"{it['number']}  —  {it['uz']}" + (f"  ·  {it['la']}" if it["la"] else "") for _, it in
             sorted(items, key=lambda kv: kv[1]["number"])]
    card = design.shape(slide, "!!answers", MSO_SHAPE.ROUNDED_RECTANGLE, bx - 0.15, by - 0.1, bw + 0.3,
                        min(bh + 0.2, 0.45 + 0.36 * len(lines)), st.panel if not st.light else "FFFFFF", alpha=0.97,
                        round_=0.08, line=st.accent, line_pt=1)
    frame = card.text_frame
    frame.margin_left = frame.margin_right = Inches(0.2)
    frame.margin_top = Inches(0.12)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.word_wrap = True
    for n, line in enumerate(lines):
        para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        para.space_after = Pt(4)
        run = para.add_run()
        run.text = line
        run.font.name, run.font.size, run.font.bold = st.body_font, Pt(14), n >= 0
        run.font.color.rgb = _rgb(st.ink if not st.light else st.ink)
    return card


def _overlap(a, b, pad: float = 0.0) -> bool:
    return a[0] < b[2] - pad and b[0] < a[2] - pad and a[1] < b[3] - pad and b[1] < a[3] - pad


def quality(lay, model_box, labels: List[Tuple[str, dict]], text_boxes, focus_box=None) -> List[str]:
    """Sifat nazorati (slayd yig'ilgach): yorliqlar bir-biriga, matnga tegmasin, model matn ustiga chiqmasin;
    fokusdagi organ slayddan chiqib ketmasin va juda kichik bo'lmasin, yorliq chizig'i model ustiga tushsin."""
    issues = []
    if focus_box is not None:
        fx0, fy0, fx1, fy1 = focus_box
        if fx0 < -0.05 or fy0 < -0.05 or fx1 > W + 0.05 or fy1 > H + 0.05:
            issues.append("fokus slayddan chiqib ketgan")
        rx, ry, rw, rh = lay.region
        if max(fx1 - fx0, fy1 - fy0) < 0.3 * min(rw, rh):
            issues.append("fokus juda kichik")
    mx0, my0, mx1, my1 = model_box
    for k, it in labels:
        if not (mx0 - 0.1 <= it["ax"] <= mx1 + 0.1 and my0 - 0.1 <= it["ay"] <= my1 + 0.1):
            issues.append(f"yorliq chizig'i modeldan tashqarida: {k}")
    rects = [(k, (it["x"], it["y"] - it["h"] / 2, it["x"] + it["w"], it["y"] + it["h"] / 2)) for k, it in labels]
    for i in range(len(rects)):
        for j in range(i + 1, len(rects)):
            if _overlap(rects[i][1], rects[j][1], 0.02):
                issues.append(f"yorliqlar ustma-ust: {rects[i][0]} / {rects[j][0]}")
    for k, r in rects:
        for name, t in text_boxes:
            if _overlap(r, t, 0.05):
                issues.append(f"yorliq matnga tegadi: {k} / {name}")
    for name, t in text_boxes:
        if _overlap(model_box, t, 0.35):
            issues.append(f"model matnga tegadi: {name}")
    return issues


def _send_back(slide) -> None:
    """Bezaklar, panel va soya — eng orqada (fon ustida, 3D model ortida)."""
    tree = slide.shapes._spTree
    anchor = tree.find(qn("p:grpSpPr"))
    back = []
    for el in list(tree):
        names = _CNVPR_NAME(el)
        if names and (names[0].startswith("!!deco_") or names[0] in ("!!panel", "!!shadow")):
            back.append(el)
    for el in reversed(back):
        anchor.addnext(el)


def _model_name(scene_key: str, group: str, main: str) -> str:
    # Morph faqat bir xil nomli modellarni bog'laydi. Boshqa sahnadagi model (boshqa chegara qutisi)
    # bilan bog'lanmasligi uchun nomda sahna ham bor; asosiy sahnada — qisqa nom.
    return f"!!m3d_{group}" if scene_key == main else f"!!m3d_{scene_key}_{group}"


def _explode_offsets(cam: Camera, groups: dict, show: Sequence[str], frame: Frame, factor: float
                     ) -> Dict[str, Tuple[float, float]]:
    """Ajralish (portlash): har organ umumiy markazdan o'z tomoniga suriladi (dyuymda)."""
    if not factor or len(show) < 2:
        return {}
    centers = {g: cam.project(groups[g]["samples"], frame.rot)[0].mean(axis=0) for g in show}
    mid = np.mean(list(centers.values()), axis=0)
    out = {}
    for g, c in centers.items():
        d = (c - mid) * frame.size
        n = float(np.hypot(*d)) or 1.0
        push = factor * (0.6 + n)                 # yaqin turganlar ham ajralsin
        out[g] = (float(d[0] / n * push), float(d[1] / n * push))
    return out


def build(plan: dict, out_path: str, *, assets_dir: str = ASSETS_DIR, scenes: Optional[Dict[str, dict]] = None,
          share: bool = True) -> str:
    """Reja → PPTX. Slayd o'z sahnasini `scene` bilan tanlashi mumkin (sukut — rejaning `scene`).

    scenes — xotirada tayyorlangan sahnalar (avtomatik tizim): {kalit: {bbox_min, bbox_max, source,
    groups: {g: {samples, glb}}, labels}}; berilmasa — `assets_dir` dagi tayyor sahnalar o'qiladi.
    share — bir xil GLB slaydlar orasida bitta faylga yoziladi (taqdimot hajmi bir necha barobar kichik).

    Rejada `style` (design.STYLES) va har slaydda ixtiyoriy: `layout` (design.LAYOUTS), `deco` (bezak varianti),
    `explode` (organlar ajralishi, 0 — yo'q), `enter`/`exit` (uchib kirish/chiqish tomoni).
    """
    main = plan["scene"]
    st = design.style(plan.get("style"))
    seed = int(plan.get("seed") or 0)
    ready = scenes or {}
    scenes = {}
    cams: Dict[str, Camera] = {}
    glbs: Dict[Tuple[str, str], bytes] = {}
    parts: Optional[Dict] = {} if share else None

    def use(key: str) -> dict:
        if key not in scenes:
            scenes[key] = ready[key] if key in ready else load_scene(key, assets_dir)
            cams[key] = Camera(scenes[key]["bbox_min"], scenes[key]["bbox_max"])
            for g, info in scenes[key]["groups"].items():
                if "glb" in info:
                    glbs[(key, g)] = info["glb"]
                else:
                    with open(os.path.join(scenes[key]["dir"], info["file"]), "rb") as fh:
                        glbs[(key, g)] = fh.read()
        return scenes[key]

    # 1-bosqich: har slaydning sahnasi, ramkasi va ko'rinadigan organlari.
    slides = plan["slides"]
    plans = []
    # Joyida almashish (yordamchi slayd) shu slayd kamerasini oladi: unda qo'shiladigan organlar ham matnga
    # tegmasligi uchun kadr oldindan ular bilan birga hisoblanadi.
    later: Dict[int, List[str]] = {}
    for j, s in enumerate(slides):
        if s.get("frame_of") is not None:
            root = s["frame_of"]
            while slides[root].get("frame_of") is not None:
                root = slides[root]["frame_of"]
            later.setdefault(root, []).extend(s.get("show") or [])
    for index, data in enumerate(slides):
        key = data.get("scene") or main
        scene, cam = use(key), cams[key]
        groups = scene["groups"]
        lay = design.layout(data.get("layout"))
        show = [g for g in (data.get("show") or list(groups)) if g in groups]
        focus = [g for g in (data.get("focus") or show) if g in show] or show
        labels = list(data.get("labels") or [])
        region = lay.region if labels else lay.wide
        points = data.get("focus_points") or [p for g in focus for p in groups[g]["samples"]]
        if data.get("frame_of") is not None:         # kamera aynan oldingi slayddagidek (joyida almashish)
            frame = plans[data["frame_of"]]["frame"]
            lay = plans[data["frame_of"]]["layout"]
        else:
            rot = data.get("rot") or (0, 0, 0)
            everything = [p for g in show for p in groups[g]["samples"]]
            extra = [p for g in later.get(index, []) if g in groups and g not in show for p in groups[g]["samples"]]
            frame = fit_frame(cam, points, rot, region, lay.fill)
            # Tayanch suyaklar (bo'g'im slaydida) yaqinlashish chegarasini belgilamaydi: bo'g'imga yaqindan qaraladi.
            whole = fit_frame(cam, data.get("whole_points") or everything, rot, region, lay.fill)
            if frame.size > whole.size * lay.zoom_limit:     # haddan tashqari yaqin: qo'shni organlar kesiladi
                f0, g0, f1, g1 = cam.extent(points, rot)
                fx, fy = frame.to_slide((f0 + f1) / 2, (g0 + g1) / 2)
                size = whole.size * lay.zoom_limit
                frame = Frame(fx - (f0 + f1) / 2 * size, fy - (g0 + g1) / 2 * size, size, rot)
            if data.get("zoom"):                         # harakat: yaqinlashish / uzoqlashish (fokus atrofida)
                f0, g0, f1, g1 = cam.extent(points, rot)
                fx, fy = frame.to_slide((f0 + f1) / 2, (g0 + g1) / 2)
                size = min(frame.size * float(data["zoom"]), MAX_FRAME, whole.size * lay.zoom_limit * 1.3)
                frame = Frame(fx - (f0 + f1) / 2 * size, fy - (g0 + g1) / 2 * size, size, rot)
            frame = keep_clear(cam, frame, points, everything + extra, lay.area)
        show = [g for g in show if g in focus or not _too_big(cam, frame, groups[g]["samples"])]
        offsets = _explode_offsets(cam, groups, show, frame, float(data.get("explode") or 0))
        plans.append({"key": key, "scene": scene, "cam": cam, "frame": frame, "show": show, "layout": lay,
                      "labels": labels, "offsets": offsets,
                      "names": {_model_name(key, g, main): (key, g) for g in show}})

    # 2-bosqich: slaydlar. Har o'tishda hamma narsa harakatlanadi: model, matn, bezaklar; ketayotgan organ
    # slayddan uchib chiqadi, keladigani chetdan uchib kiradi, yorliqlar yon tomondan suzib kiradi.
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    built = []
    quality_log: List[Tuple[int, List[str]]] = []
    total = sum(1 for s in slides if not s.get("auto"))
    for index, (data, cur) in enumerate(zip(slides, plans)):
        key, scene, cam, frame, show = cur["key"], cur["scene"], cur["cam"], cur["frame"], cur["show"]
        lay, offsets = cur["layout"], cur["offsets"]
        slide = prs.slides.add_slide(blank)
        design.background(slide, st)
        groups = scene["groups"]
        vis_points = [p for g in show for p in groups[g]["samples"]]
        u0, v0, u1, v1 = cam.extent(vis_points, frame.rot)
        model_box = (*frame.to_slide(u0, v0), *frame.to_slide(u1, v1))
        variant = data.get("deco") if data.get("deco") is not None else seed * 31 + index
        design.draw_deco(slide, design.deco_specs(st, lay, int(variant), model_box))
        design.ground_shadow(slide, st, model_box)
        # Yordamchi (avtomatik o'tadigan) slaydlar raqamga kirmaydi: u keyingi slaydning raqamini ko'rsatadi.
        number = sum(1 for s in slides[:index] if not s.get("auto"))
        _text_column(slide, data, number, total, scene["source"], st, lay)

        placed = set(cur["names"])
        # fade — bu slaydga o'tishda organlar uchib kelmaydi, o'z joyida asta almashadi (Morph o'zi so'ndiradi).
        if index and not data.get("fade"):           # oldingi slayddan ketayotganlar
            before = plans[index - 1]
            way = (data.get("exit") or "up") if before["key"] == key else (data.get("exit") or "left")
            out = before["frame"].moved_out(way)
            for name, (k, g) in before["names"].items():
                if name not in placed:
                    add_model3d(slide, glbs[(k, g)], name, out.left, out.top, out.size, out.rot, parts,
                                lights=st.lights)
                    placed.add(name)
        if index + 1 < len(plans) and not slides[index + 1].get("fade"):   # keyingilari — hozircha tashqarida
            after = plans[index + 1]
            way = slides[index + 1].get("enter") or ("down" if after["key"] == key else "right")
            incoming = after["frame"].moved_out(way)
            for name, (k, g) in after["names"].items():
                if name not in placed:
                    add_model3d(slide, glbs[(k, g)], name, incoming.left, incoming.top, incoming.size,
                                incoming.rot, parts, lights=st.lights)
                    placed.add(name)
        # Har model alohida qatlam: PowerPoint ular orasida chuqurlikni solishtirmaydi. Uzoqdagisi avval
        # (orqaga), yaqinidagisi keyin qo'yiladi.
        def depth(item):
            k, g = item[1]
            return float(np.median(cams[k].project(scenes[k]["groups"][g]["samples"], frame.rot)[1]))
        for name, (k, g) in sorted(cur["names"].items(), key=depth):
            ox, oy = offsets.get(g, (0.0, 0.0))
            add_model3d(slide, glbs[(k, g)], name, frame.left + ox, frame.top + oy, frame.size, frame.rot, parts,
                        lights=st.lights)

        placed_labels = _labels(cam, scene, frame, cur["labels"], show, model_box, st, lay, offsets)
        if data.get("quiz"):                          # savol slaydi: nomlar o'rniga raqamlar
            for n, (_, item) in enumerate(sorted(placed_labels, key=lambda kv: (kv[1]["side"] != "left",
                                                                                kv[1]["y"])), 1):
                item["number"] = n
        steps, fly = [], []
        for i, (lkey, item) in enumerate(placed_labels):
            drawn = _draw_label(slide, st, lkey, item)
            extras = _connector(slide, st, lkey, item["start"], (item["ax"], item["ay"]))
            delay = 250 + (i // 3) * 220
            if index:                                # kartochka Morph bilan yon tomondan suzib keladi
                fly.append((lkey, item))
                steps += [(s, "fade", delay + 500) for s in extras]
            else:
                steps += [(s, "fade", delay) for s in drawn + extras]
        if fly and built:
            prev_slide, prev_names = built[-1]
            for lkey, item in fly:
                if f"lbl_{lkey}" in prev_names:      # oldingi slaydda ham bor — Morph o'zi suradi
                    continue
                dx = -(item["x"] + item["w"] + 0.4) if item["side"] == "left" else (W - item["x"] + 0.4)
                _draw_label(prev_slide, st, lkey, item, dx)
        bx, by, bw, bh = lay.body
        tx, ty, tw, th = lay.title
        f0, g0, f1, g1 = cam.extent(data.get("focus_points") or vis_points, frame.rot)
        focus_box = (*frame.to_slide(f0, g0), *frame.to_slide(f1, g1))
        issues = quality(lay, model_box, placed_labels, [("sarlavha", (tx, ty, tx + tw, ty + th)),
                                                          ("matn", (bx, by, bx + bw, by + bh))],
                         None if data.get("explode") else focus_box)
        if issues and not data.get("auto"):
            log.info("med3d sifat (%d-slayd, %s): %s", index + 1, lay.key, "; ".join(issues))
        quality_log.append((index, issues))
        if data.get("quiz") and placed_labels:
            _answers_timing(slide, steps, [(_answer_card(slide, st, lay, placed_labels), "zoom", 0)])
        elif steps:
            add_entrances(slide, steps)
        _send_back(slide)
        if index:
            # Yordamchi (joyida almashish, yig'ilish) o'tishlar qisqaroq — ketma-ket kelganda harakat sust tuyulmasin.
            add_morph(slide, HELPER_MS if data.get("auto") else MORPH_MS)
            if data.get("auto"):                     # yordamchi slayd: o'tish tugashi bilan keyingisiga o'tadi
                for node in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}transition"):
                    node.set("advClick", "0")
                    node.set("advTm", "0")
        built.append((slide, {f"lbl_{k}" for k, _ in placed_labels}))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    LAST_QUALITY[:] = quality_log
    return out_path


LAST_QUALITY: List[Tuple[int, List[str]]] = []     # oxirgi `build` ning sifat hisoboti (test va log uchun)


DEMO_PLAN = {
    "scene": "skull",
    "slides": [
        {"title": "Bosh suyagi",
         "text": "Bosh suyagi miya qutisi va yuz suyaklaridan iborat. U miyani himoya qiladi, sezgi "
                 "a'zolari va chaynov apparatiga tayanch bo'ladi. Pastki jag' — bosh suyagining yagona "
                 "harakatchan suyagi.",
         "show": ["cranium", "mandible"], "rot": [0, 0, 0],
         "labels": ["frontal", "parietal_r", "temporal_l", "zygomatic_r", "nasal", "maxilla_l", "mandible"]},
        {"title": "Pastki jag' suyagi",
         "text": "Pastki jag' chakka suyagi bilan bo'g'im hosil qiladi va chaynashda harakatlanadi. "
                 "Yuqoridan qaraganda tishlar taqasimon ravoq bo'lib joylashgan: oldinda kesuv tishlar, "
                 "yonlarida qoziq va oziq tishlar.",
         "show": ["mandible"], "rot": [70, 0, 0], "exit": "up",
         "labels": ["incisors_lower", "canine_lower_l", "premolars_lower_l", "molars_lower_r"]},
        {"title": "Yon tomondan ko'rinish",
         "text": "Qismlar qayta birlashganda bosh suyagining yon tomoni ko'rinadi: chakka suyagi, ensa "
                 "suyagi va pastki jag'ning chakka-pastki jag' bo'g'imi.",
         "show": ["cranium", "mandible"], "rot": [0, 90, 0],
         "labels": ["frontal", "parietal_r", "temporal_r", "occipital", "zygomatic_r", "mandible"]},
    ],
}
