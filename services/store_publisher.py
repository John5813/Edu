"""Tayyor ishni mijoz izidan tozalab, do'kon katalogiga qo'yadi.

Oqim: tozalash → ko'rgazma rasmlari → yopiq Telegram kanaliga yuklash →
katalogga yozuv. Faylning o'zi hech qachon saytda turmaydi; sayt faqat
rasmlarni ko'rsatadi, fayl esa kanaldan `file_id` orqali olinadi.
"""

import logging
import os
import re
import shutil
import uuid

from config import STORE_PREVIEW_DIR, STORE_PREVIEW_LIMIT, STORE_VAULT_CHAT_ID, TEMP_DIR

logger = logging.getLogger(__name__)

# "Tayyorladi:", "Автор —", "Prepared by" kabi satrlar. Mijoz ismi noma'lum
# bo'lganda ham shu satrlar orqali topiladi.
_LABEL = (
    r"tayyorladi|bajardi|topshirdi|muallif|ism[\s\-]*familiya|ismi|talaba|"
    r"выполнил(?:а)?|подготовил(?:а)?|автор|студент|"
    r"prepared\s+by|author|submitted\s+by|student"
)
_LABEL_ONLY_RE = re.compile(rf"^\s*(?:{_LABEL})\s*[:\-–—]?\s*$", re.IGNORECASE)
_LABEL_VALUE_RE = re.compile(rf"^\s*(?:{_LABEL})\s*[:\-–—]\s*(.+)$", re.IGNORECASE)


def _looks_like_person(text: str) -> bool:
    """Ism-familiyaga o'xshaydimi.

    Kitob muallifi haqidagi jumla tasodifan o'chib ketmasligi uchun:
    faqat raqamsiz, tinish belgisiz, 1-4 so'zli qisqa matn ism deb qaraladi.
    """
    value = text.strip(" .,;")
    if not value or len(value) > 60:
        return False
    if any(ch.isdigit() for ch in value):
        return False
    words = value.split()
    return 1 <= len(words) <= 4 and all(w.replace("'", "").replace("`", "").isalpha() for w in words)


def _iter_shapes(shapes):
    """Guruhlangan shakllar ichidagilarni ham qaytaradi."""
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            try:
                yield from _iter_shapes(shape.shapes)
            except Exception:
                pass


def _blank_paragraph(para) -> None:
    for run in para.runs:
        run.text = ""


def _blank_paragraph_all(text_frame) -> None:
    for para in text_frame.paragraphs:
        _blank_paragraph(para)


def _scrub_text_frame(text_frame, name_re) -> int:
    """Mijozga ishora qiluvchi satrlarni bo'shatadi, sonini qaytaradi."""
    removed = 0
    after_label = False
    for para in text_frame.paragraphs:
        text = para.text
        if not text.strip():
            continue
        label_only = bool(_LABEL_ONLY_RE.match(text))
        hit = False
        if name_re is not None and name_re.search(text):
            hit = True
        elif label_only:
            hit = True
        else:
            match = _LABEL_VALUE_RE.match(text)
            if match and _looks_like_person(match.group(1)):
                hit = True
            elif after_label and _looks_like_person(text):
                # "Tayyorladi:" alohida satrda, ism esa undan keyin turadi.
                hit = True
        if hit:
            _blank_paragraph(para)
            removed += 1
        after_label = label_only
    return removed


def _scrub_table_row(cells, name_re) -> int:
    """Jadval satrini tozalaydi.

    "Talaba | Ism Familiya" ko'rinishida yorliq bir katakda, ism yonidagida
    turadi — shuning uchun kataklar bir-biriga qarab baholanadi.
    """
    before = [cell.text for cell in cells]
    removed = sum(_scrub_text_frame(cell.text_frame, name_re) for cell in cells)
    if not any(_LABEL_ONLY_RE.match(text) for text in before):
        return removed
    for cell, text in zip(cells, before):
        if _LABEL_ONLY_RE.match(text):
            continue
        if cell.text.strip() and _looks_like_person(cell.text):
            _blank_paragraph_all(cell.text_frame)
            removed += 1
    return removed


def anonymize_pptx(src_path: str, customer_name: str = "", out_path: str = "") -> str:
    """Nusxani tozalab, yangi fayl yo'lini qaytaradi. Asl fayl tegilmaydi."""
    from pptx import Presentation

    if not out_path:
        out_path = os.path.join(TEMP_DIR, f"store_{uuid.uuid4().hex[:8]}.pptx")
        os.makedirs(TEMP_DIR, exist_ok=True)

    name_re = None
    cleaned_name = (customer_name or "").strip()
    if len(cleaned_name) >= 3:
        # Ism bo'laklari alohida ham qidiriladi: taqdimotda faqat familiya
        # yozilgan bo'lishi mumkin.
        parts = [re.escape(p) for p in cleaned_name.split() if len(p) >= 3]
        if parts:
            name_re = re.compile("|".join(parts), re.IGNORECASE)

    prs = Presentation(src_path)

    removed = 0
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                removed += _scrub_text_frame(shape.text_frame, name_re)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    removed += _scrub_table_row(list(row.cells), name_re)
        # Slayd izohlari — ekranda ko'rinmaydi, shuning uchun eng oson
        # e'tibordan chetda qoladigan joy.
        if getattr(slide, "has_notes_slide", False):
            try:
                _blank_paragraph_all(slide.notes_slide.notes_text_frame)
            except Exception:
                pass

    # Fayl xossalari: Word/PowerPoint "Muallif" maydonini o'zi to'ldiradi.
    props = prs.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.category = ""
    props.keywords = ""
    props.subject = ""
    props.identifier = ""
    props.revision = 1

    prs.save(out_path)
    logger.info("Tozalandi: %s → %s (%s satr olib tashlandi)",
                os.path.basename(src_path), os.path.basename(out_path), removed)
    return out_path


def render_previews(pptx_path: str, public_code: str, limit: int = 0) -> int:
    """Birinchi slaydlarni JPG qilib saqlaydi, nechtasi saqlangani qaytadi.

    Faqat bir qismi chiqariladi: ko'rgazma xaridning o'rnini bosmasligi kerak.
    """
    from PIL import Image

    from services.premium_presentation.qa import discard_images, pptx_to_images

    limit = limit or STORE_PREVIEW_LIMIT
    out_dir = os.path.join(STORE_PREVIEW_DIR, public_code)
    os.makedirs(out_dir, exist_ok=True)

    images = pptx_to_images(pptx_path)
    if not images:
        logger.error("Ko'rgazma rasmlari yaratilmadi: %s", pptx_path)
        return 0

    saved = 0
    try:
        for source in images[:limit]:
            try:
                with Image.open(source) as img:
                    img = img.convert("RGB")
                    if img.width > 1000:
                        height = round(img.height * 1000 / img.width)
                        img = img.resize((1000, height), Image.LANCZOS)
                    saved += 1
                    img.save(os.path.join(out_dir, f"{saved}.jpg"),
                             format="JPEG", quality=82, optimize=True)
            except Exception as exc:
                logger.warning("Ko'rgazma rasmi tayyorlanmadi (%s): %s", source, exc)
    finally:
        discard_images(images)

    return saved


def discard_previews(public_code: str) -> None:
    """Katalogdan olib tashlangan ishning rasmlarini o'chiradi."""
    shutil.rmtree(os.path.join(STORE_PREVIEW_DIR, public_code), ignore_errors=True)


async def publish_pptx(
    bot,
    pptx_path: str,
    title: str,
    price: int,
    *,
    customer_name: str = "",
    description: str = "",
    category: str = "",
    language: str = "uz",
    keywords: str = "",
) -> dict:
    """Ishni tozalab, omborga yuklab, katalogga qo'yadi.

    Katalogga yozuv faqat fayl omborga yetib borgandan keyin qo'shiladi —
    aks holda saytda yuklab bo'lmaydigan ish paydo bo'lardi.
    """
    from aiogram.types import FSInputFile
    from pptx import Presentation

    from database.database import Database

    if not STORE_VAULT_CHAT_ID:
        raise RuntimeError(
            "STORE_VAULT_CHAT_ID sozlanmagan — ombor kanali ko'rsatilishi kerak"
        )

    public_code = await Database.generate_store_code()
    cleaned_path = anonymize_pptx(pptx_path, customer_name)

    try:
        slide_count = len(Presentation(cleaned_path).slides)
    except Exception:
        slide_count = 0

    try:
        preview_count = render_previews(cleaned_path, public_code)
        if not preview_count:
            raise RuntimeError("ko'rgazma rasmlari tayyorlanmadi")

        message = await bot.send_document(
            chat_id=STORE_VAULT_CHAT_ID,
            document=FSInputFile(cleaned_path, filename=f"{public_code}.pptx"),
            caption=f"{public_code} — {title}",
            parse_mode=None,
        )
        if not message.document:
            raise RuntimeError("ombor javobida hujjat yo'q")

        await Database.create_store_item(
            public_code=public_code,
            title=title,
            file_id=message.document.file_id,
            price=price,
            description=description,
            category=category,
            language=language,
            keywords=keywords,
            slide_count=slide_count,
            preview_count=preview_count,
        )
    except Exception:
        # Yarim qolgan nashrdan rasm qolib ketmasin.
        discard_previews(public_code)
        raise
    finally:
        if cleaned_path != pptx_path:
            try:
                os.remove(cleaned_path)
            except OSError:
                pass

    logger.info("Katalogga qo'shildi: %s — %s", public_code, title)
    return {
        "public_code": public_code,
        "title": title,
        "price": price,
        "slide_count": slide_count,
        "preview_count": preview_count,
    }
