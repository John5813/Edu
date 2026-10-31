"""3D taqdimotni rejadan yig'adi: uslub, fon bezaklari, katalogdagi haqiqiy 3D modellar va Morph.

Morph sahnasi: bir model qo'shni slaydlarda turganda (nomi `!!m_<id>`) PowerPoint uni yangi joyiga suzib
olib o'tadi va aylantiradi (har ko'rinishda boshqa burchak); sarlavha (`!!title`), bo'lim nomi (`!!kicker`),
fon bezaklari (`!!deco_N`) va model soyasi ham birga suriladi. Matn bandlari, raqam kartalari esa o'tishdan
keyin navbat bilan chiqadi.

Har varoq boshqacha: 14 kompozitsiya (paragraflar, vaqt chizig'i, taqqoslash, kartalar, bosqichlar, iqtibos,
savol, raqamlar, katta model, uchlik, ro'yxat...). Shrift joyga moslanadi; sig'magani ma'ruzachi izohiga o'tadi.

Katalog modellari aniq sxema emas — qismlari yorliq-chiziq bilan ko'rsatilmaydi (3D Anatomiyadan farqi).
Telefon va 3D ni bilmaydigan dasturlar uchun har modelning o'rnida uning rasmi (fon olib tashlangan) turadi.
"""
import io
import logging
import os
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from services import katalog3d
from services.med3d import design, glb
from services.med3d.pptx3d import CT_GLB, RT_MODEL3D, _PNG, _model_xml, _next_id
from services.premium_presentation.slide_anim import P_NS
from services.pro3d.build import add_entrances, add_morph, fit_size

from .uslub import STYLES

log = logging.getLogger(__name__)

W, H = design.W, design.H
MORPH_MS = 1700
TURN_MS = int(os.getenv("KAT3D_TURN_SECONDS") or 40) * 1000   # slaydda turgan model o'z joyida shuncha vaqtda 1 marta aylanadi
BASE_ROT = (14.0, -30.0, 0.0)              # preview rasmidagi burchakka yaqin: biroz yuqoridan, yon tomondan
SWING = (0.0, 28.0, -22.0, 16.0, -32.0, 10.0, 34.0, -14.0)   # har keyingi ko'rinishda Y bo'yicha burilish
NOTES_HEAD = {"uz": "3D modellar mualliflari", "ru": "Авторы 3D-моделей", "en": "3D model credits",
              "kk": "3D модельдер авторлары"}


class Ctx:
    """Bitta taqdimot yig'ilishi: uslub, GLB va zaxira rasmlar keshi, har model necha marta chiqqani."""

    def __init__(self, prs, st, plan: dict):
        self.prs, self.st, self.plan = prs, st, plan
        self.parts: Dict[str, Part] = {}
        self.data: Dict[str, Tuple[bytes, np.ndarray, np.ndarray]] = {}
        self.fallback: Dict[str, bytes] = {}
        self.seen: Dict[str, int] = {}
        self.cur = 0
        self.placed: Dict[int, Dict[str, tuple]] = {}      # slayd → {model: (cx, cy, size, rot)}
        self.ghosts: Dict[int, set] = {}
        self.spins: Dict[int, List[int]] = {}               # slayd → o'z joyida aylanadigan 3D ob'ektlar id si
        self.credits: List[str] = []
        self.total = len(plan["slides"])

    def glb(self, mid: str):
        if mid not in self.data:
            path = katalog3d.model_yoli(mid)
            if not path:
                raise FileNotFoundError(f"3D model yo'q: {mid}")
            with open(path, "rb") as fh:
                raw = fh.read()
            lo, hi = glb.bounds(raw)
            self.data[mid] = (raw, lo, hi)
        return self.data[mid]

    def picture(self, mid: str) -> bytes:
        if mid not in self.fallback:
            self.fallback[mid] = _cutout(katalog3d.preview_yoli(mid), katalog3d.maska_yoli(mid))
        return self.fallback[mid]


def _cutout(path: Optional[str], mask: Optional[str] = None) -> bytes:
    """Preview rasmidan shaffof PNG — 3D ko'rinmaydigan dasturda (telefon, WPS) model o'rnida turadi.

    Asosiy yo'l — renderer chizgan shaffoflik niqobi (`mask/<id>.png`, xuddi shu burchak va kadr). Niqob
    bo'lmasa: fon va soya rasm chetidan "suv quyish" bilan olib tashlanadi."""
    if not path:
        return _PNG
    from PIL import Image, ImageDraw, ImageFilter

    with Image.open(path) as im:
        rgb = np.asarray(im.convert("RGB")).astype(np.int16)
    h, w = rgb.shape[:2]
    bg = np.array([242, 242, 240], dtype=np.int16)
    if mask:
        with Image.open(mask) as m:
            alpha = np.asarray(m.convert("L").resize((w, h)))
        a = alpha.astype(np.float32)[..., None] / 255.0
        # chetdagi piksellardan kulrang fon "aralashmasi" ayiriladi — to'q fonda oq hoshiya qolmaydi
        clean = np.where(a > 0.04, (rgb - (1.0 - a) * bg) / np.maximum(a, 0.04), rgb)
        out = Image.fromarray(np.dstack([np.clip(clean, 0, 255).astype(np.uint8), alpha]), "RGBA")
        buf = io.BytesIO()
        out.save(buf, "PNG", optimize=True)
        return buf.getvalue()
    dist = np.abs(rgb - bg).max(axis=2)
    chroma = rgb.max(axis=2) - rgb.min(axis=2)
    cand = (dist <= 55) & (chroma <= 12)
    cand[:max(6, h // 48), :] = True                 # tepadagi yuklanish chizig'i ham fon
    mask = Image.fromarray(np.where(cand, 255, 0).astype(np.uint8), "L").copy()   # copy: numpy xotirasi o'qishga'
    for xy in ((0, 0), (w - 1, 0), (0, h - 1), (w - 1, h - 1), (w // 2, 0), (w // 2, h - 1)):
        if mask.getpixel(xy) == 255:
            ImageDraw.floodfill(mask, xy, 128)
    alpha = np.where(np.asarray(mask) == 128, 0, 255).astype(np.uint8)
    alpha = np.asarray(Image.fromarray(alpha, "L").filter(ImageFilter.GaussianBlur(0.8)))
    out = Image.fromarray(np.dstack([rgb.astype(np.uint8), alpha]), "RGBA")
    buf = io.BytesIO()
    out.save(buf, "PNG", optimize=True)
    return buf.getvalue()


# ───────────────────────────────────────────── chizish

def _place(ctx: Ctx, slide, mid: str, cx: float, cy: float, size: float, rot) -> int:
    """3D ob'ekt (am3d:model3d) — GLB taqdimotda bir marta, hamma slaydlar unga havola qiladi. Ob'ekt id si."""
    raw, lo, hi = ctx.glb(mid)
    pkg = slide.part.package
    part = ctx.parts.get(mid)
    if part is None:
        part = Part(pkg.next_partname("/ppt/media/model3d%d.glb"), CT_GLB, pkg, raw)
        ctx.parts[mid] = part
    rid_model = slide.part.relate_to(part, RT_MODEL3D)
    _, rid_png = slide.part.get_or_add_image_part(io.BytesIO(ctx.picture(mid)))
    x, y = cx - size / 2, cy - size / 2
    tree = slide.shapes._spTree
    sid = _next_id(tree)
    node = etree.fromstring(_model_xml(sid, "!!m_" + mid[:24], int(Inches(x)), int(Inches(y)),
                                       int(Inches(size)), int(Inches(size)), rid_model, rid_png, lo, hi, rot,
                                       ctx.st.lights))
    for sppr in node.iter(qn("p:spPr")):          # zaxira ko'rinish ramkasiz
        ln = etree.SubElement(sppr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
    tree.append(node)
    return sid


def _model(ctx: Ctx, slide, mid: str, cx: float, cy: float, size: float, turn: float = 0.0) -> Tuple:
    """3D modelni (cx, cy) markazli `size` dyuymli kvadrat ramkaga qo'yadi. Qaytaradi: (chap, tepa, o'ng, past).
    Har keyingi ko'rinishda model boshqa burchakda — Morph uni slayddan slaydga aylantirib olib o'tadi."""
    k = ctx.seen.get(mid, 0)
    ctx.seen[mid] = k + 1
    rot = (BASE_ROT[0], BASE_ROT[1] + SWING[k % len(SWING)] + turn, BASE_ROT[2])
    ctx.spins.setdefault(ctx.cur, []).append(_place(ctx, slide, mid, cx, cy, size, rot))
    ctx.placed.setdefault(ctx.cur, {})[mid] = (cx, cy, size, rot)
    m = katalog3d.katalog()[mid]
    if m["attribution_full"] not in ctx.credits:
        ctx.credits.append(m["attribution_full"])
    x, y = cx - size / 2, cy - size / 2
    return (x + size * 0.12, y + size * 0.12, x + size * 0.88, y + size * 0.88)


# Slayddan ketayotgan model keyingi slaydda chetdan tashqarida, aylangan va kichraygan holda turadi — Morph
# uni aylantirib uchirib chiqaradi; yangi model esa oldingi slaydda chetda turadi — aylanib uchib kiradi.
SPIN_OUT = (24.0, 150.0, 18.0)
SPIN_IN = (-18.0, -135.0, -12.0)


def _offstage(cx: float, cy: float, size: float, side: str, scale: float) -> Tuple[float, float, float]:
    s = size * scale
    if side == "right":
        return W + s * 0.7, cy + 0.6, s
    if side == "left":
        return -s * 0.7, cy + 0.6, s
    if side == "down":
        return cx + 0.8, H + s * 0.7, s
    return cx - 0.8, -s * 0.7, s


def _choreography(ctx: Ctx, slides) -> int:
    """Kirish/chiqish "arvohlari": qo'shni slaydda yo'q modellar uchun slayddan tashqaridagi nusxa."""
    added = 0
    n = len(slides)
    for i in range(n - 1):
        here, nxt = ctx.placed.get(i, {}), ctx.placed.get(i + 1, {})
        ghosts_next = ctx.ghosts.setdefault(i + 1, set())
        ghosts_here = ctx.ghosts.setdefault(i, set())
        for k, (mid, (cx, cy, size, rot)) in enumerate(here.items()):
            if mid in nxt or mid in ghosts_next:
                continue
            side = ("right" if cx > W / 2 else "left") if (i + k) % 3 else ("down" if cy > H / 2 else "up")
            gx, gy, gs = _offstage(cx, cy, size, side, 0.55)
            spin = tuple(r + d for r, d in zip(rot, SPIN_OUT))
            _place(ctx, slides[i + 1], mid, gx, gy, gs, spin)
            ghosts_next.add(mid)
            added += 1
        for k, (mid, (cx, cy, size, rot)) in enumerate(nxt.items()):
            if mid in here or mid in ghosts_here:
                continue
            side = "right" if cx > W / 2 else "left"
            side = side if (i + k) % 2 else ("down" if cy > H / 2 else "up")
            gx, gy, gs = _offstage(cx, cy, size, side, 0.7)
            spin = tuple(r + d for r, d in zip(rot, SPIN_IN))
            _place(ctx, slides[i], mid, gx, gy, gs, spin)
            ghosts_here.add(mid)
            added += 1
    return added


def _caption(ctx: Ctx, slide, mids: Sequence[str], x: float, y: float, w: float, align=PP_ALIGN.LEFT,
             name: str = "credit") -> None:
    """Muallif va litsenziya (CC BY uchun majburiy) — kichik, xira yozuv."""
    cat = katalog3d.katalog()
    text = "   ".join(dict.fromkeys(cat[m]["attribution"] for m in mids))
    if text:
        design.textbox(slide, name, x, y, w, 0.26, text, size=8, color=ctx.st.muted, font=ctx.st.body_font,
                       align=align, smallest=6.5)


def _bullets(ctx: Ctx, slide, name: str, points: Sequence[str], x, y, w, h, size: float, smallest: float,
             color: Optional[str] = None, fit: bool = True):
    """Haqiqiy PowerPoint ro'yxati (belgisi urg'u rangida). `fit` — shrift qutiga moslanadi (aks holda
    `size` o'zi, oldindan `_fit_points` bilan hisoblangan)."""
    st = ctx.st
    if fit:
        size = fit_size("\n".join(points), w - 0.3, h - 0.12 * len(points), size, smallest, 1.24)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, text in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(size * 0.75)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.3))))
        pPr.set("indent", str(-int(Inches(0.3))))
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", st.accent)
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = st.body_font
        r.font.color.rgb = design.rgb(color or st.body)
    return tb


HEAD_EM = 0.6                 # sarlavha shriftida o'rtacha harf eni (em) — Georgia/Segoe Semibold qalin
BODY_EM = 0.5                 # Segoe UI ~0.46 em + so'z ko'chishidagi bo'sh joy
SERIF_EM = 0.53               # Georgia matni kengroq
VALUE_EM = 0.7                # qalin raqamlar (ayniqsa serif) o'rtacha harfdan kengroq


def _lines(text: str, w: float, size: float, em: float) -> int:
    per_line = max(1, int(w * 72 / (size * em)))
    words, lines, cur = str(text or "").split(), 1, 0
    for word in words:
        add = len(word) + (1 if cur else 0)
        if cur and cur + add > per_line:
            lines, cur = lines + 1, len(word)
        else:
            cur += add
    return lines


def _fit(text: str, w: float, h: float, start: float, smallest: float, em: float, line: float = 1.08) -> float:
    size = start
    while size > smallest and _lines(text, w, size, em) * size * line / 72 > h:
        size -= 1
    return size


def _kicker(ctx: Ctx, slide, text: str, x, y, w, align=PP_ALIGN.LEFT):
    return design.textbox(slide, "!!kicker", x, y, w, 0.36, text, size=13, color=ctx.st.accent,
                          font=ctx.st.head_font, bold=True, caps=True, align=align, smallest=10)


def _title(ctx: Ctx, slide, text: str, x, y, w, h, size: float, smallest: float, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP) -> float:
    """Sarlavha; qaytaradi — haqiqiy pastki cheti (dyuym), undan keyingi matn shunga qarab suriladi."""
    st = ctx.st
    text = text.upper() if st.title_caps else text
    em = HEAD_EM * (1.12 if st.title_caps else 1.0)
    size = _fit(text, w, h, size, smallest, em)
    design.textbox(slide, "!!title", x, y, w, h, text, size=size, color=st.ink, font=st.head_font,
                   bold=True, align=align, anchor=anchor, line=1.02)
    return y + _lines(text, w, size, em) * size * 1.12 / 72


def _number(ctx: Ctx, slide, index: int) -> None:
    design.textbox(slide, "!!num", W - 1.55, 0.46, 0.95, 0.3, f"{index + 1:02d} / {ctx.total:02d}", size=10,
                   color=ctx.st.muted, font=ctx.st.body_font, align=PP_ALIGN.RIGHT)


def _panel(ctx: Ctx, slide, name: str, x, y, w, h):
    st = ctx.st
    return design.shape(slide, name, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, st.panel,
                        alpha=st.panel_alpha, round_=0.07)


def _decorate(ctx: Ctx, slide, index: int, box) -> None:
    design.background(slide, ctx.st)
    design.draw_deco(slide, design.deco_specs(ctx.st, None, index + 1, box))


# ───────────────────────────────────────────── kompozitsiyalar

def _muqova(ctx: Ctx, slide, s: dict, i: int) -> List:
    mids = s["models"]
    _decorate(ctx, slide, i, (6.6, 0.6, 12.4, 6.8))
    box = _model(ctx, slide, mids[0], 9.55, 3.7, 6.3)
    design.ground_shadow(slide, ctx.st, box)
    if len(mids) > 1:
        _model(ctx, slide, mids[1], 11.75, 5.85, 2.0, turn=20)        # ramkali uslublarda ham ichkarida
    _kicker(ctx, slide, s["kicker"] or ctx.plan.get("topic", ""), 0.8, 1.85, 6.0)
    bottom = _title(ctx, slide, s["title"] or ctx.plan.get("topic", ""), 0.8, 2.3, 6.2, 2.75, 54, 30)
    if s["subtitle"]:
        top = max(4.6, bottom + 0.35)
        design.textbox(slide, "sub", 0.8, top, 5.9, 6.75 - top, s["subtitle"], size=19, color=ctx.st.body,
                       font=ctx.st.body_font, smallest=13)
    _caption(ctx, slide, mids, 0.8, H - 0.5, 6.2)
    return []


def _em(font: str) -> float:
    return SERIF_EM if "Georgia" in font else BODY_EM


def _body_h(points: Sequence[str], w: float, size: float, em: float = BODY_EM) -> float:
    """`_bullets` ro'yxatining taxminiy balandligi (dyuym): qator oralig'i 1.12, punktdan keyin 0.75 qator."""
    return sum(_lines(p, w - 0.3, size, em) * size * 1.12 * 1.22 + size * 0.75 for p in points) / 72


def _para_h(text: str, w: float, size: float, em: float = BODY_EM) -> float:
    return _lines(text, w, size, em) * size * 1.15 * 1.22 / 72 if text else 0.0


def _fit_points(s: dict, points: Sequence[str], w: float, room: float, body: Tuple[float, float],
                lead: str = "", lead_w: float = 0.0, em: float = BODY_EM) -> Tuple[float, List[str]]:
    """Kirish matni va punktlar `room` balandlikka sig'adigan shrift. Eng kichik shriftda ham sig'masa —
    oxirgi punktlar ma'ruzachi izohiga ko'chadi (matn yo'qolmaydi). (o'lcham, punktlar) qaytaradi."""
    points = list(points)

    def need(size: float) -> float:
        h = _body_h(points, w, size, em) if points else 0.0
        return h + (_para_h(lead, lead_w or w, size * 1.06, em) + 0.18 if lead else 0.0)

    size = body[0]
    while size > body[1] and need(size) > room:
        size -= 0.5
    while len(points) > 1 and need(size) > room:
        moved = points.pop()
        s["notes"] = ((s.get("notes") or "") + "\n• " + moved).strip()
    return size, points


def _lead(ctx: Ctx, slide, i: int, text: str, x, y, w, size: float):
    h = _para_h(text, w, size, _em(ctx.st.body_font))
    return design.textbox(slide, f"lead_{i}", x, y, w, h + 0.05, text, size=size, color=ctx.st.ink,
                          font=ctx.st.body_font, line=1.15), h


def _stack(ctx: Ctx, slide, s: dict, i: int, x: float, w: float, head: Tuple[float, float],
           body: Tuple[float, float], points: Sequence[str], panel: bool = False, cy: float = 3.75,
           top: float = 0.75, bottom: float = 6.95) -> List:
    """Bo'lim nomi + sarlavha + kirish matni + ro'yxat bitta blok: model markaziga (cy) nisbatan vertikal
    markazda; shrift mavjud joyga moslanadi. `panel` — ro'yxat yarim shaffof kartada."""
    st = ctx.st
    title = s["title"].upper() if st.title_caps else s["title"]
    em = HEAD_EM * (1.12 if st.title_caps else 1.0)
    tsize = _fit(title, w, 2.2, head[0], head[1], em)
    th = _lines(title, w, tsize, em) * tsize * 1.12 / 72
    pad = 0.3 if panel else 0.0
    bw = w - 2 * pad
    gap = 0.4 if panel else 0.32
    lead = s.get("subtitle") or ""
    room = bottom - top - 0.4 - th - gap - (2 * pad * 0.8 if panel and points else 0.0)
    bem = _em(st.body_font)
    size, points = _fit_points(s, points, bw, room, body, lead, w, bem)
    lsize = size * 1.06
    lh = _para_h(lead, w, lsize, bem) + 0.18 if lead else 0.0
    bh = _body_h(points, bw, size, bem) if points else 0.0
    total = 0.4 + th + gap + lh + (bh + 2 * pad * 0.8 if points else 0.0)
    y0 = min(max(top, cy - total / 2), max(top, bottom - total))
    _kicker(ctx, slide, s["kicker"], x, y0, w)
    y = _title(ctx, slide, s["title"], x, y0 + 0.4, w, 2.2, tsize, tsize) + gap
    steps = []
    if lead:
        box, _ = _lead(ctx, slide, i, lead, x, y, w, lsize)
        steps.append((box, "fade", 150))
        y += lh
    if points:
        if panel:
            steps.append((_panel(ctx, slide, f"panel_{i}", x - 0.15, y, w + 0.3, bh + 2 * pad * 0.8), "fade", 250))
            y += pad * 0.8
        box = _bullets(ctx, slide, f"body_{i}", points, x + pad, y, bw, bh + 0.1, size, body[1], fit=False)
        steps.append((box, "float", 400 if panel else 300))
    return steps


def _side(ctx: Ctx, slide, s: dict, i: int, left_text: bool) -> List:
    tx, tw = (0.7, 6.2) if left_text else (6.4, 6.25)
    mx = 10.25 if left_text else 3.1
    _decorate(ctx, slide, i, (mx - 2.5, 1.1, mx + 2.5, 6.5))
    box = _model(ctx, slide, s["models"][0], mx, 3.8, 5.2)
    design.ground_shadow(slide, ctx.st, box)
    steps = _stack(ctx, slide, s, i, tx, tw, (34, 24), (18, 12), s["points"][:5])
    _caption(ctx, slide, s["models"][:1], mx - 2.6, H - 0.5, 5.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _qahramon(ctx: Ctx, slide, s: dict, i: int) -> List:
    _decorate(ctx, slide, i, (6.2, 0.4, 12.9, 7.1))
    box = _model(ctx, slide, s["models"][0], 9.6, 3.75, 6.8)
    design.ground_shadow(slide, ctx.st, box)
    steps = _stack(ctx, slide, s, i, 0.75, 5.3, (42, 26), (17, 11.5), s["points"][:3], panel=True)
    _caption(ctx, slide, s["models"][:1], 6.6, H - 0.45, 6.2, align=PP_ALIGN.RIGHT)
    _number(ctx, slide, i)
    return steps


def _raqamlar(ctx: Ctx, slide, s: dict, i: int) -> List:
    st = ctx.st
    _decorate(ctx, slide, i, (0.6, 2.2, 5.4, 6.9))
    box = _model(ctx, slide, s["models"][0], 3.0, 4.6, 4.5)
    design.ground_shadow(slide, st, box)
    _kicker(ctx, slide, s["kicker"], 0.75, 0.75, 7.5)
    bottom = _title(ctx, slide, s["title"], 0.75, 1.15, 11.8, 1.2, 34, 24)
    stats = s["stats"][:4]
    n = len(stats)
    cols, gap = 2, 0.25
    rows = 1 if n <= 2 else 2
    x0, cw_all = 5.85, 6.95
    cw = (cw_all - gap) / cols
    if s["points"]:                                  # ostida izoh punktlari — kartalar ixchamroq
        ch = 1.25 if rows == 2 else 1.5
    else:
        ch = 2.3 if rows == 1 else 1.9
    top = max(2.15, bottom + 0.3)
    steps = []
    for k, item in enumerate(stats):
        r, c = divmod(k, cols)
        x, y = x0 + c * (cw + gap), top + r * (ch + gap)
        w = cw_all if (n == 3 and k == 2) else cw      # uchta raqamda oxirgisi butun qatorga
        card = _panel(ctx, slide, f"card_{i}_{k}", x, y, w, ch)
        one_line = (w - 0.5) * 72 / (max(1, len(item["value"])) * VALUE_EM)
        val = design.textbox(slide, f"val_{i}_{k}", x + 0.25, y + 0.14, w - 0.5, ch * 0.5, item["value"],
                             size=max(18, min(40, int(one_line), int(ch * 0.5 * 72 / 1.25))), color=st.accent,
                             font=st.head_font, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        lab = design.textbox(slide, f"lab_{i}_{k}", x + 0.25, y + 0.2 + ch * 0.5, w - 0.5, ch * 0.42,
                             item["label"], size=13, color=st.body, font=st.body_font, smallest=10)
        steps += [(card, "zoom", 200 + 180 * k), (val, "fade", 260 + 180 * k), (lab, "fade", 300 + 180 * k)]
    if s["points"]:
        y = top + rows * ch + (rows - 1) * gap + 0.3
        size, points = _fit_points(s, s["points"][:3], cw_all, 6.95 - y, (16, 11), em=_em(st.body_font))
        steps.append((_bullets(ctx, slide, f"body_{i}", points, x0, y, cw_all, 6.95 - y, size, 11, fit=False),
                      "float", 300 + 180 * n))
    _caption(ctx, slide, s["models"][:1], 0.75, H - 0.45, 4.6, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _uchlik(ctx: Ctx, slide, s: dict, i: int) -> List:
    st = ctx.st
    _decorate(ctx, slide, i, (0.8, 1.8, 12.5, 5.4))
    _kicker(ctx, slide, s["kicker"], 0.75, 0.75, 11.8)
    _title(ctx, slide, s["title"], 0.75, 1.15, 11.8, 1.0, 32, 22)
    steps = []
    for k, item in enumerate(s["items"][:3]):
        cx = 2.45 + k * 4.215
        box = _model(ctx, slide, item["model"], cx, 3.35, 3.3, turn=-10 + 10 * k)
        if k == 1:
            design.ground_shadow(slide, st, box)
        name = design.textbox(slide, f"name_{i}_{k}", cx - 1.95, 5.0, 3.9, 0.42, item["name"], size=19,
                              color=st.accent, font=st.head_font, bold=True, align=PP_ALIGN.CENTER, smallest=13)
        steps.append((name, "fade", 200 + 200 * k))
        if item["text"]:
            txt = design.textbox(slide, f"text_{i}_{k}", cx - 1.95, 5.45, 3.9, 1.5, item["text"], size=15,
                                 color=st.body, font=st.body_font, align=PP_ALIGN.CENTER, smallest=10.5)
            steps.append((txt, "fade", 280 + 200 * k))
    _caption(ctx, slide, [it["model"] for it in s["items"][:3]], 0.75, H - 0.4, 11.0, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _jadval(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Taqqoslash jadvali: sarlavha qatori urg'u rangida, qatorlar alohida kartalar (navbat bilan chiqadi);
    «belgilangan» qator (`table.mark`, odatda O'zbekiston) ramka bilan ajratiladi. Model o'ngda."""
    st = ctx.st
    _decorate(ctx, slide, i, (9.7, 2.0, 12.9, 6.2))
    bottom = _head(ctx, slide, s, 0.75, 11.8, (34, 24), 1.2)
    box = _model(ctx, slide, s["models"][0], 11.35, 4.4, 3.4)
    design.ground_shadow(slide, st, box)
    tab = s["table"]
    head, rows = tab["head"], tab["rows"]
    mark = tab.get("mark", -1)
    x0, wall = 0.6, 9.15
    top = max(2.25, bottom + 0.35)
    n = len(rows)
    rh = min(0.95, (6.95 - top - 0.62) / max(n, 1) - 0.1)
    first = 0.3 if len(head) > 3 else 0.34                   # birinchi ustun ulushi (nomlar)
    widths = [wall * first] + [wall * (1 - first) / (len(head) - 1)] * (len(head) - 1)
    steps = []
    hb = design.shape(slide, f"thead_{i}", MSO_SHAPE.ROUNDED_RECTANGLE, x0, top, wall, 0.52, st.accent, round_=0.25)
    steps.append((hb, "fade", 150))
    x = x0
    for k, (h, w) in enumerate(zip(head, widths)):
        steps.append((design.textbox(slide, f"th_{i}_{k}", x + 0.18, top + 0.06, w - 0.3, 0.4, h, size=15, color=st.bg,
                                     font=st.head_font, bold=True, anchor=MSO_ANCHOR.MIDDLE, smallest=11), "fade", 180))
        x += w
    for r, row in enumerate(rows):
        y = top + 0.62 + r * (rh + 0.1)
        hi = r == mark
        card = design.shape(slide, f"trow_{i}_{r}", MSO_SHAPE.ROUNDED_RECTANGLE, x0, y, wall, rh, st.panel,
                            alpha=st.panel_alpha, round_=0.18, line=st.accent if hi else None, line_pt=2.0 if hi else 0)
        steps.append((card, "fade", 300 + 160 * r))
        x = x0
        for k, (cell, w) in enumerate(zip(row, widths)):
            tb = design.textbox(slide, f"td_{i}_{r}_{k}", x + 0.18, y + 0.05, w - 0.3, rh - 0.1, cell,
                                size=17 if k == 0 else 15, color=st.accent if (k == 0 and hi) else (st.ink if k == 0 else st.body),
                                font=st.head_font if k == 0 else st.body_font, bold=k == 0, anchor=MSO_ANCHOR.MIDDLE,
                                smallest=10.5)
            steps.append((tb, "fade", 330 + 160 * r))
            x += w
    _caption(ctx, slide, s["models"][:1], 9.75, H - 0.45, 3.3, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


# Xulosada modellar chapda (to'r), punktlar o'ngda katta panelda — oldingi slaydlardagi modellar shu yerga
# Morph bilan uchib yig'iladi.
_YAKUN_GRID = {1: ((3.05, 4.45, 4.3),),
               2: ((2.15, 3.55, 2.9), (4.0, 5.3, 2.9)),
               3: ((1.85, 3.3, 2.45), (4.25, 3.3, 2.45), (3.05, 5.55, 2.45)),
               4: ((1.85, 3.25, 2.3), (4.25, 3.25, 2.3), (1.85, 5.6, 2.3), (4.25, 5.6, 2.3))}


def _yakun(ctx: Ctx, slide, s: dict, i: int) -> List:
    mids = s["models"][:4]
    n = len(mids)
    _decorate(ctx, slide, i, (0.6, 1.9, 5.5, 6.9))
    _kicker(ctx, slide, s["kicker"], 0.75, 0.75, 11.8, align=PP_ALIGN.CENTER)
    _title(ctx, slide, s["title"], 0.75, 1.15, 11.83, 0.95, 34, 24, align=PP_ALIGN.CENTER)
    if not s["points"]:                              # matnsiz — modellar qatorda, markazda
        size = min(3.3, 11.6 / max(n, 1))
        span = size * n + 0.25 * (n - 1)
        for k, mid in enumerate(mids):
            _model(ctx, slide, mid, W / 2 - span / 2 + size / 2 + k * (size + 0.25), 4.3, size, turn=8 * (k - n / 2))
        _caption(ctx, slide, mids, 0.75, H - 0.4, 11.0, align=PP_ALIGN.CENTER)
        _number(ctx, slide, i)
        return []
    for k, (mid, (cx, cy, size)) in enumerate(zip(mids, _YAKUN_GRID[max(1, n)])):
        _model(ctx, slide, mid, cx, cy, size, turn=10 * (k - n / 2))
    x, w, top, bottom = 6.05, 6.75, 2.35, 6.9
    pad = 0.32
    em = _em(ctx.st.body_font)
    size, points = _fit_points(s, s["points"][:4], w - 2 * pad, bottom - top - 2 * pad, (17, 11.5), em=em)
    bh = _body_h(points, w - 2 * pad, size, em)
    ph = bh + 2 * pad
    y = top + max(0.0, (bottom - top - ph) / 2)
    panel = _panel(ctx, slide, f"panel_{i}", x, y, w, ph)
    body = _bullets(ctx, slide, f"body_{i}", points, x + pad, y + pad, w - 2 * pad, bh + 0.1, size, 11.5, fit=False)
    _caption(ctx, slide, mids, 0.75, H - 0.4, 11.0, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return [(panel, "fade", 200), (body, "float", 400)]


# ───────────────────────────────────────────── matnning boshqa ko'rinishlari (har varoq boshqacha)

def _head(ctx: Ctx, slide, s: dict, x: float, w: float, size: Tuple[float, float] = (34, 24), h: float = 1.25,
          align=PP_ALIGN.LEFT) -> float:
    """Tepada bo'lim nomi va sarlavha; sarlavhaning pastki chetini qaytaradi."""
    _kicker(ctx, slide, s["kicker"], x, 0.75, w, align=align)
    return _title(ctx, slide, s["title"], x, 1.15, w, h, size[0], size[1], align=align)


def _matn(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Ro'yxatsiz matn: 1-2 paragraf, chetida urg'u chizig'i; birinchisi yorqinroq."""
    st = ctx.st
    em = _em(st.body_font)
    left_text = i % 2 == 1
    tx, tw = (0.75, 6.3) if left_text else (6.3, 6.3)
    mx = 10.25 if left_text else 3.05
    _decorate(ctx, slide, i, (mx - 2.5, 1.1, mx + 2.5, 6.5))
    box = _model(ctx, slide, s["models"][0], mx, 3.8, 5.0)
    design.ground_shadow(slide, st, box)
    title = s["title"].upper() if st.title_caps else s["title"]
    hem = HEAD_EM * (1.12 if st.title_caps else 1.0)
    tsize = _fit(title, tw, 2.2, 36, 24, hem)
    th = _lines(title, tw, tsize, hem) * tsize * 1.12 / 72
    paras = list(s["paragraphs"][:2])
    pw = tw - 0.32
    room = 6.95 - 0.75 - 0.4 - th - 0.38

    def need(size: float) -> float:
        return sum(_para_h(p, pw, size * (1.05 if k == 0 else 1.0), em) + 0.2 for k, p in enumerate(paras))

    size = 19.0
    while size > 12 and need(size) > room:
        size -= 0.5
    while len(paras) > 1 and need(size) > room:
        s["notes"] = ((s.get("notes") or "") + "\n" + paras.pop()).strip()
    total = 0.4 + th + 0.38 + need(size)
    y0 = min(max(0.75, 3.8 - total / 2), max(0.75, 6.95 - total))
    _kicker(ctx, slide, s["kicker"], tx, y0, tw)
    y = _title(ctx, slide, s["title"], tx, y0 + 0.4, tw, 2.2, tsize, tsize) + 0.38
    steps = [(design.shape(slide, f"bar_{i}", MSO_SHAPE.RECTANGLE, tx, y + 0.06, 0.05, need(size) - 0.3,
                           st.accent), "fade", 150)]
    for k, text in enumerate(paras):
        psize = size * (1.05 if k == 0 else 1.0)
        h = _para_h(text, pw, psize, em)
        tb = design.textbox(slide, f"para_{i}_{k}", tx + 0.32, y, pw, h + 0.05, text, size=psize,
                            color=st.ink if k == 0 else st.body, font=st.body_font, line=1.15)
        steps.append((tb, "float", 250 + 250 * k))
        y += h + 0.2
    _caption(ctx, slide, s["models"][:1], mx - 2.6, H - 0.5, 5.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _xronologiya(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Vaqt chizig'i: sanalar chiziq ustida, voqealar ostida; model tepada o'ngda."""
    st = ctx.st
    _decorate(ctx, slide, i, (9.1, 0.5, 12.8, 3.9))
    box = _model(ctx, slide, s["models"][0], 10.9, 2.15, 3.1)
    design.ground_shadow(slide, st, box)
    _head(ctx, slide, s, 0.75, 8.6, (36, 24), 1.9)
    events = s["events"][:5]
    n = max(1, len(events))
    x0, x1, ly = 0.85, 12.5, 4.45
    sp = (x1 - x0) / n
    steps = [(design.shape(slide, f"line_{i}", MSO_SHAPE.RECTANGLE, x0, ly - 0.015, x1 - x0, 0.03, st.accent,
                           alpha=0.55), "fade", 150)]
    for k, ev in enumerate(events):
        cx = x0 + sp * (k + 0.5)
        dot = design.shape(slide, f"dot_{i}_{k}", MSO_SHAPE.OVAL, cx - 0.12, ly - 0.12, 0.24, 0.24, st.accent)
        one_line = (sp - 0.2) * 72 / (max(1, len(ev["date"])) * VALUE_EM)
        date = design.textbox(slide, f"date_{i}_{k}", cx - sp / 2 + 0.1, ly - 0.78, sp - 0.2, 0.55, ev["date"],
                              size=max(12, min(26, int(one_line))), color=st.accent, font=st.head_font, bold=True,
                              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        txt = design.textbox(slide, f"ev_{i}_{k}", cx - sp / 2 + 0.12, ly + 0.32, sp - 0.24, 6.95 - ly - 0.32,
                             ev["text"], size=16, color=st.body, font=st.body_font, align=PP_ALIGN.CENTER,
                             smallest=10.5)
        steps += [(dot, "zoom", 300 + 300 * k), (date, "fade", 340 + 300 * k), (txt, "fade", 380 + 300 * k)]
    _caption(ctx, slide, s["models"][:1], 6.6, H - 0.42, 5.9, align=PP_ALIGN.RIGHT)
    _number(ctx, slide, i)
    return steps


def _taqqoslash(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Ikki ustunli taqqoslash; model o'rtada, ikki tomon orasida."""
    st = ctx.st
    em = _em(st.body_font)
    _decorate(ctx, slide, i, (4.9, 2.2, 8.4, 6.7))
    bottom = _head(ctx, slide, s, 0.75, 11.8, (34, 24), 1.2)
    box = _model(ctx, slide, s["models"][0], W / 2, 4.55, 3.1)
    design.ground_shadow(slide, st, box)
    top = max(2.3, bottom + 0.35)
    w, pad, name_h = 4.55, 0.3, 0.55
    cols = s["columns"][:2]
    fits = [_fit_points(s, c["points"], w - 2 * pad, 6.9 - top - 2 * pad - name_h, (17, 11), em=em) for c in cols]
    size = min(f[0] for f in fits)
    ph = max(_body_h(f[1], w - 2 * pad, size, em) for f in fits) + 2 * pad + name_h
    steps = []
    for k, (col, (_, pts)) in enumerate(zip(cols, fits)):
        x = 0.6 if k == 0 else W - 0.6 - w
        panel = _panel(ctx, slide, f"col_{i}_{k}", x, top, w, ph)
        name = design.textbox(slide, f"colname_{i}_{k}", x + pad, top + pad - 0.05, w - 2 * pad, 0.45, col["name"],
                              size=20, color=st.accent, font=st.head_font, bold=True, smallest=13)
        body = _bullets(ctx, slide, f"colbody_{i}_{k}", pts, x + pad, top + pad + name_h, w - 2 * pad,
                        ph - 2 * pad - name_h + 0.1, size, 11, fit=False)
        steps += [(panel, "fade", 200 + 450 * k), (name, "fade", 260 + 450 * k), (body, "float", 320 + 450 * k)]
    _caption(ctx, slide, s["models"][:1], W / 2 - 1.6, H - 0.45, 3.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _kartalar(ctx: Ctx, slide, s: dict, i: int) -> List:
    """3-4 karta (sarlavha + matn), model chapda."""
    st = ctx.st
    _decorate(ctx, slide, i, (0.6, 2.2, 5.3, 6.8))
    bottom = _head(ctx, slide, s, 0.75, 11.8, (34, 24), 1.2)
    box = _model(ctx, slide, s["models"][0], 2.95, 4.5, 4.3)
    design.ground_shadow(slide, st, box)
    cards = s["cards"][:4]
    n = len(cards)
    cols = 2 if n == 4 else 1
    rows = (n + cols - 1) // cols
    x0, wall, gap = 5.55, 7.2, 0.22
    top = max(2.3, bottom + 0.32)
    cw = (wall - gap * (cols - 1)) / cols
    chh = (6.92 - top - gap * (rows - 1)) / rows
    steps = []
    for k, card in enumerate(cards):
        r, c = divmod(k, cols)
        x, y = x0 + c * (cw + gap), top + r * (chh + gap)
        panel = _panel(ctx, slide, f"card_{i}_{k}", x, y, cw, chh)
        mark = design.shape(slide, f"mark_{i}_{k}", MSO_SHAPE.RECTANGLE, x + 0.28, y + 0.2, 0.45, 0.05, st.accent)
        head = design.textbox(slide, f"ctitle_{i}_{k}", x + 0.28, y + 0.32, cw - 0.56, 0.42, card["title"], size=18,
                              color=st.ink, font=st.head_font, bold=True, smallest=12)
        txt = design.textbox(slide, f"ctext_{i}_{k}", x + 0.28, y + 0.78, cw - 0.56, chh - 0.9, card["text"],
                             size=15, color=st.body, font=st.body_font, smallest=10.5)
        steps += [(panel, "zoom", 200 + 220 * k), (mark, "fade", 240 + 220 * k), (head, "fade", 260 + 220 * k),
                  (txt, "fade", 300 + 220 * k)]
    _caption(ctx, slide, s["models"][:1], 0.75, H - 0.45, 4.4, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _bosqichlar(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Raqamlangan bosqichlar (jarayon), ular orasida tutashtiruvchi chiziq; model o'ngda."""
    st = ctx.st
    _decorate(ctx, slide, i, (7.9, 1.2, 12.8, 6.6))
    box = _model(ctx, slide, s["models"][0], 10.35, 4.0, 4.9)
    design.ground_shadow(slide, st, box)
    bottom = _head(ctx, slide, s, 0.75, 6.9, (34, 24), 1.25)
    items = s["steps"][:5]
    n = max(1, len(items))
    top = max(2.25, bottom + 0.3)
    row = (6.95 - top) / n
    steps = []
    if n > 1:
        steps.append((design.shape(slide, f"rail_{i}", MSO_SHAPE.RECTANGLE, 1.06, top + 0.35, 0.025,
                                   row * (n - 1), st.accent, alpha=0.4), "fade", 150))
    for k, item in enumerate(items):
        y = top + k * row
        num = design.shape(slide, f"num_{i}_{k}", MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, y + 0.05, 0.58, 0.58,
                           st.accent, round_=0.25)
        digit = design.textbox(slide, f"digit_{i}_{k}", 0.78, y + 0.05, 0.58, 0.58, str(k + 1), size=20, color=st.bg,
                               font=st.head_font, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        head = design.textbox(slide, f"stitle_{i}_{k}", 1.62, y, 5.9, 0.42, item["title"], size=18, color=st.ink,
                              font=st.head_font, bold=True, smallest=12)
        txt = design.textbox(slide, f"stext_{i}_{k}", 1.62, y + 0.43, 5.9, max(0.35, row - 0.5), item["text"],
                             size=15, color=st.body, font=st.body_font, smallest=10.5)
        steps += [(num, "zoom", 250 + 260 * k), (digit, "fade", 270 + 260 * k), (head, "fade", 300 + 260 * k),
                  (txt, "fade", 340 + 260 * k)]
    _caption(ctx, slide, s["models"][:1], 7.75, H - 0.5, 5.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _fikr(ctx: Ctx, slide, s: dict, i: int) -> List:
    """Asosiy fikr yoki iqtibos: katta qo'shtirnoq, yirik kursiv matn, muallif va izoh; model chapda katta."""
    st = ctx.st
    em = _em(st.body_font)
    _decorate(ctx, slide, i, (0.8, 0.9, 6.3, 6.8))
    box = _model(ctx, slide, s["models"][0], 3.5, 3.85, 5.9)
    design.ground_shadow(slide, st, box)
    x, w = 6.95, 5.75
    qsize = _fit(s["quote"], w, 2.1 if s["text"] else 2.7, 28, 18, HEAD_EM * 0.95)   # izohga ham joy qolsin
    qh = _lines(s["quote"], w, qsize, HEAD_EM * 0.95) * qsize * 1.12 * 1.2 / 72
    src_h = 0.5 if s["source"] else 0.0
    text = s["text"]
    tsize = 16.0
    while text and tsize > 11.5 and _para_h(text, w, tsize, em) > 6.9 - 0.9 - 1.35 - qh - src_h - 0.45:
        tsize -= 0.5
    th = _para_h(text, w, tsize, em) if text else 0.0
    total = 0.45 + 0.9 + qh + src_h + (0.45 + th if text else 0.0)
    y = min(max(0.75, 3.8 - total / 2), max(0.75, 6.95 - total))
    _kicker(ctx, slide, s["kicker"] or s["title"], x, y, w)
    mark = design.textbox(slide, f"qmark_{i}", x - 0.08, y + 0.2, 1.3, 1.1, "“", size=96, color=st.accent,
                          font="Georgia", bold=True)
    y += 0.45 + 0.9
    quote = design.textbox(slide, "!!title", x, y, w, qh + 0.1, s["quote"], size=qsize, color=st.ink,
                           font=st.head_font, italic=True, line=1.12)   # sarlavha Morph bilan iqtibosga aylanadi
    steps = [(mark, "zoom", 150), (quote, "fade", 300)]
    y += qh + 0.1
    if s["source"]:
        steps.append((design.textbox(slide, f"src_{i}", x, y, w, 0.36, "— " + s["source"], size=15, color=st.accent,
                                     font=st.head_font, bold=True), "fade", 550))
        y += src_h
    if text:
        steps.append((design.shape(slide, f"rule_{i}", MSO_SHAPE.RECTANGLE, x, y + 0.08, 1.1, 0.035, st.accent),
                      "fade", 650))
        steps.append((design.textbox(slide, f"qtext_{i}", x, y + 0.35, w, th + 0.05, text, size=tsize, color=st.body,
                                     font=st.body_font, line=1.15), "float", 750))
    _caption(ctx, slide, s["models"][:1], 0.9, H - 0.5, 5.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return steps


def _savol(ctx: Ctx, slide, s: dict, i: int) -> List:
    """«Bilasizmi?»: savol belgisi nishoni, yirik savol, javob kartada; model o'ngda."""
    st = ctx.st
    em = _em(st.body_font)
    _decorate(ctx, slide, i, (7.2, 1.0, 12.7, 6.7))
    box = _model(ctx, slide, s["models"][0], 9.95, 3.9, 5.5)
    design.ground_shadow(slide, st, box)
    x, w, pad = 0.75, 6.2, 0.32
    hem = HEAD_EM * (1.12 if st.title_caps else 1.0)
    qtext = s["question"].upper() if st.title_caps else s["question"]
    qsize = _fit(qtext, w, 2.1, 32, 22, hem)
    qh = _lines(qtext, w, qsize, hem) * qsize * 1.12 / 72
    room = 6.95 - 0.75 - 1.15 - qh - 0.35 - 2 * pad
    asize = 17.5
    while asize > 12 and _para_h(s["answer"], w - 2 * pad, asize, em) > room:
        asize -= 0.5
    ah = _para_h(s["answer"], w - 2 * pad, asize, em)
    total = 1.15 + qh + 0.35 + ah + 2 * pad
    y = min(max(0.75, 3.85 - total / 2), max(0.75, 6.95 - total))
    badge = design.shape(slide, f"badge_{i}", MSO_SHAPE.OVAL, x, y, 0.85, 0.85, st.accent)
    mark = design.textbox(slide, f"badgeq_{i}", x, y, 0.85, 0.85, "?", size=32, color=st.bg, font=st.head_font,
                          bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _kicker(ctx, slide, s["kicker"], x + 1.05, y + 0.24, w - 1.05)
    y += 1.15
    question = design.textbox(slide, "!!title", x, y, w, qh + 0.1, qtext, size=qsize, color=st.ink,
                              font=st.head_font, bold=True, line=1.02)
    y += qh + 0.35
    panel = _panel(ctx, slide, f"panel_{i}", x - 0.1, y, w + 0.2, ah + 2 * pad)
    answer = design.textbox(slide, f"answer_{i}", x + pad - 0.1, y + pad, w - 2 * pad + 0.2, ah + 0.05, s["answer"],
                            size=asize, color=st.body, font=st.body_font, line=1.15)
    _caption(ctx, slide, s["models"][:1], 7.35, H - 0.5, 5.2, align=PP_ALIGN.CENTER)
    _number(ctx, slide, i)
    return [(badge, "zoom", 150), (mark, "fade", 200), (question, "fade", 300), (panel, "fade", 550),
            (answer, "float", 650)]


def _turntable(slide, spids: Sequence[int], delay: int = 400) -> None:
    """PowerPoint'ning 3D «Turntable» effekti: slayddagi har 3D model o'z joyida juda sekin, slayd turguncha
    aylanadi (emph, presetID 37, `3d.view.rotation.y`). Kirish effektlari bilan bir guruhda, o'tishdan keyin."""
    from services.premium_presentation import slide_anim

    if not spids or TURN_MS <= 0:
        return
    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        timing = etree.fromstring(slide_anim._timing([], []))
        ext = slide._element.find(qn("p:extLst"))
        ext.addprevious(timing) if ext is not None else slide._element.append(timing)
    group = next(c for c in timing.iter(qn("p:cTn")) if c.get("id") == "4").find(qn("p:childTnLst"))
    nid = max(int(c.get("id")) for c in timing.iter(qn("p:cTn"))) + 1
    for k, spid in enumerate(spids):
        turn = 21600000 if k % 2 == 0 else -21600000           # qo'shni modellar qarama-qarshi tomonga
        group.append(etree.fromstring(
            f'<p:par {P_NS}><p:cTn id="{nid}" presetID="37" presetClass="emph" presetSubtype="128" '
            f'repeatCount="indefinite" fill="hold" grpId="0" nodeType="withEffect"><p:stCondLst>'
            f'<p:cond delay="{delay}"/></p:stCondLst><p:childTnLst><p:animRot by="{turn}"><p:cBhvr>'
            f'<p:cTn id="{nid + 1}" dur="{TURN_MS}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            f'<p:attrNameLst><p:attrName>3d.view.rotation.y</p:attrName></p:attrNameLst></p:cBhvr>'
            f'</p:animRot></p:childTnLst></p:cTn></p:par>'))
        nid += 2


LAYOUT_FN = {"muqova": _muqova, "chap": lambda c, sl, s, i: _side(c, sl, s, i, True),
             "ong": lambda c, sl, s, i: _side(c, sl, s, i, False), "qahramon": _qahramon,
             "raqamlar": _raqamlar, "uchlik": _uchlik, "yakun": _yakun, "matn": _matn,
             "xronologiya": _xronologiya, "taqqoslash": _taqqoslash, "kartalar": _kartalar,
             "bosqichlar": _bosqichlar, "fikr": _fikr, "savol": _savol, "jadval": _jadval}


def build(plan: dict, out_path: str) -> str:
    st = STYLES.get(plan.get("style") or "") or STYLES["vitrina"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    ctx = Ctx(prs, st, plan)
    slides = []
    for i, s in enumerate(plan["slides"]):
        slide = prs.slides.add_slide(blank)
        slides.append(slide)
        ctx.cur = i
        steps = LAYOUT_FN.get(s["layout"], LAYOUT_FN["chap"])(ctx, slide, s, i)
        add_morph(slide, MORPH_MS)
        if steps:
            add_entrances(slide, steps)
        _turntable(slide, ctx.spins.get(i, []))
        notes = s.get("notes") or ""
        if i == len(plan["slides"]) - 1 and ctx.credits:
            head = NOTES_HEAD.get((plan.get("language") or "uz")[:2]) or NOTES_HEAD["en"]
            notes = (notes + "\n\n" if notes else "") + head + ":\n" + "\n".join(ctx.credits)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    ghosts = _choreography(ctx, slides)
    prs.save(out_path)
    log.info("3D katalog: %d slayd, %d model, %d kirish/chiqish, uslub %s, %.1f MB", len(plan["slides"]),
             len(ctx.parts), ghosts, st.key, os.path.getsize(out_path) / 1e6)
    return out_path
