"""Qozoq tilidagi hujjatlar (referat, mustaqil ish, kurs ishi, taqdimot...).

Kod bazasidagi barcha so'rovlar va qolipchalar uch tilda (o'zbek, rus,
ingliz) yozilgan. Qozoq tilini ularning har biriga qo'shish o'rniga qozoqcha
hujjat RUSCHA yo'l bilan yasaladi, lekin:

  1. modelga har so'rovda "matnni QOZOQ tilida (kirill) yozing" qoidasi
     qo'shiladi (`with_rule`, uz_script.with_rule kabi);
  2. modeldan va koddan kelgan qat'iy ruscha yorliqlar ("Введение",
     "Заключение", "Рисунок 1"...) qozoqchaga o'giriladi (`translate`),
     tayyor DOCX/PPTX dagi qolganlari ham (`localize_docx`, `localize_pptx`).

Til kodi "kk". Servis metodlari `aware` / `wrap_class` bilan o'raladi: ularga
"kk" kelsa, ichkarida "ru" ishlaydi va qoida yoqiladi, natija o'giriladi.
"""
import contextvars
import functools
import inspect
import logging
import re
from typing import Any, Dict, Optional

logger = logging.getLogger(__name__)

LANG = "kk"

# Qozoq alifbosiga xos harflar (rus va o'zbek kirillida bunday emas).
KAZAKH_LETTERS = set("әіңүұөһӘІҢҮҰӨҺ")

_ACTIVE: contextvars.ContextVar = contextvars.ContextVar("kazakh_doc", default=False)


def active() -> bool:
    """Shu jarayonda hujjat qozoq tilida yozilmoqdami."""
    return _ACTIVE.get()


def use(flag: bool = True):
    """Shu jarayondagi barcha model so'rovlariga qozoq qoidasini qo'shadi. `reset` uchun token."""
    return _ACTIVE.set(bool(flag))


def reset(token) -> None:
    try:
        _ACTIVE.reset(token)
    except (ValueError, LookupError):
        pass


# ───────────────────────────────────────────────────────── modelga qoida

RULE = (
    "LANGUAGE RULE (highest priority): write ALL generated text in KAZAKH (қазақ тілі), "
    "Cyrillic alphabet with the Kazakh letters ә, ғ, қ, ң, ө, ұ, ү, һ, і. "
    "The instructions above are written in Russian only for convenience: if they say "
    "to write in Russian, write in Kazakh instead. Titles, headings, plan items, "
    "captions, table cells, chart labels and references text must be Kazakh as well. "
    "Do not write Russian sentences and do not mix Uzbek words in. "
    "JSON keys, layout names and technical identifiers stay exactly as specified in the "
    "instructions; internationally written abbreviations (AI, PDF) and proper names may stay. "
    "Use natural academic Kazakh, e.g. «Кіріспе», «Қорытынды», «Пайдаланылған әдебиеттер»."
)


# Ruscha so'rov ichidagi "ruscha yozing" iboralari qozoqchaga almashtiriladi: model
# bir-biriga zid ikki ko'rsatma ("rus tilida" va "qozoq tilida") olmasin.
_PROMPT_SWAPS = (
    (re.compile(r"на русском языке"), "на казахском языке"),
    (re.compile(r"на русском"), "на казахском"),
    (re.compile(r"русском языке"), "казахском языке"),
    (re.compile(r"русский язык"), "казахский язык"),
    (re.compile(r"по-русски"), "по-казахски"),
    (re.compile(r"\bрусский\b"), "казахский"),
    (re.compile(r"\bRussian\b"), "Kazakh"),
)


def localize_prompt(text: str) -> str:
    for pattern, repl in _PROMPT_SWAPS:
        text = pattern.sub(repl, text)
    return text


def with_rule(messages: list) -> list:
    """Xabarlar ro'yxatiga qozoq qoidasini qo'shadi (qoida yoqilmagan bo'lsa — o'zgarmaydi)."""
    if not active() or not messages:
        return messages
    out = [dict(m) for m in messages]
    for message in out:
        if isinstance(message.get("content"), str):
            message["content"] = localize_prompt(message["content"])
    for message in reversed(out):
        if message.get("role") == "user" and isinstance(message.get("content"), str):
            if RULE not in message["content"]:
                message["content"] = message["content"] + "\n\n" + RULE
            return out
    return [{"role": "system", "content": RULE}] + out


def target_name(russian: str = "русском", default: Optional[str] = None) -> str:
    """Prompt ichidagi "русском" o'rniga: qozoq rejimida "казахском"."""
    return "казахском" if active() else (default if default is not None else russian)


# ───────────────────────────────────────────────────────── yorliqlar (RU → KK)

# Kalit — ruscha yorliq (kichik harf, bo'sh joysiz), qiymat — qozoqcha. Katta-kichik
# harfi yorliqning o'ziga moslanadi (BARCHASI KATTA → BARCHASI KATTA).
_PHRASES = {
    # tuzilma
    "введение": "Кіріспе", "заключение": "Қорытынды", "план": "Жоспар",
    "план презентации": "Презентация жоспары", "тезисы": "Тезистер",
    "содержание": "Мазмұны", "оглавление": "Мазмұны", "агенда": "Жоспар",
    "использованная литература": "Пайдаланылған әдебиеттер",
    "список использованной литературы": "Пайдаланылған әдебиеттер тізімі",
    "список литературы": "Әдебиеттер тізімі", "литература": "Әдебиеттер",
    "аннотация": "Аңдатпа", "ключевые слова": "Түйінді сөздер",
    "приложения": "Қосымшалар", "приложение": "Қосымша",
    "глава": "Тарау", "параграф": "Параграф", "раздел": "Бөлім",
    "заключение и рекомендации": "Қорытынды және ұсыныстар",
    "обзор литературы": "Әдебиеттерге шолу", "методология": "Әдіснама",
    "результаты и обсуждение": "Нәтижелер және талқылау",
    "практические рекомендации": "Практикалық ұсынымдар",
    "практическая часть": "Практикалық бөлім",
    "цель работы": "Жұмыстың мақсаты",
    "краткие понятия по теме": "Тақырып бойынша қысқаша түсініктер",
    "глоссарий": "Терминдер сөздігі",
    "глоссарий (словарь терминов)": "Глоссарий (терминдер сөздігі)",
    "шаг": "Қадам", "формулы:": "Формулалар:", "пример:": "Мысал:",
    "анализ:": "Талдау:", "statistika i faktlar:": "Статистика және деректер:",
    "статистика и факты:": "Статистика және деректер:",
    "спасибо за внимание!": "Назарларыңызға рахмет!",
    "спасибо за внимание": "Назарларыңызға рахмет",
    "(продолжение)": "(жалғасы)",
    # титул беті
    "тема": "Тақырыбы", "по предмету": "пәні бойынша",
    "выполнил": "Орындаған", "выполнил(а)": "Орындаған",
    "принял": "Қабылдаған", "принял(а)": "Қабылдаған",
    "научный руководитель": "Ғылыми жетекші",
    "курс": "курс", "факультет": "факультеті", "студент": "студенті",
    "кафедра": "кафедрасы", "университет": "университеті",
    "направление": "бағыты",
    "республика узбекистан": "Қазақстан Республикасы",
    "министерство высшего образования, науки и инноваций": "Ғылым және жоғары білім министрлігі",
    "студента ________ курса ________ группы": "________ курс ________ топ студенті",
    "реферат": "Реферат", "самостоятельная работа": "Өзіндік жұмыс",
    "курсовая работа": "Курстық жұмыс", "дипломная работа": "Дипломдық жұмыс",
    "магистерская диссертация": "Магистрлік диссертация",
    "выпускная квалификационная работа": "Бітіру біліктілік жұмысы",
    "специальная разработка": "Арнайы әзірлеме",
    "ташкент": "Астана",
    # кіріспе құрылымы (курстық, дипломдық)
    "1. предмет дипломной работы.": "1. Дипломдық жұмыстың пәні.",
    "2. объект дипломной работы.": "2. Дипломдық жұмыстың объектісі.",
    "3. степень изученности темы.": "3. Тақырыптың зерттелу дәрежесі.",
    "4. цель дипломной работы.": "4. Дипломдық жұмыстың мақсаты.",
    "5. задачи дипломной работы:": "5. Дипломдық жұмыстың міндеттері:",
    "6. структура дипломной работы.": "6. Дипломдық жұмыстың құрылымы.",
    "1. предмет курсовой работы.": "1. Курстық жұмыстың пәні.",
    "2. объект курсовой работы.": "2. Курстық жұмыстың объектісі.",
    "4. цель курсовой работы.": "4. Курстық жұмыстың мақсаты.",
    "5. задачи курсовой работы:": "5. Курстық жұмыстың міндеттері:",
    "6. структура курсовой работы.": "6. Курстық жұмыстың құрылымы.",
    "актуальность темы.": "Тақырыптың өзектілігі.",
    "предмет курсовой работы.": "Курстық жұмыстың пәні.",
    "объект курсовой работы.": "Курстық жұмыстың объектісі.",
    "цель курсовой работы.": "Курстық жұмыстың мақсаты.",
    "структура курсовой работы.": "Курстық жұмыстың құрылымы.",
    "1. обоснование темы и её актуальность:": "1. Тақырыптың негіздемесі және өзектілігі:",
    "2. объект исследования:": "2. Зерттеу объектісі:",
    "3. предмет исследования:": "3. Зерттеу пәні:",
    "4. цель и задачи исследования:": "4. Зерттеудің мақсаты мен міндеттері:",
    "5. научная новизна:": "5. Ғылыми жаңалығы:",
    "6. основные вопросы и гипотезы исследования:": "6. Зерттеудің негізгі сұрақтары мен болжамдары:",
    "7. обзор литературы по теме исследования:": "7. Зерттеу тақырыбы бойынша әдебиеттерге шолу:",
    "8. описание применённой методики:": "8. Қолданылған әдістеменің сипаттамасы:",
    "10. описание структуры работы:": "10. Жұмыс құрылымының сипаттамасы:",
    "1. актуальность исследования:": "1. Зерттеудің өзектілігі:",
    "4. цель исследования:": "4. Зерттеудің мақсаты:",
    "5. задачи исследования:": "5. Зерттеудің міндеттері:",
    "6. методы исследования:": "6. Зерттеу әдістері:",
    "7. научная новизна:": "7. Ғылыми жаңалығы:",
    "8. структура работы:": "8. Жұмыс құрылымы:",
    # анықтама тізімдері
    "нормативно-правовые акты": "Нормативтік-құқықтық актілер",
    "учебники и пособия": "Оқулықтар мен оқу құралдары",
    "научные статьи": "Ғылыми мақалалар",
    "интернет-ресурсы": "Интернет-ресурстар",
    "рисунок": "Сурет", "таблица": "Кесте", "формула": "Формула",
    # Жоба жұмысы (services/project_work/specs.py)
    "расчёты и формулы": "Есептеулер және формулалар",
    "смета и ресурсы": "Смета және ресурстар",
    "план работ (этапы)": "Жұмыс жоспары (кезеңдер)",
    "прогноз и эффективность": "Болжам және тиімділік",
    "анализ рисков": "Тәуекелдерді талдау",
    "ожидаемые результаты": "Күтілетін нәтижелер",
    "структура расходов и себестоимость": "Шығындар құрылымы және өзіндік құн",
    "маркетинговый прогноз и план продаж": "Маркетингтік болжам және сату жоспары",
    "анализ точки безубыточности": "Залалсыздық нүктесін талдау",
    "денежный поток и срок окупаемости": "Ақша ағыны және өтелу мерзімі",
    "актуальность проблемы и цель проекта": "Мәселенің өзектілігі және жобаның мақсаты",
    "план реализации проекта": "Жобаны іске асыру жоспары",
    "бюджет проекта и ресурсы": "Жоба бюджеті және ресурстар",
    "риски и управление ими": "Тәуекелдер және оларды басқару",
    "ожидаемые результаты и эффективность": "Күтілетін нәтижелер және тиімділік",
    "прогноз и расчёт эффективности": "Болжам және тиімділікті есептеу",
    "расчёты и их обоснование": "Есептеулер және олардың негіздемесі",
    "структура расходов и расчёт себестоимости": "Шығындар құрылымы және өзіндік құнды есептеу",
    "точка безубыточности и финансовая устойчивость": "Залалсыздық нүктесі және қаржылық тұрақтылық",
    "денежный поток и окупаемость инвестиций": "Ақша ағыны және инвестицияның өтелуі",
    "структура проекта": "Жоба құрылымы",
    "экономика и бизнес": "Экономика және бизнес",
    "техника, инженерия, строительство": "Техника, инженерия, құрылыс",
    "it и программное обеспечение": "IT және бағдарламалық қамтамасыз ету",
    "педагогика и образование": "Педагогика және білім беру",
    "социальная и гуманитарная сфера": "Әлеуметтік және гуманитарлық сала",
    "медицина и биология": "Медицина және биология",
    "сельское хозяйство": "Ауыл шаруашылығы",
    "другая сфера (ai составит сам)": "Басқа сала (ЖИ өзі құрастырады)",
    "ии выберет сам по теме": "ЖИ тақырып бойынша өзі таңдайды",
}

# "Рисунок 3", "Таблица 2", "Формула 5" → "3-сурет", "2-кесте", "Формула 5"
_NUMBERED = (
    (re.compile(r"^Рисунок\s+(\d+)(.*)$", re.DOTALL), r"\1-сурет\2"),
    (re.compile(r"^Таблица\s+(\d+)(.*)$", re.DOTALL), r"\1-кесте\2"),
    (re.compile(r"^Формула\s+(\d+)(.*)$", re.DOTALL), r"Формула \1\2"),
    (re.compile(r"^(ГЛАВА|Глава)\s+([IVXLC\d]+)(.*)$", re.DOTALL), lambda m: f"{m.group(2)} ТАРАУ{m.group(3)}"
     if m.group(1).isupper() else f"{m.group(2)} тарау{m.group(3)}"),
    (re.compile(r"^([IVXLC\d]+)\s+(ГЛАВА|Глава)(.*)$", re.DOTALL), lambda m: f"{m.group(1)} ТАРАУ{m.group(3)}"
     if m.group(2).isupper() else f"{m.group(1)} тарау{m.group(3)}"),
)

_RUSSIAN = re.compile(r"[А-Яа-яЁё]")


def _match_case(source: str, target: str) -> str:
    letters = [c for c in source if c.isalpha()]
    if len(letters) > 1 and all(c.isupper() for c in letters):
        return target.upper()
    if letters and letters[0].islower() and target[:1].isupper():
        return target[:1].lower() + target[1:]
    return target


def _lookup(core: str) -> Optional[str]:
    key = re.sub(r"\s+", " ", core).lower()
    found = _PHRASES.get(key)
    if found is not None:
        return found
    base = re.match(r"^(.*?)([.:]*)$", key, re.DOTALL)
    found = _PHRASES.get(base.group(1)) if base else None
    if found is not None and base.group(2) and not found.endswith(tuple(".:")):
        found += base.group(2)          # "Тема:" → "Тақырыбы:" (belgi saqlanadi)
    return found


def translate(text: str) -> str:
    """Ruscha qat'iy yorliqni qozoqchaga o'giradi. Yorliq bo'lmasa — matn o'zgarmaydi.

    Faqat BUTUN matn yorliqqa teng bo'lsagina o'giriladi (so'z ichidagi bo'lak emas),
    shuning uchun modelning qozoqcha matniga tegilmaydi.
    """
    if not isinstance(text, str) or not text or not _RUSSIAN.search(text):
        return text
    core = text.strip()
    lead, trail = text[:len(text) - len(text.lstrip())], text[len(text.rstrip()):]
    found = _lookup(core)
    if found is not None:
        return lead + _match_case(core, found) + trail
    for pattern, repl in _NUMBERED:
        if pattern.match(core):
            return lead + pattern.sub(repl, core) + trail
    return text


_NUMBER_PREFIX = re.compile(r"^(\s*\d+(?:\.\d+)*\.?\s+)(\S.*)$", re.DOTALL)
_UNDERSCORE_LABEL = re.compile(r"^(_+\s*)(\S.*)$", re.DOTALL)
_COLON_LABEL = re.compile(r"^([^:\t\n]{1,45})(:[\s\S]*)$")
_TOC_LINE = re.compile(r"^(.*?)(\s*[\t.…·]+\s*\d+\s*)$", re.DOTALL)


def translate_loose(text: str) -> str:
    """`translate` + titul varag'idagi "____ УНИВЕРСИТЕТ", "Выполнил(а):\t___", "Введение ....3" shakllari."""
    new = translate(text)
    if new != text or not isinstance(text, str):
        return new
    numbered = _NUMBER_PREFIX.match(text)
    if numbered:
        head = translate(numbered.group(2))
        if head != numbered.group(2):
            return numbered.group(1) + head
    for pattern, rebuild in (
        (_UNDERSCORE_LABEL, lambda m, t: m.group(1) + t),
        (_COLON_LABEL, lambda m, t: t + m.group(2)),
        (_TOC_LINE, lambda m, t: t + m.group(2)),
    ):
        match = pattern.match(text)
        if match:
            head = translate(match.group(2) if pattern is _UNDERSCORE_LABEL else match.group(1))
            if head != (match.group(2) if pattern is _UNDERSCORE_LABEL else match.group(1)):
                return rebuild(match, head)
    return text


def translate_tree(value: Any) -> Any:
    """Model/kod natijasidagi (dict, list, str) butun-yorliq matnlarni o'giradi."""
    if isinstance(value, str):
        return translate(value)
    if isinstance(value, list):
        return [translate_tree(item) for item in value]
    if isinstance(value, tuple):
        return tuple(translate_tree(item) for item in value)
    if isinstance(value, dict):
        return {key: (item if key in _SKIP_KEYS else translate_tree(item)) for key, item in value.items()}
    return value


# Texnik maydonlar: ularning matni emas, mazmuni o'zgarmasligi kerak.
_SKIP_KEYS = {"layout", "image_prompt", "prompt", "icon", "kind", "style", "type", "language"}


def translate_values(texts: Dict[str, str]) -> Dict[str, str]:
    """Qolip lug'ati (titul varag'i yorliqlari) qiymatlarini o'giradi."""
    return {key: translate(val) if isinstance(val, str) else val for key, val in texts.items()}


# ───────────────────────────────────────────────────────── tayyor fayl

def localize_docx(path: str) -> int:
    """DOCX ichidagi ruscha qat'iy yorliqlarni qozoqchaga o'giradi. O'zgargan bo'laklar soni."""
    from docx import Document

    document = Document(path)
    changed = 0

    def fix_paragraph(paragraph) -> None:
        nonlocal changed
        full = paragraph.text
        if not _RUSSIAN.search(full or ""):
            return
        before = changed
        for run in paragraph.runs:
            new = translate_loose(run.text)
            if new != run.text:
                run.text = new
                changed += 1
        if changed != before:
            return
        # Yorliq bir necha run'ga bo'lingan bo'lishi mumkin: butun xatboshi bo'yicha.
        new = translate_loose(full)
        if new != full and paragraph.runs:
            paragraph.runs[0].text = new
            for run in paragraph.runs[1:]:
                run.text = ""
            changed += 1

    def walk(container) -> None:
        for paragraph in container.paragraphs:
            fix_paragraph(paragraph)
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    walk(cell)

    walk(document)
    for section in document.sections:
        for part in (section.header, section.footer):
            walk(part)
    if changed:
        document.save(path)
    return changed


def localize_pptx(path: str) -> int:
    """PPTX ichidagi ruscha qat'iy yorliqlarni qozoqchaga o'giradi. O'zgargan bo'laklar soni."""
    from pptx import Presentation

    presentation = Presentation(path)
    changed = 0

    def fix_frame(frame) -> None:
        nonlocal changed
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                new = translate_loose(run.text)
                if new != run.text:
                    run.text = new
                    changed += 1

    def walk(shapes) -> None:
        for shape in shapes:
            if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
                walk(shape.shapes)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                fix_frame(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        fix_frame(cell.text_frame)

    for slide in presentation.slides:
        walk(slide.shapes)
    if changed:
        presentation.save(path)
    return changed


def localize_file(path: Any) -> Any:
    """Natija fayl yo'li bo'lsa (DOCX/PPTX) — ruscha yorliqlarni o'giradi. Yo'lni qaytaradi."""
    if isinstance(path, str) and path.lower().endswith((".docx", ".pptx")):
        try:
            if path.lower().endswith(".docx"):
                localize_docx(path)
            else:
                localize_pptx(path)
        except Exception as exc:
            logger.warning("Qozoqcha yorliqlar o'girilmadi (%s): %s", path, exc)
    return path


# ───────────────────────────────────────────────────────── servis metodlarini o'rash

_LANG_PARAMS = ("language", "lang", "doc_lang", "target_lang", "target_language")


def detect_letters(text: str) -> bool:
    """Matnda qozoq alifbosiga xos harf bormi."""
    return any(char in KAZAKH_LETTERS for char in str(text or ""))


def aware(func, post=None):
    """Metodni qozoqcha rejimga moslaydi.

    `language` ("kk") yoki `content["language"] == "kk"` kelsa: ichkarida "ru"
    ishlaydi, model so'rovlariga qozoq qoidasi qo'shiladi, natijadagi qat'iy
    ruscha yorliqlar o'giriladi va natijadagi "language" yana "kk" bo'ladi.
    """
    signature = inspect.signature(func)
    lang_names = [name for name in _LANG_PARAMS if name in signature.parameters]

    def prepare(args, kwargs):
        try:
            bound = signature.bind_partial(*args, **kwargs)
        except TypeError:
            return args, kwargs, False, []
        restore = []
        kazakh = False
        for name in lang_names:
            if bound.arguments.get(name) == LANG:
                bound.arguments[name] = "ru"
                kazakh = True
        for value in bound.arguments.values():
            if isinstance(value, dict) and value.get("language") == LANG:
                value["language"] = "ru"
                restore.append(value)
                kazakh = True
            elif not isinstance(value, (str, bytes, int, float, list, tuple, dict, type(None))) \
                    and getattr(value, "language", None) == LANG:
                object.__setattr__(value, "language", "ru")      # dataclass (ProjectContent)
                restore.append(value)
                kazakh = True
        if kazakh:
            return bound.args, bound.kwargs, True, restore
        return args, kwargs, False, []

    def put_back(restore):
        for value in restore:
            if isinstance(value, dict):
                value["language"] = LANG
            else:
                object.__setattr__(value, "language", LANG)

    def finish(result, restore):
        put_back(restore)
        result = translate_tree(result)
        if post is not None:
            result = post(result)
        if isinstance(result, dict) and result.get("language") == "ru":
            result["language"] = LANG
        elif not isinstance(result, (str, bytes, int, float, list, tuple, dict, type(None))) \
                and getattr(result, "language", None) == "ru":
            object.__setattr__(result, "language", LANG)         # ProjectContent
        return result

    if inspect.iscoroutinefunction(func):
        @functools.wraps(func)
        async def async_wrapper(*args, **kwargs):
            args, kwargs, kazakh, restore = prepare(args, kwargs)
            if not kazakh:
                return await func(*args, **kwargs)
            token = use(True)
            try:
                result = await func(*args, **kwargs)
            finally:
                reset(token)
                put_back(restore)
            return finish(result, restore)
        return async_wrapper

    @functools.wraps(func)
    def sync_wrapper(*args, **kwargs):
        args, kwargs, kazakh, restore = prepare(args, kwargs)
        if not kazakh:
            return func(*args, **kwargs)
        token = use(True)
        try:
            result = func(*args, **kwargs)
        finally:
            reset(token)
            put_back(restore)
        return finish(result, restore)
    return sync_wrapper


def wrap_class(cls, skip=(), file_methods=(), file_prefix: str = "") -> None:
    """Sinfning `language` (yoki `content["language"]`) qabul qiladigan metodlarini o'raydi.

    `file_methods` — natijasi DOCX/PPTX yo'li bo'lgan metodlar: tayyor fayldagi
    ruscha yorliqlar qozoqchaga o'giriladi.
    """
    for name, func in list(vars(cls).items()):
        if name in skip or name.startswith("__") or not inspect.isfunction(func):
            continue
        params = inspect.signature(func).parameters
        takes_language = any(p in params for p in _LANG_PARAMS)
        takes_content = "content" in params
        if not (takes_language or takes_content):
            continue
        post = localize_file if (name in file_methods or (file_prefix and name.startswith(file_prefix))) else None
        setattr(cls, name, aware(func, post))
