"""AI-driven editing of an existing DOCX the user uploads.

The client describes the changes in free text ("3-jadvalni o'chir", "10-betdan
15-betgacha qayta yoz", "bu yerga rasm qo'sh"). The model turns that into a
plan of addressable operations against the document's own blocks, the plan is
priced here — never by the model — and applied with python-docx.
"""

import asyncio
import json
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from config import (
    FILE_EDIT_BASE_PRICE,
    FILE_EDIT_MAX_PRICE,
    FILE_EDIT_OP_PRICES,
    FILE_EDIT_PRICE_STEP,
    TEMP_DIR,
)

logger = logging.getLogger(__name__)

# How much of the document is shown to the model when planning. Large enough
# for a 40-page course work, small enough to stay inside the context window.
_MAX_PLAN_CHARS = 60_000
_BLOCK_PREVIEW_CHARS = 400
_MAX_OPERATIONS = 60

# A DOCX stores no pagination — Word computes it when the file is opened — but
# clients ask for changes by page ("4-varaqqa 5 ta snoska qo'sh"), so the model
# needs a page number per block or it can only guess.
#
# Word does leave `w:lastRenderedPageBreak` markers where it last paginated;
# those are real and preferred. Without them the page is estimated from the
# A4 / Times New Roman 14pt / 1.5-spacing layout these documents are built with.
_CHARS_PER_LINE = 75
_LINES_PER_PAGE = 34
_TABLE_ROW_LINES = 1.5


@dataclass
class EditOperation:
    op: str
    block_id: Optional[int] = None
    text: str = ""
    caption: str = ""
    headers: List[str] = field(default_factory=list)
    rows: List[List[str]] = field(default_factory=list)
    image_prompt: str = ""


@dataclass
class EditPlan:
    operations: List[EditOperation]
    summary: str
    price: int
    truncated: bool = False


@dataclass
class EditResult:
    """What actually landed in the file, as opposed to what was planned."""
    path: str
    applied: List[EditOperation]
    failed: List[EditOperation]


class FileEditNotSupported(Exception):
    """The uploaded file cannot be edited by the AI path."""


class FileEditService:
    def __init__(self, ai_service, together_service=None, document_service=None):
        self.ai = ai_service
        self.together = together_service
        self.documents = document_service

    # ---------------------------------------------------------------- reading

    @staticmethod
    def _iter_blocks(doc):
        """Yield paragraphs and tables in the order they appear in the body."""
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                yield Table(child, doc)

    @staticmethod
    def _page_numbers(blocks: list) -> List[int]:
        """Page number per block, from Word's own marks when the file carries them."""
        rendered = any(
            block._element.findall(f".//{qn('w:lastRenderedPageBreak')}")
            for block in blocks
        )

        pages = []
        page = 1
        lines_used = 0.0
        for block in blocks:
            pages.append(page)

            if rendered:
                # Word paginated this file itself; trust its marks.
                page += len(block._element.findall(f".//{qn('w:lastRenderedPageBreak')}"))
                continue

            if block._element.findall(f".//{qn('w:br')}[@{qn('w:type')}='page']"):
                page += 1
                lines_used = 0.0
                continue

            if isinstance(block, Table):
                lines_used += len(block.rows) * _TABLE_ROW_LINES
            else:
                lines_used += max(1, -(-len(block.text) // _CHARS_PER_LINE))

            while lines_used >= _LINES_PER_PAGE:
                lines_used -= _LINES_PER_PAGE
                page += 1

        return pages

    @classmethod
    def _describe(cls, path: str) -> tuple:
        """Return (doc, blocks, listing, truncated) for the document at `path`."""
        doc = Document(path)
        blocks = list(cls._iter_blocks(doc))
        pages = cls._page_numbers(blocks)

        lines = []
        used = 0
        truncated = False
        for index, block in enumerate(blocks, start=1):
            page = pages[index - 1]
            if isinstance(block, Table):
                headers = [cell.text.strip() for cell in block.rows[0].cells] if block.rows else []
                line = (
                    f"[{index}] (bet {page}) JADVAL {len(block.rows)}x{len(block.columns)}"
                    f" | ustunlar: {', '.join(headers)[:200]}"
                )
            else:
                text = block.text.strip()
                if not text:
                    continue
                preview = text[:_BLOCK_PREVIEW_CHARS]
                if len(text) > _BLOCK_PREVIEW_CHARS:
                    preview += " …"
                line = f"[{index}] (bet {page}) {preview}"

            if used + len(line) > _MAX_PLAN_CHARS:
                truncated = True
                break
            lines.append(line)
            used += len(line)

        return doc, blocks, "\n".join(lines), truncated

    # --------------------------------------------------------------- planning

    async def plan(self, path: str, instruction: str, language: str) -> EditPlan:
        if not path.lower().endswith(".docx"):
            raise FileEditNotSupported(path)

        _, _, listing, truncated = await asyncio.to_thread(self._describe, path)
        raw = await self._request_plan(listing, instruction, language, truncated)
        operations = self._parse_operations(raw)
        summary = str(raw.get("summary") or "").strip()
        return EditPlan(
            operations=operations,
            summary=summary,
            price=self.price_of(operations),
            truncated=truncated,
        )

    async def _request_plan(self, listing: str, instruction: str, language: str, truncated: bool) -> Dict:
        lang_name = {"uz": "o'zbek", "ru": "русский", "en": "English"}.get(language, "o'zbek")
        truncation_note = (
            "\nNOTE: the listing below is cut short; only plan operations for blocks you can see."
            if truncated
            else ""
        )
        prompt = f"""You edit a Word document on behalf of its owner.

Every editable block of the document is listed below with its id in square
brackets, followed by the page it falls on. Paragraphs show their text, tables
show their size and column headers.{truncation_note}

Clients name pages, not block ids ("add 5 footnotes to page 4"). Use the
"(bet N)" marker to find the blocks on the page they mean, then address those
blocks by id. When a request names a page, every operation for it must target
a block carrying that page number.

DOCUMENT BLOCKS:
{listing}

THE OWNER'S REQUEST (verbatim):
{instruction}

Produce the operations that carry out this request. Rules:
- Address blocks only by the ids shown above; never invent an id.
- Cover the whole request. If it names a range, emit one operation per block in it.
- Write every piece of new prose in {lang_name}, matching the document's tone.
- Plain text only: no markdown, no special characters.
- If the request is impossible or unrelated to this document, return an empty
  operations list and explain why in the summary.
- At most {_MAX_OPERATIONS} operations.

Available operations:
{{"op": "replace_text", "block_id": 12, "text": "new full paragraph text"}}
{{"op": "insert_paragraph", "block_id": 12, "text": "text of the new paragraph inserted after block 12"}}
{{"op": "delete_block", "block_id": 5}}
{{"op": "insert_table", "block_id": 20, "headers": ["A", "B"], "rows": [["1", "2"]], "caption": "1-jadval. Nomi"}}
{{"op": "insert_image", "block_id": 8, "image_prompt": "English description of the picture to generate", "caption": "1-rasm. Nomi"}}
{{"op": "add_footnote", "block_id": 14, "text": "footnote text shown at the bottom of the page"}}

Respond with JSON only:
{{"summary": "one or two sentences in {lang_name} describing what you will change",
  "operations": [...]}}"""

        response = await self.ai._make_request(
            messages=[
                {"role": "system", "content": "You are a precise document editor. Respond with valid JSON only."},
                {"role": "user", "content": prompt},
            ],
            max_tokens=8000,
            temperature=0.3,
        )
        return self._parse_json(response)

    @staticmethod
    def _parse_json(response: str) -> Dict:
        text = response.strip()
        if text.startswith("```"):
            text = re.sub(r"^```(?:json)?\s*", "", text)
            text = re.sub(r"\s*```$", "", text)
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            match = re.search(r"\{.*\}", text, re.DOTALL)
            if match:
                return json.loads(match.group(0))
            raise

    @staticmethod
    def _parse_operations(raw: Dict) -> List[EditOperation]:
        allowed = {
            "replace_text",
            "insert_paragraph",
            "delete_block",
            "insert_table",
            "insert_image",
            "add_footnote",
        }
        operations = []
        for item in (raw.get("operations") or [])[:_MAX_OPERATIONS]:
            if not isinstance(item, dict):
                continue
            op = str(item.get("op", "")).strip()
            if op not in allowed:
                logger.warning("Dropping unknown edit operation: %s", op)
                continue
            try:
                block_id = int(item["block_id"])
            except (KeyError, TypeError, ValueError):
                logger.warning("Dropping edit operation without a usable block_id: %s", op)
                continue
            operations.append(
                EditOperation(
                    op=op,
                    block_id=block_id,
                    text=str(item.get("text") or ""),
                    caption=str(item.get("caption") or ""),
                    headers=[str(h) for h in (item.get("headers") or [])],
                    rows=[[str(c) for c in row] for row in (item.get("rows") or []) if isinstance(row, list)],
                    image_prompt=str(item.get("image_prompt") or ""),
                )
            )
        return operations

    # ---------------------------------------------------------------- pricing

    @staticmethod
    def price_of(operations: List[EditOperation]) -> int:
        """Price the plan from the work it contains, rounded to a whole step."""
        if not operations:
            return 0
        total = FILE_EDIT_BASE_PRICE + sum(
            FILE_EDIT_OP_PRICES.get(operation.op, FILE_EDIT_OP_PRICES["replace_text"])
            for operation in operations
        )
        step = FILE_EDIT_PRICE_STEP
        total = ((total + step - 1) // step) * step
        return min(total, FILE_EDIT_MAX_PRICE)

    # --------------------------------------------------------------- applying

    async def apply(self, path: str, plan: EditPlan, language: str = "uz") -> "EditResult":
        """Apply `plan` to the document, reporting what landed and what did not."""
        images = await self._render_images(plan)
        try:
            return await asyncio.to_thread(self._apply_sync, path, plan, images)
        finally:
            for image_path in images.values():
                try:
                    os.remove(image_path)
                except OSError:
                    pass

    async def _render_images(self, plan: EditPlan) -> Dict[int, str]:
        """Generate every picture the plan asks for, keyed by operation index."""
        wanted = [
            (index, operation)
            for index, operation in enumerate(plan.operations)
            if operation.op == "insert_image" and operation.image_prompt
        ]
        if not wanted or not self.together:
            return {}

        results = await asyncio.gather(
            *(self.together.generate_image(operation.image_prompt, aspect_ratio="4:3") for _, operation in wanted),
            return_exceptions=True,
        )
        images = {}
        for (index, _), result in zip(wanted, results):
            if isinstance(result, str) and result and os.path.exists(result):
                images[index] = result
            else:
                logger.warning("Image generation failed for operation %s: %s", index, result)
        return images

    def _apply_sync(self, path: str, plan: EditPlan, images: Dict[int, str]) -> EditResult:
        doc = Document(path)
        blocks = list(self._iter_blocks(doc))

        applied: List[EditOperation] = []
        failed: List[EditOperation] = []
        for index, operation in enumerate(plan.operations):
            block = self._block_at(blocks, operation.block_id)
            if block is None:
                logger.warning("Edit operation %s targets missing block %s", operation.op, operation.block_id)
                failed.append(operation)
                continue
            try:
                self._apply_one(doc, block, operation, images.get(index))
                applied.append(operation)
            except Exception as e:
                logger.error("Edit operation %s on block %s failed: %s", operation.op, operation.block_id, e)
                failed.append(operation)

        if not applied:
            raise RuntimeError("no edit operation could be applied")

        os.makedirs(TEMP_DIR, exist_ok=True)
        base = os.path.splitext(os.path.basename(path))[0]
        out_path = os.path.join(
            TEMP_DIR, f"{base}_edited_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        )
        doc.save(out_path)
        logger.info("Applied %s/%s edit operations → %s", len(applied), len(plan.operations), out_path)
        return EditResult(path=out_path, applied=applied, failed=failed)

    @staticmethod
    def _block_at(blocks: list, block_id: Optional[int]):
        if not block_id or block_id < 1 or block_id > len(blocks):
            return None
        return blocks[block_id - 1]

    def _apply_one(self, doc, block, operation: EditOperation, image_path: Optional[str]) -> None:
        if operation.op == "delete_block":
            element = block._element
            element.getparent().remove(element)
            return

        if operation.op == "replace_text":
            if isinstance(block, Table):
                raise ValueError("replace_text targets a table")
            self._set_paragraph_text(block, operation.text)
            return

        if operation.op == "add_footnote":
            if isinstance(block, Table):
                raise ValueError("add_footnote targets a table")
            if self.documents is None:
                raise RuntimeError("document service unavailable for footnotes")
            # _add_word_footnote_xml allocates the next free id itself.
            self.documents._add_footnote(block, operation.text, 1)
            return

        if operation.op == "insert_paragraph":
            new_paragraph = self._new_body_paragraph(doc, operation.text)
            block._element.addnext(new_paragraph._element)
            return

        if operation.op == "insert_table":
            self._insert_table(doc, block, operation)
            return

        if operation.op == "insert_image":
            if not image_path:
                raise RuntimeError("image was not generated")
            self._insert_image(doc, block, operation, image_path)
            return

    @staticmethod
    def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
        """Overwrite a paragraph's text while keeping its first run's formatting."""
        runs = paragraph.runs
        if runs:
            runs[0].text = text
            for run in runs[1:]:
                run._element.getparent().remove(run._element)
        else:
            paragraph.add_run(text)

    def _new_body_paragraph(self, doc, text: str) -> Paragraph:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        return paragraph

    def _insert_table(self, doc, block, operation: EditOperation) -> None:
        if not operation.headers and not operation.rows:
            raise ValueError("table has neither headers nor rows")

        columns = len(operation.headers) or max(len(row) for row in operation.rows)
        table = doc.add_table(rows=1 + len(operation.rows), cols=columns)
        table.style = "Table Grid"

        header_cells = table.rows[0].cells
        for column, title in enumerate(operation.headers[:columns]):
            header_cells[column].text = title
            for paragraph in header_cells[column].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        for row_index, row in enumerate(operation.rows, start=1):
            cells = table.rows[row_index].cells
            for column, value in enumerate(row[:columns]):
                cells[column].text = value

        anchor = block._element
        anchor.addnext(table._element)

        if operation.caption:
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption.add_run(operation.caption)
            caption_run.font.italic = True
            caption_run.font.size = Pt(12)
            caption_run.font.name = "Times New Roman"
            anchor.addnext(caption._element)

    def _insert_image(self, doc, block, operation: EditOperation, image_path: str) -> None:
        anchor = block._element

        if operation.caption:
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption.add_run(operation.caption)
            caption_run.font.italic = True
            caption_run.font.size = Pt(12)
            caption_run.font.name = "Times New Roman"
            anchor.addnext(caption._element)

        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(image_path, width=Inches(5.0))
        anchor.addnext(picture._element)


_file_edit_service_instance = None


def get_file_edit_service() -> FileEditService:
    """Return the shared FileEditService singleton."""
    global _file_edit_service_instance
    if _file_edit_service_instance is None:
        from services.ai_service import get_ai_service
        from services.document_service import get_document_service
        from services.together_service import get_together_service

        try:
            together = get_together_service()
        except Exception as e:
            logger.warning("Together service unavailable, images disabled in file edits: %s", e)
            together = None

        _file_edit_service_instance = FileEditService(
            ai_service=get_ai_service(),
            together_service=together,
            document_service=get_document_service(),
        )
        logger.info("FileEditService singleton created")
    return _file_edit_service_instance
