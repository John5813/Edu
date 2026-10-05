"""O'zbekcha yozuv: lotin va kirill.

Taqdimot mavzusi kirillda yozilsa, ilgari sarlavhalar, "Rahmat" va jadval
nomlari lotinda, qolgan matn kirillda chiqib, bir slaydning o'zida ikki
yozuv aralashib ketardi. Endi mijoz yozuvni tanlaydi, matnni modelga
shu yozuvda yozdirishga harakat qilamiz, tayyor faylda esa qolib ketgan
boshqa yozuvdagi so'zlarni shu yozuvga o'giramiz.

O'girish rus tilini emas, faqat o'zbek imlosini nazarda tutadi.
"""
import contextvars
import re
from typing import Any, Optional

LATIN = "latin"
CYRILLIC = "cyrillic"
SCRIPTS = (LATIN, CYRILLIC)

# Premium oqimda til kodi shu ko'rinishda ham keladi.
UZ_CYRILLIC_LANG = "uz-cyrl"

_CYRILLIC_RE = re.compile(r"[Ѐ-ӿ]")
_APOSTROPHES = "ʻʼ’‘`´"

# ── kirill → lotin
_C2L = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "ё": "yo", "ж": "j", "з": "z",
    "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p",
    "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "x", "ц": "ts", "ч": "ch",
    "ш": "sh", "щ": "shch", "ъ": "'", "ы": "i", "ь": "", "э": "e", "ю": "yu", "я": "ya",
    "ў": "o'", "қ": "q", "ғ": "g'", "ҳ": "h", "ҷ": "j", "ӯ": "o'",
}
_C_VOWELS = set("аеёиоуыэюяўАЕЁИОУЫЭЮЯЎ")

# ── lotin → kirill
_L_SIMPLE = {
    "a": "а", "b": "б", "d": "д", "f": "ф", "g": "г", "h": "ҳ", "i": "и", "j": "ж",
    "k": "к", "l": "л", "m": "м", "n": "н", "o": "о", "p": "п", "q": "қ", "r": "р",
    "s": "с", "t": "т", "u": "у", "v": "в", "w": "в", "x": "х", "y": "й", "z": "з",
    "c": "к",
}
_L_VOWELS = set("aeiou")


def has_cyrillic(text: str) -> bool:
    return bool(_CYRILLIC_RE.search(str(text or "")))


def _case(source: str, target: str, upper_all: bool) -> str:
    """Birinchi harf katta bo'lsa — natijaning birinchi harfi ham katta."""
    if not target:
        return target
    if source.isupper():
        return target.upper() if upper_all else target[0].upper() + target[1:]
    return target


# ───────────────────────────────────────────────────────── kirill → lotin

def _cyrillic_word_to_latin(word: str) -> str:
    out = []
    n = len(word)
    for i, char in enumerate(word):
        low = char.lower()
        if low == "е":
            prev = word[i - 1] if i else ""
            # So'z boshida, unli yoki ayiruv belgisidan keyin "ye"
            piece = "ye" if (not prev or prev in _C_VOWELS or prev in "ъЪ") else "e"
        elif low in _C2L:
            piece = _C2L[low]
        else:
            out.append(char)
            continue
        # Katta harf: "Ш" → "Sh"; butun so'z katta bo'lsa "SH"
        if char.isupper():
            following_upper = i + 1 < n and word[i + 1].isupper()
            piece = piece.upper() if (following_upper or (i and word[i - 1].isupper())) and len(piece) > 1 \
                else piece[:1].upper() + piece[1:]
        out.append(piece)
    return "".join(out)


def to_latin(text: str) -> str:
    """Kirilldagi so'zlarni o'zbek lotiniga o'giradi (lotin qismi tegilmaydi)."""
    return re.sub(r"[Ѐ-ӿ]+", lambda m: _cyrillic_word_to_latin(m.group(0)), str(text or ""))


# ───────────────────────────────────────────────────────── lotin → kirill

_LATIN_WORD = re.compile(r"[A-Za-z](?:[A-Za-z]|['" + re.escape(_APOSTROPHES) + r"](?=[A-Za-z]))*")


def _latin_word_to_cyrillic(word: str) -> str:
    text = re.sub("[" + re.escape(_APOSTROPHES) + "]", "'", word)
    low = text.lower()
    out = []
    i, n = 0, len(text)
    while i < n:
        two = low[i:i + 2]
        char = text[i]
        piece, step = None, 1
        nxt = low[i + 1:i + 2]
        after = low[i + 2:i + 3]
        if two in ("o'", "g'"):
            piece, step = ("ў" if two == "o'" else "ғ"), 2
        elif two == "sh":
            piece, step = "ш", 2
        elif two == "ch":
            piece, step = "ч", 2
        elif two == "ts":
            piece, step = "ц", 2
        elif two == "yo" and after != "'":
            piece, step = "ё", 2
        elif two == "yu":
            piece, step = "ю", 2
        elif two == "ya":
            piece, step = "я", 2
        elif two == "ye":
            piece, step = "е", 2
        elif low[i] == "e":
            prev = low[i - 1] if i else ""
            piece = "э" if (not prev or prev == "'") else "е"
        elif low[i] == "'":
            # Tutuq belgisi: "ma'no" → "маъно", "e'tibor" → "эътибор"
            piece = "ъ" if (i and i + 1 < n) else ""
        elif low[i] in _L_SIMPLE:
            piece = _L_SIMPLE[low[i]]
        else:
            piece = char
        source = text[i:i + step]
        if piece and source[:1].isupper():
            all_upper = (n > step and text[i:].isupper()) or (i and text[i - 1].isupper()) \
                or (i + step < n and text[i + step].isupper())
            piece = piece.upper() if (all_upper and len(piece) > 1) else piece[:1].upper() + piece[1:]
        out.append(piece)
        i += step
    return "".join(out)


_UNITS = {"mg", "kg", "km", "sm", "ml", "mm", "cm", "hz", "kw", "mb", "gb", "kb", "tb", "ph", "gr"}
_MATH_NEIGHBOUR = re.compile(r"[=+\-*/^()<>\d]")


def _keep_latin(token: str, before: str, after: str) -> bool:
    """Kirillga o'tkazilmaydigan lotincha bo'laklar: qisqartma, o'zgaruvchi, birlik."""
    if len(token) == 1:
        # "x = 3", "2y" — o'zgaruvchi; "u" (o'zbekcha olmosh) esa matn ichida o'giriladi.
        near = (before.rstrip(" ")[-1:] if before.endswith(" ") or not before else before[-1:]) + \
               (after.lstrip(" ")[:1])
        return bool(_MATH_NEIGHBOUR.search(near)) or token.isupper()
    if len(token) == 2 and token.lower() in _UNITS:
        return True                       # mg, kg, sm
    if token.isupper() and len(token) <= 6:
        return True                       # AI, PDF, DNK, WHO
    if any(char.isdigit() for char in (before[-1:] + after[:1])):
        return True                       # H2O, B12, COVID19
    return False


# Havola, pochta va sayt nomlari o'zgarmaydi.
_PROTECTED = (r"\$[^$]*\$|&#?\w+;|https?://\S+|www\.\S+|[\w.+-]+@[\w-]+\.[\w.]+|"
              r"\b[\w-]+\.(?:uz|com|ru|org|net|me|edu|gov)\b\S*")


def to_cyrillic(text: str) -> str:
    """Lotindagi so'zlarni o'zbek kirillchasiga o'giradi (kirill qismi tegilmaydi).

    `$...$` ichidagi formulalar, HTML belgilar (&amp;), qisqartmalar, havolalar va qisqa belgilar
    (x, y, mg) o'zgarmaydi.
    """
    source = str(text or "")
    result = []
    for index, part in enumerate(re.split("(" + _PROTECTED + ")", source)):
        if index % 2:                     # himoyalangan bo'lak
            result.append(part)
            continue
        result.append(_LATIN_WORD.sub(
            lambda m: m.group(0) if _keep_latin(m.group(0), part[:m.start()], part[m.end():])
            else _latin_word_to_cyrillic(m.group(0)), part))
    return "".join(result)


# ───────────────────────────────────────────────────────── umumiy

def convert(text: str, script: Optional[str]) -> str:
    """Matnni tanlangan yozuvga keltiradi (yozuv noma'lum bo'lsa — o'zgarmaydi)."""
    if script == LATIN:
        return to_latin(text)
    if script == CYRILLIC:
        return to_cyrillic(text)
    return text


def convert_tree(value: Any, script: Optional[str], skip_keys=("layout", "image_prompt", "prompt",
                                                              "icon", "kind", "style", "type")) -> Any:
    """JSON ko'rinishidagi mazmun (dict/list/str) ichidagi hamma matnni o'giradi."""
    if script not in SCRIPTS:
        return value
    if isinstance(value, str):
        return convert(value, script)
    if isinstance(value, list):
        return [convert_tree(item, script, skip_keys) for item in value]
    if isinstance(value, dict):
        return {key: (item if key in skip_keys else convert_tree(item, script, skip_keys))
                for key, item in value.items()}
    return value


def script_of_language(language: str) -> Optional[str]:
    """Premium til kodidan yozuv: "uz" → lotin, "uz-cyrl" → kirill, boshqasi — yo'q."""
    if language == "uz":
        return LATIN
    if language == UZ_CYRILLIC_LANG:
        return CYRILLIC
    return None


# ───────────────────────────────────────────────────────── modelga qoida

def prompt_rule(script: Optional[str]) -> str:
    """Modelga yuboriladigan yozuv qoidasi (inglizcha: model uni aniqroq bajaradi)."""
    if script == CYRILLIC:
        return ("SCRIPT RULE: write every Uzbek word in the UZBEK CYRILLIC alphabet "
                "(ў, қ, ғ, ҳ), for example «Таҳлил», «Ўзбекистон», «Эътибор». Never use Latin "
                "letters for Uzbek words and never mix the two alphabets. JSON keys, layout "
                "names and technical identifiers stay exactly as specified in the instructions; "
                "internationally written abbreviations (AI, PDF) may stay.")
    if script == LATIN:
        return ("SCRIPT RULE: write every Uzbek word in the UZBEK LATIN alphabet "
                "(oʻ, gʻ written as o', g'). Never use Cyrillic letters, even if the topic was "
                "given in Cyrillic: transliterate it into Latin.")
    return ""


_current: contextvars.ContextVar = contextvars.ContextVar("uz_script", default="")


def use(script: Optional[str]):
    """Shu jarayon ichidagi barcha model so'rovlariga yozuv qoidasini qo'shadi.

    Qaytgan belgi bilan `reset` chaqiriladi.
    """
    return _current.set(script if script in SCRIPTS else "")


def reset(token) -> None:
    try:
        _current.reset(token)
    except Exception:
        pass


def current() -> str:
    return _current.get()


def with_rule(messages: list) -> list:
    """Xabarlar ro'yxatiga joriy yozuv qoidasini qo'shadi (qoida yo'q bo'lsa — o'zgarmaydi)."""
    rule = prompt_rule(current())
    if not rule or not messages:
        return messages
    out = [dict(m) for m in messages]
    for message in reversed(out):
        if message.get("role") == "user" and isinstance(message.get("content"), str):
            message["content"] = message["content"] + "\n\n" + rule
            return out
    return [{"role": "system", "content": rule}] + out


# ───────────────────────────────────────────────────────── fayl va HTML

def normalize_pptx(path: str, script: Optional[str]) -> int:
    """PPTX ichidagi hamma matnni tanlangan yozuvga keltiradi. O'zgargan bo'laklar soni."""
    if script not in SCRIPTS:
        return 0
    from pptx import Presentation

    prs = Presentation(path)
    changed = 0

    def fix_frame(frame) -> None:
        nonlocal changed
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                new = convert(run.text, script)
                if new != run.text:
                    run.text = new
                    changed += 1

    def walk(shapes) -> None:
        for shape in shapes:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):          # guruh
                walk(shape.shapes)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                fix_frame(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        fix_frame(cell.text_frame)

    for slide in prs.slides:
        walk(slide.shapes)
        if slide.has_notes_slide:
            fix_frame(slide.notes_slide.notes_text_frame)
    if changed:
        prs.save(path)
    return changed


_TEXT_NODE = re.compile(r">([^<>]+)<")
_ATTR_PLAIN = re.compile(r'(data-(?:labels|unit|xlabel)=")([^"]*)(")')
_ATTR_SERIES = re.compile(r'(data-series=")([^"]*)(")')


def _series_names(value: str, script: str) -> str:
    """"Nomi: 1,2,3|Boshqa: 4,5" — faqat nomlar o'giriladi, formula va raqamlar emas."""
    rows = []
    for row in value.split("|"):
        name, sep, rest = row.partition(":")
        rows.append(convert(name, script) + sep + rest)
    return "|".join(rows)


def normalize_html(html: str, script: Optional[str]) -> str:
    """Slayd HTML'idagi ko'rinadigan matnlarni tanlangan yozuvga keltiradi.

    Atributlardan faqat diagramma yorliqlari, birligi va qator nomlari o'giriladi;
    rasm tavsifi (`data-prompt`), ikonka nomi va formulalar tegilmaydi.
    """
    if script not in SCRIPTS or not html:
        return html
    html = _TEXT_NODE.sub(lambda m: ">" + convert(m.group(1), script) + "<", html)
    html = _ATTR_PLAIN.sub(lambda m: m.group(1) + convert(m.group(2), script) + m.group(3), html)
    html = _ATTR_SERIES.sub(lambda m: m.group(1) + _series_names(m.group(2), script) + m.group(3), html)
    return html
