"""Yuklangan fayldan matn olish — botni yiqitmaydigan yo'l bilan.

Bir mijoz 2000 betlik kitob tashlaganda bot butunlay o'lib qolardi. Sabab
hostning kuchsizligida emas, o'qish usulida edi:

- `pdf2docx` PDF ning sahifa maketini qayta quradi (jadval, ustun, shrift).
  Bu matn o'qishdan o'nlab barobar qimmat, va bizga maket kerak emas.
- U butun hujjatni bir yo'la o'qirdi, holbuki prompt baribir 60 000 belgida
  kesiladi — qolgani behuda sarflanardi.
- Ish oddiy tredda bajarilardi. Tredni to'xtatib bo'lmaydi va unga xotira
  chegarasi qo'yib bo'lmaydi; xotira tugasa butun protsess — ya'ni bot — o'ladi.

Shuning uchun bu yerda matn betma-bet olinadi, kerakli hajmga yetganda
o'qish to'xtaydi, va og'ir ishlar navbatma-navbat bajariladi.
"""

import asyncio
import logging
import os
import zipfile

logger = logging.getLogger(__name__)

# Telegram botlarga 20 MB dan katta fayl bermaydi, shuning uchun undan
# past chegara qo'yiladi — aks holda tekshiruv hech qachon ishlamaydi.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# Prompt baribir shu atrofda kesiladi; bundan ortig'ini o'qish behuda.
MAX_EXTRACT_CHARS = 60_000
MAX_PDF_PAGES = 60
MAX_PPTX_SLIDES = 120

# Ziq siqilgan fayl ochilganda o'nlab barobar shishishi mumkin ("zip bomba").
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024

# Bir vaqtda bitta og'ir o'qish. O'n mijoz birdan kitob tashlasa, o'nta
# og'ir ish parallel ketmaydi.
_EXTRACTION_SLOT = asyncio.Semaphore(1)


class SourceTooLarge(Exception):
    """Fayl yuklab olinmasdan oldin rad etildi."""

    def __init__(self, size_bytes: int):
        self.size_mb = size_bytes / (1024 * 1024)
        super().__init__(f"fayl juda katta: {self.size_mb:.1f} MB")


class SourceUnreadable(Exception):
    """Fayl o'qildi, lekin ishlatishga yaroqli matn chiqmadi."""


class Extract:
    """O'qilgan matn va undan qancha qismi olingani."""

    def __init__(self, text: str, used_units: int = 0, total_units: int = 0, unit: str = ""):
        self.text = text
        self.used_units = used_units
        self.total_units = total_units
        self.unit = unit

    @property
    def is_partial(self) -> bool:
        return bool(self.total_units) and self.used_units < self.total_units

    @property
    def words(self) -> int:
        return len(self.text.split())


def check_size(file_size: int | None) -> None:
    """Yuklashdan OLDIN chaqiriladi — `message.document.file_size` xabar bilan keladi."""
    if file_size and file_size > MAX_UPLOAD_BYTES:
        raise SourceTooLarge(file_size)


async def read(path: str, file_name: str) -> Extract:
    """Fayldan matn oladi. Og'ir ish navbat bilan, alohida tredda bajariladi."""
    async with _EXTRACTION_SLOT:
        return await asyncio.to_thread(read_sync, path, file_name)


def read_sync(path: str, file_name: str) -> Extract:
    lowered = (file_name or path).lower()
    if lowered.endswith(".pdf"):
        extract = _read_pdf(path)
    elif lowered.endswith(".docx"):
        extract = _read_docx(path)
    elif lowered.endswith(".pptx"):
        extract = _read_pptx(path)
    else:
        raise SourceUnreadable(f"qo'llab-quvvatlanmaydigan tur: {file_name}")

    if extract.words < 20:
        raise SourceUnreadable("fayldan yetarli matn ajratilmadi")
    return extract


# ────────────────────────────────────────────────────────────────────── PDF

def _open_pdf(path: str):
    """PyMuPDF. Yangi nomi `pymupdf`, eski `fitz` — ikkalasi ham qo'llanadi."""
    try:
        import pymupdf
        return pymupdf.open(path)
    except ImportError:
        import fitz
        return fitz.open(path)


def _read_pdf(path: str) -> Extract:
    """PyMuPDF bilan betma-bet. `pdf2docx` bu yerda ishlatilmaydi — u maket
    quradi, bizga esa faqat matn kerak."""

    parts: list = []
    size = 0
    pages_read = 0
    with _open_pdf(path) as pdf:
        total = pdf.page_count
        for index in range(min(total, MAX_PDF_PAGES)):
            pages_read = index + 1
            try:
                text = pdf[index].get_text().strip()
            except Exception as e:
                logger.warning("PDF %s-betini o'qib bo'lmadi: %s", index + 1, e)
                continue
            if not text:
                continue
            parts.append(text)
            size += len(text)
            if size >= MAX_EXTRACT_CHARS:
                break

    return Extract("\n\n".join(parts)[:MAX_EXTRACT_CHARS], pages_read, total, "bet")


# ───────────────────────────────────────────────────────────────────── DOCX

def _guard_zip(path: str, inner: str) -> int:
    """Ochilgandagi hajmni tekshiradi va ichki fayl hajmini qaytaradi."""
    with zipfile.ZipFile(path) as archive:
        total = sum(item.file_size for item in archive.infolist())
        if total > MAX_UNCOMPRESSED_BYTES:
            raise SourceTooLarge(total)
        try:
            return archive.getinfo(inner).file_size
        except KeyError:
            raise SourceUnreadable(f"{inner} topilmadi")


def _read_docx(path: str) -> Extract:
    """`word/document.xml` ni oqim sifatida o'qiydi.

    python-docx butun XML daraxtini xotiraga yig'adi — katta hujjatda bu
    fayl hajmidan o'nlab barobar ko'p joy egallaydi. Bu yerda xatboshilar
    birma-bir o'qiladi va o'qilgani darhol xotiradan chiqariladi.
    """
    from lxml import etree

    _guard_zip(path, "word/document.xml")
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    parts: list = []
    size = 0
    paragraphs = 0
    with zipfile.ZipFile(path) as archive, archive.open("word/document.xml") as stream:
        context = etree.iterparse(stream, events=("end",), tag=f"{namespace}p")
        for _, element in context:
            paragraphs += 1
            text = "".join(element.itertext()).strip()
            if text:
                parts.append(text)
                size += len(text)

            # O'qib bo'lingan tugunni darhol bo'shatamiz.
            element.clear()
            parent = element.getparent()
            if parent is not None:
                while element.getprevious() is not None:
                    del parent[0]

            if size >= MAX_EXTRACT_CHARS:
                break

    return Extract("\n".join(parts)[:MAX_EXTRACT_CHARS], paragraphs, paragraphs, "xatboshi")


# ───────────────────────────────────────────────────────────────────── PPTX

def _read_pptx(path: str) -> Extract:
    from pptx import Presentation

    _guard_zip(path, "ppt/presentation.xml")
    presentation = Presentation(path)
    slides = presentation.slides

    parts: list = []
    size = 0
    used = 0
    for number, slide in enumerate(slides, start=1):
        if number > MAX_PPTX_SLIDES:
            break
        used = number
        lines = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if lines:
            block = f"[{number}-slayd] " + "\n".join(lines)
            parts.append(block)
            size += len(block)
            if size >= MAX_EXTRACT_CHARS:
                break

    return Extract("\n\n".join(parts)[:MAX_EXTRACT_CHARS], used, len(slides), "slayd")


def probe_pdf_pages(path: str) -> int:
    """Betlar sonini arzon usulda oladi — matn o'qilmaydi."""
    try:
        with _open_pdf(path) as pdf:
            return pdf.page_count
    except Exception as e:
        logger.warning("PDF betlari sanalmadi: %s", e)
        return 0


def cleanup(path: str) -> None:
    if not path:
        return
    try:
        os.remove(path)
    except OSError:
        pass
