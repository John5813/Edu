"""PPTX matnini lotin ↔ kirill (o'zbek yozuvi) o'tkazish — yashirin /abs xizmati.

Faqat ko'rinadigan matn (slayd, jadval, guruh, diagramma yorliqlari, izohlar) o'giriladi.
Rasm, joylashuv, shrift, rang, animatsiya va boshqa hamma narsa o'z holicha qoladi.
"""
import re
from dataclasses import dataclass
from typing import Callable, Iterator, List, Optional

from services import uz_script

_LATIN_LETTER = re.compile(r"[A-Za-z]")
_CYR_LETTER = re.compile(r"[Ѐ-ӿ]")
_UZ_ONLY = re.compile(r"[ўқғҳЎҚҒҲ]")
_RU_ONLY = re.compile(r"[ыщьёЫЩЬЁ]")

_ENGLISH = {"the", "and", "of", "to", "is", "are", "in", "for", "with", "that", "this", "on", "by",
            "from", "as", "it", "be", "was", "an", "or", "at", "your", "our", "we", "you", "not"}
_UZBEK = {"va", "bilan", "uchun", "bu", "bir", "ham", "yoki", "ni", "ning", "dan", "ga", "lar",
          "emas", "kerak", "bo'lib", "qilish"}
_WORD = re.compile(r"[A-Za-z']+")


class AbsError(Exception):
    """Foydalanuvchiga ko'rsatiladigan xato (kod: no_text | russian | bad_file)."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass
class AbsResult:
    script: str          # natija yozuvi (uz_script.LATIN | CYRILLIC)
    runs: int            # o'zgargan matn bo'laklari
    skipped: int         # inglizcha deb tegilmagan bo'laklar


def _looks_english(text: str) -> bool:
    words = [w.lower() for w in _WORD.findall(text)]
    if len(words) < 2 or "'" in text or "ʻ" in text or "ʼ" in text:
        return False
    english = sum(1 for w in words if w in _ENGLISH)
    uzbek = sum(1 for w in words if w in _UZBEK)
    return english >= 2 and uzbek == 0 and english * 3 >= len(words)


def _iter_frames(prs) -> Iterator:
    """Presentation ichidagi hamma matn ramkalari (slayd, guruh, jadval, izoh)."""
    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):
                yield from walk(shape.shapes)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                yield shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

    for slide in prs.slides:
        yield from walk(slide.shapes)
        if slide.has_notes_slide:
            yield slide.notes_slide.notes_text_frame


def _iter_runs(prs) -> Iterator:
    for frame in _iter_frames(prs):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                yield run


def _chart_nodes(prs) -> List:
    """Diagramma ichidagi matn tugunlari: toifa/qator nomlari va sarlavhalar (keshdagi)."""
    nodes: list = []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):
                walk(shape.shapes)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                try:
                    root = shape.chart._chartSpace
                except Exception:
                    continue
                for element in root.iter():
                    tag = element.tag if isinstance(element.tag, str) else ""
                    if tag.endswith("}v") and element.text and element.getparent() is not None \
                            and element.getparent().getparent() is not None \
                            and element.getparent().getparent().tag.endswith("}strCache"):
                        nodes.append(element)
                    elif tag.endswith("}t") and element.text and tag.startswith("{http://schemas.openxmlformats.org/drawingml"):
                        nodes.append(element)

    for slide in prs.slides:
        walk(slide.shapes)
    return nodes


def detect_target(prs) -> str:
    """Lotin matn ko'p bo'lsa — kirill, kirill ko'p bo'lsa — lotin."""
    latin = cyr = 0
    uz_marks = ru_marks = 0
    for run in _iter_runs(prs):
        text = run.text or ""
        if _looks_english(text):
            continue
        latin += len(_LATIN_LETTER.findall(text))
        found = _CYR_LETTER.findall(text)
        cyr += len(found)
        uz_marks += len(_UZ_ONLY.findall(text))
        ru_marks += len(_RU_ONLY.findall(text))
    if latin + cyr == 0:
        raise AbsError("no_text")
    if cyr > latin:
        # Rus tilidagi taqdimotni o'zbek lotiniga o'girish ma'nosiz — rad etamiz.
        if ru_marks >= 3 and ru_marks > uz_marks:
            raise AbsError("russian")
        return uz_script.LATIN
    return uz_script.CYRILLIC


def convert_pptx(src: str, dst: str, script: Optional[str] = None) -> AbsResult:
    """src dagi matnni o'girib, dst ga saqlaydi. script berilmasa — avtomatik aniqlanadi."""
    from pptx import Presentation

    try:
        prs = Presentation(src)
    except Exception as exc:
        raise AbsError("bad_file") from exc

    target = script or detect_target(prs)
    changed = skipped = 0

    def fix(text: str) -> str:
        nonlocal skipped
        if target == uz_script.CYRILLIC and _looks_english(text):
            skipped += 1
            return text
        return uz_script.convert(text, target)

    for run in _iter_runs(prs):
        new = fix(run.text or "")
        if new != run.text:
            run.text = new
            changed += 1
    for node in _chart_nodes(prs):
        new = fix(node.text)
        if new != node.text:
            node.text = new
            changed += 1

    prs.save(dst)
    return AbsResult(script=target, runs=changed, skipped=skipped)
