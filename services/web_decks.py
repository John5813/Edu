"""Saytda yaratilgan zamonaviy taqdimotning saqlangan nusxasi: sahifalar, suratlar, tarix.

Mijoz taqdimotni yuklab olmasdan avval brauzerda varaqlab ko'radi va istalgan sahifani AI ga qayta
yozdiradi. Buning uchun taqdimotning yakuniy HTML sahifalari (PPTX bilan bir xil) va har sahifaning
surati diskda saqlanadi. Hammasi buyurtma fayli bilan birga `JOB_TTL_HOURS` dan keyin o'chadi.

Papka: `<RESULTS_DIR>/decks/<job_id>/` — `deck.json` (sahifalar, reja, tarix) va `p01.jpg` ... suratlar.
"""
import asyncio
import io
import json
import logging
import os
import re
import shutil
import tempfile
import time
from typing import Dict, List, Optional

log = logging.getLogger(__name__)

PREVIEW_W = 1280          # brauzerga beriladigan surat kengligi (16:9)
PREVIEW_H = 720
_locks: Dict[str, asyncio.Lock] = {}
_ID = re.compile(r"^[0-9a-f]{32}$")


def root() -> str:
    from services import web_jobs

    return os.path.join(web_jobs.RESULTS_DIR, "decks")


def _dir(job_id: str) -> str:
    if not _ID.match(str(job_id or "")):
        raise ValueError("noto'g'ri buyurtma raqami")
    return os.path.join(root(), job_id)


def lock(job_id: str) -> asyncio.Lock:
    """Bitta taqdimotni bir vaqtda bitta ish o'zgartiradi."""
    return _locks.setdefault(job_id, asyncio.Lock())


def exists(job_id: str) -> bool:
    try:
        return os.path.isfile(os.path.join(_dir(job_id), "deck.json"))
    except ValueError:
        return False


def load(job_id: str) -> Optional[dict]:
    try:
        with open(os.path.join(_dir(job_id), "deck.json"), encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return None


def save(job_id: str, deck: dict) -> None:
    folder = _dir(job_id)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, "deck.json")
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as handle:
        json.dump(deck, handle, ensure_ascii=False)
    os.replace(tmp, path)


def preview_path(job_id: str, number: int) -> str:
    return os.path.join(_dir(job_id), f"p{int(number):02d}.jpg")


def _jpeg(png_path: str, out_path: str) -> None:
    from PIL import Image

    image = Image.open(png_path).convert("RGB")
    image = image.resize((PREVIEW_W, PREVIEW_H), Image.LANCZOS)
    tmp = out_path + ".tmp"
    image.save(tmp, "JPEG", quality=82, optimize=True)
    os.replace(tmp, out_path)


def store_shots(job_id: str, shots: Dict[int, str]) -> int:
    """`{sahifa indeksi (0 dan): png yo'li}` → `pNN.jpg`. Saqlangan suratlar soni."""
    os.makedirs(_dir(job_id), exist_ok=True)
    done = 0
    for index, path in shots.items():
        try:
            _jpeg(path, preview_path(job_id, index + 1))
            done += 1
        except Exception as exc:
            log.warning("%s: %d-sahifa surati saqlanmadi: %s", job_id, index + 1, exc)
    return done


def titles_of(pages: List[str]) -> List[str]:
    from services.premium_presentation import deck_logic, html_slides

    return [deck_logic.title_of(html_slides.source_of(page) or page) for page in pages]


def create(job_id: str, params: dict, data: dict) -> None:
    """Yangi yaratilgan taqdimotni saqlaydi (`data` — `pipeline.build_deck(deck_out=...)` natijasi)."""
    pages = list(data.get("pages") or [])
    if not pages:
        return
    shots = {}
    folder = data.get("shots_dir") or ""
    for index in range(len(pages)):
        path = os.path.join(folder, f"shot_{index + 1:02d}.png")
        if folder and os.path.exists(path):
            shots[index] = path
    deck = {
        "topic": params.get("topic", ""), "language": params.get("language", "uz"), "style": params.get("style", ""),
        "volume": params.get("volume", "kop"), "theme_key": data.get("theme_key", ""), "author": params.get("author", ""),
        "preferences": params.get("preferences", ""), "source_text": (params.get("source_text") or "")[:4000],
        "level": 2, "family": data.get("family", ""), "outline": data.get("outline") or [],
        "pages": pages, "titles": titles_of(pages), "version": 1, "history": [], "created": time.time(),
        "design_seed": data.get("design_seed")}
    store_shots(job_id, shots)
    save(job_id, deck)
    if folder:
        shutil.rmtree(folder, ignore_errors=True)


def count(deck: dict) -> int:
    return int(deck.get("count") or len(deck.get("pages") or []))


def public(job_id: str, deck: dict) -> dict:
    """Brauzerga beriladigan qisqa ko'rinish (sahifalar HTML isiz)."""
    count_ = count(deck)
    titles = list(deck.get("titles") or [])
    shown = [os.path.exists(preview_path(job_id, n + 1)) for n in range(count_)]
    editable = not deck.get("view_only")
    charts = [0] * count_
    if editable:
        from services.premium_presentation import manual_edit
        charts = [manual_edit.count_charts(deck, n) for n in range(count_)]
    plan = editable and bool(deck.get("has_plan", count_ > 3))
    return {"count": count_, "version": int(deck.get("version") or 1), "editable": editable, "has_plan": plan,
            "slides": [{"n": n + 1, "title": titles[n] if n < len(titles) else "", "image": shown[n],
                        "charts": charts[n]} for n in range(count_)],
            "history": [h for h in list(deck.get("history") or []) if not h.get("manual")][-30:]}


# ─────────────────────────────────────────── qo'lda tahrirlash (bepul): matn, diagramma, tartib

_pptx_gen: Dict[str, int] = {}


async def manual(job_id: str, target: str, op: str, version: int, **kw) -> dict:
    """Qo'lda o'zgarishni qo'llab saqlaydi; PPTX biroz keyin fonda qayta yig'iladi. Yangi `deck` ni qaytaradi."""
    from services.premium_presentation import manual_edit

    async with lock(job_id):
        deck = load(job_id)
        if not deck:
            raise manual_edit.ManualEditError("Taqdimot topilmadi.")
        if deck.get("view_only"):
            raise manual_edit.ManualEditError("Bu taqdimotni tahrirlab bo'lmaydi.")
        if int(version or 0) != int(deck.get("version") or 1):
            raise manual_edit.ManualEditError("Taqdimot boshqa joyda o'zgargan. Sahifani yangilab, qaytadan urinib ko'ring.")
        deck.setdefault("has_plan", len(deck["pages"]) > 3)
        if op == "text":
            result = await asyncio.to_thread(manual_edit.apply_text, deck, kw["index"], kw["edits"], kw["count"])
        elif op == "chart":
            result = await asyncio.to_thread(manual_edit.apply_chart, deck, kw["index"], kw["k"], kw["spec"])
        elif op == "order":
            result = await asyncio.to_thread(manual_edit.apply_order, deck, kw["order"])
        else:
            raise manual_edit.ManualEditError("Noma'lum amal.")
        try:
            if op == "order":
                _reorder(job_id, deck, result)
            else:
                pages = list(deck["pages"])
                for index, page in result["pages"].items():
                    pages[index] = page
                deck["pages"] = pages
                await asyncio.to_thread(store_shots, job_id, result["shots"])
                if result.get("title_changed"):
                    outline = list(deck.get("outline") or [])
                    if kw["index"] < len(outline) and isinstance(outline[kw["index"]], dict):
                        outline[kw["index"]] = {**outline[kw["index"]], "title": result["title"]}
                        deck["outline"] = outline
            deck.update(titles=titles_of(deck["pages"]), version=int(deck.get("version") or 1) + 1, pptx_stale=True)
            deck.setdefault("history", []).append({"index": (kw.get("index") or 0) + 1, "instruction": op,
                                                   "manual": True, "at": time.time()})
            save(job_id, deck)
        finally:
            if result.get("work"):
                shutil.rmtree(result["work"], ignore_errors=True)
    schedule_pptx(job_id, target)
    return deck


def _reorder(job_id: str, deck: dict, result: dict) -> None:
    """Sahifalar yangi tartibda: suratlar ham shu tartibga ko'chiriladi, reja (bo'lsa) yangisi."""
    mapping = result["mapping"]
    old = {}
    for index in set(i for i in mapping if i is not None):
        try:
            with open(preview_path(job_id, index + 1), "rb") as handle:
                old[index] = handle.read()
        except OSError:
            pass
    outline = list(deck.get("outline") or [])
    for position, index in enumerate(mapping):
        path = preview_path(job_id, position + 1)
        if index is not None and index in old:
            with open(path + ".tmp", "wb") as handle:
                handle.write(old[index])
            os.replace(path + ".tmp", path)
    store_shots(job_id, result["shots"])
    for extra in range(len(mapping), len(deck["pages"])):
        try:
            os.remove(preview_path(job_id, extra + 1))
        except OSError:
            pass
    blank = {"title": "", "brief": "", "category": ""}
    deck["outline"] = [dict(outline[i]) if i is not None and i < len(outline) else
                       dict(outline[1]) if i is None and len(outline) > 1 else dict(blank) for i in mapping]
    deck["pages"] = result["pages"]


def schedule_pptx(job_id: str, target: str, delay: float = 5.0) -> None:
    """Ketma-ket tahrirlarda PPTX bir marta, oxirgi o'zgarishdan keyin yig'iladi."""
    _pptx_gen[job_id] = _pptx_gen.get(job_id, 0) + 1
    gen = _pptx_gen[job_id]

    async def later():
        await asyncio.sleep(delay)
        if _pptx_gen.get(job_id) == gen:
            try:
                await ensure_pptx(job_id, target)
            except Exception as exc:
                log.warning("%s: PPTX qayta yig'ilmadi: %s", job_id, exc)

    asyncio.get_running_loop().create_task(later())


async def ensure_pptx(job_id: str, target: str) -> None:
    """Qo'lda tahrirdan keyin PPTX hali yig'ilmagan bo'lsa — hozir yig'adi (yuklab olishdan oldin ham)."""
    from services.premium_presentation import slide_edit

    async with lock(job_id):
        deck = load(job_id)
        if not deck or not deck.get("pptx_stale") or not target:
            return
        work = tempfile.mkdtemp(prefix="manualpptx_")
        try:
            path = await asyncio.to_thread(slide_edit.build_pptx, deck["pages"], work, deck)
            tmp = target + ".new"
            shutil.copyfile(path, tmp)
            os.replace(tmp, target)
        finally:
            shutil.rmtree(work, ignore_errors=True)
        deck["pptx_stale"] = False
        save(job_id, deck)


def _pptx_titles(pptx_path: str) -> List[str]:
    from pptx import Presentation

    titles = []
    for slide in Presentation(pptx_path).slides:
        text = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            text = slide.shapes.title.text_frame.text
        if not text.strip():
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    text = shape.text_frame.text
                    break
        titles.append(re.sub(r"\s+", " ", text).strip()[:80])
    return titles


def pptx_previews(pptx_path: str, out_dir: str) -> List[str]:
    """PPTX → PDF (LibreOffice) → har sahifa PNG. Qaytaradi: PNG yo'llari (tartib bilan)."""
    import subprocess

    import pymupdf

    os.makedirs(out_dir, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{out_dir}/profile", "--headless", "--convert-to", "pdf",
                    "--outdir", out_dir, pptx_path], check=True, capture_output=True, timeout=240)
    pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
    paths = []
    with pymupdf.open(pdf) as doc:
        for number, page in enumerate(doc, 1):
            zoom = PREVIEW_W / max(page.rect.width, 1)
            path = os.path.join(out_dir, f"v{number:02d}.png")
            page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).save(path)
            paths.append(path)
    return paths


def create_view(job_id: str, params: dict, pptx_path: str) -> int:
    """Oddiy (orqa fonli) taqdimot: faqat ko'rish. Sahifalarni LibreOffice chizadi; qayta yozish yo'q."""
    import tempfile

    work = tempfile.mkdtemp(prefix="deckview_")
    try:
        paths = pptx_previews(pptx_path, work)
        if not paths:
            return 0
        done = store_shots(job_id, {index: path for index, path in enumerate(paths)})
        titles = []
        try:
            titles = _pptx_titles(pptx_path)
        except Exception as exc:
            log.debug("Slayd sarlavhalari olinmadi: %s", exc)
        save(job_id, {"view_only": True, "topic": params.get("topic", ""), "count": len(paths), "pages": [],
                      "titles": (titles + [""] * len(paths))[:len(paths)], "version": 1, "history": [],
                      "created": time.time()})
        return done
    finally:
        shutil.rmtree(work, ignore_errors=True)


def drop(job_id: str) -> None:
    try:
        shutil.rmtree(_dir(job_id), ignore_errors=True)
    except ValueError:
        pass
    _locks.pop(job_id, None)


async def purge_orphans() -> int:
    """Buyurtmasi o'chirilgan (muddati o'tgan) taqdimot papkalarini o'chiradi."""
    from database import web_store

    base = root()
    if not os.path.isdir(base):
        return 0
    removed = 0
    for name in os.listdir(base):
        if not _ID.match(name):
            continue
        if await web_store.get_job(name) is None:
            drop(name)
            removed += 1
    return removed
