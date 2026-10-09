""""Rahmat" sahifasi: tayyor taqdimot oxiriga harakatlanuvchi personaj, ovoz va matn animatsiyasi qo'shadi.

Nega video emas: PowerPoint videoni pleer qilib ko'rsatadi — o'zi boshlanmaydi, telefonda bosilsa ekranni
qoplab oladi. Shuning uchun personaj — GIF (rasm kabi turadi va o'zi harakatlanadi), ovoz — slayddan
tashqaridagi ko'rinmas audio, yozuvlar — PowerPoint'ning o'z animatsiyalari (navbat bilan paydo bo'ladi).

GIF shaffofligi qirrali bo'lgani uchun personaj sahifaning o'ng panelida, panel rangidagi fon bilan
chiziladi: panel va GIF foni aynan bir xil rang — chegara ko'rinmaydi.
"""
import logging
import os
import shutil
import subprocess
import tempfile
import zipfile
from typing import Optional, Tuple

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from services.premium_presentation.designer.kit import MIDDLE, SANS, fill, grad, line, mix, nofill, oval, px, rect, text
from services.premium_presentation.designer.styles import WHITE, Palette

from .catalog import Animation

log = logging.getLogger(__name__)

GIF_FPS = 12
GIF_HEIGHT = 600        # sahifada 900 px ga cho'ziladi
PANEL_X, PANEL_W = 1060, 860
CHAR_H = 900

TEXTS = {
    "uz": ("SO'NGGI SO'Z", "E'tiboringiz uchun rahmat!", "Savollaringiz bo'lsa, marhamat", "Muallif"),
    "uz-cyrl": ("СЎНГГИ СЎЗ", "Эътиборингиз учун раҳмат!", "Саволларингиз бўлса, марҳамат", "Муаллиф"),
    "ru": ("В ЗАВЕРШЕНИЕ", "Спасибо за внимание!", "Буду рад ответить на вопросы", "Автор"),
    "en": ("FINAL WORD", "Thank you for your attention!", "Questions are welcome", "Author"),
    "kk": ("СОҢҒЫ СӨЗ", "Назарларыңызға рахмет!", "Сұрақтарыңыз болса, марқабат", "Автор"),
}


def texts(language: str) -> Tuple[str, str, str, str]:
    from services import uz_script

    if language == uz_script.UZ_CYRILLIC_LANG:
        return TEXTS["uz-cyrl"]
    return TEXTS.get(language) or TEXTS["en"]


def _run(cmd) -> None:
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if res.returncode != 0:
        raise RuntimeError("ffmpeg: " + res.stderr[-400:])


def _size(path: str) -> Tuple[int, int]:
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height",
                          "-of", "csv=p=0", path], capture_output=True, text=True, timeout=30).stdout
    w, h = out.strip().split(",")[:2]
    return int(w), int(h)


def make_gif(anim: Animation, panel: str, out: str) -> str:
    """Personaj GIF i panel rangidagi fonda. Qaytaradi: GIF dagi haqiqiy fon rangi (palitra uni biroz siljitadi)."""
    packed = anim.path("packed.mp4")
    pw, ph = _size(packed)
    w = int(pw // 2 * GIF_HEIGHT / ph) // 2 * 2
    graph = (f"[0:v]fps={GIF_FPS},split[l][r];[l]crop=iw/2:ih:0:0[c];[r]crop=iw/2:ih:iw/2:0,format=gray[m];"
             f"[c][m]alphamerge,scale={w}:{GIF_HEIGHT}:flags=lanczos[ca];"
             f"color=c=0x{panel}:s={w}x{GIF_HEIGHT}:r={GIF_FPS}[bg];[bg][ca]overlay=shortest=1,"
             "split[a][b];[a]palettegen=max_colors=128:stats_mode=diff[p];"
             "[b][p]paletteuse=dither=none:diff_mode=rectangle")
    _run(["ffmpeg", "-v", "error", "-y", "-i", packed, "-filter_complex", graph, "-loop", "0", out])
    with Image.open(out) as im:
        colours = im.convert("RGB").getcolors(1 << 16) or [(0, (0, 0, 0))]
    r, g, b = max(colours)[1]                  # eng ko'p uchraydigan rang — fon
    return f"{r:02X}{g:02X}{b:02X}"


def make_audio(anim: Animation, out: str) -> Optional[str]:
    try:
        _run(["ffmpeg", "-v", "error", "-y", "-i", anim.path("packed.mp4"), "-vn", "-c:a", "copy", out])
    except RuntimeError as exc:
        log.info("Animatsiyada ovoz yo'q (%s): %s", anim.key, exc)
        return None
    return out if os.path.getsize(out) > 1000 else None


def _duration_ms(path: str) -> int:
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", path],
                         capture_output=True, text=True).stdout.strip()
    try:
        return int(float(out) * 1000)
    except ValueError:
        return 10000


P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'


def _fade(cid: int, spid: int, delay: int, dur: int = 600) -> str:
    """Kirish animatsiyasi: "Paydo bo'lish" (fade), boshqasi bilan birga, kechikish bilan."""
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" '
            f'nodeType="withEffect"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid + 1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'</p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility'
            f'</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid + 2}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')


def _timing(steps, audio_id: Optional[int], audio_ms: int) -> str:
    """Slayd ochilishi bilan (bosish shart emas) yozuvlar navbat bilan paydo bo'ladi, ovoz boshlanadi."""
    inner, cid = [], 5
    if audio_id:
        inner.append(f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" '
                     f'nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                     f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{cid + 1}" dur="{audio_ms}" '
                     f'fill="hold"/><p:tgtEl><p:spTgt spid="{audio_id}"/></p:tgtEl></p:cBhvr></p:cmd>'
                     f'</p:childTnLst></p:cTn></p:par>')
        cid += 2
    for spid, delay in steps:
        inner.append(_fade(cid, spid, delay))
        cid += 3
    media = ""
    if audio_id:
        media = (f'<p:audio><p:cMediaNode vol="80000"><p:cTn id="{cid}" fill="hold" display="0"><p:stCondLst>'
                 f'<p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{audio_id}"/></p:tgtEl>'
                 f'</p:cMediaNode></p:audio>')
    builds = "".join(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>' for spid, _ in steps)
    return (f'<p:timing {P_NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            f'<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{"".join(inner)}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            f'</p:seq>{media}</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{builds}</p:bldLst></p:timing>')


def _fix_audio_rels(path: str, rels_name: str) -> None:
    """python-pptx audio aloqasini "video" deb yozadi — PowerPoint uni ovoz deb tanishi uchun almashtiramiz.

    Faqat "Rahmat" sahifasining aloqalari o'zgaradi: unda video yo'q, bitta "video" aloqa — shu ovoz.
    """
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == rels_name:
                data = data.replace(b'officeDocument/2006/relationships/video"',
                                    b'officeDocument/2006/relationships/audio"')
            zout.writestr(item, data)
    shutil.move(tmp, path)


def add_thanks_slide(pptx_in: str, pptx_out: str, anim: Animation, theme, language: str = "uz",
                     author: str = "") -> str:
    """`pptx_in` oxiriga "Rahmat" sahifasini qo'shib `pptx_out` ga yozadi."""
    pal = Palette.of(theme)
    panel = mix(WHITE, pal.accent, 0.1)
    work = tempfile.mkdtemp(prefix="thanks_")
    try:
        gif = os.path.join(work, "thanks_character.gif")
        panel = make_gif(anim, panel, gif)
        audio = make_audio(anim, os.path.join(work, "thanks_audio.m4a"))

        prs = Presentation(pptx_in)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        base = rect(s, 0, 0, 1920, 1080)
        grad(base, [(0, mix(pal.dark, pal.accent, 0.3)), (100, pal.deep)], angle=45)
        for r, a in ((520, 30), (400, 22), (280, 14)):
            ring = oval(s, 0, 1080, r)
            nofill(ring)
            line(ring, pal.glowing, 3, a)
        glow = oval(s, 900, 150, 200)
        grad(glow, [(0, pal.accent, 40), (100, pal.accent, 0)], radial=True)
        fill(rect(s, PANEL_X, 0, PANEL_W, 1080), panel)

        label, title, sub, author_word = texts(language)
        steps = []
        t_label = text(s, 150, 250, 820, 60, [[(label, 28, pal.glowing, True, SANS, 300)]])
        bar = rect(s, 150, 320, 140, 8)
        fill(bar, pal.accent)
        t_title = text(s, 150, 360, 860, 330, [[(title, 84, WHITE, True, pal.head_font)]], anchor=MIDDLE,
                       line_spacing=1.05)
        t_sub = text(s, 150, 720, 820, 70, [[(sub, 34, pal.on_dark)]])
        t_author = text(s, 150, 900, 820, 50, [[(f"{author_word}: {author}", 22, pal.glowing)]]) if author else None

        with Image.open(gif) as im:
            gw, gh = im.size
        cw = int(gw * CHAR_H / gh)
        if cw > PANEL_W - 40:                      # juda keng personaj panelga sig'sin
            cw = PANEL_W - 40
        ch = int(gh * cw / gw)
        s.shapes.add_picture(gif, px(PANEL_X + (PANEL_W - cw) / 2), px(1080 - ch - 30), px(cw), px(ch))

        audio_id = None
        if audio:
            poster = os.path.join(work, "thanks_icon.png")
            Image.new("RGBA", (8, 8), (0, 0, 0, 0)).save(poster)
            # Slayddan tashqarida: belgisi ko'rinmaydi, ovoz esa slayd ochilganda o'zi boshlanadi.
            mov = s.shapes.add_movie(audio, px(-300), px(40), px(120), px(120), poster_frame_image=poster,
                                     mime_type="audio/mp4")
            node = mov._element.nvPicPr.nvPr.find(qn("a:videoFile"))
            node.tag = qn("a:audioFile")
            audio_id = mov.shape_id

        delay = 200
        for shape, gap in ((t_label, 300), (bar, 300), (t_title, 900), (t_sub, 500), (t_author, 0)):
            if shape is not None:
                steps.append((shape.shape_id, delay))
                delay += gap
        old = s._element.find(qn("p:timing"))
        new = etree.fromstring(_timing(steps, audio_id, _duration_ms(audio) if audio else 0))
        if old is not None:
            old.addprevious(new)
            old.getparent().remove(old)
        else:
            s._element.append(new)
        prs.save(pptx_out)
        if audio:
            part = s.part.partname                 # /ppt/slides/slideN.xml
            _fix_audio_rels(pptx_out, f"ppt/slides/_rels/{os.path.basename(part)}.rels")
        return pptx_out
    finally:
        shutil.rmtree(work, ignore_errors=True)
