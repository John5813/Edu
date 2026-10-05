"""Oddiy taqdimot: 1, 2, 3 raqamli doiralar o'rniga ikonkalar.

    python test_ikonka_raqam.py
"""
import asyncio, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

from pptx import Presentation
from services.document_service import DocumentService
from services.template_service import TemplateService

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

LAYOUTS = {"four_cards": 4, "numbered_list": 4, "bullets_icons": 4, "timeline": 4,
           "staircase": 4, "icon_row": 4, "image_left_bullets": 4, "image_right_cards": 3}

def slides():
    out = [{"title": "Mavzu", "layout": "cover", "content": ""}]
    for name, n in LAYOUTS.items():
        out.append({"title": f"{name} sarlavhasi", "layout": name,
                    "items": [{"head": h, "text": "Bu yerda mavzuga oid to'liq izoh matni yoziladi. Ikkinchi gap ham bor.", "value": "20" + str(20 + i)}
                              for i, h in enumerate(["Ta'lim", "Salomatlik", "Texnologiya", "Iqtisodiyot"][:n])]})
    out.append({"title": "Xulosa", "layout": "conclusion", "content": "Xulosa matni."})
    return out

async def build(use_icons):
    svc = DocumentService.__new__(DocumentService)
    svc.documents_dir = tempfile.mkdtemp(); svc.together = None
    svc.use_icons = use_icons; svc._last_used_icons = set()
    return await svc.create_presentation_with_smart_images("Sinov", {"slides": slides()}, "Talaba", "uz",
                                                           TemplateService(), "template_20")

def names(path):
    prs = Presentation(path)
    return [(i, sh.name, sh.text_frame.text if sh.has_text_frame else "") for i, s in enumerate(prs.slides, 1) for sh in s.shapes]

on = names(asyncio.run(build(True)))
numbers = [x for x in on if x[1] in ("badge_text", "fixed:badge_text", "num", "fixed:num")]
icons = [x for x in on if "icon" in x[1]]
check("ikonkali: raqamli doira yo'q", not numbers, numbers[:3])
check("ikonkali: ikonkalar qo'yilgan", len(icons) >= 15, len(icons))

off = names(asyncio.run(build(False)))
check("ikonkasiz tanlansa raqamlar qoladi (mijoz tanlovi)", any(x[1] in ("badge_text", "fixed:badge_text") for x in off))

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
