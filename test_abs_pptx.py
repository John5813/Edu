"""Yashirin /abs xizmati: PPTX matnini lotin ↔ kirill o'girish, boshqa hech narsaga tegmasdan.

    python test_abs_pptx.py
"""
import asyncio, os, sys, tempfile, zipfile, hashlib

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from services import uz_script
from services.pptx_script import convert_pptx, AbsError

TMP = tempfile.mkdtemp()

def build(path, title, body, cell, note, english="The market is growing and the demand is high"):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = body
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(23)
    tb = s.shapes.add_table(2, 2, Inches(1), Inches(3.5), Inches(6), Inches(1.5)).table
    tb.cell(0, 0).text = cell
    tb.cell(1, 1).text = "2024"
    eng = s.shapes.add_textbox(Inches(1), Inches(5.2), Inches(6), Inches(0.6))
    eng.text_frame.text = english
    url = s.shapes.add_textbox(Inches(1), Inches(6), Inches(6), Inches(0.5))
    url.text_frame.text = "https://example.com/info"
    cd = CategoryChartData(); cd.categories = [title.split()[0], "B"]; cd.add_series("S", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(6.6), Inches(3), Inches(1), cd)
    s.notes_slide.notes_text_frame.text = note
    prs.save(path)

def texts(path):
    out = []
    prs = Presentation(path)
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame: out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [c.text for r in sh.table.rows for c in r.cells]
        out.append(sl.notes_slide.notes_text_frame.text)
    return out

def geometry(path):
    prs = Presentation(path)
    return (prs.slide_width, prs.slide_height,
            [(sh.shape_type, sh.left, sh.top, sh.width, sh.height) for sl in prs.slides for sh in sl.shapes])

print("1) Lotin → kirill")
src = os.path.join(TMP, "lat.pptx"); out = os.path.join(TMP, "lat_out.pptx")
build(src, "Iqtisodiy o'sish omillari", "Bozor iqtisodiyoti va xalq farovonligi", "Ko'rsatkich", "Shaxsiy izoh bilan")
r = convert_pptx(src, out)
t = texts(out)
check("yo'nalish kirill", r.script == uz_script.CYRILLIC)
check("sarlavha", "Иқтисодий ўсиш омиллари" in t, t)
check("matn", any("Бозор иқтисодиёти" in x for x in t), t)
check("jadval katagi", "Кўрсаткич" in t, t)
check("izoh", any("Шахсий изоҳ" in x for x in t), t)
check("raqam o'zgarmadi", "2024" in t)
check("inglizcha tegilmadi", "The market is growing and the demand is high" in t, t)
check("havola tegilmadi", "https://example.com/info" in t)
check("slayd o'lchami va shakllar joyi bir xil", geometry(src) == geometry(out))
check("shrift o'lchami saqlandi", Presentation(out).slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size == Pt(23))
with zipfile.ZipFile(out) as z:
    cx = z.read([n for n in z.namelist() if n.startswith("ppt/charts/chart")][0]).decode("utf8")
check("diagramma toifasi o'girildi", "Иқтисодий" in cx)

print("2) Kirill → lotin")
back = os.path.join(TMP, "back.pptx")
r2 = convert_pptx(out, back)
t2 = texts(back)
check("yo'nalish lotin", r2.script == uz_script.LATIN)
check("sarlavha qaytdi", "Iqtisodiy o'sish omillari" in t2, t2)
check("jadval qaytdi", "Ko'rsatkich" in t2, t2)

print("3) Rus va matnsiz fayllar")
ru = os.path.join(TMP, "ru.pptx")
build(ru, "Факторы экономического роста", "Рыночная экономика и благосостояние народа", "Показатель", "Личные заметки", english="Ы")
try:
    convert_pptx(ru, os.path.join(TMP, "ru_out.pptx")); check("rus rad etildi", False)
except AbsError as e:
    check("rus rad etildi", e.code == "russian", e.code)
empty = os.path.join(TMP, "empty.pptx"); p = Presentation(); p.slides.add_slide(p.slide_layouts[6]); p.save(empty)
try:
    convert_pptx(empty, os.path.join(TMP, "e_out.pptx")); check("matnsiz rad etildi", False)
except AbsError as e:
    check("matnsiz rad etildi", e.code == "no_text", e.code)
bad = os.path.join(TMP, "bad.pptx"); open(bad, "wb").write(b"not a pptx")
try:
    convert_pptx(bad, os.path.join(TMP, "b_out.pptx")); check("buzuq fayl rad etildi", False)
except AbsError as e:
    check("buzuq fayl rad etildi", e.code == "bad_file", e.code)

print("4) Bot oqimi")
from bot.handlers import abs_converter as ac
from bot.states import AbsStates

class FakeState:
    def __init__(self): self.s = None
    async def clear(self): self.s = None
    async def set_state(self, s): self.s = s

class FakeBot:
    async def get_file(self, fid):
        class F: file_path = src
        return F()
    async def download_file(self, path, dest):
        import shutil; shutil.copy(path, dest)

class Doc:
    def __init__(self, name, size=1000): self.file_name = name; self.file_id = "x"; self.file_size = size

class Msg:
    def __init__(self, doc=None): self.document = doc; self.bot = FakeBot(); self.sent = []; self.docs = []
    async def answer(self, text, **kw):
        self.sent.append(text)
        class M:
            async def delete(self_): pass
        return M()
    async def answer_document(self, f, caption=None, **kw): self.docs.append((f.filename, caption))

async def flow():
    st = FakeState(); m = Msg()
    await ac.abs_start(m, st, "uz")
    check("fayl so'raladi", st.s == AbsStates.waiting_for_file and m.sent)
    m2 = Msg(Doc("a.docx")); await ac.abs_file(m2, st, "uz")
    check("docx rad etildi", m2.sent == [ac._t("uz", "not_pptx")] and st.s is not None)
    m3 = Msg(Doc("a.pptx", 50 * 1024 * 1024)); await ac.abs_file(m3, st, "uz")
    check("katta fayl rad etildi", m3.sent == [ac._t("uz", "too_big")])
    m4 = Msg(Doc("Taqdimot.pptx")); await ac.abs_file(m4, st, "uz")
    check("fayl qaytarildi", m4.docs and m4.docs[0][0] == "Taqdimot_kirill.pptx", m4.docs)
    check("holat tozalandi", st.s is None)
    left = [f for f in os.listdir("temp") if f.startswith("abs_")] if os.path.isdir("temp") else []
    check("vaqtinchalik fayllar o'chdi", not left, left)
asyncio.run(flow())

src_main = open("main.py", encoding="utf8").read()
check("router ulangan (documents/start dan oldin)", "abs_converter.router" in src_main and src_main.index("abs_converter.router") < src_main.index("documents.router"))
check("menyu/yordamda ko'rinmaydi", "/abs" not in open("translations.py", encoding="utf8").read())

print("\nXATO:" if FAILS else "\nHAMMASI YAXSHI", FAILS or "")
sys.exit(1 if FAILS else 0)
