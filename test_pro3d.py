"""3D Pro taqdimot: fonsiz 3D obyekt, yorliq-chiziqlar va PowerPoint Morph o'tishlari.

- AI rejasi tekshiriladi va tuzatiladi (muqova/xulosa, "focus" oldingi obyektni oladi, kam elementli
  joylashuv "points" ga aylanadi, so'z chegaralari);
- fon olib tashlanadi (bu yerda zaxira usul — model yuklanmaydi), shaffof chetlar kesiladi;
- yorliq nuqtasi obyekt ustiga tushadi (vision javobi shaffof joyni ko'rsatsa ham);
- PPTX: har slaydda !!panel/!!obj, ikkinchisidan boshlab Morph (Fallback — Fade), yorliqlar animatsiyasi;
- to'liq oqim soxta AI va soxta rasm xizmati bilan; rasm chizilmasa — xato (pul qaytariladi).

    python test_pro3d.py
"""
import asyncio, json, os, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")
os.environ["PRO3D_CUTOUT_MODEL"] = os.path.join(tempfile.gettempdir(), "yoq_model.onnx")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.ns import qn

import services.pro3d as pro3d
from services.pro3d import build, content, cutout, images
from services.premium_presentation import llm_client

cutout._session_failed = True          # testda 180 MB model yuklanmasin — zaxira usul
TMP = tempfile.mkdtemp()


def sample(path, color=(200, 60, 80), bg=(255, 255, 255)):
    """Bir xil fonda "3D obyekt": soyali doira va tayoqcha."""
    im = Image.new("RGB", (512, 512), bg)
    d = ImageDraw.Draw(im)
    for r in range(150, 0, -6):
        k = r / 150
        d.ellipse((256 - r, 220 - r, 256 + r, 220 + r), fill=tuple(int(c * (0.55 + 0.45 * k)) for c in color))
    d.rectangle((236, 360, 276, 470), fill=(90, 90, 110))
    im.save(path)
    return path


print("1) Reja (normalize)")
raw = {"palette": "emerald", "render_style": "realistic", "slides": [
    {"layout": "callouts", "title": "Yurak", "object": {"id": "a", "prompt": "human heart"},
     "callouts": [{"label": "Chap qorincha", "text": "so'z " * 40, "part": "left ventricle"},
                  {"label": "Aorta", "text": "Qonni tanaga olib chiqadi", "part": "aorta"}]},
    {"layout": "callouts", "title": "Bitta yorliq", "object": {"id": "b", "prompt": "valve"},
     "callouts": [{"label": "Klapan", "text": "Qon orqaga qaytmaydi", "part": "valve"}]},
    {"layout": "focus", "title": "Yaqindan", "object": {"id": "zzz", "prompt": "other"}, "lead": "L", "text": "T"},
    {"layout": "nomalum", "title": "X", "object": {"id": "c", "prompt": "lungs"}, "points": [{"head": "A", "text": "B"}]},
    {"layout": "stats", "title": "Raqamlar", "object": {"id": "d", "prompt": "d"}, "stats": []},
    {"layout": "points", "title": "Oxiri", "object": {"id": "e", "prompt": "e"}, "points": [{"head": "1", "text": "Bir"}]},
]}
deck = content.normalize(raw, "Yurak", 6)
layouts = [s["layout"] for s in deck["slides"]]
check("birinchisi muqova, oxirgisi xulosa", layouts[0] == "cover" and layouts[-1] == "conclusion", layouts)
check("bitta yorliqli 'callouts' → 'points'", layouts[1] == "points" and deck["slides"][1]["points"], deck["slides"][1])
check("'focus' oldingi slayd obyektini oladi (Morph yaqinlashuvi)", deck["slides"][2]["object"]["id"] == "b",
      deck["slides"][2]["object"])
check("noma'lum joylashuv → 'points'", layouts[3] == "points", layouts)
check("bo'sh 'stats' → 'points'", layouts[4] == "points", layouts)
check("xulosa muqova obyektini oladi", deck["slides"][-1]["object"]["id"] == deck["slides"][0]["object"]["id"])
check("palitra va uslub", deck["palette"] == "emerald" and deck["style"] == "realistic")
bad = content.normalize({"palette": "pushti", "render_style": "x", "slides": raw["slides"]}, "Yurak", 6)
check("noma'lum palitra/uslub — sukut", bad["palette"] == "indigo" and bad["style"] == "realistic")
long = content.normalize({"slides": [{"layout": "cover", "title": "T", "object": {"id": "a", "prompt": "p"}},
                                     {"layout": "callouts", "title": "C", "object": {"id": "b", "prompt": "p"},
                                      "callouts": [{"label": "A", "text": "so'z " * 40, "part": "a"},
                                                   {"label": "B", "text": "qisqa", "part": "b"}]},
                                     {"layout": "conclusion", "title": "X", "points": ["a"]}]}, "T", 3)
check("yorliq matni 18 so'zgacha qisqaradi", len(long["slides"][1]["callouts"][0]["text"].split()) <= 18,
      long["slides"][1]["callouts"][0]["text"])
alias = content.normalize({"slides": [
    {"layout": "cover", "title": "T", "object": {"id": "a", "prompt": "p"}},
    {"layout": "timeline", "title": "Bosqichlar", "object": {"id": "b", "prompt": "p"},
     "items": [{"year": "1628", "description": "Garvey qon aylanishini tushuntirdi"},
               {"stage": "1967", "event": "Birinchi yurak ko'chirib o'tkazildi"}]},
    {"layout": "points", "title": "P", "object": {"id": "c", "prompt": "p"},
     "items": [{"title": "Himoya", "description": "Miyani asraydi"}, "Oddiy satr"]},
    {"layout": "conclusion", "title": "X", "points": [{"head": "Bir"}, "Ikki"]}]}, "T", 4)
check("boshqa nomli maydonlar ham olinadi (items/year/description)",
      alias["slides"][1]["layout"] == "timeline" and alias["slides"][1]["steps"][0]["when"] == "1628"
      and "Garvey" in alias["slides"][1]["steps"][0]["text"], alias["slides"][1])
check("punktlar: title/description va oddiy satr", [p["head"] for p in alias["slides"][2]["points"]] == ["Himoya", ""]
      and alias["slides"][2]["points"][1]["text"] == "Oddiy satr", alias["slides"][2]["points"])
check("xulosa punktlari: dict va satr", alias["slides"][3]["points"] == ["Bir", "Ikki"], alias["slides"][3]["points"])
objs = content.objects(deck)
check("noyob obyektlar (takror chizilmaydi)", len(objs) == len(set(objs)) and "zzz" not in objs, objs)

seen = {}
def fake_plan(system, user, **kw):
    seen["user"] = user
    return raw
orig = llm_client._call_openrouter
llm_client._call_openrouter = fake_plan
try:
    planned = content.plan("Yurak", 6, "ru")
finally:
    llm_client._call_openrouter = orig
check("reja so'rovida taqdimot tili va slaydlar soni", "русском" in seen["user"] and "exactly 6" in seen["user"],
      seen["user"][:300])
check("reja tekshirildi", planned["slides"][0]["layout"] == "cover")

print("\n2) Fon olib tashlash (zaxira usul) va kesish")
src = sample(os.path.join(TMP, "obj.png"))
cut, method = cutout.cut(Image.open(src))
check("fon olib tashlandi", method == "oq fon" and cutout.is_transparent(cut), method)
alpha = np.asarray(cut.split()[-1])
check("burchak shaffof, markaz to'liq", alpha[5, 5] == 0 and alpha[220, 256] > 240, (alpha[5, 5], alpha[220, 256]))
trimmed = cutout.trim(cut)
check("shaffof chetlar kesildi", trimmed.width < 512 and trimmed.height < 512, trimmed.size)
white = Image.new("RGB", (300, 300), (255, 255, 255))
same, method = cutout.cut(white)
check("bo'sh (faqat oq) rasm — fon olib tashlanmaydi", method == "", method)
pic = images._finish(src)
check("Picture: shaffof PNG va kichik niqob", pic and pic.transparent and pic.alpha is not None and os.path.exists(pic.path))

print("\n2b) Rangli fon (chroma key): oq obyekt, ichki teshik, soya")
green = Image.new("RGB", (512, 512), (0, 177, 64))
d = ImageDraw.Draw(green)
d.ellipse((96, 96, 416, 416), fill=(246, 244, 238))          # oq obyekt (oq fonda adashib ketardi)
d.ellipse((216, 216, 296, 296), fill=(0, 177, 64))            # ichki teshik — fon ko'rinadi
d.ellipse((150, 440, 360, 480), fill=(0, 92, 33))             # to'q yashil soya
d.ellipse((20, 20, 30, 30), fill=(250, 250, 250))             # mayda dog' (chiqindi)
cut, method = cutout.cut(green, "green")
a = np.asarray(cut.split()[-1])
check("yashil fon usuli ishladi", method == "yashil fon", method)
check("oq obyekt to'liq qoldi", a[150, 256] > 240 and a[400, 256] > 240, (a[150, 256], a[400, 256]))
check("ichki teshik shaffof", a[256, 256] < 10, a[256, 256])
check("soya va burchak shaffof", a[460, 256] < 60 and a[5, 5] == 0, (a[460, 256], a[5, 5]))
check("mayda dog' o'chdi", a[25, 25] < 10, a[25, 25])
edge = np.asarray(cut.convert("RGB"))[256, 97]
check("chetda yashil nur qolmadi (despill)", edge[1] <= max(edge[0], edge[2]) + 2, edge)
leafy = Image.new("RGB", (512, 512), (255, 0, 255))
ImageDraw.Draw(leafy).ellipse((96, 96, 416, 416), fill=(40, 170, 60))
cut, method = cutout.cut(leafy, "magenta")
check("yashil obyekt pushti fonda butun qoladi", method == "pushti fon" and np.asarray(cut.split()[-1])[256, 256] > 240,
      method)
check("fon rangi: odatda yashil, yashil obyektga pushti",
      images.key_for("human skull") == "green" and images.key_for("fresh green lettuce leaves") == "magenta")
cut, method = cutout.cut(Image.open(src), "green")
check("model yashil fon bermasa — oddiy usulga o'tadi", method == "oq fon", method)

print("\n3) Yorliq nuqtalari")
def fake_request(kind, payload, **kw):
    check("vision so'rovida rasm bor", payload["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image"))
    return {"choices": [{"message": {"content": json.dumps({"points": [{"n": 1, "x": 500, "y": 400},
                                                                        {"n": 2, "x": 0, "y": 0}]})}}]}
orig_req = llm_client._request
llm_client._request = fake_request
try:
    found = images.locate(pic, "a red ball", ["top", "corner"])
finally:
    llm_client._request = orig_req
check("ikki nuqta qaytdi", len(found) == 2 and all(found), found)
h, w = pic.alpha.shape
x, y = found[1]
check("shaffof burchakni ko'rsatgan nuqta obyekt ustiga tushdi", pic.alpha[min(h - 1, int(y * h)), min(w - 1, int(x * w))] > 96,
      found[1])
llm_client._request = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("tarmoq yo'q"))
try:
    check("vision ishlamasa — None lar (zaxira chiziq chetidan)", images.locate(pic, "x", ["a", "b"]) == [None, None])
finally:
    llm_client._request = orig_req
edges = images.edge_points(pic, ["left", "right"], [0.4, 0.4])
check("zaxira nuqtalar obyekt ustida", all(pic.alpha[int(y * h), min(w - 1, int(x * w))] > 96 for x, y in edges), edges)

print("\n4) PPTX: Morph, nomlar, animatsiya")
full = content.normalize({"palette": "graphite", "slides": [
    {"layout": "cover", "title": "Inson yuragi", "subtitle": "Tana nasosi", "object": {"id": "a", "prompt": "heart"}},
    {"layout": "callouts", "title": "Tuzilishi", "object": {"id": "a", "prompt": "heart"},
     "callouts": [{"label": "Aorta", "text": "Eng katta arteriya", "part": "aorta"},
                  {"label": "Qorincha", "text": "Qonni haydaydi", "part": "ventricle"},
                  {"label": "Bo'lmacha", "text": "Qonni qabul qiladi", "part": "atrium"}]},
    {"layout": "points", "title": "Vazifasi", "object": {"id": "b", "prompt": "b"},
     "points": [{"head": "Nasos", "text": "Qonni haydaydi"}, {"head": "Ritm", "text": "Daqiqasiga 70 marta"}]},
    {"layout": "focus", "title": "Klapan", "object": {"id": "b"}, "lead": "Bir tomonlama", "text": "Qon orqaga qaytmaydi"},
    {"layout": "table", "title": "Qon tomirlari", "object": {"id": "c", "prompt": "c"}, "columns": ["Tomir", "Vazifa"],
     "rows": [["Arteriya", "Yurakdan"], ["Vena", "Yurakka"], ["Kapillyar", "Almashinuv"]]},
    {"layout": "stats", "title": "Raqamlar", "object": {"id": "c"}, "text": "Matn",
     "stats": [{"value": "70", "unit": "", "label": "Urish", "text": "daqiqada"}]},
    {"layout": "timeline", "title": "Tarix", "object": {"id": "d", "prompt": "d"},
     "steps": [{"when": "1628", "text": "Garvey"}, {"when": "1967", "text": "Ko'chirib o'tkazish"}]},
    {"layout": "conclusion", "title": "Xulosa", "points": ["Bir", "Ikki", "Uch"]}]}, "Yurak", 8)
pictures = {k: images._finish(sample(os.path.join(TMP, f"{k}.png"), c))
            for k, c in (("a", (190, 40, 60)), ("b", (60, 120, 200)), ("c", (40, 160, 90)))}   # "d" chizilmagan
path = build.build(full, pictures, {1: [(0.5, 0.2), None, (0.4, 0.6)]}, author="Ali", year="2026",
                   out_path=os.path.join(TMP, "deck.pptx"))
prs = Presentation(path)
check("8 ta slayd, 16:9", len(prs.slides) == 8 and abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01)
names = [[sh.name for sh in s.shapes] for s in prs.slides]
check("har slaydda !!panel va !!title", all("!!panel" in n and "!!title" in n for n in names), names[0])
check("obyekt bor slaydlarda !!obj_<id> (chizilmagan 'd' — obyektsiz)",
      all((f"!!obj_{s['object']['id']}" in n) == (s["object"]["id"] in pictures) for n, s in zip(names, full["slides"])),
      names)
import lxml.etree as _et
def timing_of(slide):
    node = slide._element.find(qn("p:timing"))
    return _et.tostring(node).decode() if node is not None else ""
def spid(slide, name):
    return next(sh.shape_id for sh in slide.shapes if sh.name == name)
slides_list = list(prs.slides)
# 1-slayd (muqova, obyekt "a"): yangi obyekt — uchib kiradi; 2-slayd ham "a" — Morph olib o'tadi, uchmaydi.
cover_t, call_t = timing_of(slides_list[0]), timing_of(slides_list[1])
check("yangi obyekt yondan uchib, kattalashib kiradi", f'spid="{spid(slides_list[0], "!!obj_a")}"' in cover_t
      and 'presetID="2"' in cover_t and "#ppt_w*0.75" in cover_t, cover_t[:400])
check("o'sha obyekt keyingi slaydda uchmaydi (Morph siljitadi), faqat suzib turadi",
      'presetID="2"' not in call_t and 'presetClass="path"' in call_t and 'repeatCount="indefinite"' in call_t)
check("hamma obyektli slaydda suzish harakati", all('presetClass="path"' in timing_of(s) for s, d in
                                                    zip(slides_list, full["slides"]) if d["object"]["id"] in pictures))
check("bir slaydda !! nom takrorlanmaydi", all(len([x for x in n if x.startswith("!!")]) ==
                                              len({x for x in n if x.startswith("!!")}) for n in names))
morphs = []
for s in prs.slides:
    xml = s._element
    alt = xml.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
    morphs.append(alt is not None and b"morph" in __import__("lxml.etree", fromlist=["etree"]).tostring(alt))
check("Morph: 1-slayddan tashqari hammasida", morphs == [False] + [True] * 7, morphs)
fallback_ok = all(b"<p:fade/>" in __import__("lxml.etree", fromlist=["etree"]).tostring(s._element) for s in list(prs.slides)[1:])
check("Morph yo'q PowerPoint uchun Fade zaxirasi", fallback_ok)
order_ok = True
for s in prs.slides:
    tags = [__import__("lxml.etree", fromlist=["etree"]).QName(c).localname for c in s._element]
    if "timing" in tags and "AlternateContent" in tags:
        order_ok &= tags.index("AlternateContent") < tags.index("timing")
check("XML tartibi: o'tish animatsiyadan oldin", order_ok)
callout = list(prs.slides)[1]
check("yorliqlar: 3 ta chiziq, 3 ta nuqta, animatsiya bor",
      sum(1 for n in names[1] if n.startswith("call_line")) == 3 and sum(1 for n in names[1] if n.startswith("call_dot")) == 3
      and callout._element.find(qn("p:timing")) is not None, names[1])
inside = all(0 <= sh.left and sh.left + sh.width <= prs.slide_width * 1.001 and 0 <= sh.top
             and sh.top + sh.height <= prs.slide_height * 1.001
             for sh in callout.shapes if sh.name.startswith(("call_pill", "call_text")))
check("yorliqlar varaqdan chiqmaydi", inside)
table = [sh for sh in list(prs.slides)[4].shapes if sh.has_table]
check("jadval: 4 qator (sarlavha + 3)", table and len(table[0].table.rows) == 4)
texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
check("muallif va yil muqovada", "Ali" in texts and "2026" in texts)
check("sahifa raqami", "02 / 08" in texts and "08 / 08" in texts)

print("\n5) To'liq oqim (soxta AI va rasm xizmati)")
class FakeService:
    def __init__(self, fail=False):
        self.fail, self.prompts = fail, []
    async def generate_with_models(self, prompt, models, stem="", clean=True):
        self.prompts.append((prompt, models))
        if self.fail:
            return None
        return sample(os.path.join(TMP, f"raw_{len(self.prompts)}.png"))

import services.together_service as ts
orig_get = ts.get_together_service
for fail in (False, True):
    service = FakeService(fail)
    ts.get_together_service = lambda: service
    llm_client._call_openrouter = lambda s, u, **k: raw
    llm_client._request = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("vision yo'q"))
    stages = []
    try:
        result = asyncio.run(pro3d.build_deck("Yurak", 6, language="uz", author="Ali",
                                              stage_cb=lambda n, i: stages.append(n)))
        error = None
    except Exception as exc:
        result, error = None, exc
    finally:
        ts.get_together_service = orig_get
        llm_client._call_openrouter = orig
        llm_client._request = orig_req
    if not fail:
        check("taqdimot tayyor: yo'l, 6 slayd, rasmlar", result and os.path.exists(result[0]) and result[1] == 6
              and result[2] == len(content.objects(content.normalize(raw, "Yurak", 6))), (result, error))
        check("bosqichlar: reja → rasmlar → yig'ish", stages[0] == "plan" and "images" in stages and stages[-1] == "render",
              stages)
        prompt, models = service.prompts[0]
        check("rasm so'rovi: yashil fon, bitta obyekt, ta'limiy, Together modellari",
              "chroma-key green" in prompt and "single isolated object" in prompt.lower() and "non-graphic" in prompt
              and models and "FLUX.2-pro" in models[0] and not any("imagen" in m for m in models), (prompt, models))
    else:
        check("rasm chizilmasa — xato (mijozga pul qaytariladi)", isinstance(error, RuntimeError)
              and "3D" in str(error), error)

print("\n6) Together: moderatsiya rad etsa qayta urilmaydi")
from services.together_service import TogetherImageService
svc = TogetherImageService.__new__(TogetherImageService)
calls = []
class Refusing:
    class images:
        @staticmethod
        def generate(**kw):
            calls.append(kw["model"])
            raise RuntimeError("Error code: 400 - {'message': 'Invalid content detected. The generated content was "
                               "flagged and rejected by Black Forest Labs's content moderation system.', "
                               "'code': 'content_policy_violation'}")
svc.client = Refusing
svc.image_semaphore = asyncio.Semaphore(2)          # __init__ chaqirilmagan: xizmat holati qo'lda
svc._cooldown_until = 0.0
import time as _time
t0 = _time.monotonic()
result = asyncio.run(svc._call_model("a skull", "black-forest-labs/FLUX.2-pro"))
check("moderatsiya: bitta urinish, kutishsiz, keyingi modelga o'tiladi",
      result is None and calls == ["black-forest-labs/FLUX.2-pro"] and _time.monotonic() - t0 < 1, (calls, result))
check("moderatsiya modelni 'ishlamaydi' deb belgilamaydi", "black-forest-labs/FLUX.2-pro" not in ts._unavailable)

print("\n7) Narx")
check("10 slayd — 15 000 so'm (zamonaviydan 3 baravar)", pro3d.price_for(10) == 15000, pro3d.price_for(10))
check("narx 500 ga yaxlit", all(pro3d.price_for(n) % 500 == 0 for n in pro3d.SLIDE_OPTIONS))

print("\n8) Obyekt atrofida effekt yo'q (halqa-izlar qolmasin)")
from services.pro3d import images as _img, content as _content
got = _img.clean_object("A stylized, glowing 3D rendering of a human heart, with subtle energy lines emanating, translucent")
check("«glowing» va «energy lines» olib tashlanadi, obyekt qoladi",
      "glow" not in got.lower() and "energy" not in got.lower() and "human heart" in got, got)
check("oddiy tavsifga tegilmaydi", _img.clean_object("a vintage microscope with brass details") == "a vintage microscope with brass details")
check("rasm so'rovida «no glow, no halo»", "no glow, no halo" in _img.prompt_for("heart", "realistic"))
check("rejada effekt va kelajak yillari taqiqlangan", "no glow" in _content.PROMPT.lower() and "NEVER future years" in _content.PROMPT)

print("\nNATIJA:", "HAMMASI O'TDI" if not FAILS else f"{len(FAILS)} ta xato: {FAILS}")
sys.exit(1 if FAILS else 0)
