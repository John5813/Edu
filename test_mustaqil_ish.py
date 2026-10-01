"""Mustaqil ish: qo'shimchasiz — oddiy matn; qo'shimchalar mavzuga qarab rejalashtiriladi va BEPUL.

    python test_mustaqil_ish.py
"""
import asyncio, os, sys, tempfile
from unittest.mock import AsyncMock, patch

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

from docx import Document
from services import document_service as ds

service = ds.DocumentService.__new__(ds.DocumentService)
service.documents_dir = tempfile.mkdtemp()
service.temp_dir = tempfile.mkdtemp()
service._last_used_icons = set()
service.together = AsyncMock()                 # chaqirilsa — test yiqiladi

content = {
    "language": "uz", "author_name": "Ali",
    "sections": [{"title": f"Bo'lim {i}", "content": f"Bu {i}-bo'lim matni. " * 40} for i in range(7)],
    "references": ["Ahmad S. Ufq. T., 1998."],
    "table_data_3": {"headers": ["a", "b"], "rows": [["1", "2"]] * 5},   # eski kontent shu kalitni bergan bo'lishi mumkin
}

async def build(extras, planned=None):
    calls = {"extras": [], "visuals": []}
    async def fake_extras(self, doc, title, topic, lang, ex, **kw):
        calls["extras"].append((title, list(ex)))
    async def fake_visual(self, doc, item, lang):
        calls["visuals"].append(item["kind"])
        doc.add_paragraph(f"[{item['kind']}] {item.get('title','')}")
    fake_ai = AsyncMock()
    fake_ai.plan_document_visuals = AsyncMock(return_value=planned or [])
    with patch.object(ds.DocumentService, "_add_section_extras", fake_extras), \
         patch.object(ds.DocumentService, "_add_planned_visual", fake_visual), \
         patch("services.ai_service.get_ai_service", return_value=fake_ai), \
         patch.object(ds.DocumentService, "_create_independent_work_title_page", AsyncMock()):
        path = await service.create_independent_work("Mavzu", dict(content), extras=extras)
    return Document(path), calls, fake_ai

doc, calls, ai = asyncio.run(build(None))
check("qo'shimchasiz: jadval yo'q", len(doc.tables) == 0 and not calls["visuals"] and not calls["extras"], (len(doc.tables), calls))
check("qo'shimchasiz: rasm ham chaqirilmaydi (together ishlatilmadi)", not service.together.method_calls and not service.together.mock_calls)
check("qo'shimchasiz: vizual rejasi so'ralmaydi", not ai.plan_document_visuals.await_count)
check("qo'shimchasiz: matn to'liq (6 bo'lim + xulosa)", sum("-bo'lim matni" in p.text for p in doc.paragraphs) >= 6)

plan = [{"subsection": "2", "kind": "chart", "title": "Diagramma", "explanation": "x"},
        {"subsection": "4", "kind": "table", "title": "Jadval", "explanation": "y"},
        {"subsection": "4", "kind": "formula", "title": "F", "explanation": "z"},
        {"subsection": "1", "kind": "formula", "title": "Formula", "explanation": "z"}]
doc, calls, ai = asyncio.run(build(["tables", "statistics"], plan))
check("jadval+statistika tanlansa: vizual reja so'raladi", ai.plan_document_visuals.await_count == 1)
kw = ai.plan_document_visuals.await_args.kwargs
check("diagramma va jadval soni so'raldi, formula 0", kw["charts"] >= 1 and kw["tables"] >= 1 and kw["formulas"] == 0, kw)
check("rejadagi diagramma va jadval qo'yildi, formula (tanlanmagan) qo'yilmadi", sorted(calls["visuals"]) == ["chart", "table"], calls["visuals"])
check("rejalashtirilgan extra'lar qat'iy sikl bo'yicha qayta qo'yilmadi", all("tables" not in e and "statistics" not in e for _, e in calls["extras"]), calls["extras"])

doc, calls, ai = asyncio.run(build(["images", "scheme"], plan))
check("faqat rasm/sxema tanlansa vizual reja so'ralmaydi, eski yo'l ishlaydi", not ai.plan_document_visuals.await_count and calls["extras"], calls)

doc, calls, ai = asyncio.run(build(["tables"], []))
check("reja bo'sh kelsa — eski sikl bo'yicha jadval qoladi", any("tables" in e for _, e in calls["extras"]), calls["extras"])

# ── narx: mustaqil ishda qo'shimchalar bepul
from bot import keyboards
from bot.handlers import documents as docs
from config import EXTRAS_PRICES
free_kb = keyboards.get_extras_keyboard("uz", ["images", "tables"], 5000, free=True)
paid_kb = keyboards.get_extras_keyboard("uz", ["images", "tables"], 5000)
texts = lambda kb: [b.text for row in kb.inline_keyboard for b in row]
check("bepul panel: har qo'shimchada 'bepul', jami o'zgarmaydi", all("bepul" in t for t in texts(free_kb)[:-1]) and "5,000" in texts(free_kb)[-1], texts(free_kb))
check("pullik panel (referat) avvalgidek", "+2,000" in " ".join(texts(paid_kb)) and f"{5000 + EXTRAS_PRICES['images'] + EXTRAS_PRICES['tables']:,}" in texts(paid_kb)[-1], texts(paid_kb))
check("faqat mustaqil ish bepul", docs._extras_are_free("independent_work") and not docs._extras_are_free("referat") and not docs._extras_are_free("mahsus_ishlanma"))

# tasdiqlashda yakuniy narx: mustaqil ishda asosiy narx, boshqasida qo'shimchalar bilan
from types import SimpleNamespace
from unittest.mock import MagicMock
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage

async def confirm(free):
    state = FSMContext(MemoryStorage(), StorageKey(bot_id=1, chat_id=1, user_id=1))
    await state.update_data(selected_extras=["images", "tables", "scheme"], base_price=5000, extras_free=free)
    cb = MagicMock(); cb.answer = AsyncMock(); cb.message = MagicMock()
    cb.message.edit_reply_markup = AsyncMock(); cb.message.answer = AsyncMock()
    user = SimpleNamespace(balance=100000)
    await docs.handle_extras_confirm(cb, state, None, "uz", user)
    return (await state.get_data())["price"]

check("mustaqil ish: 3 ta qo'shimcha tanlansa ham narx 5 000 (bepul)", asyncio.run(confirm(True)) == 5000)
check("referat: qo'shimchalar narxga qo'shiladi", asyncio.run(confirm(False)) == 5000 + EXTRAS_PRICES["images"] + EXTRAS_PRICES["tables"] + EXTRAS_PRICES["scheme"])

print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
