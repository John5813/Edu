"""Saytda taqdimotni yuklab olmasdan varaqlab ko'rish va bitta sahifani AI ga qayta yozdirish (900 so'm).

Haqiqiy baza (vaqtinchalik) va haqiqiy brauzer (sahifalar chiziladi, PPTX yig'iladi); faqat AI soxta.

    python test_sayt_dekk.py
"""
import asyncio, json, os, sys, tempfile, time, types

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import config
import database.database as dbmod
from database.database import Database, init_db
from database import web_store
from services import web_jobs, web_kinds, web_decks  # noqa
from services.premium_presentation import html_render, html_slides, llm_client, slide_edit, themes
import webapp
from aiohttp.test_utils import TestClient, TestServer

TMP = tempfile.mkdtemp()
dbmod.DATABASE_FILE = os.path.join(TMP, "test.db")
web_jobs.RESULTS_DIR = os.path.join(TMP, "results")


class FakeBot:
    def __init__(self): self.docs = []
    async def send_document(self, chat_id, document, **kw): self.docs.append(chat_id); return types.SimpleNamespace()
    async def send_message(self, *a, **kw): return types.SimpleNamespace(message_id=1)
BOT = FakeBot()
webapp.BOT, webapp.BOT_USERNAME = BOT, "Edufayl_bot"


def slide(title, text, dark=False):
    if dark:
        return f'<section class="slide dark"><div class="body"><h1 class="title big">{title}</h1><p class="lead">{text}</p></div></section>'
    return (f'<section class="slide"><div class="head"><h2 class="title">{title}</h2><div class="rule"></div></div>'
            f'<div class="body"><p class="lead">{text}</p><div class="list"><div class="item"><span class="item-dot"></span>'
            f'<div class="item-text">{text} birinchi band</div></div><div class="item"><span class="item-dot"></span>'
            f'<div class="item-text">{text} ikkinchi band</div></div></div></div></section>')

TITLES = ["Raqamli iqtisodiyot", "Taqdimot rejasi", "Kirish va asosiy tushunchalar", "Rivojlanish bosqichlari",
          "Hozirgi holat va natijalar", "Xulosa va takliflar"]
BODIES = [slide(TITLES[0], "Mavzu haqida", dark=True), slide(TITLES[1], "Reja")] + [slide(t, "Matn " + t) for t in TITLES[2:]]
OUTLINE = [{"title": t, "brief": f"{t} haqida bir gap", "category": "kartalar"} for t in TITLES]


def fake_write_slides(topic, count, theme, language="uz", level=2, preferences="", source_text="", author="",
                      progress_cb=None, outline_out=None):
    if outline_out is not None:
        outline_out.update(family="umumiy", outline=OUTLINE)
    return html_slides.build_pages(BODIES, theme, language)
html_slides.write_slides = fake_write_slides

STATE = {"writer": "chart", "planner": "chart", "calls": []}

def fake_openrouter(system, user, temperature=0.7, max_tokens=1000, **kw):
    STATE["calls"].append(("json", system[:30]))
    if "art director" in system:
        if STATE["planner"] == "broken":
            raise RuntimeError("model javob bermadi")
        return {"category": "diagramma", "title": "Raqamli xizmatlar ulushi", "brief": "Xizmatlar ulushi", "chart_kind": "halqa"}
    return {"ok": True, "kind": "donut", "labels": ["Onlayn", "Offlayn", "Aralash"],
            "series": [{"name": "", "values": [55, 30, 15]}], "unit": "%", "xlabel": "",
            "source": "Statistika agentligi, 2024", "approx": False, "forecast": False, "reason": ""}
llm_client._call_openrouter = fake_openrouter

def fake_text(system, user, temperature=0.7, max_tokens=1000, accept=None, **kw):
    STATE["calls"].append(("text", user))
    if STATE["writer"] == "empty":
        return ""
    if STATE["writer"] == "plain":      # diagramma so'ralgan, lekin model uni yozmagan
        return slide("Raqamli xizmatlar ulushi", "Oddiy matn")
    return slide("Raqamli xizmatlar ulushi", "Xizmatlar ulushi o'smoqda")
llm_client._call_openrouter_text = fake_text


async def wait_job(client, job_id, limit=240.0):
    end = time.time() + limit
    data = {}
    while time.time() < end:
        data = await (await client.get(f"/api/v1/jobs/{job_id}")).json()
        if data["job"]["status"] in ("done", "failed"):
            return data
        await asyncio.sleep(0.3)
    return data


async def login(client, origin, telegram_id):
    start = await (await client.post("/api/v1/auth/start", headers=origin)).json()
    await web_store.confirm_login(start["token"], telegram_id)
    await client.get("/api/v1/auth/poll", params={"token": start["token"]})


async def main():
    if not html_render.available():
        check("brauzer o'rnatilgan", False, "Chromium kerak")
        return
    from webapp.server import create_web_app
    await init_db()
    await Database.create_user(100, "ali", "Ali", "uz")
    await Database.create_user(200, "vali", "Vali", "uz")
    await Database.update_user_balance(100, 20_000)
    server = TestServer(create_web_app())
    client = TestClient(server)
    await client.start_server()
    origin = {"Origin": f"http://{server.host}:{server.port}"}
    await login(client, origin, 100)

    print("1) Taqdimot yaratiladi va varaqlash uchun saqlanadi")
    r = await client.post("/api/v1/jobs", json={"kind": "premium_presentation", "params": {
        "topic": "Raqamli iqtisodiyot", "slide_count": 4, "author": "Ali Valiyev"}}, headers=origin)
    job_id = (await r.json())["job"]["id"]
    job = await wait_job(client, job_id)
    check("taqdimot tayyor", job["job"]["status"] == "done", job)
    info = await (await client.get(f"/api/v1/jobs/{job_id}/deck")).json()
    check("varaqlash mavjud: 6 sahifa, hammasida surat", info["available"] and info["deck"]["count"] == 6
          and all(s["image"] for s in info["deck"]["slides"]), info)
    check("sarlavhalar to'g'ri", [s["title"] for s in info["deck"]["slides"]] == TITLES,
          [s["title"] for s in info["deck"]["slides"]])
    check("narx 900 so'm", info["price"] == 900 and info["busy"] is None, info)
    r = await client.get(f"/api/v1/jobs/{job_id}/slide/3")
    img = await r.read()
    check("sahifa surati JPEG, 1280x720", r.status == 200 and img[:2] == b"\xff\xd8", r.status)
    from PIL import Image
    import io
    check("o'lcham 16:9", Image.open(io.BytesIO(img)).size == (1280, 720), Image.open(io.BytesIO(img)).size)
    r = await client.get(f"/api/v1/jobs/{job_id}/slide/99")
    check("bo'lmagan sahifa — 404", r.status == 404, r.status)
    file_before = open((await web_store.get_job(job_id))["result_path"], "rb").read()
    docs_before = len(BOT.docs)

    print("2) Boshqa odam ko'ra olmaydi")
    other = TestClient(TestServer(create_web_app())); await other.start_server()
    oorigin = {"Origin": f"http://{other.server.host}:{other.server.port}"}
    await login(other, oorigin, 200)
    r = await other.get(f"/api/v1/jobs/{job_id}/deck")
    check("begona taqdimot topilmaydi", r.status == 404, r.status)
    r = await other.get(f"/api/v1/jobs/{job_id}/slide/1")
    check("begona surat ham berilmaydi", r.status == 404, r.status)
    r = await other.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 3, "instruction": "diagramma"}, headers=oorigin)
    check("begona taqdimotni qayta yozdirib bo'lmaydi", r.status == 404, r.status)
    await other.close()

    print("3) Sahifani qayta yozdirish: tekshiruvlar")
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 2, "instruction": "boshqacha qil"}, headers=origin)
    check("reja sahifasi (2) o'zgartirilmaydi", r.status == 400 and (await r.json())["code"] == "plan_slide", r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 99, "instruction": "diagramma qil"}, headers=origin)
    check("bo'lmagan sahifa rad etildi", r.status == 400, r.status)
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 3, "instruction": "  "}, headers=origin)
    check("iltimossiz so'rov rad etildi", r.status == 400, r.status)
    balance0 = (await Database.get_user(100)).balance

    print("4) Sahifani doirasimon diagrammali qilib berish")
    STATE["calls"].clear()
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite",
                          json={"index": 4, "instruction": "Sahifani doirasimon diagrammali qilib ber"}, headers=origin)
    body = await r.json()
    check("buyurtma qabul qilindi, 900 so'm yechildi", r.status == 200 and body["balance"] == balance0 - 900, (r.status, body))
    rid = body["job"]["id"]
    r2 = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 5, "instruction": "rasm qo'y"}, headers=origin)
    check("bir vaqtda ikkinchi o'zgartirish rad etildi", r2.status == 409, r2.status)
    done = await wait_job(client, rid)
    check("qayta yozish tugadi", done["job"]["status"] == "done", done)
    info2 = await (await client.get(f"/api/v1/jobs/{job_id}/deck")).json()
    deck = info2["deck"]
    check("versiya oshdi, tarixga yozildi", deck["version"] == 2 and deck["history"] and deck["history"][-1]["index"] == 4
          and "doirasimon" in deck["history"][-1]["instruction"], deck)
    check("4-sahifa sarlavhasi yangilandi", deck["slides"][3]["title"] == "Raqamli xizmatlar ulushi", deck["slides"][3])
    check("boshqa sahifalar o'z joyida", [s["title"] for i, s in enumerate(deck["slides"]) if i not in (1, 3)]
          == [t for i, t in enumerate(TITLES) if i not in (1, 3)])
    stored = web_decks.load(job_id)
    check("sahifada haqiqiy ma'lumotli diagramma (halqa)", 'data-kind="donut"' in html_slides.source_of(stored["pages"][3])
          and "<svg" in stored["pages"][3], html_slides.source_of(stored["pages"][3])[:300])
    check("diagramma manbasi AI bergan ma'lumotdan", "Statistika agentligi" in stored["pages"][3])
    plan_titles = [t for t in ("Kirish", "Rivojlanish", "Raqamli xizmatlar", "Hozirgi") if t in stored["pages"][1]]
    check("reja sahifasi yangi sarlavhalardan qayta yig'ildi", "Raqamli xizmatlar ulushi" in stored["pages"][1]
          and "Rivojlanish bosqichlari" not in stored["pages"][1], plan_titles)
    r = await client.get(f"/api/v1/jobs/{job_id}/slide/4")
    check("4-sahifa surati yangilandi", r.status == 200 and (await r.read()) != img)
    file_after = open((await web_store.get_job(job_id))["result_path"], "rb").read()
    check("'Yuklab olish' fayli yangilandi", file_after != file_before and file_after[:2] == b"PK")
    from pptx import Presentation
    prs = Presentation((await web_store.get_job(job_id))["result_path"])
    texts = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    check("PPTX da 6 slayd, yangi sarlavha bor", len(prs.slides) == 6 and "Raqamli xizmatlar ulushi" in texts, len(prs.slides))
    jobs = await (await client.get("/api/v1/jobs")).json()
    check("qayta yozish 'Hujjatlarim' ro'yxatini to'ldirmaydi", all(j["kind"] != "slide_rewrite" for j in jobs["jobs"]))
    check("Telegramga qayta yuborilmadi", len(BOT.docs) == docs_before, (len(BOT.docs), docs_before))
    me = await (await client.get("/api/v1/me")).json()
    check("hisobdan aynan 900 so'm yechilgan", me["user"]["balance"] == balance0 - 900, me["user"]["balance"])
    check("AI ga iltimos va qo'shni sahifalar berildi", any("doirasimon" in c[1] and "Neighbours" in c[1] for c in STATE["calls"] if c[0] == "text"),
          [c for c in STATE["calls"] if c[0] == "text"][:1])

    print("5) Reja AI dan kelmasa ham iltimos kalit so'zlardan tushuniladi")
    STATE["planner"] = "broken"; STATE["writer"] = "plain"
    balance1 = (await Database.get_user(100)).balance
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 5, "instruction": "doirasimon diagramma qo'shib ber"}, headers=origin)
    done = await wait_job(client, (await r.json())["job"]["id"])
    stored = web_decks.load(job_id)
    check("model diagramma yozmagan bo'lsa, tizim uni qo'yadi", done["job"]["status"] == "done"
          and 'data-kind="donut"' in html_slides.source_of(stored["pages"][4]), done["job"])
    check("hisobdan yana 900 so'm", (await Database.get_user(100)).balance == balance1 - 900)

    print("6) Xato bo'lsa pul qaytadi, taqdimot o'zgarmaydi")
    STATE["planner"] = "chart"; STATE["writer"] = "empty"
    balance2 = (await Database.get_user(100)).balance
    version = web_decks.load(job_id)["version"]
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 5, "instruction": "ustunli diagramma bilan ko'rsat"}, headers=origin)
    done = await wait_job(client, (await r.json())["job"]["id"])
    check("AI javob bermasa ish muvaffaqiyatsiz", done["job"]["status"] == "failed" and "qaytarildi" in done["job"]["error"], done["job"])
    check("pul to'liq qaytdi", (await Database.get_user(100)).balance == balance2)
    check("taqdimot o'zgarmadi", web_decks.load(job_id)["version"] == version)

    print("7) Balans yetmasa")
    await Database.update_user_balance(100, -(await Database.get_user(100)).balance + 500)
    r = await client.post(f"/api/v1/jobs/{job_id}/rewrite", json={"index": 5, "instruction": "rasm va matn qo'y"}, headers=origin)
    check("balans yetmasa 402", r.status == 402 and (await r.json())["code"] == "no_balance", r.status)


    print("9) Oddiy (orqa fonli) taqdimot ham ochib ko'riladi, lekin qayta yozdirib bo'lmaydi")
    await Database.update_user_balance(100, 20_000)
    import services.ai_service as ais, services.document_service as dss
    from pptx import Presentation
    class FakeAI:
        async def generate_presentation_in_batches(self, topic, count, lang):
            return {"slides": [{"layout": "content", "title": "T", "content": "M"}]}
        async def generate_plan_items(self, topic, lang): return []
    class FakeDocs:
        use_icons = True
        async def create_presentation_with_template_background(self, topic, content, author, template, ts, lang, refs, plan):
            prs = Presentation()
            for title in ("Birinchi sahifa", "Ikkinchi sahifa", "Uchinchi sahifa"):
                sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text = title
            path = os.path.join(TMP, f"s{time.time_ns()}.pptx"); prs.save(path); return path
    real = (ais.get_ai_service, dss.get_document_service)
    ais.get_ai_service = lambda: FakeAI(); dss.get_document_service = lambda: FakeDocs()
    r = await client.post("/api/v1/jobs", json={"kind": "simple_presentation", "params": {
        "topic": "Amir Temur davri", "slide_count": 10, "template": "template_3", "author": "Ali"}}, headers=origin)
    sid = (await r.json())["job"]["id"]
    sdone = await wait_job(client, sid)
    check("oddiy taqdimot tayyor", sdone["job"]["status"] == "done", sdone)
    sinfo = await (await client.get(f"/api/v1/jobs/{sid}/deck")).json()
    check("3 sahifa ko'rinadi, sarlavhalari bilan", sinfo["available"] and sinfo["deck"]["count"] == 3
          and [x["title"] for x in sinfo["deck"]["slides"]] == ["Birinchi sahifa", "Ikkinchi sahifa", "Uchinchi sahifa"]
          and all(x["image"] for x in sinfo["deck"]["slides"]), sinfo)
    check("faqat ko'rish (qayta yozish yo'q)", sinfo["deck"]["editable"] is False)
    r = await client.get(f"/api/v1/jobs/{sid}/slide/2")
    check("2-sahifa surati beriladi", r.status == 200 and (await r.read())[:2] == b"\xff\xd8", r.status)
    r = await client.post(f"/api/v1/jobs/{sid}/rewrite", json={"index": 2, "instruction": "diagramma qil"}, headers=origin)
    check("oddiy taqdimotni qayta yozdirib bo'lmaydi", r.status == 400 and (await r.json())["code"] == "view_only", r.status)
    ais.get_ai_service, dss.get_document_service = real

    print("8) Tozalash")
    await web_store.update_job(job_id, finished_at=1.0)
    import sqlite3
    con = sqlite3.connect(dbmod.DATABASE_FILE); con.execute("UPDATE web_jobs SET created_at = ?", (time.time() - 100 * 3600,)); con.commit(); con.close()
    await web_jobs.purge()
    check("muddati o'tgach taqdimot nusxasi ham o'chdi", not web_decks.exists(job_id) and not os.path.isdir(os.path.join(web_decks.root(), job_id)))

    await client.close()


asyncio.run(main())
print("\n" + ("✅ hammasi o'tdi" if not FAILS else f"❌ {len(FAILS)} ta xato: {FAILS}"))
sys.exit(1 if FAILS else 0)
