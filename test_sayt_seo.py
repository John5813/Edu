"""Tayyor ishlar Google (va Google Rasmlar, Yandex) da topilishi uchun: ish sahifasi, katalog bo'limlari,
nomli rasm manzillari, sitemap, robots, reja/parcha, saytdagi qidiruv va brauzerda ishlashi.

Ishga tushirish: BOT_TOKEN=1:x python test_sayt_seo.py
"""
import asyncio
import json
import os
import re
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET

TMP = tempfile.mkdtemp(prefix="seo_")
os.environ.setdefault("BOT_TOKEN", "1:x")
os.environ["STORE_PREVIEW_DIR"] = os.path.join(TMP, "previews")
os.environ["STORE_PREVIEW_MAX"] = "4"

FAILED = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail and not cond else ""))
    if not cond:
        FAILED.append(name)


import database.database as dbmod  # noqa: E402
from database.database import Database, init_db  # noqa: E402

dbmod.DATABASE_FILE = os.path.join(TMP, "test.db")

import webapp  # noqa: E402
from aiohttp import web  # noqa: E402
from aiohttp.test_utils import TestClient, TestServer  # noqa: E402
from PIL import Image  # noqa: E402

from services import store_seo  # noqa: E402
from webapp import store  # noqa: E402

webapp.BOT_USERNAME = "Edufayl_bot"


def make_previews(code, count, size=(1000, 563)):
    folder = os.path.join(os.environ["STORE_PREVIEW_DIR"], code)
    os.makedirs(folder, exist_ok=True)
    for n in range(1, count + 1):
        Image.new("RGB", size, (30 * n % 255, 90, 160)).save(os.path.join(folder, f"{n}.jpg"))
    Image.new("RGB", (480, 270), (10, 90, 160)).save(os.path.join(folder, "thumb.jpg"))


def make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("______ fakulteti")
    doc.add_paragraph().add_run("KURS ISHI").bold = True
    doc.add_paragraph("Mavzu: Bank tizimi va uning ahamiyati")
    doc.add_paragraph("Bajardi: ____")
    doc.add_paragraph().add_run("MUNDARIJA").bold = True
    for line in ("KIRISH\t3", "I BOB. BANK TIZIMINING MOHIYATI\t5", "1.1 Tijorat banklari\t5",
                 "1.2 Markaziy bank va pul-kredit siyosati\t9", "XULOSA\t20", "FOYDALANILGAN ADABIYOTLAR\t22"):
        doc.add_paragraph(line)
    doc.add_paragraph().add_run("KIRISH").bold = True
    doc.add_paragraph("Kirish matni. " * 20)
    doc.add_paragraph().add_run("I BOB. BANK TIZIMINING MOHIYATI").bold = True
    doc.add_paragraph().add_run("1.1 Tijorat banklari").bold = True
    doc.add_paragraph("Tijorat banklari aholi va korxonalarning bo'sh pul mablag'larini jalb qilib, ularni "
                      "kredit shaklida iqtisodiyotga yo'naltiradi va shu orqali iqtisodiy o'sishga xizmat qiladi. " * 2)
    doc.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for title, body in (("Ekologiya muammolari", "Muqova"),
                        ("Reja", "1. Havo ifloslanishi 2. Suv resurslari"),
                        ("Havo ifloslanishi va uning manbalari",
                         "Sanoat korxonalari va avtotransport chiqindilari shahar havosining asosiy "
                         "ifloslantiruvchilari hisoblanadi, bu esa aholi salomatligiga bevosita ta'sir qiladi."),
                        ("Orol dengizi fojiasi", "Qisqa matn"),
                        ("E'tiboringiz uchun rahmat", "")):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1)).text_frame.text = title
        if body:
            slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3)).text_frame.text = body
    prs.save(path)


class FakeBot:
    def __init__(self, files):
        self.files = files

    async def download(self, file_id, destination):
        if file_id not in self.files:
            raise RuntimeError("file is too big")
        shutil.copy(self.files[file_id], destination)


async def main():
    await init_db()

    # ── fayldan reja va parcha
    docx_path, pptx_path = os.path.join(TMP, "kurs.docx"), os.path.join(TMP, "deck.pptx")
    make_docx(docx_path)
    make_pptx(pptx_path)
    doc = store_seo.extract(docx_path, "Bank tizimi va uning ahamiyati")
    check("docx: reja mundarijadan olinadi (umumiy bandlarsiz)",
          doc["outline"] == ["I BOB. BANK TIZIMINING MOHIYATI", "1.1 Tijorat banklari",
                             "1.2 Markaziy bank va pul-kredit siyosati"], doc["outline"])
    check("docx: parcha asosiy qismdan", doc["excerpt"].startswith("Tijorat banklari aholi"), doc["excerpt"][:60])
    deck = store_seo.extract(pptx_path, "Ekologiya muammolari")
    check("pptx: reja slayd sarlavhalaridan (muqova, Reja, Rahmat tashlanadi)",
          deck["outline"] == ["Havo ifloslanishi va uning manbalari", "Orol dengizi fojiasi"], deck["outline"])
    check("pptx: parcha uzun matnli slayddan", deck["excerpt"].startswith("Sanoat korxonalari"), deck["excerpt"][:40])
    check("buzuq fayl xato ko'tarmaydi", store_seo.extract(os.path.join(TMP, "yoq.pptx")) == {"outline": [], "excerpt": ""})
    sample = "attached_assets/Premium_Aholi_daromadlari_1790180246958.pptx"
    if os.path.exists(sample):
        real = store_seo.extract(sample, "Aholi daromadlari: Tahlil va istiqbollar")
        check("haqiqiy premium taqdimot: reja olinadi", len(real["outline"]) >= 10 and real["excerpt"], real["outline"][:3])

    # ── manzil va yozuv
    check("slug: apostrof va kirill", store_seo.item_slug("O‘zbekiston milliy boyligi", "kurs_ishi")
          == "ozbekiston-milliy-boyligi-kurs-ishi")
    check("slug: tur ikki marta yozilmaydi", store_seo.item_slug("Kurs ishi: Bank tizimi", "kurs_ishi") == "kurs-ishi-bank-tizimi")
    check("slug: kirill mavzu", store_seo.item_slug("Ўзбекистон тарихи", "taqdimot") == "ozbekiston-tarixi-taqdimot")
    check("kirill varianti", store_seo.other_script("Bank tizimi") == "Банк тизими")
    check("lotin varianti", store_seo.other_script("Ўзбекистон тарихи") == "O'zbekiston tarixi")
    check("ruscha va inglizcha mavzuga variant yo'q",
          store_seo.other_script("Финансовая система") == "" and store_seo.other_script("Climate change and society") == "")

    # ── katalog ma'lumoti
    items = []
    for n in range(30):
        code = f"TQ{n:06d}"
        await Database.create_store_item(code, f"Iqtisodiyot mavzusi {n}", f"f{n}", 9000, category="Iqtisodiyot",
                                         work_type="taqdimot", slide_count=12, preview_count=4,
                                         outline=["Bozor iqtisodiyoti"] if n == 5 else [], excerpt="")
        make_previews(code, 4)
        items.append(code)
    await Database.create_store_item("KURS0001", "Bank tizimi va uning ahamiyati", "kf", 25000,
                                     category="Iqtisodiyot", work_type="kurs_ishi", slide_count=28,
                                     file_type="docx", preview_count=3, outline=doc["outline"], excerpt=doc["excerpt"])
    make_previews("KURS0001", 3, size=(800, 1131))
    await Database.create_store_item("KURS0002", "Moliya bozori", "kf2", 25000, category="Iqtisodiyot",
                                     work_type="kurs_ishi", slide_count=25, file_type="docx", preview_count=2)
    make_previews("KURS0002", 2, size=(800, 1131))
    await Database.create_store_item("ECO00001", "Ekologiya muammolari", "eco", 9000, category="Ekologiya",
                                     work_type="taqdimot", slide_count=5, preview_count=2)
    make_previews("ECO00001", 2)
    # Eski ish: reja hali olinmagan (seo_state 0)
    await Database.create_store_item("OLD00001", "Bozor iqtisodiyoti asoslari", "old", 9000,
                                     category="Iqtisodiyot", work_type="taqdimot", slide_count=5, preview_count=1)
    make_previews("OLD00001", 1)

    # ── eski ishlarni to'ldirish (fon vazifasi)
    pending = {r["public_code"] for r in await Database.store_items_without_seo(100)}
    check("reja yo'q ishlar topiladi (yangilari emas)", pending == {"KURS0002", "ECO00001", "OLD00001"}, pending)
    bot = FakeBot({"eco": pptx_path, "kf2": docx_path})
    done = await store_seo.backfill(bot, pause=0)
    eco = await Database.get_store_item("ECO00001")
    old = await Database.get_store_item("OLD00001")
    check("to'ldirildi: hammasi bir martadan", done == 3, done)
    check("ombordan olingan faylning rejasi yozildi",
          eco["seo_state"] == 1 and "Orol dengizi fojiasi" in store_seo.outline_of(eco) and eco["excerpt"])
    check("yuklab bo'lmagan ish belgilanadi (qayta-qayta urinilmaydi)", old["seo_state"] == 2)
    check("ikkinchi marta hech narsa qilinmaydi", await store_seo.backfill(bot, pause=0) == 0)

    # ── yangi nashr: reja darhol yoziladi, Yandex/Bing'ga xabar ketadi
    from types import SimpleNamespace

    from services import store_publisher

    pinged = []

    async def fake_ping(paths):
        pinged.append(paths)
        return True

    class VaultBot:
        async def send_document(self, chat_id, document, **kw):
            return SimpleNamespace(document=SimpleNamespace(file_id="vault-1"))

    store_publisher.STORE_VAULT_CHAT_ID = "-100"
    store_publisher.render_previews = lambda path, code: (make_previews(code, 2) or (2, 5))
    original_ping, store_seo.ping = store_seo.ping, fake_ping
    try:
        result = await store_publisher.publish_work(VaultBot(), pptx_path, "Ekologiya va inson", 9000,
                                                    work_type="taqdimot", category="Ekologiya")
        await asyncio.sleep(0.05)
    finally:
        store_seo.ping = original_ping
    new = await Database.get_store_item(result["public_code"])
    check("yangi nashr: reja va parcha darhol yoziladi",
          new["seo_state"] == 1 and "Orol dengizi fojiasi" in store_seo.outline_of(new) and new["excerpt"])
    check("yangi nashr: sahifa Yandex/Bing'ga bildiriladi",
          pinged and pinged[0][0] == f"/shop/{result['public_code']}/ekologiya-va-inson-taqdimot", pinged)

    app = web.Application()
    store.setup_store_routes(app)
    client = TestClient(TestServer(app))
    await client.start_server()
    try:
        await http_checks(client, items)
        await browser_checks(client)
    finally:
        await client.close()


async def http_checks(client, items):
    kurs = await Database.get_store_item("KURS0001")
    path = store._item_path(kurs)
    check("ish manzili mavzu nomi bilan", path == "/shop/KURS0001/bank-tizimi-va-uning-ahamiyati-kurs-ishi", path)

    r = await client.get("/shop/KURS0001", allow_redirects=False)
    check("eski qisqa manzil → 301 nomli manzilga", r.status == 301 and r.headers["Location"] == path)
    r = await client.get("/shop/KURS0001/notogri-nom", allow_redirects=False)
    check("noto'g'ri nom → 301 to'g'risiga", r.status == 301 and r.headers["Location"] == path)
    r = await client.get("/shop/XXXXXXXX/abc", allow_redirects=False)
    check("yo'q ish → 404", r.status == 404)

    r = await client.get(path)
    page = await r.text()
    check("ish sahifasi ochiladi", r.status == 200)
    title = re.search(r"<title>(.*?)</title>", page).group(1)
    check("sarlavhada mavzu, ish turi, hajm va «yuklab olish»",
          title == "Bank tizimi va uning ahamiyati — tayyor kurs ishi (28 varaq) yuklab olish", title)
    check("canonical — nomli manzil", f'rel="canonical" href="http://' in page and path in page)
    check("h1 — mavzu", "<h1>Bank tizimi va uning ahamiyati</h1>" in page)
    check("reja sahifada", "<li>1.1 Tijorat banklari</li>" in page and "<h2>Reja</h2>" in page)
    check("parcha sahifada", "Ishdan parcha" in page and "Tijorat banklari aholi" in page)
    check("ish haqida: tur nomlari, format, hajm",
          "kurs loyihasi" in page and "28 ta varaqdan iborat, DOCX formatida" in page and "Word" in page)
    check("mavzuning kirill yozuvi", "Банк тизими ва унинг аҳамияти" in page)
    desc = re.search(r'<meta name="description" content="([^"]*)"', page).group(1)
    check("meta description: tur, hajm, reja", "tayyor kurs ishi: 28 varaq, DOCX" in desc and "Reja:" in desc, desc)
    check("rasm nomli manzil va mazmunli alt",
          'src="/shop/img/KURS0001/bank-tizimi-va-uning-ahamiyati-kurs-ishi-1.jpg"' in page
          and 'alt="Bank tizimi va uning ahamiyati — kurs ishi, 1-varaq"' in page)
    check("rasm o'lchami yozilgan (sahifa sakramaydi)", 'width="800" height="1131"' in page)
    check("birinchi rasm darhol, qolganlari kechiktirib",
          page.count('fetchpriority="high"') == 1 and page.count('loading="lazy"') >= 2)
    check("og:image nomli rasm", 'property="og:image" content="http://' in page and "-kurs-ishi-1.jpg" in page)
    check("yo'l: tur va fan bo'limlariga havola",
          'href="/shop/tur/kurs-ishi"' in page and 'href="/shop/tur/kurs-ishi/iqtisodiyot"' in page)
    check("o'xshash ishlar (boshqa ishlarga havola)", "O'xshash ishlar" in page and "/shop/KURS0002/moliya-bozori-kurs-ishi" in page)
    ld = json.loads(re.search(r'<script type="application/ld\+json">(.*?)</script>', page, re.S).group(1))
    kinds = [node["@type"] for node in ld["@graph"]]
    product = ld["@graph"][0]
    check("JSON-LD: Product + BreadcrumbList + rasmlar", kinds[:2] == ["Product", "BreadcrumbList"]
          and kinds.count("ImageObject") == 3 and product["offers"]["priceCurrency"] == "UZS", kinds)
    check("JSON-LD: rasmlar nomli manzilda", all("/shop/img/KURS0001/" in u for u in product["image"]))
    crumbs = [i["name"] for i in ld["@graph"][1]["itemListElement"]]
    check("JSON-LD: yo'l", crumbs == ["Katalog", "Tayyor kurs ishlari", "Iqtisodiyot", "Bank tizimi va uning ahamiyati"], crumbs)

    # ── rasmlar
    r = await client.get("/shop/img/KURS0001/bank-tizimi-va-uning-ahamiyati-kurs-ishi-2.jpg")
    check("nomli rasm ochiladi", r.status == 200 and r.headers["Content-Type"] == "image/jpeg")
    r = await client.get("/shop/img/KURS0001/bank-tizimi-va-uning-ahamiyati-kurs-ishi-kichik.jpg")
    check("«-kichik» — katalogdagi kichik rasm", r.status == 200)
    check("raqamsiz/0-rasm → 404", (await client.get("/shop/img/KURS0001/bank-tizimi.jpg")).status == 404
          and (await client.get("/shop/img/KURS0001/bank-0.jpg")).status == 404)
    check("chegaradan tashqari raqam → 404", (await client.get("/shop/img/KURS0001/x-9.jpg")).status == 404)
    check("noto'g'ri kod → 404", (await client.get("/shop/img/..%2F..%2Fetc/x-1.jpg")).status == 404)
    check("eski rasm manzili ham ishlaydi", (await client.get("/shop/preview/KURS0001/1.jpg")).status == 200)

    # ── katalog
    r = await client.get("/shop")
    page = await r.text()
    cards = re.findall(r'<a class="card" href="(/[^"]+)"', page)
    check("katalog serverda chiziladi (JavaScriptsiz havolalar)", r.status == 200 and len(cards) == 24, len(cards))
    check("kartochka nomli manzilga", all(re.match(r"^/shop/[A-Z0-9]{8}/[a-z0-9-]+$", c) for c in cards))
    check("sahifalash havolalari", 'href="/shop?page=2"' in page and 'rel="next"' in page)
    check("tur va fan havolalari", 'href="/shop/tur/taqdimot"' in page and 'href="/shop/fan/iqtisodiyot"' in page)
    check("katalog h1 va matn", "<h1>Tayyor mavzular</h1>" in page and "prezentatsiya" in page)
    check("katalog indekslanadi", 'content="index,follow' in page)
    ld = json.loads(re.search(r'<script type="application/ld\+json">(.*?)</script>', page, re.S).group(1))
    check("katalog JSON-LD: ItemList", ld["@graph"][0]["mainEntity"]["numberOfItems"] == 35
          and len(ld["@graph"][0]["mainEntity"]["itemListElement"]) == 24)
    r = await client.get("/shop?page=2")
    page2 = await r.text()
    cards2 = re.findall(r'<a class="card" href="(/[^"]+)"', page2)
    check("2-sahifa: qolgan ishlar, o'z canonical'i", len(cards2) == 11 and "/shop?page=2" in page2
          and "2-sahifa" in page2 and not set(cards) & set(cards2))
    check("yo'q sahifa → 404", (await client.get("/shop?page=99")).status == 404)

    r = await client.get("/shop?type=kurs_ishi", allow_redirects=False)
    check("eski ?type= havola → 301 tur sahifasiga", r.status == 301 and r.headers["Location"] == "/shop/tur/kurs-ishi")
    r = await client.get("/shop?type=kurs_ishi&category=Iqtisodiyot", allow_redirects=False)
    check("eski tur+fan havola → 301", r.status == 301 and r.headers["Location"] == "/shop/tur/kurs-ishi/iqtisodiyot")
    r = await client.get("/shop/tur/kurs-ishi")
    page = await r.text()
    cards = re.findall(r'<a class="card" href="(/[^"]+)"', page)
    title = re.search(r"<title>(.*?)</title>", page).group(1)
    check("tur sahifasi: faqat kurs ishlari", r.status == 200 and len(cards) == 2 and all("kurs-ishi" in c for c in cards))
    check("tur sahifasi sarlavhasi", title == "Tayyor kurs ishlari — 2 ta mavzu, yuklab olish" and
          "<h1>Tayyor kurs ishlari</h1>" in page, title)
    check("tur sahifasida fanlar shu tur ichida", 'href="/shop/tur/kurs-ishi/iqtisodiyot"' in page)
    r = await client.get("/shop/tur/kurs-ishi/iqtisodiyot")
    page = await r.text()
    check("tur + fan sahifasi", r.status == 200 and "<h1>Tayyor kurs ishlari: Iqtisodiyot</h1>" in page
          and "Iqtisodiyot bo&#x27;yicha tayyor kurs ishlari" in page)
    r = await client.get("/shop/fan/ekologiya")
    page = await r.text()
    check("fan sahifasi", r.status == 200 and "<h1>Ekologiya: tayyor ishlar</h1>" in page
          and len(re.findall(r'<a class="card" href="/', page)) == 2)
    check("yo'q tur/fan → 404", (await client.get("/shop/tur/yoq")).status == 404
          and (await client.get("/shop/fan/yoq")).status == 404
          and (await client.get("/shop/tur/referat")).status == 404)
    r = await client.get("/shop?q=bank")
    page = await r.text()
    check("qidiruv natijasi indekslanmaydi, lekin ishlaydi",
          'content="noindex,follow' in page and "/shop/KURS0001/" in page and 'value="bank"' in page)

    # ── API va saytdagi qidiruv
    data = await (await client.get("/api/shop/items?q=bozor")).json()
    codes = [it["code"] for it in data["items"]]
    check("qidiruv rejadagi so'zni ham topadi", "TQ000005" in codes, codes)
    check("sarlavhada mos kelgan ish yuqorida", codes and codes[0] == "OLD00001", codes)
    check("API: nomli manzil va rasm", data["items"][0]["url"].startswith("/shop/OLD00001/")
          and data["items"][0]["preview"] == "/shop/img/OLD00001/bozor-iqtisodiyoti-asoslari-taqdimot-kichik.jpg")
    check("raqam bilan tugaydigan turisiz mavzu: kichik rasm adashmaydi",
          store._image_path({"public_code": "AAAA0000", "title": "Matematika 9"}) == "/shop/img/AAAA0000/matematika-9-kichik.jpg"
          and store._image_path({"public_code": "AAAA0000", "title": "Matematika 9"}, 2) == "/shop/img/AAAA0000/matematika-9-2.jpg")
    types = (await (await client.get("/api/shop/types")).json())["types"]
    check("API turlar: manzil bilan", any(t["url"] == "/shop/tur/kurs-ishi" for t in types))
    item = await (await client.get("/api/shop/items/KURS0001")).json()
    check("API ish: nomli rasmlar", item["previews"][0].endswith("-kurs-ishi-1.jpg"))

    # ── sitemap, robots, IndexNow
    r = await client.get("/sitemap.xml")
    body = await r.text()
    root = ET.fromstring(body)
    ns = {"s": "http://www.sitemaps.org/schemas/sitemap/0.9", "i": "http://www.google.com/schemas/sitemap-image/1.1"}
    locs = [u.find("s:loc", ns).text for u in root.findall("s:url", ns)]
    check("sitemap: to'g'ri XML", r.status == 200 and len(locs) == len(set(locs)))
    check("sitemap: hamma ish nomli manzil bilan", sum(1 for l in locs if re.search(r"/shop/[A-Z0-9]{8}/", l)) == 35)
    check("sitemap: tur, fan va tur+fan bo'limlari",
          any(l.endswith("/shop/tur/kurs-ishi") for l in locs) and any(l.endswith("/shop/fan/ekologiya") for l in locs)
          and any(l.endswith("/shop/tur/taqdimot/iqtisodiyot") for l in locs)
          and any(l.endswith("/shop/tur/taqdimot/ekologiya") for l in locs)
          and not any(l.endswith("/shop/tur/kurs-ishi/ekologiya") for l in locs))
    kurs_url = next(u for u in root.findall("s:url", ns) if "/shop/KURS0001/" in u.find("s:loc", ns).text)
    images = [i.find("i:loc", ns).text for i in kurs_url.findall("i:image", ns)]
    check("sitemap: Google Rasmlar uchun rasmlar", len(images) == 3 and images[0].endswith("-kurs-ishi-1.jpg"), images)
    robots = await (await client.get("/robots.txt")).text()
    check("robots: sitemap, shaxsiy sahifalar yopiq", "Sitemap: http" in robots and "Disallow: /app" in robots
          and "Disallow: /api/" in robots and "Allow: /shop" in robots and "Allow: /api/shop/" in robots)
    key = store_seo.indexnow_key()
    r = await client.get(f"/{key}.txt")
    check("IndexNow kalit fayli", r.status == 200 and (await r.text()) == key)
    check("IndexNow localhost'da yuborilmaydi", await store_seo.ping(["/shop"]) is False)


async def browser_checks(client):
    from playwright.async_api import async_playwright

    from services.premium_presentation import html_render

    if not html_render._executable():
        check("brauzer bor", False, "Chromium topilmadi")
        return
    base = str(client.make_url("/")).rstrip("/")
    async with async_playwright() as p:
        browser = await p.chromium.launch(executable_path=html_render._executable(), args=html_render._LAUNCH_ARGS)
        page = await browser.new_page(viewport={"width": 390, "height": 844})
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        await page.goto(base + "/shop", wait_until="networkidle")
        check("brauzer: 24 ta kartochka", await page.locator(".card").count() == 24)
        await page.click("#more")
        await page.wait_for_function("document.querySelectorAll('.card').length > 24")
        check("brauzer: «Yana ko'rsatish» qo'shadi, sahifalash yashiriladi",
              await page.locator(".card").count() == 35 and await page.locator("#pager").is_hidden()
              and await page.locator("#more").is_hidden())
        await page.fill("#q", "bank")
        # «bank» — KURS0001 sarlavhasida, KURS0002 rejasida (sarlavhadagisi yuqorida).
        await page.wait_for_function("document.querySelectorAll('.card').length === 2", timeout=5000)
        check("brauzer: qidiruv ishlaydi, manzil yangilanadi",
              page.url.endswith("/shop?q=bank") and "/shop/KURS0001/" in await page.locator(".card").first.get_attribute("href"))
        await page.goto(base + "/shop/tur/kurs-ishi", wait_until="networkidle")
        await page.select_option("#sort", "expensive")
        await page.wait_for_timeout(600)
        check("brauzer: saralash tur sahifasida qoladi", "/shop/tur/kurs-ishi?sort=expensive" in page.url
              and await page.locator(".card").count() == 2)
        await page.goto(base + "/shop/KURS0001/bank-tizimi-va-uning-ahamiyati-kurs-ishi", wait_until="networkidle")
        # Sarlavhadagi logotip (/static) bu sinov serverida yo'q — faqat ish rasmlari tekshiriladi.
        broken = await page.evaluate("[...document.images].filter(i => i.src.includes('/shop/') && i.complete && !i.naturalWidth).length")
        overflow = await page.evaluate("document.documentElement.scrollWidth > window.innerWidth")
        check("brauzer: ish sahifasi — rasmlar ochiladi, telefonda yonga siljimaydi", broken == 0 and not overflow,
              f"buzuq rasm: {broken}, yonga siljish: {overflow}")
        await page.screenshot(path=os.path.join(TMP, "item.png"), full_page=True)
        check("brauzer: JavaScript xatosi yo'q", not errors, errors)
        await browser.close()


asyncio.run(main())
shutil.rmtree(TMP, ignore_errors=True) if not FAILED else print("Fayllar:", TMP)
print("\nNATIJA:", "HAMMASI O'TDI" if not FAILED else f"{len(FAILED)} ta xato")
sys.exit(1 if FAILED else 0)
