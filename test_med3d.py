"""Tibbiy 3D taqdimot (med3d): haqiqiy .glb modellar, PowerPoint 3D XML, Morph va yorliqlar.

- GLB: yozish/o'qish, chegara qutisi ko'rinmas "anchor" bilan hamma guruhda bir xil;
- tayyor sahna (assets/med3d/skull): guruhlar va yorliqlar o'rnida;
- kamera: old ko'rinishda bemorning chapi slaydning o'ngida, rotx musbat — tepasi kameraga;
- PPTX: model3d bog'lanishi, model/gltf.binary turi, am3d:model3d (kamera, aylanish), Morph 2-slayddan;
- 1-slaydda bosh suyagi tepasi va jag' bir xil ramkada (aniq yopishadi), 2-slaydda tepa slayddan
  yuqoriga chiqib ketadi, jag' aylanadi; yorliqlar slayd ichida va ustma-ust tushmaydi.

    python test_med3d.py
"""
import os, sys, tempfile, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE); os.chdir(HERE)

FAILS = []
def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond: FAILS.append(name)

import numpy as np
from lxml import etree

from services.med3d import available_scenes, load_scene
from services.med3d import glb
from services.med3d.camera import Camera, IdBuffer
from services.med3d.deck import DEMO_PLAN, W, H, _connector, build
from services.med3d import design
from pptx import Presentation
from pptx.util import Inches

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "am3d": "http://schemas.microsoft.com/office/drawing/2017/model3d",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main"}
EMU = 914400

print("1) Ko'rsatkich koordinatalari")
pointer_prs = Presentation()
pointer_slide = pointer_prs.slides.add_slide(pointer_prs.slide_layouts[6])
pointer_start, pointer_end = (2.1, 3.4), (8.6, 5.2)
pointer_shapes = _connector(pointer_slide, design.style("atlas"), "anchor_test",
                             pointer_start, pointer_end)
pointer_xml = etree.fromstring(pointer_shapes[-1]._element.xml.encode())
off = pointer_xml.find(".//a:xfrm/a:off", NS)
ext = pointer_xml.find(".//a:xfrm/a:ext", NS)
line_start = (int(off.get("x")), int(off.get("y")))
line_end = (line_start[0] + int(ext.get("cx")), line_start[1] + int(ext.get("cy")))
dot = pointer_shapes[0]
check("ko'rsatkich to'g'ri chiziq bo'lib anchor koordinatasida tugaydi",
      pointer_shapes[-1].shape_type == 9
      and np.allclose(line_start, [Inches(v) for v in pointer_start])
      and np.allclose(line_end, [Inches(v) for v in pointer_end]),
      (line_start, line_end, pointer_start, pointer_end))
check("marker markazi ko'rsatkichning oxirgi nuqtasiga teng",
      np.allclose([dot.left + dot.width / 2, dot.top + dot.height / 2], [Inches(v) for v in pointer_end]))

print("1) GLB yozish va o'qish")
tri = glb.Primitive(np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0]], float), np.tile([0, 0, 1.0], (3, 1)),
                    np.array([0, 1, 2]), (0.9, 0.8, 0.7))
data = glb.write_glb([tri], anchors=([-2, -3, -4], [5, 6, 7]))
doc, binary = glb.read_glb(data)
check("GLB sarlavhasi va uzunligi to'g'ri", data[:4] == b"glTF" and len(data) % 4 == 0)
lo, hi = glb.bounds(data)
check("anchor chegara qutisini belgilaydi", list(lo) == [-2, -3, -4] and list(hi) == [5, 6, 7], (lo, hi))
pts = np.frombuffer(binary, np.float32, 6 * 3, doc["bufferViews"][0]["byteLength"] - 6 * 12).reshape(6, 3)
areas = [np.linalg.norm(np.cross(t[1] - t[0], t[2] - t[0])) / 2 for t in (pts[:3], pts[3:])]
check("anchor: ikkita haqiqiy (PowerPoint tashlab yubormaydi), lekin piksel ulushicha uchburchak",
      doc["accessors"][-1]["count"] == 9 and all(0 < a < 1e-4 * 11 ** 2 for a in areas), areas)
check("material suyak rangida", doc["materials"][0]["pbrMetallicRoughness"]["baseColorFactor"][:3] == [0.9, 0.8, 0.7])

print("2) Tayyor sahna")
check("skull sahnasi bor", "skull" in available_scenes(), available_scenes())
scene = load_scene("skull")
check("guruhlar: cranium va mandible", set(scene["groups"]) == {"cranium", "mandible"}, list(scene["groups"]))
boxes = []
for g, info in scene["groups"].items():
    with open(os.path.join(scene["dir"], info["file"]), "rb") as fh:
        raw = fh.read()
    glb.read_glb(raw)
    boxes.append(glb.bounds(raw))
    check(f"{g}.glb hajmi kichik (< 3 MB)", len(raw) < 3e6, len(raw))
check("hamma guruhning chegara qutisi bir xil (qismlar aniq yopishadi)",
      all(np.allclose(b[0], boxes[0][0]) and np.allclose(b[1], boxes[0][1]) for b in boxes))
check("chegara qutisi sahna.json bilan bir xil",
      np.allclose(boxes[0][0], scene["bbox_min"], atol=1e-4) and np.allclose(boxes[0][1], scene["bbox_max"], atol=1e-4))
check("yorliqlar o'zbekcha nom va guruhga ega",
      all(v["uz"] and v["group"] in scene["groups"] and len(v["samples"]) > 10 for v in scene["labels"].values()))

print("3) Kamera")
cam = Camera(scene["bbox_min"], scene["bbox_max"])
left_u, _ = cam.anchor(scene["labels"]["temporal_l"]["samples"], (0, 0, 0))
right_u, _ = cam.anchor(scene["labels"]["temporal_r"]["samples"], (0, 0, 0))
check("old ko'rinish: bemorning chapi slaydning o'ngida", left_u > 0 > right_u, (left_u, right_u))
_, front_v = cam.anchor(scene["labels"]["incisors_lower"]["samples"], (70, 0, 0))
_, back_v = cam.anchor(scene["labels"]["molars_lower_l"]["samples"], (70, 0, 0))
check("rotx=70: jag' yuqoridan — kesuv tishlar pastda, oziq tishlar tepada", front_v > back_v, (front_v, back_v))
u0, v0, u1, v1 = cam.extent([p for g in scene["groups"].values() for p in g["samples"]], (0, 0, 0))
check("butun model ramka ichida", -0.5 < u0 < u1 < 0.5 and -0.5 < v0 < v1 < 0.5, (u0, v0, u1, v1))

print("4) Mayda yorliq koordinatasi")
tiny_xy = np.array([(6.0 + x, 3.0 + y) for x in (0, 0.019, 0.038) for y in (0, 0.012, 0.024)])
tiny_layers = [(('scene', 'tiny'), tiny_xy[:, 0], tiny_xy[:, 1], np.ones(len(tiny_xy)))]
coarse_anchor = IdBuffer(tiny_layers, W, H, dpi=40).anchor(('scene', 'tiny'))
precise_anchor = IdBuffer(tiny_layers, W, H, dpi=80).anchor(('scene', 'tiny'))
check("yuqori aniqlikdagi z-buffer mayda organ yorlig'ini topadi",
      coarse_anchor is None and precise_anchor is not None, (coarse_anchor, precise_anchor))
check("mayda yorliq uchi aniq koordinataga yaqin",
      precise_anchor is not None and abs(precise_anchor[0] - 6.019) < 1 / 80
      and abs(precise_anchor[1] - 3.012) < 1 / 80, precise_anchor)

print("4) PPTX")
out = os.path.join(tempfile.mkdtemp(), "skull.pptx")
build(DEMO_PLAN, out)
z = zipfile.ZipFile(out)
names = z.namelist()
ct = z.read("[Content_Types].xml").decode()
glbs = [n for n in names if n.endswith(".glb")]
check("bir xil GLB slaydlar orasida bitta faylda (6 ta model → 2 ta fayl)", len(glbs) == 2, glbs)
check("GLB turi model/gltf.binary", ct.count('ContentType="model/gltf.binary"') == 2)
rels = [z.read(f"ppt/slides/_rels/slide{i}.xml.rels").decode() for i in (1, 2, 3)]
check("har slayd GLB ga bog'langan", all(r.count("relationships/model3d") == 2 for r in rels))

def frames(i):
    root = etree.fromstring(z.read(f"ppt/slides/slide{i}.xml"))
    rels = z.read(f"ppt/slides/_rels/slide{i}.xml.rels").decode()
    res = {}
    for gf in root.iterfind(".//p:graphicFrame", NS):
        name = gf.find(".//p:cNvPr", NS).get("name")
        off, ext = gf.find("p:xfrm/a:off", NS), gf.find("p:xfrm/a:ext", NS)
        model = gf.find(".//am3d:model3d", NS)
        rot = model.find(".//am3d:rot", NS)
        rid = model.get(f"{{{NS['r']}}}embed")
        res[name] = dict(x=int(off.get("x")) / EMU, y=int(off.get("y")) / EMU, size=int(ext.get("cx")) / EMU,
                         rot=tuple(int(rot.get(k)) / 60000 for k in ("ax", "ay", "az")),
                         linked=f'Id="{rid}" Type="http://schemas.microsoft.com/office/2017/06/relationships/model3d"' in rels,
                         camera=model.find("am3d:camera/am3d:perspective", NS) is not None)
    return root, res

root1, f1 = frames(1)
root2, f2 = frames(2)
root3, f3 = frames(3)
check("1-slayd: ikki model", set(f1) == {"!!m3d_cranium", "!!m3d_mandible"}, list(f1))
check("model3d GLB ga bog'langan va kamerasi bor", all(v["linked"] and v["camera"] for f in (f1, f2, f3) for v in f.values()))
check("1-slayd: tepa va jag' bir xil ramkada", f1["!!m3d_cranium"] == f1["!!m3d_mandible"])
centre = lambda f: f["x"] + f["size"] / 2   # noqa: E731
check("2-slayd: tepa slayddan yuqoriga chiqib ketadi",
      f2["!!m3d_cranium"]["y"] + f2["!!m3d_cranium"]["size"] < 0
      and abs(centre(f2["!!m3d_cranium"]) - centre(f1["!!m3d_cranium"])) < 1e-3, f2["!!m3d_cranium"])
check("2-slayd: chiqib ketayotgan tepa aylanib-buralib, kichrayib ketadi (3D ekani bilinadi)",
      f2["!!m3d_cranium"]["rot"] != f1["!!m3d_cranium"]["rot"]
      and f2["!!m3d_cranium"]["size"] < f1["!!m3d_cranium"]["size"], f2["!!m3d_cranium"])
check("2-slayd: jag' aylanib kattalashadi",
      f2["!!m3d_mandible"]["rot"][0] == 70 and f2["!!m3d_mandible"]["size"] > f1["!!m3d_mandible"]["size"])
check("3-slayd: qismlar qayta birlashadi (yon ko'rinish)",
      f3["!!m3d_cranium"] == f3["!!m3d_mandible"] and f3["!!m3d_cranium"]["rot"][1] == 90)
check("Morph 1-slaydda yo'q, 2- va 3-slaydda bor",
      not root1.findall(".//p159:morph", NS) and root2.findall(".//p159:morph", NS) and root3.findall(".//p159:morph", NS))
connector_names = [{n for n in root.xpath(".//p:cNvPr/@name", namespaces=NS)
                    if n.endswith(("_line", "_dot"))} for root in (root1, root2, root3)]
check("Morph paytida ko'rsatkich chiziqlari organ ustida sirg'almaydi",
      all(not connector_names[i] & connector_names[i + 1] for i in range(len(connector_names) - 1)))
order = [(etree.QName(el).localname, (el.xpath("./*/p:cNvPr/@name|.//p:cNvPr/@name", namespaces=NS) or [""])[0])
         for el in root1.find(".//p:spTree", NS)]
models_at = [i for i, (kind, _) in enumerate(order) if kind == "AlternateContent"]
title_at = next(i for i, (_, name) in enumerate(order) if name == "!!title")
back_at = [i for i, (_, name) in enumerate(order) if name.startswith("!!deco_") or name in ("!!panel", "!!shadow")]
check("modellar matn ortida, fon bezaklari esa modellar ortida",
      models_at and max(models_at) < title_at and back_at and max(back_at) < min(models_at), order[:8])

from pptx import Presentation
prs = Presentation(out)
for i, slide in enumerate(prs.slides, 1):
    # Keyingi slayd yorliqlari bu slaydda chetda kutib turadi (Morph ularni suzib olib kiradi) — sanalmaydi.
    pills = [s for s in slide.shapes if s.name.startswith("lbl_") and not s.name.endswith(("_line", "_dot", "_la"))
             and 0 <= s.left / EMU < W]
    want = len(DEMO_PLAN["slides"][i - 1]["labels"])
    check(f"{i}-slayd: hamma yorliq bor ({want})", len(pills) == want, len(pills))
    inside = all(0 <= p.left / EMU and p.left / EMU + p.width / EMU <= W and 0 <= p.top / EMU
                 and p.top / EMU + p.height / EMU <= H for p in pills)
    check(f"{i}-slayd: yorliqlar slayd ichida", inside)
    rects = [(p.left, p.top, p.left + p.width, p.top + p.height) for p in pills]
    overlap = any(a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]
                  for k, a in enumerate(rects) for b in rects[k + 1:])
    check(f"{i}-slayd: yorliqlar ustma-ust tushmaydi", not overlap)
    check(f"{i}-slayd: manba (CC BY) izohda, slaydda ko'rinmaydi",
          "BodyParts3D" in slide.notes_slide.notes_text_frame.text
          and not any("BodyParts3D" in (s.text_frame.text if s.has_text_frame else "") for s in slide.shapes))

print("5) Hazm organlari: 12 slayd, uchta sahna")
from services.med3d.deck import MAX_FRAME
from services.med3d.plans import DIGESTIVE_PLAN
out2 = os.path.join(tempfile.mkdtemp(), "hazm.pptx")
build(DIGESTIVE_PLAN, out2)
z = zipfile.ZipFile(out2)
check("12 slayd", sum(1 for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")) == 12)
all_frames = {i: frames(i)[1] for i in range(1, 13)}
check("ramkalar MAX_FRAME dan katta emas",
      all(v["size"] <= MAX_FRAME + 1e-3 for f in all_frames.values() for v in f.values()))
on_slide = lambda f: {n for n, v in f.items() if -v["size"] < v["x"] < W and -v["size"] < v["y"] < H}
check("og'iz slaydi o'z sahnasida (Morph boshqa sahnaga ulanmaydi)",
      on_slide(all_frames[3]) == {"!!m3d_digestive_head_ogiz"}, on_slide(all_frames[3]))
check("sahna almashganda eskilari slayddan tashqariga uchadi, yangilari tashqaridan keladi",
      "!!m3d_jigar" in all_frames[3] and "!!m3d_jigar" not in on_slide(all_frames[3])
      and "!!m3d_digestive_head_ogiz" in all_frames[2] and "!!m3d_digestive_head_ogiz" not in on_slide(all_frames[2]))
check("7-slayd: jigar yuqoriga chiqib ketadi, o't pufagi qoladi",
      all_frames[7]["!!m3d_digestive_abdomen_jigar"]["y"] + all_frames[7]["!!m3d_digestive_abdomen_jigar"]["size"] < 0
      and "!!m3d_digestive_abdomen_ot" in all_frames[7])
prs2 = Presentation(out2)
for i, slide in enumerate(prs2.slides, 1):
    pills = [s for s in slide.shapes if s.name.startswith("lbl_") and not s.name.endswith(("_line", "_dot", "_la"))
             and 0 <= s.left / EMU < W]
    if len(pills) != len(DIGESTIVE_PLAN["slides"][i - 1].get("labels") or []):
        check(f"hazm {i}-slayd: hamma yorliq bor", False, len(pills))

print("Z-Anatomy: ko'zgu nusxalar ichi-tashi teskari emas")
from services.med3d import manba as _manba
_at = _manba.atlas("male")
_bad = {"left": 0, "right": 0}
for _pid, _info in _at.by_id.items():
    _side = "left" if "-left-" in _pid else ("right" if "-right-" in _pid else "")
    if not _pid.startswith("ZA-") or not _side or "muscle" not in _pid:
        continue
    _p, _n, _i = _at._one(_info)
    _t = _i.reshape(-1, 3)
    _vol = np.einsum("ij,ij->i", _p[_t[:, 0]], np.cross(_p[_t[:, 1]], _p[_t[:, 2]])).sum()
    _bad[_side] += int(_vol < 0)
check("chap va o'ng mushaklar bir xil yo'nalishda (chap tomoni teskari emas)", _bad["left"] <= _bad["right"] + 2, _bad)

print()
print("HAMMASI O'TDI" if not FAILS else f"XATOLAR: {len(FAILS)}")
sys.exit(1 if FAILS else 0)
