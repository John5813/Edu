"""Yangi o'zbek lotin alifbosi (oʻ→ö, gʻ→ğ, sh→ş, ch→ç, tutuq ʼ): o'girish, fayllar va oqimlar.

    python test_yangi_alifbo.py
"""
import asyncio
import os
import sys
import tempfile
from types import SimpleNamespace
from unittest.mock import AsyncMock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.chdir(HERE)
os.environ.setdefault("BOT_TOKEN", "1:x")

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  XATO ") + name + ("" if cond else f" — {detail}"))
    if not cond:
        FAILS.append(name)


from services import uz_script as u

TMP = tempfile.mkdtemp()

print("1) O'girish")
pairs = {
    "ko'cha, tog', shahar, chiroq": "köça, toğ, şahar, çiroq",
    "koʻcha va togʻ": "köça va toğ",
    "O'ZBEKISTON RESPUBLIKASI": "ÖZBEKISTON RESPUBLIKASI",
    "Shahar va Choyxona": "Şahar va Çoyxona",
    "ma'no, ta'lim, e'tibor, Is'hoq": "maʼno, taʼlim, eʼtibor, Isʼhoq",
    "g'oya, go'zal, bog'": "ğoya, gözal, boğ",
    "mo''jiza": "möʼjiza",
    "mashhur, ng: tong, keng": "maşhur, ng: tong, keng",
    "Кўча, тоғ, шаҳар, чироқ, маъно": "Köça, toğ, şahar, çiroq, maʼno",
    "«Shoh asar» – O'zbekiston": "«Şoh asar» – Özbekiston",
}
for old, new in pairs.items():
    check(f"{old} → {new}", u.to_latin_new(old) == new, u.to_latin_new(old))
keep = "CH4 va H2O, $x^2 + ch$, https://shop.uz/chat, info@chat.uz, Washington, Chicago"
check("formula, havola, pochta, chet so'zlar tegilmaydi", u.to_latin_new(keep) == keep, u.to_latin_new(keep))
eng = "The children should change the school"
check("inglizcha jumla tegilmaydi", u.to_latin_new(eng) == eng, u.to_latin_new(eng))
once = u.to_latin_new("O'zbekiston shahar ma'no")
check("ikki marta qo'llansa ham o'zgarmaydi", u.to_latin_new(once) == once, once)
check("yangi → kirill", u.to_cyrillic("Özbekiston şahar çiroq toğ maʼno") == "Ўзбекистон шаҳар чироқ тоғ маъно",
      u.to_cyrillic("Özbekiston şahar çiroq toğ maʼno"))
check("yangi → eski lotin", u.to_latin_old("Özbekiston ŞAHAR çiroq toğ maʼno") == "O'zbekiston SHAHAR chiroq tog' ma'no",
      u.to_latin_old("Özbekiston ŞAHAR çiroq toğ maʼno"))
check("convert(LATIN_NEW)", u.convert("shahar", u.LATIN_NEW) == "şahar")
check("eski lotin o'zgarmadi", u.convert("Шаҳар", u.LATIN) == "Shahar")

print("2) Til kodlari")
check("uz-yangi → yangi lotin", u.script_of_language(u.UZ_NEW_LANG) == u.LATIN_NEW)
check("uz → lotin (sukut)", u.script_of_language("uz") == u.LATIN)
token = u.use(u.LATIN_NEW)
check("uz + use(LATIN_NEW) → yangi lotin", u.script_of_language("uz") == u.LATIN_NEW)
check("modelga lotin qoidasi", "UZBEK LATIN" in u.prompt_rule(u.current()))
u.reset(token)
check("reset'dan keyin oddiy lotin", u.script_of_language("uz") == u.LATIN)
check("base_language", u.base_language(u.UZ_NEW_LANG) == "uz" and u.base_language("ru") == "ru")
check("hujjat tugmasi uzn", u.doc_language_choice("uzn") == ("uz", True) and u.doc_language_choice("ru") == ("ru", False))

print("3) DOCX va PPTX")
from docx import Document
from pptx import Presentation
from pptx.util import Inches

docx_path = os.path.join(TMP, "a.docx")
doc = Document()
doc.add_paragraph("O'zbekiston shahar chiroq")
doc.add_table(rows=1, cols=1).cell(0, 0).text = "Ta'lim tizimi"
doc.sections[0].header.paragraphs[0].text = "Kurs ishi: g'oya"
doc.save(docx_path)
changed = u.normalize_file(docx_path, u.LATIN_NEW)
doc = Document(docx_path)
check("docx: asosiy matn", doc.paragraphs[0].text == "Özbekiston şahar çiroq", doc.paragraphs[0].text)
check("docx: jadval", doc.tables[0].cell(0, 0).text == "Taʼlim tizimi", doc.tables[0].cell(0, 0).text)
check("docx: kolontitul", doc.sections[0].header.paragraphs[0].text == "Kurs işi: ğoya",
      doc.sections[0].header.paragraphs[0].text)
check("docx: o'zgarganlar soni", changed >= 3, changed)

pptx_path = os.path.join(TMP, "a.pptx")
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Toshkent shahri"
prs.save(pptx_path)
u.normalize_file(pptx_path, u.LATIN_NEW)
check("pptx: sarlavha", Presentation(pptx_path).slides[0].shapes.title.text == "Toşkent şahri")
check("boshqa fayl turi tegilmaydi", u.normalize_file(os.path.join(TMP, "x.pdf"), u.LATIN_NEW) == 0)


class FakeService:
    async def create_x(self, path):
        return path

    def create_sync(self, path):
        return path


u.wrap_file_methods(FakeService, prefix="create_")


async def wrapped():
    p1 = os.path.join(TMP, "b.docx")
    d = Document(); d.add_paragraph("shahar"); d.save(p1)
    await FakeService().create_x(p1)
    check("o'ralgan metod: yangi alifbo tanlanmagan — tegilmaydi", Document(p1).paragraphs[0].text == "shahar")
    t = u.use(u.LATIN_NEW)
    try:
        await FakeService().create_x(p1)
    finally:
        u.reset(t)
    check("o'ralgan metod: yangi alifbo — o'giriladi", Document(p1).paragraphs[0].text == "şahar")
    from services.document_service import DocumentService
    check("DocumentService.create_referat o'ralgan", DocumentService.create_referat.__wrapped__ is not None)

asyncio.run(wrapped())

print("4) Bot: tugma, middleware, navbat")
from bot.keyboards import get_doc_language_keyboard
datas = [b.callback_data for r in get_doc_language_keyboard("uz").inline_keyboard for b in r]
check("hujjat tili tugmalarida yangi alifbo", "doc_lang_uzn" in datas, datas)


async def middleware():
    from bot.middlewares import UzScriptMiddleware
    seen = []

    async def handler(event, data):
        seen.append(u.current())

    class State:
        def __init__(self, d): self.d = d
        async def get_data(self): return self.d

    mw = UzScriptMiddleware()
    event = SimpleNamespace(from_user=SimpleNamespace(id=5))
    await mw(handler, event, {"state": State({"topic": "x", "uz_new": True})})
    await mw(handler, event, {"state": State({"topic": "x", "uz_new": False})})
    from bot import checkout
    checkout.remember(5, "doc", {"topic": "x", "uz_new": True})
    await mw(handler, event, {"state": State({})})
    checkout.forget(5)
    check("middleware: uz_new → yangi lotin, aks holda yo'q, to'lovdan keyin ham",
          seen == [u.LATIN_NEW, "", u.LATIN_NEW], seen)
    check("middleware'dan keyin kontekst tozalandi", u.current() == "")

    from bot.queue_service import DocumentQueue, HeavyDocTask
    t = u.use(u.LATIN_NEW)
    got = []

    async def job():
        got.append(u.current())
    task = HeavyDocTask(task_id="t1", coro_factory=job, user_telegram_id=1, chat_id=1, lang="uz",
                        doc_type="course_work", topic="x", bot=None)
    u.reset(t)
    queue = DocumentQueue()
    worker = asyncio.create_task(queue._worker())
    await queue.enqueue(task)
    await asyncio.wait_for(task.done_event.wait(), 5)
    worker.cancel()
    check("navbat: yozuv ishchiga o'tadi", got == [u.LATIN_NEW], got)

asyncio.run(middleware())


async def documents_choice():
    from bot.handlers import documents
    store = {}

    class State:
        async def update_data(self, **kw): store.update(kw)
        async def get_data(self): return dict(store)
        async def set_state(self, *_): pass

    msg = SimpleNamespace(delete=AsyncMock(), answer=AsyncMock())
    cb = SimpleNamespace(data="doc_lang_uzn", answer=AsyncMock(), message=msg)
    await documents.handle_doc_language_selection(cb, State(), "uz")
    check("hujjat: uzn → doc_language=uz, uz_new", store.get("doc_language") == "uz" and store.get("uz_new") is True, store)
    cb.data = "doc_lang_ru"
    await documents.handle_doc_language_selection(cb, State(), "uz")
    check("hujjat: boshqa til — uz_new o'chadi", store.get("doc_language") == "ru" and store.get("uz_new") is False, store)

asyncio.run(documents_choice())

print("5) Premium va «Rahmat» slaydi")
from bot.handlers import premium_presentation as pp
check("premium til tugmasi", u.UZ_NEW_LANG in pp._LANG_BUTTONS)
check("xulosadagi yozuv", pp._language_label(u.UZ_NEW_LANG) == "UZ (yangi alifbo)")
kb = pp._confirm_keyboard("uz", u.UZ_NEW_LANG)
check("xulosada tanlangan ✓", any(b.text.startswith("✓ ") and b.callback_data == f"prem_ppt_lang:{u.UZ_NEW_LANG}"
                                  for r in kb.inline_keyboard for b in r))
from services.thanks_anim.slide import texts
check("rahmat slaydi yangi alifboda", texts(u.UZ_NEW_LANG)[:2] == ("SÖNGGI SÖZ", "Eʼtiboringiz uçun rahmat!"),
      texts(u.UZ_NEW_LANG))

print("6) Sayt")
from services import web_jobs, web_kinds
check("sayt: tillar ro'yxatida", u.UZ_NEW_LANG in web_kinds.LANGUAGES and u.UZ_NEW_LANG in web_kinds.DOC_LANGUAGES)


async def site():
    seen = {}

    async def run(params, report):
        seen.update(language=params["language"], topic=params["topic"], script=u.current())
        p = os.path.join(TMP, "s.docx")
        d = Document(); d.add_paragraph("Ta'lim va shahar"); d.save(p)
        return p, "s.docx"

    kind = SimpleNamespace(run=run)
    path, _ = await web_jobs._run_kind(kind, {"language": u.UZ_NEW_LANG, "topic": "Özbekiston tarixi"}, lambda *a: None)
    check("sayt: ichkarida uz + eski lotin mavzu + yangi yozuv konteksti",
          seen == {"language": "uz", "topic": "O'zbekiston tarixi", "script": u.LATIN_NEW}, seen)
    check("sayt: tayyor fayl yangi alifboda", Document(path).paragraphs[0].text == "Taʼlim va şahar")
    check("sayt: kontekst tozalandi", u.current() == "")
    seen.clear()
    await web_jobs._run_kind(kind, {"language": "uz", "topic": "Tarix"}, lambda *a: None)
    check("sayt: oddiy uz o'zgarmaydi", seen["language"] == "uz" and seen["script"] == "", seen)

asyncio.run(site())

print("\nNATIJA:", "HAMMASI O'TDI" if not FAILS else f"{len(FAILS)} ta xato: {FAILS}")
sys.exit(1 if FAILS else 0)
